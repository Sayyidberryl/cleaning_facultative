"""
================================================================================
 SCRIPT CLEANSING DATA 3 - SUSPENSE (CEDANT : IBS)
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
INPUT_FILE = "3b. Database Suspense.xlsx"
SHEET_NAME = "Sheet1"
HEADER_ROW = 2
OUTPUT_FILE = "ibs_output_suspense_V2.xlsx"
CEDANT_FILTER = "IBS"
STATUS_FILTER = "SUSPENSE"

MAX_BREAKDOWN_CODES = 5
MAX_BREAKDOWN_INSURED = 5
TEXT_FORMAT_COLUMN_KEYWORDS = ("POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO", "CERTIFICATE")
RE_I = re.IGNORECASE

# ==============================================================================
# 1. KODE RESMI IBS & CHECKER
# ==============================================================================
_REAL_CODE_PATTERN_RE = re.compile(
    r"IBS\s*RE|IBSRESG|BSRE|MCOC|AIG[\s\-]*MAP|MD[\s\-]*C\d|11-F\d|[A-Z]{1,10}\s*IBS\s*\d{4}|OC(?:\.|\s)*\d{2}", RE_I
)

def _has_real_code(text: str) -> bool:
    return bool(_REAL_CODE_PATTERN_RE.search(text))

def _compact_if_real_code(text: str) -> str:
    if _has_real_code(text):
        return re.sub(r"[\s.:/\-,]+", "", text).upper()
    return text.upper()

# ==============================================================================
# 2. CLEANSING INSURED
# ==============================================================================
_ENTITY_TITLE_RE = re.compile(r"\b(?:P\.?T\.?|C\.?V\.?|TBK\.?|LTD\.?|PTE\.?|N\.?V\.?|BAPAK|IBU|NY\.?|NYONYA|MR\.?|MRS\.?|MS\.?)\b", RE_I)
_PERSERO_PAREN_RE = re.compile(r"\(\s*PERSERO\s*\)", RE_I)
_AMPERSAND_WHITELIST = [
    "BUDI STARCH & SWEETENER",
    "ULTRAJAYA MILK INDUSTRY & TRADING CO",
]
_LEGAL_ABBREV_PATTERNS = [r"N\.A\.", r"N\.V\."]
_ENTITY_SPLIT_QQ_SLASH_RE = re.compile(r"\bQQ\b|/", RE_I)
_ENTITY_SPLIT_COMMA_AMP_RE = re.compile(r",|&")
_COMMA_BEFORE_NA_NV_RE = re.compile(r",\s*(N\.?A\.?|N\.?V\.?)\b", RE_I)
_BOILERPLATE_PREFIX_RE = re.compile(r"^\s*PERHIMPUNAN\s+PENGHUNI\s+RUMAH\s+SUSUN\s+", RE_I)

def _strip_boilerplate_prefix(text: str) -> str:
    return _BOILERPLATE_PREFIX_RE.sub("", text)

def _protect_whitelist(text: str) -> tuple:
    placeholders = {}
    upper = text.upper()
    for i, phrase in enumerate(_AMPERSAND_WHITELIST):
        if phrase in upper:
            key = f"__WL{i}__"
            text = re.compile(re.escape(phrase), RE_I).sub(key, text)
            placeholders[key] = phrase
    return text, placeholders

def _restore_whitelist(text: str, placeholders: dict) -> str:
    for key, phrase in placeholders.items():
        text = text.replace(key, phrase)
    return text

def _split_insured_entities(raw: str) -> list:
    text = raw.strip()
    if not text:
        return []

    text = _strip_boilerplate_prefix(text)
    text, placeholders = _protect_whitelist(text)
    text = _COMMA_BEFORE_NA_NV_RE.sub(r" __COMMA_NA__\1", text)

    entities = []
    for part in _ENTITY_SPLIT_QQ_SLASH_RE.split(text):
        for sub in _ENTITY_SPLIT_COMMA_AMP_RE.split(part):
            sub = _restore_whitelist(sub.strip().replace("__COMMA_NA__", ", "), placeholders)
            if sub:
                entities.append(sub)
    return entities

def _clean_single_insured(name: str) -> str:
    text = name.strip()
    protected_abbrev = {}
    for i, pat in enumerate(_LEGAL_ABBREV_PATTERNS):
        m = re.search(pat, text, RE_I)
        if m:
            key = f"__ABBR{i}__"
            protected_abbrev[key] = m.group(0).upper()
            text = re.sub(pat, key, text, flags=RE_I)

    text = _ENTITY_TITLE_RE.sub("", text)
    text = _PERSERO_PAREN_RE.sub("", text)
    text = re.sub(r"^[-/]+|[-/]+$", "", re.sub(r"[,;:.]+", " ", text))

    for key, val in protected_abbrev.items():
        text = text.replace(key, val)

    return re.sub(r"\s+", " ", text).strip().upper()

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return []

    original = text.strip()
    entities = _split_insured_entities(original)
    if not entities:
        return []

    cleaned = [c for c in (_clean_single_insured(e) for e in entities) if c]
    if not cleaned:
        return []
    return [original] if len(cleaned) > MAX_BREAKDOWN_INSURED else cleaned

# ==============================================================================
# 3. CLEANSING POLIS, SLIP, & CERTIFICATE
# ==============================================================================
_POLIS_INSTALLMENT_SUFFIX_RE = re.compile(r"-\d{1,2}/\d{1,2}$")
_POLIS_CERTIFICATE_SUFFIX_RE = re.compile(r"-(?P<cert>(?:\d+\s*(?:-|S/?D)\s*\d+)|\d+)\s*$", RE_I)

_SLIP_NO_PREFIX_RE = re.compile(r"^NO\s+", RE_I)
_SLIP_NO_DOTTED_NUMBER_RE = re.compile(r"^NO\s+[\d.]+$", RE_I)
_SLIP_SUFFIX_RE = re.compile(r"-[A-Z]{2,}\d+$", RE_I)
_SLIP_DECLARATION_CODE_RE = re.compile(r"OC\.\d+\.\d+\.\d+\.\d+", RE_I)

def _format_certificate_string(cert_raw: str) -> str:
    if not cert_raw or pd.isna(cert_raw):
        return ""
    cert_raw = str(cert_raw).upper().replace("S/D", "-").replace("SD", "-")
    cert_raw = re.sub(r"\s+", " ", cert_raw).strip()

    range_match = re.match(r"^\s*(\d+)\s*-\s*(\d+)\s*$", cert_raw)
    if range_match:
        start_str, end_str = range_match.group(1), range_match.group(2)
        if len(start_str) > 6 or len(end_str) > 6:
            return ""

        start_num = int(start_str)
        end_num = int(end_str)

        start_pad = str(start_num).zfill(6)
        end_pad = str(end_num).zfill(6)

        if start_num > end_num:
            return ""

        diff = end_num - start_num
        if diff > 2:
            return f"{start_pad} SD {end_pad}"
        else:
            parts = [str(i).zfill(6) for i in range(start_num, end_num + 1)]
            return ", ".join(parts)
    else:
        clean_num = re.sub(r"[^\d]", "", cert_raw)
        if not clean_num or len(clean_num) > 6:
            return ""
        return clean_num.zfill(6)

def clean_polis_code(text: str) -> tuple:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return [], ""

    t = re.sub(r"\s+", " ", text.strip())

    cert_match = _POLIS_CERTIFICATE_SUFFIX_RE.search(t)
    if cert_match:
        cert_raw = cert_match.group("cert")
        cert_formatted = _format_certificate_string(cert_raw)
        polis_base = _POLIS_CERTIFICATE_SUFFIX_RE.sub("", t)

        if _has_real_code(polis_base):
            return [re.sub(r"[\s.:/\-,]+", "", polis_base).upper()], cert_formatted
        return [polis_base.upper()], cert_formatted

    t = _POLIS_INSTALLMENT_SUFFIX_RE.sub("", t)

    if t:
        if _has_real_code(t):
            return [re.sub(r"[\s.:/\-,]+", "", t).upper()], ""
        return [t.upper()], ""
    return [], ""

def clean_slip_code(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return []

    t = re.sub(r"\s+", " ", text.strip())
    t = _SLIP_SUFFIX_RE.sub("", t)

    decl_match = _SLIP_DECLARATION_CODE_RE.search(t)
    if decl_match:
        return [_compact_if_real_code(decl_match.group(0))]

    if _SLIP_NO_DOTTED_NUMBER_RE.match(t):
        t = _SLIP_NO_PREFIX_RE.sub("", t)
        t = t.replace(".", "")
        return [_compact_if_real_code(t)] if t else []

    return [_compact_if_real_code(t)] if t else []

# ==============================================================================
# 4. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_cedant(df_raw: pd.DataFrame, keyword: str) -> pd.DataFrame:
    cedant_cols = [c for c in df_raw.columns if "CEDANT" in str(c).upper()]
    if not cedant_cols:
        raise ValueError("Kolom CEDANT tidak ditemukan di source file.")

    mask = pd.Series(False, index=df_raw.index)
    for col in cedant_cols:
        mask |= df_raw[col].astype(str).str.contains(keyword, case=False, na=False)

    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant mengandung '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered

def filter_status(df_raw: pd.DataFrame, keyword: str) -> pd.DataFrame:
    status_col = next((c for c in df_raw.columns if str(c).upper().strip() == "STATUS"), None)
    if status_col is None:
        raise ValueError("Kolom STATUS tidak ditemukan di source file.")

    mask = df_raw[status_col].astype(str).str.strip().str.upper() == keyword.upper()
    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter status = '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered

# ==============================================================================
# 5. BUILD OUTPUT
# ==============================================================================
def _detect_key_columns(df: pd.DataFrame) -> tuple:
    col_insured = next((c for c in df.columns if "INSURED" in str(c).upper()), None)
    col_polis = next((c for c in df.columns if str(c).upper().strip() == "POLIS"), None)
    col_slip = next((c for c in df.columns if "SLIP" in str(c).upper()), None)

    if None in (col_insured, col_polis, col_slip):
        raise ValueError(f"Kolom kunci tidak lengkap. INSURED={col_insured}, POLIS={col_polis}, SLIP={col_slip}")
    print(f"[INFO] Kolom terdeteksi -> INSURED='{col_insured}', POLIS='{col_polis}', SLIP='{col_slip}'")
    return col_insured, col_polis, col_slip

def _breakdown_all_rows(df: pd.DataFrame, col_insured: str, col_polis: str, col_slip: str) -> tuple:
    insured_cln_all, polis_cln_all, cert_cln_all, slip_cln_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for _, row in df.iterrows():
        ins_val, pol_val, slp_val = row[col_insured], row[col_polis], row[col_slip]
        ins_cln = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []
        pol_cln, cert_formatted = clean_polis_code(str(pol_val)) if pd.notna(pol_val) else ([], "")
        slp_cln = clean_slip_code(str(slp_val)) if pd.notna(slp_val) else []

        insured_cln_all.append(ins_cln)
        polis_cln_all.append(pol_cln)
        cert_cln_all.append(cert_formatted)
        slip_cln_all.append(slp_cln)
        max_ins, max_pol = max(max_ins, len(ins_cln)), max(max_pol, len(pol_cln))
        max_slp = max(max_slp, len(slp_cln))

    return insured_cln_all, polis_cln_all, cert_cln_all, slip_cln_all, max_ins, max_pol, max_slp

def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_insured, col_polis, col_slip = _detect_key_columns(df_filtered)
    breakdown_res = _breakdown_all_rows(df_filtered, col_insured, col_polis, col_slip)
    insured_cln_all, polis_cln_all, cert_cln_all, slip_cln_all, max_ins, max_pol, max_slp = breakdown_res

    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)

    output_cols_data = {}
    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        if col == col_insured:
            for i in range(1, max_ins + 1):
                output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in insured_cln_all]
        if col == col_polis:
            if max_pol >= 1:
                output_cols_data["POLIS_CLN_1"] = [lst[0] if len(lst) >= 1 else "" for lst in polis_cln_all]
            output_cols_data["CERTIFICATE"] = cert_cln_all
            for i in range(2, max_pol + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in polis_cln_all]
        if col == col_slip:
            for i in range(1, max_slp + 1):
                output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slip_cln_all]

    df_master = pd.DataFrame(output_cols_data)
    drop_cols = [c for c in df_master.columns if "_CLN_" in c and df_master[c].astype(str).str.strip().eq("").all()]
    return df_master.drop(columns=drop_cols)

# ==============================================================================
# 6. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    workbook = load_workbook(output_path)
    worksheet = workbook.active
    text_col_indices = [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None:
                cell.value = str(cell.value)
    workbook.save(output_path)

# ==============================================================================
# 7. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print("=" * 70 + "\n CLEANSING DATA 3 - SUSPENSE | CEDANT: PT INTER BENUA REASINDO (IBS)\n" + "=" * 70)
    input_path = INPUT_FILE

    if not os.path.exists(input_path):
        candidates = [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]
        if not candidates:
            raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")
        input_path = candidates[0]
        print(f"[AUTO-DETECT] File input tidak ditemukan di path default, memakai: '{input_path}'")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    print(f"[INFO] Total baris source: {len(df_raw)}")

    df_ibs = filter_cedant(df_raw, CEDANT_FILTER)
    if df_ibs.empty:
        print("[WARNING] Tidak ada baris yang cocok dengan filter cedant. Proses dihentikan.")
        return

    df_ibs = filter_status(df_ibs, STATUS_FILTER)
    if df_ibs.empty:
        print("[WARNING] Tidak ada baris yang cocok dengan filter status. Proses dihentikan.")
        return

    df_hasil = build_output(df_ibs)
    save_with_text_format(df_hasil, OUTPUT_FILE)
    print("=" * 70 + f"\n[SUCCESS] Selesai. Total baris output: {len(df_hasil)}\n[SUCCESS] File hasil: '{OUTPUT_FILE}'\n" + "=" * 70)

if __name__ == "__main__":
    main()