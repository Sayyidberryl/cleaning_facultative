"""
Cleaning Data 1 - Facultative | Cedant: KB INS (KOOKMIN BEST INS. EX. LIG INS.)

CATATAN VERSI: v3 -- update lanjutan setelah BUG_KB_INS_pt2.md (bug report kedua). v2 sebelumnya rewrite setelah BUG_KB_INS.md (bug report dari hasil
pengerjaan v1). Semua perubahan di bawah ini dipicu oleh contoh konkret
di bug report tsb -- mayoritas contoh polis/slip & insured di bug report
justru diambil dari Data 1, jadi file ini yang paling banyak divalidasi
langsung.

CATATAN ARSITEKTUR:
- Mesin polis/slip & insured SAMA PERSIS dengan KB_INS_cleaning_osbal.py
  (lihat docstring file tsb untuk detail lengkap rule & alasannya) --
  BEDANYA: Data 1 tidak punya kolom CLSDT terpisah, jadi TIDAK ada rule
  seleksi sumber FAC vs CLSDT -- FAC_POLICY_NO/FAC_SLIP langsung dipakai
  & di-clean spt biasa.
- [BARU v2] Rule mesin polis/slip yg berubah dari v1 (identik dgn Data 2):
    * Ambang pengulangan digit ekor naik dari 8 ke 12 digit.
    * Blok sertifikat/register yang nempel via "/" atau spasi (bukan cuma
      dash rapat) sekarang dikenali sebelum tokenisasi -- lihat
      try_extract_cert_block().
    * Ekor narasi bulan+tahun (mis. "/ DESEMBER 2017") dibuang.
    * Nilai "garbled"/format asing (mis. "P0B3-8006F3-2403A-7740",
      "30101G[81900002") dibiarkan apa adanya, tidak ditokenisasi.
    * Prefiks "P<n>/" dibuang SEBELUM proses pisah spasi panjang (fix
      kasus "P1/  <basis>" dgn 2 spasi).
- [BARU v2] Rule insured yang berubah dari v1 -- lihat docstring
  KB_INS_cleaning_suspend.py bagian "Insured" untuk daftar lengkap.
- CERTIFICATE: diambil dari FAC_POLICY_NO ORI dulu, fallback ke FAC_SLIP
  ORI kalau FAC_POLICY_NO tidak mengandung pola sertifikat.
- Kolom broker: file terbaru punya 2 kolom "COMP_NAME" duplikat (satu utk
  cedant, satu utk broker) -- pandas otomatis rename yg kedua jadi
  "COMP_NAME.1". Kalau file yang dipakai beda struktur, sesuaikan
  BROKER_COL di bawah.

=====================================================================
CATATAN YANG MASIH PERLU DIKONFIRMASI:
1. Sama seperti v1: karena Data 1 tidak dibahas terpisah di dokumen
   eksplorasi awal (KB_INS - Eksplorasi Cedant.docx), rule mesin polis/
   slip di file ini diadaptasi dari pola yang sama dgn Data 2 (skema
   penomoran identik, sudah dicek silang ke data mentah). Bug report kali
   ini justru BANYAK memberi contoh nyata dari Data 1 sendiri (bukan lagi
   murni inferensi), jadi tingkat kepercayaannya sudah lebih tinggi dari
   v1 -- tapi tetap disarankan spot-check ke hasil akhir sebelum dipakai
   produksi penuh.
2. Ambang MAX_SHORT_SUFFIX=12 & threshold "narasi"/"garbled" sama seperti
   Data 2 -- lihat catatan di docstring KB_INS_cleaning_osbal.py poin 1-2.
3. [SELESAI -- lihat BUG_KB_INS_pt2.md] Trailing "/CJ INDONESIA -PASURUAN
   –JOMBANG" pada contoh insured CJ BIO sekarang DIBUANG total, dan
   "CJ BIO PASURUAN DAN JOMBANG" dipecah jadi 2 insured (CJ BIO PASURUAN,
   CJ BIO JOMBANG) -- lihat _try_dan_location_with_discard().
=====================================================================
"""

import os
import re

import pandas as pd


# ─────────────────────────────────────────────────────────────────────────────
# KONFIGURASI
# ─────────────────────────────────────────────────────────────────────────────

INPUT_FILE  = os.path.join("input", "1b. Transaksi Facul 01.01.23 - 17.08.2026.xlsx")
INPUT_SHEET = "Query result"
OUTPUT_FILE = os.path.join("output", "KbIns_output_facul.xlsx")

CEDANT_COL   = "COMP_NAME"
CEDANT_VALUE = "KOOKMIN BEST INS.  EX. LIG INS."

BROKER_COL  = "COMP_NAME.1"
POLIS_COL   = "FAC_POLICY_NO"
SLIP_COL    = "FAC_SLIP"
INSURED_COL = "FAC_INSURED"

MAX_SPLIT_COLS = 5

FULL_LEN_MIN     = 15
MAX_SHORT_SUFFIX = 12
FRESH_BASE_MIN   = 13


# ─────────────────────────────────────────────────────────────────────────────
# KAMUS POLA STANDAR KB INS
# ─────────────────────────────────────────────────────────────────────────────

KNOWN_POLIS_PATTERNS = [re.compile(r"^\d{13,19}$")]
KNOWN_SLIP_PATTERNS = [
    re.compile(r"^\d{13,19}$"),
    re.compile(r"^\d{3,6}/CN/\d{1,4}/\d{2}/\d{2}$", re.IGNORECASE),
]


def _matches_known_pattern(value: str, patterns) -> bool:
    return any(p.match(value) for p in patterns)


def refine_with_known_pattern(value: str, patterns) -> str:
    if not value:
        return value
    if _matches_known_pattern(value, patterns):
        return value
    candidates = [re.sub(r"\s+", "", value), value.strip(" .-/")]
    for c in candidates:
        if _matches_known_pattern(c, patterns):
            return c
    return value


def _refine_list(items: list, patterns) -> list:
    return [refine_with_known_pattern(v, patterns) for v in items]


def _normalize_spaces(text: str) -> str:
    return re.sub(r"\s{2,}", " ", text).strip()


def _dedup_preserve_order(items: list) -> list:
    seen, out = set(), []
    for i in items:
        if i not in seen:
            seen.add(i)
            out.append(i)
    return out


def _cap_or_join(items: list) -> list:
    if len(items) > MAX_SPLIT_COLS:
        return [",".join(items)]
    return items


# ─────────────────────────────────────────────────────────────────────────────
# KOLOM CERTIFICATE [DIROMBAK v2] -- mendukung list campuran & "S.D."
# ─────────────────────────────────────────────────────────────────────────────

_CERT_PAIR_SEP_RE = re.compile(
    r"\s*(?:S\s*\.\s*D\s*\.?|S\s*/\s*D|SD|S(?!\d)|-)\s*", re.IGNORECASE
)


def _format_cert_chunk(chunk: str):
    chunk = chunk.strip()
    if not chunk:
        return None
    parts = [p for p in _CERT_PAIR_SEP_RE.split(chunk) if p.strip()]
    if len(parts) == 2 and all(re.match(r"^\d{1,7}$", p.strip()) for p in parts):
        lo, hi = int(parts[0]), int(parts[1])
        if hi < lo:
            lo, hi = hi, lo
        count = hi - lo + 1
        if count > 3:
            return f"{str(lo).zfill(6)} SD {str(hi).zfill(6)}"
        return ", ".join(str(n).zfill(6) for n in range(lo, hi + 1))
    if len(parts) == 1 and re.match(r"^\d{1,7}$", parts[0].strip()):
        return str(int(parts[0])).zfill(6)
    return None


def build_certificate(cert_raw) -> str:
    if cert_raw is None or (isinstance(cert_raw, float) and pd.isna(cert_raw)):
        return ""
    text = str(cert_raw).strip()
    if not text or text.upper() in {"-", "NAN", "NONE", "NULL"}:
        return ""
    top_chunks = [c.strip() for c in text.split(",") if c.strip()]
    if not top_chunks:
        return ""
    formatted = []
    for c in top_chunks:
        f = _format_cert_chunk(c)
        if f is None:
            return ""
        formatted.append(f)
    return ", ".join(formatted)


# ─────────────────────────────────────────────────────────────────────────────
# DETEKSI BLOK SERTIFIKAT/REGISTER YANG NEMPEL DI POLIS/SLIP
# ─────────────────────────────────────────────────────────────────────────────

_CERT_JUNK_TAIL_RE = re.compile(r"\s*/\s*(?:VAR|VARIOUS|TBA)\s*$", re.IGNORECASE)
_CERT_BLOCK_SLASH_RE = re.compile(r"^\s*(\d{10,})\s*/\s*(.+)$")
_CERT_BLOCK_SPACE_RE = re.compile(r"^\s*(\d{10,})\s+-?\s*(.+)$")
_CERT_BLOCK_TIGHT_DASH_RE = re.compile(r"^\s*(\d{10,})\s*[-.]\s*(\d{5,7})\s*$")

# [BARU pt2] cert bisa nempel di dlm chain "+" pada polis duplikat (mis.
# "2010101102500276 + 2010101102500276-000002" -- basis sama, ekor cert di
# salah satu -- basis dianggap 1 (duplikat), cert diambil dari ekornya).
_INLINE_CERT_TAIL_RE = re.compile(r"(\d{13,19})-(\d{5,7})(?!\d)")


def _extract_inline_certs(text: str):
    found = []

    def repl(m):
        found.append(m.group(2))
        return m.group(1)

    new_text = _INLINE_CERT_TAIL_RE.sub(repl, text)
    return new_text, found


def try_extract_cert_block(primary: str):
    for rx in (_CERT_BLOCK_SLASH_RE, _CERT_BLOCK_SPACE_RE):
        m = rx.match(primary)
        if m:
            base, rest = m.groups()
            rest = _CERT_JUNK_TAIL_RE.sub("", rest.strip()).strip()
            if rest and not re.match(r"^\d{10,}", rest):
                cert = build_certificate(rest)
                if cert:
                    return base, cert
    m = _CERT_BLOCK_TIGHT_DASH_RE.match(primary)
    if m:
        base, rest = m.groups()
        cert = build_certificate(rest)
        if cert:
            return base, cert
    return None, None


def extract_certificate(polis_raw, slip_raw) -> str:
    """Ambil sertifikat dari FAC_POLICY_NO ORI dulu; fallback ke FAC_SLIP
    ORI (Data 1 tidak punya kolom CLSDT terpisah)."""
    for raw in (polis_raw, slip_raw):
        if raw is None or (isinstance(raw, float) and pd.isna(raw)):
            continue
        val = str(raw).strip()
        if not val:
            continue
        _, cert = try_extract_cert_block(val)
        if cert:
            return cert
        # [BARU pt2] fallback: cert nempel di dlm chain "+" pada polis duplikat
        _, inline_certs = _extract_inline_certs(val)
        if inline_certs:
            return build_certificate(",".join(inline_certs))
    return ""


# ─────────────────────────────────────────────────────────────────────────────
# MESIN POLIS/SLIP KB INS [DIROMBAK v2] -- identik dgn KB_INS_cleaning_osbal.py
# ─────────────────────────────────────────────────────────────────────────────

_STANDALONE_P_RE = re.compile(
    r"^\s*(?:P\d+|TBA|VAR|VARIOUS)(?:\s*[+,]\s*(?:P\d+|TBA|VAR|VARIOUS))*\s*$",
    re.IGNORECASE,
)
_P_PREFIX_RE = re.compile(r"^\s*P\d+\s*/\s*", re.IGNORECASE)
_CN_PATTERN_RE = re.compile(r"^\s*\d{3,6}\s*/\s*CN\s*/\s*\d{1,4}\s*/\s*\d{2}\s*/\s*\d{2}\s*$", re.IGNORECASE)
_JUNK_TOKEN_RE = re.compile(r"^(?:P\d+|VAR|VARIOUS|TBA)$", re.IGNORECASE)

_MONTH_RE = (
    r"JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|"
    r"NOVEMBER|DES[EA]?MBER|JANUARY|FEBRUARY|MARCH|MAY|JUNE|JULY|AUGUST|OCTOBER|DECEMBER"
)
_MONTH_YEAR_TAIL_RE = re.compile(
    r"\s*/?\s*(?:" + _MONTH_RE + r")\.?\s*(?:19|20)?\d{0,4}\s*$", re.IGNORECASE
)

_GARBLED_TOKEN_RE = re.compile(r"^(?=.*[A-Za-z])(?=.*\d).+$")

# [BARU pt2] literal whitelist "BORD <BULAN> <TAHUN>" -- dibiarkan apa adanya
_BORD_MONTH_RE = re.compile(
    r"^\s*BORD\s+(?:" + _MONTH_RE + r")\s+\d{4}\s*$", re.IGNORECASE
)


def _looks_garbled(primary: str) -> bool:
    tokens = re.split(r"[\s+\-/,]", primary)
    for t in tokens:
        t = t.strip()
        if not t:
            continue
        if _JUNK_TOKEN_RE.match(t):
            continue
        if re.fullmatch(r"S/?D\.?", t, re.IGNORECASE):
            continue
        if _GARBLED_TOKEN_RE.match(t):
            return True
    return False


def _split_long_space_tail(val: str):
    parts = re.split(r"\s{2,}", val, maxsplit=1)
    if len(parts) == 1:
        return val, None
    return parts[0].strip(), parts[1].strip()


def _clean_value(val) -> list:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return []
    val = str(val).strip()
    if not val:
        return []

    if _BORD_MONTH_RE.match(val):
        return [_normalize_spaces(val)]
    if _CN_PATTERN_RE.match(val):
        return [_normalize_spaces(val)]
    if _STANDALONE_P_RE.match(val):
        return [_normalize_spaces(val)]
    if val.upper().startswith("TBA") and not re.search(r"\d{10,}", val):
        return [_normalize_spaces(val)]

    val2 = _MONTH_YEAR_TAIL_RE.sub("", val).strip(" -/,")
    check_val = val2 if val2 else val

    if _looks_garbled(check_val):
        return [_normalize_spaces(val)]

    m = _P_PREFIX_RE.match(check_val)
    if m:
        check_val = check_val[m.end():]

    primary, tail = _split_long_space_tail(check_val)

    base, cert = try_extract_cert_block(primary)
    if base:
        return [_normalize_spaces(base)]

    # [BARU pt2] cert bisa nempel di dlm chain "+" pada polis DUPLIKAT
    # (mis. "2010101102500276 + 2010101102500276-000002") -- basis dianggap
    # 1 (duplikat, akan ke-dedup di bawah), cert diurus terpisah.
    primary_nocert, _inline_certs = _extract_inline_certs(primary)
    text = primary_nocert.strip(" -/,")

    raw_tokens = [t for t in re.split(r"\s+|[+\-/,]", text) if t.strip()]
    raw_tokens = [t.strip(".") for t in raw_tokens]
    raw_tokens = [t for t in raw_tokens if t and not _JUNK_TOKEN_RE.match(t)]
    raw_tokens = [t for t in raw_tokens if re.search(r"\d", t)]

    results, current_base = [], None
    reconstruction_ok = True
    for i, tok in enumerate(raw_tokens):
        digits_only = re.match(r"^\d+$", tok) is not None
        if not digits_only:
            results.append(tok)
            continue

        if len(tok) >= FULL_LEN_MIN:
            results.append(tok)
            current_base = tok
        elif len(tok) >= FRESH_BASE_MIN:
            # basis msh di bawah panjang penuh (13-14 digit) -- cek token
            # BERIKUTNYA (lookahead): kalau pendek (mungkin pelengkap) ->
            # tahan dulu; kalau tidak ada token berikutnya, atau berikutnya
            # JUGA basis panjang -> anggap basis ini sudah final, tampilkan.
            next_tok = raw_tokens[i + 1] if i + 1 < len(raw_tokens) else None
            if next_tok and re.match(r"^\d+$", next_tok) and len(next_tok) < FRESH_BASE_MIN:
                current_base = tok  # ditahan, menunggu pelengkap
            else:
                results.append(tok)
                current_base = tok
        elif current_base is not None and len(current_base) >= FULL_LEN_MIN and len(tok) <= MAX_SHORT_SUFFIX:
            new_val = current_base[:-len(tok)] + tok
            results.append(new_val)
            # current_base TIDAK diupdate -- token ekor berikutnya yg
            # independen tetap mengacu ke basis asli yang sama.
        elif current_base is not None and len(current_base) < FULL_LEN_MIN:
            # basis SEDANG pending (msh terpotong) -- sambung dgn token ini
            new_val = current_base + tok
            # [BARU] kalau hasil sambungan MASIH di bawah panjang basis yg
            # masuk akal (< FRESH_BASE_MIN), pola ini kemungkinan BUKAN
            # "basis terpotong" -- cuma daftar kode pendek yg kebetulan
            # nempel (mis. "052 + 057 + 058 + 053", "040 + 041"). Batalkan
            # rekonstruksi utk SELURUH nilai ini, biarkan apa adanya.
            if len(new_val) < FRESH_BASE_MIN:
                reconstruction_ok = False
                break
            results.append(new_val)
            current_base = new_val
        else:
            # token pendek tanpa basis yg bisa disambung -> jadi basis
            # pending baru (mis. prefiks pendek spt "301010" di awal chain)
            current_base = tok

    if not reconstruction_ok:
        return [_normalize_spaces(val)]

    final = [t for t in results if len(t) >= 5]
    if not final:
        return [_normalize_spaces(primary)]
    final = _dedup_preserve_order(final)

    if tail:
        tail_norm = re.sub(r"\s+", "", tail)
        if tail_norm in [re.sub(r"\s+", "", f) for f in final]:
            final = [tail_norm]

    # [BARU pt2] kalau hasil rekonstruksi > MAX_SPLIT_COLS, JANGAN dipakai
    # (jangan gabung-koma hasil rekonstruksi) -- biarkan NILAI ASLI apa
    # adanya, sesuai instruksi "perulangannya gausah dijalankan".
    if len(final) > MAX_SPLIT_COLS:
        return [_normalize_spaces(val)]

    return [_normalize_spaces(f) for f in final]


def clean_polis(val) -> list:
    return _clean_value(val)


def clean_slip(val) -> list:
    return _clean_value(val)


# ─────────────────────────────────────────────────────────────────────────────
# CLEAN INSURED [DIROMBAK v2] -- identik dgn KB_INS_cleaning_osbal.py
# ─────────────────────────────────────────────────────────────────────────────

_GELAR_PREFIX_RE = re.compile(
    r"^\s*(?:BAPAK|IBU|SDRI?|NYONYA|NY|TUAN|TN|MR|MRS|MS|DR|IR|PROF|HJ|H)\b\.?\s*",
    re.IGNORECASE,
)
_GELAR_SUFFIX_RE = re.compile(
    r"\s*,?\s*\b(?:S\.?E|S\.?H|S\.?SI|S\.?T|S\.?SOS|S\.?KOM|S\.?PD|"
    r"M\.?SI|M\.?M|M\.?H|M\.?KOM|SH|SE)\b\.?\s*$",
    re.IGNORECASE,
)
_SUSPENSE_KEEP_RE = re.compile(r"SUSPEN|\(?H\)?\s*UTANG\s*PIUTANG", re.IGNORECASE)
_SPLIT_RE = re.compile(r"\bQQ\b|\bAND\s*/\s*OR\.?\b|/{1,2}|,|–|\s-\s|\+|¿", re.IGNORECASE)

INSURED_JUNK_WORDS = frozenset({
    "PT", "CV", "TBK", "PERSERO", "LTD", "INC", "LLC",
    "AND", "OR", "THE", "OF", "AS", "VARIOUS", "VAR", "QQ", "TBA",
})

_KNOWN_LOCATION_WORDS = {"BREBES", "GROBOGAN", "PEMALANG", "PASURUAN", "JOMBANG"}
_BRANCH_MERGE_RE = re.compile(
    r"^(?:FACTORY|PABRIK|CABANG)\s+\S+$|^\S+\s+(?:FACTORY|PABRIK|CABANG)$",
    re.IGNORECASE,
)


def _is_location_merge_part(p: str) -> bool:
    up = p.strip().upper()
    if up in _KNOWN_LOCATION_WORDS:
        return True
    if _BRANCH_MERGE_RE.match(up):
        return True
    return False


_FACTORY_MULTI_RE = re.compile(
    r"^(.*?)\bFACTORY\s+([IVXLCDM0-9]+(?:\s*(?:,|&|DAN)\s*[IVXLCDM0-9]+)+)\s*$",
    re.IGNORECASE,
)


def _expand_factory_multi(val: str):
    m = _FACTORY_MULTI_RE.match(val.strip())
    if not m:
        return None
    prefix, nums_part = m.groups()
    prefix = prefix.strip()
    nums = [n.strip() for n in re.split(r"\s*(?:,|&|DAN)\s*", nums_part, flags=re.IGNORECASE) if n.strip()]
    if len(nums) < 2:
        return None
    return [f"{prefix} FACTORY {n}".strip() for n in nums]


# [BARU pt2] "PREFIX (FACTORY N1 DAN FACTORY N2)" -- beda dari
# _expand_factory_multi di atas: di sini kata "FACTORY" DIULANG utk tiap
# nomor (bukan sekali lalu daftar nomor), dan boleh dibungkus kurung.
_FACTORY_REPEATED_RE = re.compile(
    r"^(.*?)\(?\s*FACTORY\s+([IVXLCDM0-9]+)\s+DAN\s+FACTORY\s+([IVXLCDM0-9]+)\s*\)?\s*$",
    re.IGNORECASE,
)


def _expand_factory_repeated(val: str):
    m = _FACTORY_REPEATED_RE.match(val.strip())
    if not m:
        return None
    prefix, n1, n2 = m.groups()
    prefix = prefix.strip().rstrip(",").strip()
    prefix = re.sub(r",?\s*(?:PT|CV)\.?\s*$", "", prefix, flags=re.IGNORECASE).strip()
    if not prefix:
        return None
    return [f"{prefix} FACTORY {n1}".strip(), f"{prefix} FACTORY {n2}".strip()]


_DAN_LOCATION_RE = re.compile(
    r"^(.*?)\s+(" + "|".join(_KNOWN_LOCATION_WORDS) + r")\s+DAN\s+("
    + "|".join(_KNOWN_LOCATION_WORDS) + r")\s*$",
    re.IGNORECASE,
)


def _expand_dan_location(val: str):
    m = _DAN_LOCATION_RE.match(val.strip())
    if not m:
        return None
    prefix, loc1, loc2 = m.groups()
    prefix = prefix.strip()
    if not prefix:
        return None
    return [f"{prefix} {loc1}".strip(), f"{prefix} {loc2}".strip()]


def _try_dan_location_with_discard(val: str):
    """Coba _expand_dan_location pada value utuh dulu; kalau gagal, coba
    pada BAGIAN PERTAMA (sebelum '/' teratas) saja -- kalau cocok, bagian
    setelah '/' DIBUANG total (bukan dianggap insured terpisah), sesuai
    konfirmasi user utk kasus spt 'CJ BIO PASURUAN DAN JOMBANG/CJ INDONESIA
    -PASURUAN -JOMBANG' -> 2 insured (CJ BIO PASURUAN, CJ BIO JOMBANG),
    bagian '/CJ INDONESIA ...' dibuang total."""
    r = _expand_dan_location(val)
    if r is not None:
        return r
    first_part = val.split("/", 1)[0].strip()
    if first_part != val.strip():
        r2 = _expand_dan_location(first_part)
        if r2 is not None:
            return r2
    return None


_AS_ROLE_RE = re.compile(
    r"\bAS\s+(?:THE\s+)?(?:OWNER|PRINCIPAL|INSURED|CONTRACTOR|OFF-TAKER)\b\.?",
    re.IGNORECASE,
)


def _initials_windows(name: str) -> set:
    words = [w for w in re.split(r"[\s\-]+", name.upper()) if w]
    variants = set()
    n = len(words)
    for start in range(n):
        for end in range(start + 1, n + 1):
            window = words[start:end]
            variants.add("".join(w[0] for w in window if w[0].isalpha()))
    return variants


def _is_abbreviation_of(short: str, long_name: str) -> bool:
    short_clean = re.sub(r"[^A-Z0-9]", "", short.upper())
    if not short_clean or len(short_clean) < 2 or len(short_clean) > 8:
        return False
    if short_clean in _initials_windows(long_name):
        return True
    words = [w for w in re.split(r"[\s\-]+", long_name.upper()) if w]
    if any(w.startswith(short_clean) and w != short_clean for w in words):
        return True
    return False


def _strip_parenthetical_abbrev(text: str) -> str:
    def repl(m):
        inner = m.group(1).strip()
        if re.search(r"\bPT\b|\bCV\b", inner, re.IGNORECASE):
            return " "
        inner_letters = re.sub(r"[^A-Za-z]", "", inner)
        before = text[:m.start()].strip()
        if before and inner_letters and _is_abbreviation_of(inner_letters, before):
            return " "
        return m.group(0)
    return re.sub(r"\(([^()]*)\)", repl, text)


_SHORT_AMP_RE = re.compile(r"\b[A-Z]{1,2}&[A-Z]{1,2}\b")

_KNOWN_INSURED_ABBREV = {
    "SEIN": "SAMSUNG ELECTRONICS INDONESIA",
}

_PT_CV_PREFIX_RE = re.compile(r"^\s*(?:PT|CV)(?:\.|(?=\s|$))\.?\s*", re.IGNORECASE)
_PT_CV_SUFFIX_RE = re.compile(r",?\s*(?:PT|CV)\.?\s*$", re.IGNORECASE)


def _remove_polis_slip_from_insured(text: str, polis_ori, slip_ori) -> str:
    original_text = text.strip()
    for token in (polis_ori, slip_ori):
        if token is not None and not (isinstance(token, float) and pd.isna(token)):
            t = str(token).strip()
            if not t or t == "-":
                continue
            if t.upper() == original_text.upper():
                continue
            text = text.replace(t, "")
    return text


def _clean_insured_name(name: str) -> str:
    name = _normalize_spaces(name)
    name = _GELAR_PREFIX_RE.sub("", name)
    for _ in range(3):
        stripped = _GELAR_SUFFIX_RE.sub("", name).strip()
        if stripped == name:
            break
        name = stripped
    name = re.sub(r",?\s*\(\s*PERSERO\s*\)\s*", " ", name, flags=re.IGNORECASE)
    name = re.sub(r"\bTBK\b\s*,\s*PT\.?\s*", " ", name, flags=re.IGNORECASE)
    name = re.sub(r",?\s*\bTBK\b\s*", " ", name, flags=re.IGNORECASE)
    name = re.sub(r"\bLTD\b\.?\s*$", "", name, flags=re.IGNORECASE)
    name = re.sub(r"\bPTE\b\.?\s*", "", name, flags=re.IGNORECASE)
    name = _PT_CV_SUFFIX_RE.sub("", name)
    name = _PT_CV_PREFIX_RE.sub("", name)
    name = re.sub(r"(?<!\d)\.(?!\d)", " ", name)
    return _normalize_spaces(name).strip().upper()


def _dedup_abbreviations(cleaned: list) -> list:
    final = []
    for name in cleaned:
        words = name.split()
        first_word = words[0] if words else ""
        is_dup = False
        for prior in final:
            if name in _KNOWN_INSURED_ABBREV and _KNOWN_INSURED_ABBREV[name] in prior:
                is_dup = True
                break
            if _is_abbreviation_of(name, prior):
                is_dup = True
                break
            if _SHORT_AMP_RE.search(name) and prior.split() and prior.split()[0] == first_word:
                is_dup = True
                break
        if not is_dup:
            final.append(name)
    return final


def clean_insured(val, polis_ori=None, slip_ori=None) -> list:
    if val is None or (isinstance(val, float) and pd.isna(val)):
        return []
    val = str(val).strip()
    if not val:
        return []
    if _SUSPENSE_KEEP_RE.search(val):
        return [_normalize_spaces(val).upper()]

    val = _remove_polis_slip_from_insured(val, polis_ori, slip_ori)
    val = _strip_parenthetical_abbrev(val)
    val = _AS_ROLE_RE.sub(" ", val)

    for fn in (_expand_factory_multi, _expand_factory_repeated, _try_dan_location_with_discard):
        r = fn(val)
        if r is not None:
            return [_normalize_spaces(x).upper() for x in r]

    parts = [p.strip() for p in _SPLIT_RE.split(val) if p.strip()]
    cleaned = []
    for p in parts:
        p_clean = _clean_insured_name(p)
        if not p_clean:
            continue
        merge_check = p_clean.strip("() ")
        if _is_location_merge_part(merge_check) and cleaned:
            cleaned[-1] = _normalize_spaces(cleaned[-1] + " " + p_clean)
            continue
        if len(p_clean) < 2 or p_clean in INSURED_JUNK_WORDS:
            continue
        cleaned.append(p_clean)

    if not cleaned:
        fallback = _clean_insured_name(val)
        return [fallback] if fallback else []

    cleaned = _dedup_abbreviations(cleaned)
    return _cap_or_join(cleaned)


# ─────────────────────────────────────────────────────────────────────────────
# MITRA BISNIS
# ─────────────────────────────────────────────────────────────────────────────

def get_business_partner(comp_name2, comp_name) -> str:
    broker = "" if pd.isna(comp_name2) else str(comp_name2).strip()
    if not broker or broker.upper() == "DIRECT":
        result = "" if pd.isna(comp_name) else str(comp_name).strip()
    else:
        result = broker
    result = re.sub(r"\.", " ", result)
    return re.sub(r"\s+", " ", result).upper().strip()


# ─────────────────────────────────────────────────────────────────────────────
# PROSES UTAMA
# ─────────────────────────────────────────────────────────────────────────────

def _force_text_format(output_file, prefixes):
    import openpyxl
    wb = openpyxl.load_workbook(output_file)
    ws = wb.active
    header = [c.value for c in ws[1]]
    target_cols = [i for i, h in enumerate(header, start=1) if h and str(h).startswith(prefixes)]
    for col_idx in target_cols:
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                cell.number_format = "@"
    wb.save(output_file)


def _insert_clean_columns(df, all_lists, prefix, max_cols):
    added = []
    actual_max_cols = max((len(lst) for lst in all_lists if isinstance(lst, list)), default=1)
    actual_max_cols = min(actual_max_cols, max_cols)
    for i in range(1, actual_max_cols + 1):
        col_name = f"{prefix}_{i}"
        df[col_name] = [lst[i - 1] if isinstance(lst, list) and i - 1 < len(lst) else None for lst in all_lists]
        added.append(col_name)
    return added


def _insert_polis_cert_columns(df, all_polis_lists, all_cert_values, max_cols):
    added = []
    actual_max_cols = max((len(lst) for lst in all_polis_lists if isinstance(lst, list)), default=1)
    actual_max_cols = min(actual_max_cols, max_cols)
    for i in range(1, actual_max_cols + 1):
        polis_col = f"POLIS_CLEAN_{i}"
        df[polis_col] = [lst[i - 1] if isinstance(lst, list) and i - 1 < len(lst) else None for lst in all_polis_lists]
        added.append(polis_col)
        if i == 1:
            df["CERTIFICATE_1"] = list(all_cert_values)
            added.append("CERTIFICATE_1")
    return added


def process_data(input_file: str, output_file: str) -> None:
    print(f"[1/5] Membaca data dari: {input_file} (sheet: {INPUT_SHEET}) ...")
    df = pd.read_excel(input_file, sheet_name=INPUT_SHEET, header=0)
    print(f"      Total baris keseluruhan: {len(df):,}")

    if CEDANT_COL not in df.columns:
        print(f"\n[ERROR] Kolom '{CEDANT_COL}' tidak ditemukan!")
        print(f"        Kolom tersedia: {list(df.columns)}")
        return

    df[CEDANT_COL] = df[CEDANT_COL].astype(str).str.strip()
    df = df[df[CEDANT_COL] == CEDANT_VALUE].copy()
    print(f"[2/5] Filter cedant '{CEDANT_VALUE}': {len(df):,} baris ditemukan.")

    if df.empty:
        print("\n[WARN] Tidak ada data setelah filter. Proses dihentikan.")
        return

    for col in [POLIS_COL, SLIP_COL, INSURED_COL, BROKER_COL]:
        if col not in df.columns:
            print(f"\n[ERROR] Kolom '{col}' tidak ditemukan di file!")
            print(f"        Kolom tersedia: {list(df.columns)}")
            return

    business_partner_values = [get_business_partner(row[BROKER_COL], row[CEDANT_COL]) for _, row in df.iterrows()]

    df.rename(columns={INSURED_COL: "FAC_INSURED_ORI"}, inplace=True)

    insert_pos = list(df.columns).index(BROKER_COL) + 1 if BROKER_COL in df.columns else len(df.columns)
    df.insert(insert_pos, "BUSINESS_PARTNER", business_partner_values)

    print("[3/5] Menjalankan proses cleaning ...")

    all_clean_polis, all_clean_slip, all_clean_ins, all_certificate = [], [], [], []
    max_polis = max_slip = max_ins = 1

    for _, row in df.iterrows():
        c_polis = _refine_list(clean_polis(row.get(POLIS_COL, "")), KNOWN_POLIS_PATTERNS)
        c_slip  = _refine_list(clean_slip(row.get(SLIP_COL, "")), KNOWN_SLIP_PATTERNS)
        c_ins   = clean_insured(row.get("FAC_INSURED_ORI", ""), row.get(POLIS_COL, ""), row.get(SLIP_COL, ""))
        cert    = extract_certificate(row.get(POLIS_COL, ""), row.get(SLIP_COL, ""))

        max_polis = max(max_polis, len(c_polis))
        max_slip  = max(max_slip, len(c_slip))
        max_ins   = max(max_ins, len(c_ins))

        all_clean_polis.append(c_polis)
        all_clean_slip.append(c_slip)
        all_clean_ins.append(c_ins)
        all_certificate.append(cert)

    print("[4/5] Menyusun kolom output ...")

    new_columns = []
    for col in df.columns:
        new_columns.append(col)
        if col == POLIS_COL:
            new_columns += _insert_polis_cert_columns(df, all_clean_polis, all_certificate, max_polis)
        elif col == SLIP_COL:
            new_columns += _insert_clean_columns(df, all_clean_slip, SLIP_COL, max_slip)
        elif col == "FAC_INSURED_ORI":
            new_columns += _insert_clean_columns(df, all_clean_ins, "FAC_INSURED", max_ins)

    df = df[new_columns]

    os.makedirs(os.path.dirname(output_file), exist_ok=True)
    print(f"[5/5] Menyimpan hasil ke: {output_file} ...")
    df.to_excel(output_file, index=False)
    _force_text_format(output_file, prefixes=("POLIS_CLEAN_", f"{SLIP_COL}_", "FAC_INSURED_", "CERTIFICATE"))

    print(f"\n{'=' * 55}")
    print("  [OK] Selesai!")
    print(f"  Total baris output  : {len(df):,}")
    print(f"  Total kolom output  : {len(df.columns)}")
    print(f"  File disimpan di    : {output_file}")
    print(f"{'=' * 55}")


if __name__ == "__main__":
    process_data(INPUT_FILE, OUTPUT_FILE)