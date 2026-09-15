"""
Cleaning Data 3 - Suspend | Cedant: JASA R.P (PT ASURANSI JASA RAHARJA PUTRA)

CATATAN ARSITEKTUR (beda dari Data 2/Osbal):
- Sesuai dokumen eksplorasi JRP: "cleansing tanpa nomor sertifikat, matching
  data menggunakan polis ori dan polis tanpa nomor sertifikat" -- artinya di
  Data 3, sertifikat (mis. "-000023") justru DIBUANG TOTAL dari POLIS_CLEAN_1
  (kebalikan dari Data 2 yang mempertahankannya menempel). Ini KONSISTEN
  dengan mesin lama (loop strip suffix "-NNN"), jadi TIDAK perlu logika baru
  seperti di Osbal/Facul -- cukup pakai clean_polis versi lama apa adanya.
- Diverifikasi ke sample data asli: pola POLIS JRP di Data 3 memang persis
  "BASE-CERT" (base 15-19 digit + dash + cert 6 digit), dan SLIP NO sudah
  bersih (14-19 digit polos, tidak perlu proses tambahan).
- Insured: "PT. Pelabuhan Indonesia (Persero) Regional N" -- confirmed
  cukup strip PT/Persero standar, TANPA transformasi "Pelindo" atau angka
  romawi (sesuai konfirmasi user, itu cuma catatan matching manual).

UPDATE (file 3b, rule baru):
- Sheet & struktur file berubah: sheet sekarang "Sheet1" (bukan lagi
  "Detail Database"), header tetap di baris ke-3 (index 2).
- Ada kolom baru "STATUS" berisi "SUSPENSE" atau "ADJUSTED". Cleaning HANYA
  memproses baris berstatus "SUSPENSE" -- baris "ADJUSTED" DIBUANG TOTAL,
  tidak ditampilkan di output sama sekali (difilter sebelum cleaning, sama
  seperti filter cedant).

=====================================================================
UPDATE STRUKTUR KOLOM OUTPUT (terbaru) -- HANYA STRUKTUR, LOGIC CLEANING
TIDAK BERUBAH:
- Sebelumnya urutannya: POLIS_ORI -> CERTIFICATE -> POLIS_CLN.
- Sekarang kolom hasil clean polis diganti nama dari "POLIS_CLN" jadi
  "POLIS_CLEAN_1" (Data 3 memang selalu cuma 1 polis per baris, tidak
  pernah pecah jadi banyak kolom seperti Data 1/Data 2 -- jadi otomatis
  hanya ada POLIS_CLEAN_1, tidak ada POLIS_CLEAN_2 dst), dan CERTIFICATE
  (diganti nama jadi CERTIFICATE_1) dipindah ke SETELAH POLIS_CLEAN_1:
      POLIS_ORI -> POLIS_CLEAN_1 -> CERTIFICATE_1
- CERTIFICATE_1 tetap selalu dibuat walau tidak ada certificate sama
  sekali (nilainya string kosong).
=====================================================================
UPDATE SEBELUMNYA -- KOLOM CERTIFICATE (rule sama persis dgn Data 1/Data 2,
lihat build_certificate()):
- SUMBER: Data 3 TIDAK punya kolom CLSDT_* terpisah seperti Data 2 --
  sesuai docstring di atas, sertifikat di Data 3 SELALU menempel langsung
  di kolom POLIS (pola "BASE-CERT"), jadi ini SATU-SATUNYA sumber yang
  dipakai (tidak ada fallback lain karena tidak ada kolom lain yang relevan).
- CERTIFICATE diambil dari POLIS_ORI (nilai mentah, SEBELUM clean_polis
  membuang suffix cert-nya).
- Rule format (SAMA seperti Data 1 & Data 2):
    * range 2-angka ("-" / S/D / SD / S): >3 anggota -> "AWAL SD AKHIR",
      <=3 anggota -> breakdown penuh dipisah koma.
    * daftar eksplisit koma/campur S-D: >3 item -> kompres
      "ITEM_PERTAMA SD ITEM_TERAKHIR" (tanpa isi celah), <=3 -> apa adanya.
    * semua angka di-pad ke 6 digit.
    * "S/D"/"s/d"/"SD"/"S" di output SELALU distandarisasi jadi "SD".
    * pola tidak masuk akal (bukan angka murni) -> CERTIFICATE kosong.
- clean_polis() TIDAK berubah sama sekali (masih buang total suffix cert
  dari POLIS_CLEAN_1, sesuai desain asli Data 3).
=====================================================================
"""

import os
import re

import pandas as pd


# ─────────────────────────────────────────────────────────────────────────────
# KONFIGURASI
# ─────────────────────────────────────────────────────────────────────────────

INPUT_FILE  = os.path.join("input", "3b. Database Suspense 150826.xlsx")
INPUT_SHEET = "Sheet1"
HEADER_ROW  = 2          # header data ada di baris ke-3 file (index 2, 0-based)

OUTPUT_FILE = os.path.join("output", "JasaRp_output_suspend.xlsx")

CEDANT_FILTER_COL   = "CEDANT SHRT NAME"
CEDANT_FILTER_VALUE = "JASA R.P"

# Rule baru: hanya baris berstatus SUSPENSE yang dicleaning & ditampilkan;
# baris ADJUSTED dibuang total sebelum proses cleaning.
STATUS_COL          = "STATUS"
STATUS_KEEP_VALUE   = "SUSPENSE"

INSURED_COL = "INSURED"
POLIS_COL   = "POLIS"
SLIP_COL    = "SLIP NO"

MAX_SPLIT_COLS = 5


# ─────────────────────────────────────────────────────────────────────────────
# KAMUS POLA STANDAR JRP (dari Karakteristik_Polis_dan_Slip.xlsx sheet
# "Jasa Raharja" -- polis base tanpa sertifikat, 14-19 digit)
# ─────────────────────────────────────────────────────────────────────────────

KNOWN_POLIS_PATTERNS = [
    re.compile(r"^\d{14,19}$"),
]

KNOWN_SLIP_PATTERNS = [
    re.compile(r"^\d{14,19}$"),
]

# P1 / P2 / P3 dst tetap dipertahankan apa adanya (konsisten dgn cedant lain)
_POLIS_KEEP_AS_IS_RE = re.compile(r"^\s*P\d+\s*/", re.IGNORECASE)


def _matches_known_pattern(value: str, patterns) -> bool:
    return any(p.match(value) for p in patterns)


def refine_with_known_pattern(value: str, patterns) -> str:
    """Perbaikan ringan format (spasi, 0 di depan hilang, dll) SETELAH
    cleaning utama -- tidak mengubah makna/isi data."""
    if not value:
        return value
    if _matches_known_pattern(value, patterns):
        return value
    candidates = [
        re.sub(r"\s+", "", value),
        value.strip(" .-/"),
        re.sub(r"\.0+$", "", value),
        "0" + value if value.isdigit() else value,
    ]
    for c in candidates:
        if _matches_known_pattern(c, patterns):
            return c
    return value


def _normalize_spaces(text: str) -> str:
    return re.sub(r"\s{2,}", " ", text).strip()


# ─────────────────────────────────────────────────────────────────────────────
# KOLOM CERTIFICATE (rule sama persis dengan Data 1 & Data 2)
# ─────────────────────────────────────────────────────────────────────────────

_CERT_RANGE_PAIR_RE = re.compile(
    r"^\s*(\d{1,7})\s*(?:-|S\s*/\s*D|SD|S(?!\d))\s*(\d{1,7})\s*$",
    re.IGNORECASE,
)
_CERT_SPLIT_RE = re.compile(r"\s*(?:S\s*/\s*D|SD|S(?!\d)|,)\s*", re.IGNORECASE)
_CERT_TOKEN_RE = re.compile(r"^\d{1,7}$")


def build_certificate(cert_raw) -> str:
    """Format nomor sertifikat sesuai rule baru:
    - range 2-angka ("-" / S/D / SD / S): >3 anggota -> "AWAL SD AKHIR",
      <=3 anggota -> breakdown penuh dipisah koma.
    - daftar eksplisit (koma dan/atau campur S/D/SD/S, bukan range murni):
      >3 item -> dikompres "ITEM_PERTAMA SD ITEM_TERAKHIR" (tanpa isi celah),
      <=3 item -> tetap apa adanya, dipisah koma.
    - semua angka di-pad ke 6 digit.
    - "S/D"/"s/d"/"SD"/"S" di output SELALU ditulis "SD".
    - kalau pola tidak masuk akal (bukan angka murni) -> kosong.
    """
    if pd.isna(cert_raw):
        return ""
    text = str(cert_raw).strip()
    if not text or text.upper() in {"-", "NAN", "NONE", "NULL"}:
        return ""

    m = _CERT_RANGE_PAIR_RE.match(text)
    if m:
        lo, hi = int(m.group(1)), int(m.group(2))
        if hi < lo:
            lo, hi = hi, lo
        count = hi - lo + 1
        if count > 3:
            return f"{str(lo).zfill(6)} SD {str(hi).zfill(6)}"
        return ", ".join(str(n).zfill(6) for n in range(lo, hi + 1))

    raw_tokens = [t for t in _CERT_SPLIT_RE.split(text) if t.strip()]
    if not raw_tokens or not all(_CERT_TOKEN_RE.match(t.strip()) for t in raw_tokens):
        return ""
    nums = [int(t) for t in raw_tokens]
    if len(nums) > 3:
        return f"{str(nums[0]).zfill(6)} SD {str(nums[-1]).zfill(6)}"
    return ", ".join(str(n).zfill(6) for n in nums)


# Pola "base + '-' + sertifikat" yang menempel di kolom POLIS (satu-satunya
# sumber cert di Data 3, sesuai catatan arsitektur di atas).
_POLICY_CERT_TAIL_RE = re.compile(r"^\s*(?:VARIOUS\s*/\s*)?(\d{10,})\s*[-.]\s*(.+)$", re.IGNORECASE)


def extract_certificate(polis_raw) -> str:
    """Ambil blok sertifikat dari POLIS_ORI (pola base+'-'+sertifikat),
    lalu format lewat build_certificate(). Return '' kalau tidak ada pola
    sertifikat yang cocok."""
    if pd.isna(polis_raw):
        return ""
    val = str(polis_raw).strip()
    if not val:
        return ""

    m = _POLICY_CERT_TAIL_RE.match(val)
    if not m:
        return ""
    base, rest = m.groups()

    first_num_m = re.match(r"^\s*(\d+)", rest)
    if not first_num_m or not (5 <= len(first_num_m.group(1)) <= 7):
        return ""

    rest = re.sub(r"\s*/\s*P\d+\s*$", "", rest, flags=re.IGNORECASE)
    rest = re.sub(r"\s*/\s*VARIOUS\s*$", "", rest, flags=re.IGNORECASE)
    rest = rest.strip()
    if not rest or not re.search(r"\d", rest):
        return ""

    return build_certificate(rest)


# ─────────────────────────────────────────────────────────────────────────────
# FUNGSI CLEANING POLIS/SLIP (TIDAK BERUBAH -- masih buang total suffix cert)
# ─────────────────────────────────────────────────────────────────────────────

def clean_polis(val) -> str:
    """Bersihkan nomor polis. P1/P2/dst di depan dibiarkan apa adanya.
    Suffix sertifikat (mis. "-000023", "-000023/1") DIBUANG total -- beda
    dari Data 2/Osbal yang justru mempertahankannya."""
    if pd.isna(val):
        return ""
    val = str(val).strip()
    if not val:
        return ""

    if _POLIS_KEEP_AS_IS_RE.match(val):
        return val

    p = val
    while True:
        stripped = re.sub(r"-\d+(?:/\d+)?$", "", p)
        if stripped == p:
            break
        p = stripped

    return p if p else val


def clean_slip(val) -> str:
    """Slip Data 3 JRP sudah bersih (14-19 digit polos) -- cukup trim."""
    if pd.isna(val):
        return ""
    return str(val).strip()


def _remove_polis_slip_from_text(text: str, polis_ori, slip_ori) -> str:
    if pd.notna(polis_ori):
        token = str(polis_ori).strip()
        if token and token != "-":
            text = text.replace(token, "")
    if pd.notna(slip_ori):
        token = str(slip_ori).strip()
        if token and token != "-":
            text = text.replace(token, "")
    return text


# ─────────────────────────────────────────────────────────────────────────────
# CLEAN INSURED (pipeline standar -- PT/CV/Persero strip)
# ─────────────────────────────────────────────────────────────────────────────

_INSURED_TAIL_RE = re.compile(
    r"""
    \bAS\b\s+(?:THE\s+)?(?:PRINCIPAL|OFF-TAKER|MAINTENANCE|CONTRACTOR).*
  | \bBEING\s+(?:THE\s+)?(?:PRINCIPAL|OFF-TAKER).*
  | \bAND\s+ALL\s+SUBSIDIAR.*
  | \bAND\s+ITS\s+SUBSIDIAR.*
  | \bINCLUDING\s+ALL\s+SUBSIDIAR.*
  | \bINCLUDING\s+ANY\s+SUBSIDIAR.*
  | \bANY\s+SUBSIDIAR.*
  | \bCOMPRISING\s+OF.*
  | \bINSTALLMENT\b.*
  | \bRELATED\s+COMPANY\b.*
  | \bPURCHASED\s+OR\s+OTHERWISE\b.*
    """,
    re.IGNORECASE | re.VERBOSE,
)

_ATTN_TAIL_RE = re.compile(r"\s*-\s*ATTN\b.*$", re.IGNORECASE)

_INSURED_JUNK_WORDS = frozenset({
    "SHANGHAI", "PR OF CHINA", "CHINA", "INDONESIA", "JAKARTA",
    "OFFICERS", "EMPLOYEES", "ALL OTHER CONTRACTORS",
    "SUB-CONTRACTORS", "SUB CONTRACTORS",
    "COMPANIES", "AFFILIATED", "AFFILIATES",
    "SUBSIDIARY", "SUBSIDIARIES", "ANY SUBSIDIARY COMPANY",
    "RELATED COMPANY", "THE PRINCIPAL", "PRINCIPAL", "OWNER",
})


def _normalize_insured_part(p: str) -> str:
    p = _INSURED_TAIL_RE.sub("", p)
    p = re.sub(
        r"\b(?:BAPAK|IBU|SDRI?|NYONYA|NY|TUAN|MR|MRS|MS|"
        r"DR|IR|PROF|HJ|H)\b\.?\s*",
        "", p, flags=re.IGNORECASE,
    )
    p = re.sub(r"\(([^()]*)\)", "", p)
    p = re.sub(r"^(?:AND|OR)\b\s*", "", p, flags=re.IGNORECASE)
    p = re.sub(r"\s*\b(?:AND|OR)$", "", p, flags=re.IGNORECASE)
    p = re.sub(r"^[^a-zA-Z0-9(]+", "", p)
    p = re.sub(r"[^a-zA-Z0-9)]+$", "", p)
    p = re.sub(r"\s{2,}", " ", p)
    return p.strip()


def _is_valid_insured_part(p: str) -> bool:
    if len(p) <= 2:
        return False
    up = p.upper()
    if re.match(r"^[\d\/\-\.]+$", up):
        return False
    if re.search(r"\b(?:NO\.\s*\d+|BUILDING|FLOOR|ROOM|ROAD|STREET|TOWER|KAV\.?|BLOK)\b", up):
        return False
    return up not in _INSURED_JUNK_WORDS


def _merge_digit_start_parts(parts: list) -> list:
    merged = []
    for p in parts:
        p_stripped = p.strip()
        if merged and re.match(r"^\d", p_stripped):
            merged[-1] = merged[-1] + " " + p_stripped
        else:
            merged.append(p_stripped)
    return merged


_BRANCH_KEYWORD_RE = re.compile(r"^(?:KCU|KANWIL|DTC|CABANG)\b", re.IGNORECASE)


def _merge_branch_keyword_parts(parts: list) -> list:
    merged = []
    for p in parts:
        p_stripped = p.strip()
        if merged and _BRANCH_KEYWORD_RE.match(p_stripped):
            merged[-1] = merged[-1] + " " + p_stripped
        else:
            merged.append(p_stripped)
    return merged


def _name_initials_variants(name: str) -> set:
    words = [w for w in re.split(r"[\s\-]+", name.upper()) if w]
    if not words:
        return set()
    variants = {"".join(w[0] for w in words if w[0].isalpha())}
    if len(words) >= 2:
        variants.add("".join(w[0] for w in words[:-1] if w[0].isalpha()) + words[-1])
    return variants


def _is_abbreviation_of(short: str, long_name: str) -> bool:
    short_clean = re.sub(r"[^A-Z0-9]", "", short.upper())
    if len(short_clean) < 2 or len(short_clean) > 8:
        return False
    return short_clean in _name_initials_variants(long_name)


def _strip_trailing_abbrev_word(name: str, priors: list) -> str:
    words = name.split()
    if len(words) >= 2 and any(_is_abbreviation_of(words[-1], p) for p in priors):
        return " ".join(words[:-1])
    return name


def clean_insured(val, polis_ori, slip_ori) -> list:
    if pd.isna(val):
        return []
    val = str(val).strip()
    if not val:
        return []

    val = _remove_polis_slip_from_text(val, polis_ori, slip_ori)
    val = _ATTN_TAIL_RE.sub("", val)

    val = re.sub(r"\(\s*[\d\.\/\-]+\s*\)", "", val)
    val = re.sub(r"\b(?:AND|AN|OR)\s*/\s*(?:AND|OR)\b", ",", val, flags=re.IGNORECASE)
    val = re.sub(r"\bAND\s+OR\b", ",", val, flags=re.IGNORECASE)
    val = re.sub(r"\bCO\.,?\s*LTD\.?\b", ",", val, flags=re.IGNORECASE)
    for pattern in [r"\bTBK\.?\b", r"\(PERSERO\)", r"\bPERSERO\b", r"\bLTD\.?\b"]:
        val = re.sub(pattern, "", val, flags=re.IGNORECASE)

    split_pattern = r"\bQQ\b|/|,|\d+\.|\bPT\.?\b|\bCV\.?\b|:|;"
    parts = re.split(split_pattern, val, flags=re.IGNORECASE)
    parts = _merge_digit_start_parts(parts)
    parts = _merge_branch_keyword_parts(parts)

    cleaned = []
    for p in parts:
        p = _normalize_insured_part(p)
        if _is_valid_insured_part(p):
            cleaned.append(p.replace(".", "").upper())

    if not cleaned:
        fallback = _normalize_insured_part(val).replace(".", "").upper()
        return [fallback] if fallback else []

    final = []
    for name in cleaned:
        name = _strip_trailing_abbrev_word(name, final)
        if any(_is_abbreviation_of(name, prior) for prior in final):
            continue
        final.append(name)
    cleaned = final

    if len(cleaned) > MAX_SPLIT_COLS:
        return [", ".join(cleaned)]

    return cleaned


# ─────────────────────────────────────────────────────────────────────────────
# PROSES UTAMA
# ─────────────────────────────────────────────────────────────────────────────

def _force_text_format(output_file, target_cols_names):
    import openpyxl
    wb = openpyxl.load_workbook(output_file)
    ws = wb.active
    header = [c.value for c in ws[1]]
    target_idx = [i for i, h in enumerate(header, start=1) if h in target_cols_names]
    for col_idx in target_idx:
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                cell.number_format = "@"
    wb.save(output_file)


def process_data(input_file: str, output_file: str) -> None:
    print(f"[1/5] Membaca data dari: {input_file} (sheet: {INPUT_SHEET}) ...")
    df = pd.read_excel(input_file, sheet_name=INPUT_SHEET, header=HEADER_ROW)
    print(f"      Total baris keseluruhan: {len(df):,}")

    if CEDANT_FILTER_COL not in df.columns:
        print(f"\n[ERROR] Kolom '{CEDANT_FILTER_COL}' tidak ditemukan!")
        print(f"        Kolom tersedia: {list(df.columns)}")
        return

    df[CEDANT_FILTER_COL] = df[CEDANT_FILTER_COL].astype(str).str.strip()
    df = df[df[CEDANT_FILTER_COL] == CEDANT_FILTER_VALUE].copy()
    print(f"[2/5] Filter cedant '{CEDANT_FILTER_VALUE}': {len(df):,} baris ditemukan.")

    if df.empty:
        print("\n[WARN] Tidak ada data setelah filter. Proses dihentikan.")
        return

    if STATUS_COL not in df.columns:
        print(f"\n[ERROR] Kolom '{STATUS_COL}' tidak ditemukan!")
        print(f"        Kolom tersedia: {list(df.columns)}")
        return

    before_status = len(df)
    df[STATUS_COL] = df[STATUS_COL].astype(str).str.strip()
    df = df[df[STATUS_COL] == STATUS_KEEP_VALUE].copy()
    print(f"[2b/5] Filter status '{STATUS_KEEP_VALUE}': {len(df):,} baris "
          f"(dibuang {before_status - len(df):,} baris ADJUSTED).")

    if df.empty:
        print("\n[WARN] Tidak ada data setelah filter status. Proses dihentikan.")
        return

    for col in [INSURED_COL, POLIS_COL, SLIP_COL]:
        if col not in df.columns:
            print(f"\n[ERROR] Kolom '{col}' tidak ditemukan di file!")
            print(f"        Kolom tersedia: {list(df.columns)}")
            return

    df.rename(columns={
        INSURED_COL: "INSURED_ORI",
        POLIS_COL:   "POLIS_ORI",
        SLIP_COL:    "SLIP_NO_ORI",
    }, inplace=True)

    print("[3/5] Menjalankan proses cleaning ...")

    polis_cln, slip_cln, all_certificate, all_insured_parts = [], [], [], []
    max_insured_parts = 1

    for _, row in df.iterrows():
        p_ori = row.get("POLIS_ORI", "")
        s_ori = row.get("SLIP_NO_ORI", "")
        i_ori = row.get("INSURED_ORI", "")

        polis_cln.append(refine_with_known_pattern(clean_polis(p_ori), KNOWN_POLIS_PATTERNS))
        slip_cln.append(refine_with_known_pattern(clean_slip(s_ori), KNOWN_SLIP_PATTERNS))
        all_certificate.append(extract_certificate(p_ori))

        parts = clean_insured(i_ori, p_ori, s_ori)
        max_insured_parts = max(max_insured_parts, len(parts))
        all_insured_parts.append(parts)

    print("[4/5] Menyusun kolom output ...")

    # Rename: POLIS_CLN -> POLIS_CLEAN_1, CERTIFICATE -> CERTIFICATE_1.
    # Data 3 selalu cuma 1 polis per baris (clean_polis mengembalikan
    # string tunggal, bukan list), jadi tidak ada POLIS_CLEAN_2 dst.
    df["POLIS_CLEAN_1"] = polis_cln
    df["CERTIFICATE_1"] = all_certificate
    df["SLIP_CLEAN_1"]   = slip_cln

    insured_cols = []
    for i in range(1, max_insured_parts + 1):
        col_name = f"INSURED_{i}"
        df[col_name] = [p[i - 1] if i - 1 < len(p) else None for p in all_insured_parts]
        insured_cols.append(col_name)

    new_columns = []
    for col in df.columns:
        if col in ("POLIS_CLEAN_1", "CERTIFICATE_1", "SLIP_CLEAN_1") or col in insured_cols:
            continue
        new_columns.append(col)
        if col == "INSURED_ORI":
            new_columns += insured_cols
        elif col == "POLIS_ORI":
            # Urutan baru: POLIS_ORI -> POLIS_CLEAN_1 -> CERTIFICATE_1
            new_columns.append("POLIS_CLEAN_1")
            new_columns.append("CERTIFICATE_1")
        elif col == "SLIP_NO_ORI":
            new_columns.append("SLIP_CLEAN_1")

    df = df[new_columns]

    os.makedirs(os.path.dirname(output_file), exist_ok=True)
    print(f"[5/5] Menyimpan hasil ke: {output_file} ...")
    df.to_excel(output_file, index=False)
    _force_text_format(output_file, target_cols_names={"POLIS_ORI", "POLIS_CLEAN_1", "CERTIFICATE_1", "SLIP_NO_ORI", "SLIP_CLEAN_1"})

    print(f"\n{'=' * 55}")
    print("  [OK] Selesai!")
    print(f"  Total baris output  : {len(df):,}")
    print(f"  Total kolom output  : {len(df.columns)}")
    print(f"  File disimpan di    : {output_file}")
    print(f"{'=' * 55}")


if __name__ == "__main__":
    process_data(INPUT_FILE, OUTPUT_FILE)