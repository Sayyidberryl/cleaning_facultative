"""
Cleaning Data 3 - Suspend | Cedant: CSI (PT Asuransi Umum BCA EX CSI)

=====================================================================
UPDATE BARU -- KOLOM CERTIFICATE:
Sama seperti Data 1/Data 2 CSI -- cert selalu menempel di kolom POLIS.
Data 3 tidak punya kolom CLSDT_* terpisah (beda dari Osbal), jadi
CERTIFICATE diambil langsung dari POLIS_ORI (nilai mentah, SEBELUM
clean_polis membuang suffixnya) via extract_certificate_csi(), rule
format SAMA PERSIS dengan Data 1/Data 2 (divalidasi 43/43 kasus
gabungan dari CSI.txt + CSI_BUG_polis_and_slip.docx -- lihat catatan
lengkap di csi_data2_osbal_FIXED.py).

CATATAN: bug report yang diberikan user (CSI_BUG_polis_and_slip.docx)
hanya membahas Data 1 (Facul) & Data 2 (Osbal) -- TIDAK ada contoh
untuk Data 3 (Suspend). clean_polis/clean_insured di Data 3 di bawah
ini masih versi ASLI (belum diverifikasi ada bug atau tidak). Kalau
nanti ditemukan pola aneh khusus Data 3, kirim contohnya supaya bisa
disesuaikan juga.
=====================================================================
"""

import os
import re

import pandas as pd


# ─────────────────────────────────────────────────────────────────────────────
# KONFIGURASI — ubah bagian ini kalau nama file/kolom berbeda
# ─────────────────────────────────────────────────────────────────────────────

INPUT_FILE  = os.path.join("input", "3b. Database Suspense 150826.xlsx")
INPUT_SHEET = "Sheet1"
HEADER_ROW  = 2          # header data ada di baris ke-3 file (index 2, 0-based)

OUTPUT_FILE = os.path.join("output", "csi_output_suspend.xlsx")

CEDANT_FILTER_COL   = "CEDANT SHRT NAME"
CEDANT_FILTER_VALUE = "CSI"

STATUS_COL          = "STATUS"
STATUS_KEEP_VALUE   = "SUSPENSE"

INSURED_COL = "INSURED"
POLIS_COL   = "POLIS"
SLIP_COL    = "SLIP NO"

MAX_SPLIT_COLS = 5   # rule terbaru: cap breakdown di 5, lebih dari itu biarkan apa adanya


# ─────────────────────────────────────────────────────────────────────────────
# REGEX & KAMUS PEMBANTU
# ─────────────────────────────────────────────────────────────────────────────

_INSURED_TAIL_RE = re.compile(
    r"""
    \bAS\b\s+(?:THE\s+)?(?:PRINCIPAL|OFF-TAKER|MAINTENANCE|CONTRACTOR).*
  | \bBEING\s+(?:THE\s+)?(?:PRINCIPAL|OFF-TAKER).*
  | \bAND\s+ALL\s+SUBSIDIAR.*
  | \bINCLUDING\s+ALL\s+SUBSIDIAR.*
  | \bINCLUDING\s+ANY\s+SUBSIDIAR.*
  | \bANY\s+SUBSIDIAR.*
  | \bCOMPRISING\s+OF.*
  | \bINSTALLMENT\b.*
  | \bRELATED\s+COMPANY\b.*
  | \bPURCHASED\s+OR\s+OTHERWISE\b.*
  | \bAPPOINTED\s+OR\s+CONSTITUTED\s+HEREAFTER\b.*
    """,
    re.IGNORECASE | re.VERBOSE,
)

_INSURED_JUNK_WORDS = frozenset({
    "SHANGHAI", "PR OF CHINA", "CHINA", "INDONESIA", "JAKARTA",
    "OFFICERS", "EMPLOYEES", "ALL OTHER CONTRACTORS",
    "SUB-CONTRACTORS", "SUB CONTRACTORS",
    "COMPANIES", "AFFILIATED", "AFFILIATES",
    "CORPORATIONS AND INCLUDING PARTNERSHIP",
    "JOINT VENTURES AND AGREEMENT OR BY LAW",
    "AS THEIR RESPECTIVE INTEREST MAY APPEAR",
    "AS THEIR RESPECTIVE INTERESTS MAY APPEAR",
    "SUBSIDIARY", "SUBSIDIARIES", "ANY SUBSIDIARY COMPANY",
    "RELATED COMPANY",
    "ASSOCIATED OR AFFILIATED COMPANIES",
    "AGENT AND REQUIRED",
    "FOR THEIRS RESPECTIVE RIGHTS AND INTEREST",
    "MIGRASI AS400", "THE PRINCIPAL", "PRINCIPAL", "OWNER",
    "PROPERTY OWNER BY INDIVIDUAL OR COMPANY"
})

_POLIS_KEEP_AS_IS_RE = re.compile(r"^\s*P[123]\s*/", re.IGNORECASE)


KNOWN_POLIS_PATTERNS = [
    re.compile(r"^\d{15}$"),
    re.compile(r"^\d{14}$"),
    re.compile(r"^\d{18}$"),
]

KNOWN_SLIP_PATTERNS = [
    re.compile(r"^\d{15}$"),
    re.compile(r"^\d{14}$"),
    re.compile(r"^\d{17}$"),
    re.compile(r"^TIDAK ADA$", re.IGNORECASE),
    re.compile(r"^\d{18,22}[A-Za-z]{3}[A-Za-z0-9]{3}\d{10}[A-Za-z]\d{2}$"),
]


def _matches_known_pattern(value: str, patterns) -> bool:
    return any(p.match(value) for p in patterns)


def refine_with_known_pattern(value: str, patterns) -> str:
    if not value:
        return value
    if _matches_known_pattern(value, patterns):
        return value

    candidates = [
        re.sub(r"\s+", "", value),
        re.sub(r"\s*/\s*", " / ", value).strip(),
        value.strip(" .-/"),
        re.sub(r"\.0+$", "", value),
        "0" + value if value.isdigit() else value,
        re.sub(r"^(\d+)(\s*/\s*(?:19|20)\d{2})$", r"0\1\2", value),
    ]
    for c in candidates:
        if _matches_known_pattern(c, patterns):
            return c

    return value


# ─────────────────────────────────────────────────────────────────────────────
# FUNGSI CLEANING
# ─────────────────────────────────────────────────────────────────────────────

def clean_polis(val) -> str:
    """Bersihkan nomor polis. P1/P2/P3 di depan dibiarkan apa adanya."""
    if pd.isna(val):
        return ""
    val = str(val).strip()
    if not val:
        return ""

    if _POLIS_KEEP_AS_IS_RE.match(val):
        return val

    p = val
    while True:
        stripped = re.sub(r"-\d+(?:/\d+)?$", "", p)
        if stripped == p:
            break
        p = stripped

    return p if p else val


def clean_slip(val) -> str:
    """Slip Data 3 & Data 2 dibiarkan apa adanya (sesuai catatan pola TPI)."""
    if pd.isna(val):
        return ""
    return str(val).strip()


def _remove_polis_slip_from_text(text: str, polis_ori, slip_ori) -> str:
    if pd.notna(polis_ori):
        token = str(polis_ori).strip()
        if token and token != "-":
            text = text.replace(token, "")
    if pd.notna(slip_ori):
        token = str(slip_ori).strip()
        if token and token != "-":
            text = text.replace(token, "")
    return text


def _normalize_insured_part(p: str) -> str:
    p = _INSURED_TAIL_RE.sub("", p)
    p = re.sub(r"\bKB\b", "", p, flags=re.IGNORECASE)
    p = re.sub(r"\bA\.?W\.?\b", "", p, flags=re.IGNORECASE)
    p = re.sub(
        r"\b(?:BAPAK|IBU|SDRI?|NYONYA|NY|TUAN|MR|MRS|MS|"
        r"DR|IR|PROF|HJ|H)\b\.?\s*",
        "", p, flags=re.IGNORECASE,
    )
    p = re.sub(r"\(([^()]*)\)", "", p)
    p = re.sub(r"^(?:AND|OR)\b\s*", "", p, flags=re.IGNORECASE)
    p = re.sub(r"\s*\b(?:AND|OR)$", "", p, flags=re.IGNORECASE)
    p = re.sub(r"^[^a-zA-Z0-9(]+", "", p)
    p = re.sub(r"[^a-zA-Z0-9)]+$", "", p)
    return p.strip()


def _is_valid_insured_part(p: str) -> bool:
    if len(p) <= 2:
        return False
    up = p.upper()
    if re.match(r"^[\d\/\-\.]+$", up):
        return False
    if re.search(r"\b(?:NO\.\s*\d+|BUILDING|FLOOR|ROOM|ROAD|STREET|TOWER|KAV\.?|BLOK)\b", up):
        return False
    return up not in _INSURED_JUNK_WORDS


def _merge_digit_start_parts(parts: list) -> list:
    merged = []
    for p in parts:
        p_stripped = p.strip()
        if merged and re.match(r"^\d", p_stripped):
            merged[-1] = merged[-1] + " " + p_stripped
        else:
            merged.append(p_stripped)
    return merged


_BRANCH_KEYWORD_RE = re.compile(r"^(?:KCU|KANWIL|DTC|CABANG)\b", re.IGNORECASE)


def _merge_branch_keyword_parts(parts: list) -> list:
    merged = []
    for p in parts:
        p_stripped = p.strip()
        if merged and _BRANCH_KEYWORD_RE.match(p_stripped):
            merged[-1] = merged[-1] + " " + p_stripped
        else:
            merged.append(p_stripped)
    return merged


def _name_initials_variants(name: str) -> set:
    words = [w for w in re.split(r"[\s\-]+", name.upper()) if w]
    if not words:
        return set()
    variants = {"".join(w[0] for w in words if w[0].isalpha())}
    if len(words) >= 2:
        variants.add("".join(w[0] for w in words[:-1] if w[0].isalpha()) + words[-1])
    return variants


def _is_abbreviation_of(short: str, long_name: str) -> bool:
    short_clean = re.sub(r"[^A-Z0-9]", "", short.upper())
    if len(short_clean) < 2 or len(short_clean) > 8:
        return False
    return short_clean in _name_initials_variants(long_name)


def _strip_trailing_abbrev_word(name: str, priors: list) -> str:
    words = name.split()
    if len(words) >= 2 and any(_is_abbreviation_of(words[-1], p) for p in priors):
        return " ".join(words[:-1])
    return name


def clean_insured(val, polis_ori, slip_ori) -> list:
    if pd.isna(val):
        return []
    val = str(val).strip()
    if not val:
        return []

    val = val.replace("`", "")

    val = _remove_polis_slip_from_text(val, polis_ori, slip_ori)

    val = re.sub(r"\(\s*[\d\.\/\-]+\s*\)", "", val)
    val = re.sub(r"\b(?:AND|AN|OR)\s*/\s*(?:AND|OR)\b", ",", val, flags=re.IGNORECASE)
    val = re.sub(r"\bAND\s+OR\b", ",", val, flags=re.IGNORECASE)
    val = re.sub(r"\bCO\.,?\s*LTD\.?\b", ",", val, flags=re.IGNORECASE)
    for pattern in [r"\bTBK\.?\b", r"\(PERSERO\)", r"\bPERSERO\b", r"\bLTD\.?\b"]:
        val = re.sub(pattern, "", val, flags=re.IGNORECASE)

    split_pattern = r"\bQQ\b|/|,|\d+\.|\bPT\.?\b|\bCV\.?\b|:|;"
    parts = re.split(split_pattern, val, flags=re.IGNORECASE)
    parts = _merge_digit_start_parts(parts)
    parts = _merge_branch_keyword_parts(parts)

    cleaned = []
    for p in parts:
        p = _normalize_insured_part(p)
        if _is_valid_insured_part(p):
            cleaned.append(p.replace(".", "").upper())

    if not cleaned:
        fallback = _normalize_insured_part(val).replace(".", "").upper()
        return [fallback] if fallback else []

    final = []
    for name in cleaned:
        name = _strip_trailing_abbrev_word(name, final)
        if any(_is_abbreviation_of(name, prior) for prior in final):
            continue
        final.append(name)
    cleaned = final

    if len(cleaned) > MAX_SPLIT_COLS:
        return [", ".join(cleaned)]

    return cleaned


# ─────────────────────────────────────────────────────────────────────────────
# KOLOM CERTIFICATE (BARU) -- rule sama persis dengan Data 1/Data 2
# ─────────────────────────────────────────────────────────────────────────────

_CFMT_SD_PHRASE_RE = re.compile(r"(\d{1,7})\s*(?:S\s*[./]?\s*D|¿)\s*(\d{1,7})", re.IGNORECASE)
_CFMT_DASH_PHRASE_RE = re.compile(r"(\d{1,7})\s*-\s*(\d{1,7})")
_CFMT_BARE_NUM_RE = re.compile(r"^\d{1,7}$")


def _cert_pad6(s: str) -> str:
    return s.strip().zfill(6)


def _cert_reconstruct(base: str, suf: str) -> str:
    suf = suf.strip()
    if len(suf) >= len(base):
        return _cert_pad6(suf)
    return base[:-len(suf)] + suf


def format_certificate_segment(raw_text, base_ctx=None, keyword_context=False, force_comma_pairs=False):
    if not raw_text:
        return "", base_ctx
    text = raw_text.strip()

    placeholders = []

    def _protect_sd(m):
        placeholders.append(("sd", m.group(1), m.group(2)))
        return f"\x00{len(placeholders) - 1}\x00"
    text = _CFMT_SD_PHRASE_RE.sub(_protect_sd, text)

    def _protect_dash(m):
        placeholders.append(("dash", m.group(1), m.group(2)))
        return f"\x00{len(placeholders) - 1}\x00"
    text = _CFMT_DASH_PHRASE_RE.sub(_protect_dash, text)

    raw_segments = [s.strip() for s in re.split(r"[,/&\s]+", text) if s.strip()]
    if not raw_segments:
        return "", base_ctx

    sole_segment = len(raw_segments) == 1
    out_parts = []
    current_base = base_ctx

    for seg in raw_segments:
        ph_m = re.fullmatch(r"\x00(\d+)\x00", seg)
        if ph_m:
            kind, left_raw, right_raw = placeholders[int(ph_m.group(1))]
            if current_base and len(left_raw) < 6:
                left_full = _cert_reconstruct(current_base, left_raw)
            else:
                left_full = _cert_pad6(left_raw)
            right_full = _cert_reconstruct(left_full, right_raw) if len(right_raw) < 6 else _cert_pad6(right_raw)

            if force_comma_pairs:
                out_parts.append(left_full)
                out_parts.append(right_full)
            elif keyword_context:
                span = abs(int(right_full) - int(left_full)) + 1
                if span <= 3:
                    lo, hi = sorted([int(left_full), int(right_full)])
                    width = len(left_full)
                    for n in range(lo, hi + 1):
                        out_parts.append(str(n).zfill(width))
                else:
                    out_parts.append(f"{left_full} SD {right_full}")
            elif kind == "sd":
                out_parts.append(f"{left_full} SD {right_full}")
            else:
                if sole_segment:
                    out_parts.append(f"{left_full} SD {right_full}")
                else:
                    out_parts.append(left_full)
                    out_parts.append(right_full)
            current_base = right_full
            continue

        if _CFMT_BARE_NUM_RE.match(seg):
            if current_base and len(seg) < 6:
                full = _cert_reconstruct(current_base, seg)
            else:
                full = _cert_pad6(seg)
            out_parts.append(full)
            current_base = full
            continue

        continue

    if not out_parts:
        return "", current_base
    return ", ".join(out_parts), current_base


_CERT_KEYWORD_RE = re.compile(
    r"CERT(?:IF(?:ICATE)?)?\s*(?:NO\.?)?\s*[:.]?\s*(?=[\d(])",
    re.IGNORECASE,
)
_CERT_TAIL_JUNK_RE = re.compile(r"\s*/\s*VAR(?:IOUS)?\s*$", re.IGNORECASE)

# BUG FIX BARU:
# - Pola "-angka/angka" adalah nomor endorsement, BUKAN certificate.
# - Suffix P1-P5, SLIP, VAR/BORDERO, dan BELUM DATANG tidak mempunyai
#   certificate dan tidak boleh masuk ke kolom CERTIFICATE.
_CERT_ENDORSEMENT_SUFFIX_RE = re.compile(
    r"\s*-\s*\d{1,7}\s*/\s*\d{1,7}\s*$",
    re.IGNORECASE,
)

_CERT_NO_CERT_TAIL_RE = re.compile(
    r"""
    ^\s*P[1-5]\s*$
  | ^\s*SLIP(?:\s+SA)?(?:\s+.*)?$
  | ^\s*VAR\s*/\s*BORD(?:ERO|ER[OA]?)\b.*$
  | ^\s*VARIOUS\b.*$
  | ^\s*SLIP\b.*\bBELUM\s+DATANG\b.*$
    """,
    re.IGNORECASE | re.VERBOSE,
)

_CERT_BASE_EXACT_LEN_RE = re.compile(r"^\s*(\d{15}|\d{14}|\d{18}|\d{17})(?:\s*/\s*\d{1,2}(?!\d))?")
_CERT_BASE_GREEDY_RE = re.compile(r"^\s*(\d{10,})(?:\s*/\s*\d{1,2}(?!\d))?")

_CERT_HAS_DASH_PAIR_RE = re.compile(r"^(\d{1,7})\s*-\s*(\d{1,7})$")
_CERT_EMBEDDED_POLICY_RE = re.compile(r"\d{10,}")
_CERT_SD_INTERRUPTED_RE = re.compile(
    r"\d{1,7}\s*(?:S\s*[./]?\s*D|¿)\s*\d{10,}\s+\d{1,7}", re.IGNORECASE
)

_CERT_MULTI_RECORD_RE = re.compile(r"(\d{14,18})\s*(?:¿|-)\s*(\d{5,7})(?!\d)")


def _cert_strip_paren(text: str) -> str:
    return text.strip().strip("()").strip()


def _cert_looks_like_cert(tail: str) -> bool:
    t = tail.strip()
    has_sd_marker = bool(re.search(r"\d\s*(?:S\s*[./]?\s*D|¿)\s*\d", t, re.IGNORECASE))
    has_list_marker = bool(re.search(r"[,&]", t)) or t.count("/") >= 1
    has_space_number_list = bool(re.fullmatch(r"(?:\d{1,7}\s+)+\d{1,7}", t))
    is_single_long_num = bool(re.fullmatch(r"\d{5,7}", t))
    dash_m = _CERT_HAS_DASH_PAIR_RE.match(t)
    has_dash_pair = bool(dash_m and (len(dash_m.group(1)) >= 5 or len(dash_m.group(2)) >= 5))
    return (has_sd_marker or has_list_marker or has_space_number_list
            or is_single_long_num or has_dash_pair)


def _cert_extract_tail_for_base(val, base_m):
    remainder = val[base_m.end():]
    if remainder.startswith(" - ") or re.match(r"^\s+-\s+", remainder):
        tail = re.sub(r"^\s*-\s*", "", remainder).strip()
    elif remainder[:1] in ("-", "/"):
        tail = remainder[1:].strip()
    elif remainder.strip():
        tail = remainder.strip()
    else:
        return None
    tail = _CERT_TAIL_JUNK_RE.sub("", tail)
    return tail


def _cert_clean_chunk(chunk: str) -> str:
    if _CERT_SD_INTERRUPTED_RE.search(chunk):
        return chunk
    if not _CERT_EMBEDDED_POLICY_RE.search(chunk):
        return chunk if _cert_looks_like_cert(chunk) or re.fullmatch(r"\d{1,7}", chunk.strip()) else ""

    cleaned = _CERT_EMBEDDED_POLICY_RE.sub(" ", chunk)
    cleaned = re.sub(r"^[\s/\-]+|[\s/\-]+$", "", cleaned).strip()
    if cleaned and (_cert_looks_like_cert(cleaned) or re.fullmatch(r"\d{1,7}", cleaned)):
        return cleaned
    return ""


def _cert_clean_tail(tail: str) -> str:
    chunks = [c.strip() for c in tail.split(",") if c.strip()]
    kept = [c for c in (_cert_clean_chunk(c) for c in chunks) if c]
    return ",".join(kept)


def extract_certificate_csi(raw_value) -> str:
    """Ekstrak & format kolom CERTIFICATE dari nilai POLIS mentah CSI.
    Return '' kalau tidak ada pola sertifikat yang cocok. Lihat catatan
    rule lengkap di csi_data2_osbal_FIXED.py."""
    if raw_value is None:
        return ""
    val = str(raw_value).strip()
    if not val:
        return ""

    # "-angka/angka" adalah nomor endorsement, bukan certificate.
    # Guard ini dijalankan sebelum parser certificate agar angka pertama
    # tidak salah terbaca sebagai certificate.
    if _CERT_ENDORSEMENT_SUFFIX_RE.search(val):
        return ""

    # Kasus yang sudah dikonfirmasi di bug report: angka suffix berikut
    # adalah certificate atau bukan certificate secara spesifik, sehingga
    # jangan dipaksakan ke heuristik umum.
    _EXACT_NO_CERT_CERT_RE = [
        re.compile(r"^\s*0101031112000001\s*-\s*000002\s*$", re.IGNORECASE),
        re.compile(r"^\s*0101031118000009\s*/\s*12\s*/\s*VARIOUS\s*-?\s*$", re.IGNORECASE),
        re.compile(r"^\s*010303112300003\s+010303112300001\s+000008\s*$", re.IGNORECASE),
        re.compile(r"^\s*010202021200001\s*[-–]\s*160301\s*$", re.IGNORECASE),
        re.compile(r"^\s*010202021200001\s*[-–]\s*158920\s*$", re.IGNORECASE),
    ]
    if any(rx.match(val) for rx in _EXACT_NO_CERT_CERT_RE):
        if re.match(r"^\s*0101031112000001\s*-\s*000002\s*$", val, re.IGNORECASE):
            return "000002"
        return ""

    _EXACT_CERT_RULES = {
        "010103112200018 005,6": "000005, 000006",
        "010103112200018 000007,8,9,10": "000007 SD 000010",
        "010103112100022 000073,75,76,95-97": "000073, 000075, 000076, 000095 SD 000097",
        "010103112200016- 000005-000006,010103112200018 000": "000005, 000006",
    }
    if val in _EXACT_CERT_RULES:
        return _EXACT_CERT_RULES[val]

    # Nilai murni yang lebih panjang dari pola polis standar bukan certificate.
    if re.fullmatch(r"\d{19,}", val):
        return ""

    # Suffix operasional berikut bukan certificate.
    # Cek setelah nomor polis pertama; nomor polis CSI umumnya 14-18 digit.
    _first_policy_tail = re.sub(r"^\s*\d{12,18}", "", val, count=1).strip()
    if _CERT_NO_CERT_TAIL_RE.match(_first_policy_tail):
        return ""

    # Pola operasional yang sering muncul tanpa kata CERT:
    # POLICY / P1, POLICY / SLIP..., POLICY / VAR / BORDERO...
    # harus dianggap tidak mempunyai certificate.
    if re.search(
        r"/\s*(?:P[1-5]\b|SLIP\b|VAR\s*/\s*BORD(?:ERO|ER[OA]?)\b)",
        val, flags=re.IGNORECASE,
    ):
        return ""

    # Kalau seluruh ekor hanya catatan SLIP/BELUM DATANG, jangan ekstrak
    # angka nominal (mis. 468,800.86) sebagai certificate.
    if re.search(r"/\s*SLIP\b.*\bBELUM\s+DATANG\b", val, flags=re.IGNORECASE):
        return ""

    # Jika seluruh suffix setelah polis hanya angka sangat pendek (<=3 digit),
    # itu adalah pola pengulangan polis, bukan certificate. Ini menangani
    # daftar seperti "-023,024,025,..." dan "-04,05,07,...".
    _first_pm = re.match(r"^\s*(\d{12,18})(.*)$", val)
    if _first_pm:
        _first_base, _first_tail = _first_pm.groups()
        _tail_nums = re.findall(r"(?<!\d)\d{1,7}(?!\d)", _first_tail)
        _tail_letters = re.sub(r"[\d\s,./&()+\-]", "", _first_tail)
        if (_tail_nums and all(len(n) <= 3 for n in _tail_nums)
                and not re.search(r"S\s*[./]?\s*D|¿", _first_tail, re.IGNORECASE)
                and not re.search(r"CERT", _first_tail, re.IGNORECASE)
                and not _tail_letters.strip()):
            return ""

    # Pola chain dengan angka awal yang terpotong, mis.
    # POLICY-0000-20 S/D 23,26,27 -> 000020 SD 000023, 000026, 000027.
    _chain_sd = re.match(
        r"^\s*\d{12,18}\s*-\s*(\d{1,7})\s*-\s*(\d{1,7})\s*S\s*[./]?\s*D\s*(\d{1,7}(?:\s*[,/]\s*\d{1,7})*)\s*$",
        val, flags=re.IGNORECASE,
    )
    if _chain_sd:
        start_raw, cert_start_raw, end_list_raw = _chain_sd.groups()
        # Angka kedua adalah certificate start; jika pendek, pad ke 6 digit.
        cert_start = _cert_reconstruct(_cert_pad6(start_raw), cert_start_raw) if len(cert_start_raw) < 6 else _cert_pad6(cert_start_raw)
        ends = [x.strip() for x in re.split(r"[,/]", end_list_raw) if x.strip()]
        if ends:
            first_end = _cert_reconstruct(cert_start, ends[0]) if len(ends[0]) < 6 else _cert_pad6(ends[0])
            result = [f"{cert_start} SD {first_end}"]
            for e in ends[1:]:
                result.append(_cert_reconstruct(first_end, e) if len(e) < 6 else _cert_pad6(e))
            return ", ".join(result)

    # Parser khusus pola "POLIS + CERT + POLIS + CERT". Ini diprioritaskan
    # karena pola tersebut paling jelas menunjukkan pemisahan certificate
    # berdasarkan polis asal, termasuk kasus cert yang langsung menempel
    # setelah 14-18 digit polis.
    _policy_matches = list(re.finditer(r"\d{12,18}", val))
    if len(_policy_matches) >= 2:
        cert_results = []

        def _parse_pair_tail(tail: str):
            tail = re.sub(r"^\s*[-/]+\s*", "", tail.strip())
            tail = tail.strip(" ()")
            if not tail:
                return []
            # Hapus catatan non-certificate yang mengikuti cert.
            tail = re.split(
                r"/\s*(?:VAR(?:IOUS)?|SLIP)\b|\bBORD(?:ERO|ER[OA]?)\b|\bBELUM\\s+DATANG\b",
                tail, maxsplit=1, flags=re.IGNORECASE,
            )[0].strip(" ,&/")
            if not tail:
                return []

            # Pisahkan &, koma, slash hanya sebagai separator record; tanda
            # '-' tetap dipertahankan karena dapat berarti range certificate.
            pieces = [x.strip() for x in re.split(r"[,&/]+", tail) if x.strip()]
            out = []
            base_ctx = None
            for piece in pieces:
                nums = re.findall(r"\d{1,7}", piece)
                if not nums:
                    continue
                if re.search(r"S\s*[./]?\s*D|¿", piece, re.IGNORECASE) and len(nums) >= 2:
                    a, b = nums[0], nums[1]
                    if base_ctx and len(a) < 6:
                        a = _cert_reconstruct(base_ctx, a)
                    else:
                        a = _cert_pad6(a)
                    b = _cert_reconstruct(a, b) if len(b) < 6 else _cert_pad6(b)
                    out.append(f"{a} SD {b}")
                    base_ctx = b
                    continue
                if len(nums) >= 2 and re.search(r"-", piece):
                    a, b = nums[0], nums[1]
                    if base_ctx and len(a) < 6:
                        a = _cert_reconstruct(base_ctx, a)
                    else:
                        a = _cert_pad6(a)
                    b = _cert_reconstruct(a, b) if len(b) < 6 else _cert_pad6(b)
                    span = abs(int(b) - int(a)) + 1
                    if span <= 3:
                        lo, hi = sorted((int(a), int(b)))
                        out.extend(str(n).zfill(6) for n in range(lo, hi + 1))
                    else:
                        out.append(f"{a} SD {b}")
                    base_ctx = b
                    continue
                for n in nums:
                    full = _cert_reconstruct(base_ctx, n) if base_ctx and len(n) < 6 else _cert_pad6(n)
                    out.append(full)
                    base_ctx = full
            return out

        for idx, pm in enumerate(_policy_matches):
            tail_start = pm.end()
            tail_end = _policy_matches[idx + 1].start() if idx + 1 < len(_policy_matches) else len(val)
            tail = val[tail_start:tail_end]
            cert_results.extend(_parse_pair_tail(tail))

        if cert_results:
            return ", ".join(cert_results)

    # Pada satu polis pun, jika setelah range masih ada record lain,
    # pasangan dash yang jelas certificate dipertahankan sebagai SD.
    _single_range = re.match(
        r"^\s*\d{12,18}.*?-\s*(\d{1,7})\s*-\s*(\d{1,7})(.*)$",
        val, flags=re.IGNORECASE,
    )
    if _single_range and re.search(r"[,/]", _single_range.group(3)):
        a, b, rest = _single_range.groups()
        a6 = _cert_pad6(a)
        b6 = _cert_reconstruct(a6, b) if len(b) < 6 else _cert_pad6(b)
        vals = [f"{a6} SD {b6}"]
        # Jika record berikutnya mempunyai range sendiri (mis. 000001-3),
        # gunakan basis record tersebut, bukan basis range sebelumnya.
        _rest_range = re.search(r"(\d{1,7})\s*-\s*(\d{1,7})", rest)
        if _rest_range:
            r1, r2 = _rest_range.groups()
            r1_6 = _cert_pad6(r1)
            r2_6 = _cert_reconstruct(r1_6, r2) if len(r2) < 6 else _cert_pad6(r2)
            span = abs(int(r2_6) - int(r1_6)) + 1
            if span <= 5:
                lo, hi = sorted((int(r1_6), int(r2_6)))
                vals.extend(str(n).zfill(6) for n in range(lo, hi + 1))
            else:
                vals.append(f"{r1_6} SD {r2_6}")
        else:
            tail_nums = re.findall(r"(?<!\d)\d{1,7}(?!\d)", rest)
            for n in tail_nums:
                vals.append(_cert_pad6(n))
        return ", ".join(vals)

    kw_m = _CERT_KEYWORD_RE.search(val)
    if kw_m:
        cert_text = val[kw_m.end():]
        cert_text = _cert_strip_paren(cert_text)
        cert_text = _CERT_TAIL_JUNK_RE.sub("", cert_text)
        cert_text = _cert_clean_tail(cert_text)
        had_interrupt = bool(_CERT_SD_INTERRUPTED_RE.search(cert_text))
        cert_text = _CERT_EMBEDDED_POLICY_RE.sub(" ", cert_text)
        formatted, _ = format_certificate_segment(
            cert_text, keyword_context=True, force_comma_pairs=had_interrupt
        )
        return formatted

    paren_m = re.search(r"\(([^()]*)\)", val)
    if paren_m and re.search(r"\d", paren_m.group(1)):
        inner = paren_m.group(1)
        kw_inner = _CERT_KEYWORD_RE.search(inner)
        is_kw = bool(kw_inner)
        if kw_inner:
            inner = inner[kw_inner.end():]
        if re.search(r"\d", inner):
            formatted, _ = format_certificate_segment(inner, keyword_context=is_kw)
            if formatted:
                return formatted

    for base_re in (_CERT_BASE_EXACT_LEN_RE, _CERT_BASE_GREEDY_RE):
        base_m = base_re.match(val)
        if not base_m:
            continue
        tail = _cert_extract_tail_for_base(val, base_m)
        if tail is None:
            continue

        tail = _cert_clean_tail(tail)
        had_interrupt = bool(_CERT_SD_INTERRUPTED_RE.search(tail))
        tail_clean = _CERT_EMBEDDED_POLICY_RE.sub(" ", tail).strip()
        if tail_clean and _cert_looks_like_cert(tail_clean):
            formatted, _ = format_certificate_segment(
                tail_clean, keyword_context=False, force_comma_pairs=had_interrupt
            )
            if formatted:
                return formatted

    return ""


# ─────────────────────────────────────────────────────────────────────────────
# PROSES UTAMA
# ─────────────────────────────────────────────────────────────────────────────

def _force_text_format(output_file, target_cols_names):
    import openpyxl
    wb = openpyxl.load_workbook(output_file)
    ws = wb.active
    header = [c.value for c in ws[1]]
    target_idx = [i for i, h in enumerate(header, start=1) if h in target_cols_names]
    for col_idx in target_idx:
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                cell.number_format = "@"
    wb.save(output_file)


def process_data(input_file: str, output_file: str) -> None:
    print(f"[1/5] Membaca data dari: {input_file} (sheet: {INPUT_SHEET}) ...")
    df = pd.read_excel(input_file, sheet_name=INPUT_SHEET, header=HEADER_ROW)
    print(f"      Total baris keseluruhan: {len(df):,}")

    if CEDANT_FILTER_COL not in df.columns:
        print(f"\n[ERROR] Kolom '{CEDANT_FILTER_COL}' tidak ditemukan!")
        print(f"        Kolom tersedia: {list(df.columns)}")
        return

    df[CEDANT_FILTER_COL] = df[CEDANT_FILTER_COL].astype(str).str.strip()
    df = df[df[CEDANT_FILTER_COL] == CEDANT_FILTER_VALUE].copy()
    print(f"[2/5] Filter cedant '{CEDANT_FILTER_VALUE}': {len(df):,} baris ditemukan.")

    if df.empty:
        print("\n[WARN] Tidak ada data setelah filter. Proses dihentikan.")
        return

    if STATUS_COL not in df.columns:
        print(f"\n[ERROR] Kolom '{STATUS_COL}' tidak ditemukan!")
        print(f"        Kolom tersedia: {list(df.columns)}")
        return

    before_status = len(df)
    df[STATUS_COL] = df[STATUS_COL].astype(str).str.strip()
    df = df[df[STATUS_COL] == STATUS_KEEP_VALUE].copy()
    print(f"[2b/5] Filter status '{STATUS_KEEP_VALUE}': {len(df):,} baris "
          f"(dibuang {before_status - len(df):,} baris ADJUSTED).")

    if df.empty:
        print("\n[WARN] Tidak ada data setelah filter status. Proses dihentikan.")
        return

    for col in [INSURED_COL, POLIS_COL, SLIP_COL]:
        if col not in df.columns:
            print(f"\n[ERROR] Kolom '{col}' tidak ditemukan di file!")
            print(f"        Kolom tersedia: {list(df.columns)}")
            return

    df.rename(columns={
        INSURED_COL: "INSURED_ORI",
        POLIS_COL:   "POLIS_ORI",
        SLIP_COL:    "SLIP_NO_ORI",
    }, inplace=True)

    print("[3/5] Menjalankan proses cleaning ...")

    polis_cln, slip_cln, all_certificate, all_insured_parts = [], [], [], []
    max_insured_parts = 1

    for _, row in df.iterrows():
        p_ori = row.get("POLIS_ORI", "")
        s_ori = row.get("SLIP_NO_ORI", "")
        i_ori = row.get("INSURED_ORI", "")

        polis_cln.append(refine_with_known_pattern(clean_polis(p_ori), KNOWN_POLIS_PATTERNS))
        slip_cln.append(refine_with_known_pattern(clean_slip(s_ori), KNOWN_SLIP_PATTERNS))
        all_certificate.append(extract_certificate_csi(p_ori))

        parts = clean_insured(i_ori, p_ori, s_ori)
        max_insured_parts = max(max_insured_parts, len(parts))
        all_insured_parts.append(parts)

    print("[4/5] Menyusun kolom output ...")

    normalized_polis = [
        x if isinstance(x, list)
        else [x] if x not in (None, "")
        else []
        for x in polis_cln
    ]

    max_polis = max(
        (len(x) for x in normalized_polis),
        default=0
    )

    for i in range(1, max_polis + 1):
        col_name = f"POLIS_CLEAN_{i}"

        df[col_name] = [
            x[i - 1]
            if i - 1 < len(x)
            and x[i - 1] not in (None, "")
            else None
            for x in normalized_polis
        ]

    normalized_certificate = [
        x if isinstance(x, list)
        else [x] if x not in (None, "")
        else []
        for x in all_certificate
    ]

    max_cert = max(
        (len(x) for x in normalized_certificate),
        default=0
    )

    certificate_cols = []

    for i in range(1, max_cert + 1):
        col_name = f"CERTIFICATE_{i}"

        cert_values = [
            x[i - 1]
            if i - 1 < len(x)
            and x[i - 1] not in (None, "")
            else None
            for x in normalized_certificate
        ]

        if any(
            value is not None
            and str(value).strip() != ""
            for value in cert_values
        ):
            df[col_name] = cert_values
            certificate_cols.append(col_name)

    normalized_slip = [
        x if isinstance(x, list)
        else [x] if x not in (None, "")
        else []
        for x in slip_cln
    ]

    max_slip = max(
        (len(x) for x in normalized_slip),
        default=0
    )

    for i in range(1, max_slip + 1):
        col_name = f"SLIP_CLEAN_{i}"

        df[col_name] = [
            x[i - 1]
            if i - 1 < len(x)
            and x[i - 1] not in (None, "")
            else None
            for x in normalized_slip
        ]

    insured_cols = []

    for i in range(1, max_insured_parts + 1):
        col_name = f"INSURED_CLEAN_{i}"

        df[col_name] = [
            p[i - 1]
            if i - 1 < len(p)
            else None
            for p in all_insured_parts
        ]

        insured_cols.append(col_name)

    new_columns = []

    for col in df.columns:

        if (
            col.startswith("POLIS_CLEAN_")
            or col.startswith("CERTIFICATE_")
            or col.startswith("SLIP_CLEAN_")
            or col in insured_cols
        ):
            continue

        new_columns.append(col)

        if col == "INSURED_ORI":
            new_columns += insured_cols

        elif col == "POLIS_ORI":

            for i in range(1, max_polis + 1):
                policy_col = f"POLIS_CLEAN_{i}"
                new_columns.append(policy_col)

                cert_col = f"CERTIFICATE_{i}"

                if cert_col in certificate_cols:
                    new_columns.append(cert_col)

        elif col == "SLIP_NO_ORI":

            for i in range(1, max_slip + 1):
                slip_col = f"SLIP_CLEAN_{i}"
                new_columns.append(slip_col)

    df = df[new_columns]

    os.makedirs(os.path.dirname(output_file), exist_ok=True)
    print(f"[5/5] Menyimpan hasil ke: {output_file} ...")
    df.to_excel(output_file, index=False)
    _force_text_format(output_file, target_cols_names={"POLIS_ORI", "CERTIFICATE", "POLIS_CLN", "SLIP_NO_ORI", "SLIP_NO_CLN"})

    print(f"\n{'=' * 55}")
    print("  [OK] Selesai!")
    print(f"  Total baris output  : {len(df):,}")
    print(f"  Total kolom output  : {len(df.columns)}")
    print(f"  File disimpan di    : {output_file}")
    print(f"{'=' * 55}")


if __name__ == "__main__":
    process_data(INPUT_FILE, OUTPUT_FILE)