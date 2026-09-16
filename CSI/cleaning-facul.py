import os
import re

import pandas as pd


# ─────────────────────────────────────────────────────────────────────────────
# KONFIGURASI — ubah bagian ini kalau nama file/kolom berbeda
# ─────────────────────────────────────────────────────────────────────────────

INPUT_FILE  = os.path.join("input", "1b. Transaksi Facul 01.01.23 - 17.08.2026.xlsx")
OUTPUT_FILE = os.path.join("output", "csi_output_facul.xlsx")

CEDANT_COL   = "COMP_NAME"
CEDANT_VALUE = "PT. ASURANSI UMUM BCA EX CSI"

BROKER_COL  = "COMP_NAME.1"
POLIS_COL   = "FAC_POLICY_NO"
SLIP_COL    = "FAC_SLIP"
INSURED_COL = "FAC_INSURED"

MAX_SPLIT_COLS = 5   # rule terbaru: cap breakdown di 5, lebih dari itu biarkan apa adanya


# ─────────────────────────────────────────────────────────────────────────────
# REGEX & KAMUS PEMBANTU
# ─────────────────────────────────────────────────────────────────────────────

POLIS_EXCEPTION_RE = re.compile(
    r"""
    MOP\s*MARINE
  | (?:LINE\s*SLIP|LINESLIP)
  | \b(?:P1|P2|P3|P73)\s*CANCEL
  | \b(?:P1|P2|P3|P73)\b
  | \bCANCEL\b
  | PENYELESAIAN
  | (?:HUTANG|UTANG)\s*PIUTANG
    """,
    re.IGNORECASE | re.VERBOSE,
)

SLIP_EXCEPTION_RE = re.compile(
    r"""
    \bSUMMARY\b
  | \bBORDER[OA]\b
  | \bBORDRO\b
  | \bSINGGLESHIPMENT\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

INSURED_JUNK_WORDS = frozenset({
    "PT", "CV", "TBK", "PERSERO", "LTD", "INC", "LLC",
    "AND", "OR", "THE", "OF", "AS",
    "NON FOOD", "DIV", "VARIOUS", "ATTACHMENT",
})

INSURED_SUFFIX_RE = re.compile(
    r"""
    ,?\s*\bTBK\b(?:\s*,?\s*PT\.?)?
  | ,?\s*\(\s*PERSERO\s*\)\s*
  | ,?\s*\bPERSERO\b\s*
  | ,?\s*\bLTD\.?\b\s*
  | ,?\s*\bPT\.?\s*$
  | ,?\s*\bCV\.?\s*$
    """,
    re.IGNORECASE | re.VERBOSE,
)

_SLIP_NOISE_RE = re.compile(
    r"""
    \b(?:JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS
        |SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\b
  | \b(?:JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST
        |SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)\b
  | \b(?:IDR|USD|ENG)\b
  | \b(?:P1|P2|P3|P73)\b
  | \bNEW\b
  | \bVARIOUS\b
  | \b20[0-9]{2}\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

_SLIP_TOKEN_RE = re.compile(r"[A-Z0-9][A-Z0-9\-]{6,}", re.IGNORECASE)


# ─────────────────────────────────────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def _normalize_spaces(text: str) -> str:
    return re.sub(r"\s{2,}", " ", text).strip()


def _is_valid_polis_token(tok: str) -> bool:
    tok = tok.strip()
    return len(tok) >= 5 and not re.search(r"TBA$", tok, re.IGNORECASE) and bool(re.search(r"\d", tok))


def _extract_polis_tokens(text: str) -> list:
    tokens = []
    for block in re.split(r"\s{2,}", text.strip()):
        for tok in re.split(r"\+|\s+", block.strip()):
            tok = tok.strip().strip("-/")
            if _is_valid_polis_token(tok):
                tokens.append(tok)
    return tokens


def _expand_dash_chain(segment: str) -> list:
    results = []
    current_base = None
    for part in segment.split("-"):
        part = part.strip()
        if not part:
            continue
        if len(part) >= 8 and re.search(r"\d", part):
            current_base = part
            results.append(part)
        elif current_base and len(part) >= 2 and re.match(r"^\d+$", part) and len(part) < len(current_base):
            n = len(part)
            results.append(current_base[:-n] + part)
        else:
            if _is_valid_polis_token(part):
                results.append(part)
    return results or ([segment] if _is_valid_polis_token(segment) else [])


def _is_valid_slip_token(tok: str) -> bool:
    tok = tok.strip()
    return len(tok) >= 7 and bool(re.search(r"\d", tok)) and not re.match(r"^\d{4}$", tok)


def _strip_slip_suffix(tok: str) -> str:
    m = re.match(r"^(.+?)-(\d{1,6})$", tok)
    if m and len(m.group(1)) > len(m.group(2)):
        return m.group(1)
    return tok


def _extract_slip_tokens(text: str) -> list:
    results = []
    for block in re.split(r"\s{2,}", text.strip()):
        clean = _SLIP_NOISE_RE.sub(" ", block)
        clean = re.sub(r"^[\s\-/+,]+|[\s\-/+,]+$", "", clean).strip()
        for cand in _SLIP_TOKEN_RE.findall(clean):
            cand = cand.strip("-")
            parts = cand.split("-")
            if len(parts) == 2:
                a, b = parts
                if _is_valid_slip_token(a) and _is_valid_slip_token(b) and abs(len(a) - len(b)) <= 2:
                    results.extend([a, b])
                    continue
            if _is_valid_slip_token(cand):
                results.append(_strip_slip_suffix(cand))
    return results


_INSURED_PAREN_JUNK = {"PERSERO", "TBK", "SMALL SHIP FLEET"}


def _handle_insured_paren(match: "re.Match") -> str:
    content_ = match.group(1).strip()
    if content_.upper() in _INSURED_PAREN_JUNK or not content_:
        return ""
    kept = content_.split(",")[0].strip()
    return " " + kept + " " if kept else ""


def _merge_digit_start_parts(parts: list) -> list:
    merged = []
    for p in parts:
        p_stripped = p.strip()
        if merged and re.match(r"^\d", p_stripped):
            merged[-1] = merged[-1] + " " + p_stripped
        else:
            merged.append(p_stripped)
    return merged


_BRANCH_KEYWORD_RE = re.compile(r"^(?:(?:TBK|PT\.?|CV\.?)\s+)*(?:KCU|KANWIL|DTC|CABANG)\b", re.IGNORECASE)


def _merge_branch_keyword_parts(parts: list) -> list:
    merged = []
    for p in parts:
        p_stripped = p.strip()
        if merged and _BRANCH_KEYWORD_RE.match(p_stripped):
            merged[-1] = merged[-1] + " " + p_stripped
        else:
            merged.append(p_stripped)
    return merged


def _name_initials_variants(name: str) -> set:
    words = [w for w in re.split(r"[\s\-]+", name.upper()) if w]
    if not words:
        return set()
    variants = {"".join(w[0] for w in words if w[0].isalpha())}
    if len(words) >= 2:
        variants.add("".join(w[0] for w in words[:-1] if w[0].isalpha()) + words[-1])
    return variants


def _is_abbreviation_of(short: str, long_name: str) -> bool:
    short_clean = re.sub(r"[^A-Z0-9]", "", short.upper())
    if len(short_clean) < 2 or len(short_clean) > 8:
        return False
    return short_clean in _name_initials_variants(long_name)


def _strip_trailing_abbrev_word(name: str, priors: list) -> str:
    words = name.split()
    if len(words) >= 2 and any(_is_abbreviation_of(words[-1], p) for p in priors):
        return " ".join(words[:-1])
    return name


def _clean_insured_name(name: str) -> str:
    name = _normalize_spaces(name)
    name = _INSURED_PREFIX_RE.sub("", name)
    name = _INSURED_GELAR_RE.sub("", name)
    name = re.sub(r"\bPTE\.?\b", "", name, flags=re.IGNORECASE)
    name = _INSURED_BOILERPLATE_RE.sub("", name)
    name = re.sub(r"\(([^()]*)\)", _handle_insured_paren, name)
    name = name.replace("(", " ").replace(")", " ")
    name = re.sub(r"/", " ", name)
    for _ in range(3):
        cleaned = _normalize_spaces(INSURED_SUFFIX_RE.sub("", name).strip().strip(","))
        if cleaned == name:
            break
        name = cleaned
    name = name.replace(".", "")
    name = _normalize_spaces(name).strip()
    name = name.upper()
    return name


def _cap_or_join(items: list) -> list:
    if len(items) > MAX_SPLIT_COLS:
        return [",".join(items)]
    return items


# ─────────────────────────────────────────────────────────────────────────────
# KAMUS POLA STANDAR
# ─────────────────────────────────────────────────────────────────────────────

KNOWN_POLIS_PATTERNS = [
    re.compile(r"^\d{15}$"),
    re.compile(r"^\d{14}$"),
    re.compile(r"^\d{18}$"),
]

KNOWN_SLIP_PATTERNS = [
    re.compile(r"^\d{15}$"),
    re.compile(r"^\d{14}$"),
    re.compile(r"^\d{17}$"),
    re.compile(r"^TIDAK ADA$", re.IGNORECASE),
]


def _matches_known_pattern(value: str, patterns) -> bool:
    return any(p.match(value) for p in patterns)


def refine_with_known_pattern(value: str, patterns) -> str:
    if not value:
        return value
    if _matches_known_pattern(value, patterns):
        return value

    candidates = [
        re.sub(r"\s+", "", value),
        re.sub(r"\s*/\s*", " / ", value).strip(),
        value.strip(" .-/"),
        re.sub(r"\.0+$", "", value),
        "0" + value if value.isdigit() else value,
        re.sub(r"^(\d+)(\s*/\s*(?:19|20)\d{2})$", r"0\1\2", value),
    ]
    for c in candidates:
        if _matches_known_pattern(c, patterns):
            return c

    return value


def _refine_list(items: list, patterns) -> list:
    return [refine_with_known_pattern(v, patterns) for v in items]


# ─────────────────────────────────────────────────────────────────────────────
# MESIN "REPETITION CHAIN"
# ─────────────────────────────────────────────────────────────────────────────

_MONTH_NAMES_RE = (
    r"JANUARI|FEBRUARI|SEPTEMBER|NOVEMBER|DESEMBER|OKTOBER|AGUSTUS|MARET"
    r"|JANUARY|FEBRUARY|SEPTEMBER|NOVEMBER|DECEMBER|OCTOBER|AUGUST|MARCH"
    r"|APRIL|JUNI|JULI|JUNE|JULY|MEI|MAY|MARHC"
    r"|SEPT|AUGUS|AGUS|OKT|DES|JAN|FEB|MAR|APR|JUN|JUL|AUG|SEP|OCT|NOV|DEC"
)

_JUNK_PHRASE_RE = re.compile(
    r"""
    SLIP\s+FOR\s+NEXT\s+\d+\s+MONTHS
  | TANPA\s+INDUK\s+SLIP
  | TANPA\s+SLIP\s+INDUK
  | SLIP\s+MASIH\s+PAKAI\s+SLIP\s+SEMENTARA
  | PAKAI\s+SLIP\s+SEMENTARA
  | SUPEN\s+SLIP
  | SUSPEN\s+SLIP
  | PAYMENT\s+NUMBER\s*:?
  | BINDER\s+NO\.?
  | REN\.?\s*POL\.?\s*NO\.?
  | REPLACEMENT
  | FINAL\s+ADJUSTMENT
  | ACTUAL\s+PREMIUM\s+FOR
  | MINDEPPREM\.?\s*FOR
  | MINDEP\s+PREMI(?:UM)?
  | MINDEP\s+FOR
  | \bMINDEP\b
  | \bMIN\.?
  | \bPERIOD\b
  | \bFOR\b
  | DATE\s+\d{1,2}\.\d{1,2}\.\d{2,4}
  | (?:19|20)\d{2}\s*-\s*\d{2}\b
  | \.?\s*\+?\s*SLIP\s+END\s*\d*(?:\s*\+\s*\d+\.?)*
  | \.?\s*\+?\s*END\.?\s*\d*(?:\s*\+\s*\d+\.?)*
  | DE(?:KLARASI|CLARASI)\s*:?
  | DECL\.?\s*:?
  | INTER\s+ISLAND
  | \bCOL\.?\s*NUMBER\b
  | \bCREDIT\s+NOTE\b
  | \bNPWN\b
  | BIND\.?\s*NO\.?
  | \bPN\.?\s+(?=\S)
  | \bR\.?\s*PREMIUM\b
  | \bMENJADI\b
  | \bS\s*[./]?\s*D\b
  | \bADDITIONAL\b
  | \bAS\s+PER\s+LIST\b
  | AS\s+PER\s+LIST(?:\s+ATTACHED)?
  | \bATTACHED\b
  | \bATTACHMENT\b
  | PROD\.JULY\s+2018
  | \bPOI\b
  | \bCANCEL\b
  | \bVARIOUS\b
  | \bVARI\b
  | \bVARIUOS\b
  | \bMACH\b
  | \bLOCAL\s+SHIPMENT\b
  | \bCERTIFICATE\b
  | \bCERTIF\b
  | \bCERT\b
  | \bREIN\b
  | (?:ADDENDUM\s+)?MOP\s+NO\.?
  | \bADD\b
  | \bBORDERO\b
  | \bBELUM\s+DATANG\b
  | \bMB\b
  | \b(?:IDR|USD|EUR|GBP|CNY|SGD|JPY|AUD|HKD|MYR|CHF|THB)\b
  | \bSLIP\s+BLM\s+DATANG\s+SEMUA\b
  | \(\s*SLIP\s+1\s+LAGI\s+BLM\s+ADA\s*\)
  | \bDEPOSIT\b.*?\bFOR\b
  | AND\s+AC\.?\s*PREMI\.?\s*FOR
  | \bFROM\b
  | \bLIHAT\s+PAGE\s*\d*\b
  | \bEX\s*:?
  | \bLOL\b
  | \bPA\b
  | TANPA\s+ADA\s+SLIP\s+INDUK
  | \bIDR\b
  | \bUSD\b
  | \bCN\b\s*:?
  | \bSLIP\b
  | \bID\s*:
  | \bAND\b
  | \bPR\.
  | \bACTUAL\s+PREMIUM\b
  | \bEQ\b
  | \bNO\b\s*:?
  | \bMONTHS\b
  | \bPOLIS\b
  | \bVARIOUS\b
    """
    + r"| (?:" + _MONTH_NAMES_RE + r")\.?\s*\d{0,4}",
    re.IGNORECASE | re.VERBOSE,
)

_STANDALONE_NARRATIVE_RE = re.compile(r"SUSPEN", re.IGNORECASE)

_STANDALONE_P_RE = re.compile(
    r"^\s*(?:P[12345]|TBA|VAR)(?:\s*[+,]\s*(?:P[12345]|TBA|VAR))*\s*$", re.IGNORECASE
)

_YEAR_RE = re.compile(r"/\s*(?:19|20)\d{2}\b")
_YEAR_2DIGIT_RE = re.compile(r"/\s*\d{2}$")
_DOT_YEAR_RE = re.compile(r"^(.*)\.(?:19|20)?\d{2,4}$")

# BUG FIX #1: glue vs replace sekarang dinamis (lihat _classify_and_reconstruct)
_MIN_COMPLETE_BASE_LEN = 14
_MAX_SUFFIX_LEN = 7

_CURRENCY_CODES = {"USD", "IDR", "EUR", "GBP", "CNY", "SGD", "JPY",
                    "AUD", "HKD", "MYR", "CHF", "THB"}

_LITERAL_WHITELIST = {
    "PENYELESAIAN SUSPENSE 2020 CSI / BCA",
    "*FACCODE EKOSONG*",
    "P1 APRIL PERIOD",
    "VARIOUS / LIHAT ELO",
    "P1 REVISED",
    "VARIUOS / LIHAT ELO",
    "JANUARY2023/CNY",
    "DEKLARASI AGUSTUS 2023 - USD",
    "DTC 02 - 05 2025 + P2",
    "PROD.MAY 2017",
    "FABRUARY 2020/USD",
    "AGUSTUS 2021/EUR",
    "AGUSTUS 2021/IDR/CN 002583/CN/0101/08/21/VARIOUS",
    "DEKLARASI JUNI 2023 - EUR",
    "FEB.2023/EUR",
    "DEKLARASI AGUSTUS 2023 - GBP",
    "BORD NOVEMBER 2025",
    "BORD APRIL 2026",
    "265.266",
    "1562.266",
    "702703-071072",
    "#SPN-2014#",
    "102.01.08.0011.F/REIN",
    "134-081-093-092",
    "14-04-000852 S/D 14-04-000862/END.1,2-5,6-9",
    "3042.227.3045.221.3048.212",
    "010-008-012-016",
    "BORD OCTOBER 2025",
    "BORD NOVEMBER 2025",
    "BORD AGUSTUS 2025",
    "BORD APRIL 2026",
    "BORD DESEMBER 2025",
    "BORD SEPTEMBER 2025",    "401.01.11.000001.F/REIN",
    "401.01.09.000001.F/REIN",
    "401.01.10.001.F/REIN",
    "401.01.10.000002.F/REIN",
    "011-006",
    "210&116",
    "614&615-627&629",
    "138-139-055-056",
    "CERTIFICATE : 000014 S/D 000019",
    "*FACCODE KOSONG*",
    "4259..3040.3766.1415.0855.1010",
    "3042.3045.3048.227.221.212",
    "JUNE 2017/CERT.00006-000009&000018-000022,000027",
    "BORDERO BCA INS DJARUM",
    "BORD JANUARY 2026",
    "BORD FEBRUARY 2026",
    "BORD MARET 2026",    "00014&02739-00002&00041",
    "000002&00041",
    "006-009-012-015-019",
    "006-009-012-015-019-007-044",
    "008 + 002 + 004 + 006",
    "001 + 011 + 004",
    "008 + 011 + 005",
    "004.012.007.032.040.043",
    "008 + 011 + 002 + 002",
    "008+007+009+008+010+017+011+09+012+010+013+011+014+012+015+013+016+014+017+015+018+016",
    "008+007+009+008+010+017+011+009+012+010+013+011+014+012+015+013+016+014+017+015+018+016",
    "006+957+392+071+074+106+109+403+126",
    "005 + 053 + 009 + 057 + 011 + 060 + 013 + 062",
    "1508&0008",
    "010101091800390/2/1 / 010101221800250/49/48",
    "001 + 00138 + 002 + 139 + 003 + 140 + 141",
}


def _is_standalone_exception(val: str) -> bool:
    if val.strip() in _LITERAL_WHITELIST:
        return True
    if _STANDALONE_CERTIFICATE_RE.match(val.strip()):
        return True
    return bool(_STANDALONE_P_RE.match(val)) or bool(_STANDALONE_NARRATIVE_RE.search(val))


def _is_non_cert_standalone(val: str) -> bool:
    """BUG FIX #7b: value yang sudah dikenali sbg literal whitelist / gabungan
    P1-P5/TBA/VAR / narasi SUSPEN (bukan pola sertifikat CERTIFICATE...) tidak
    boleh menghasilkan CERTIFICATE. Beda dari _is_standalone_exception():
    sengaja TIDAK termasuk _STANDALONE_CERTIFICATE_RE, supaya value seperti
    "CERTIFICATE NO 000023 S/D 000026/31" tetap bisa jatuh ke fallback
    extract_certificate_csi()."""
    v = val.strip()
    if v in _LITERAL_WHITELIST:
        return True
    if _STANDALONE_P_RE.match(v):
        return True
    if _STANDALONE_NARRATIVE_RE.search(v):
        return True
    return False


_CERT_SD_RANGE_RE = re.compile(
    r"-?\s*(?<!\d)\d{1,7}\s*(?:S\s*[./]?\s*D|¿)\s*\d{1,7}\b(?:\s*[,&]\s*\d{1,7}\b)*",
    re.IGNORECASE,
)

_STANDALONE_CERTIFICATE_RE = re.compile(
    r"^CERT(?:IF(?:ICATE)?)?\b", re.IGNORECASE
)


_SLIP_AMOUNT_NOTE_RE = re.compile(
    r"/?\s*SLIP\s+[\d.,]+\s*BELUM\s+DATANG",
    re.IGNORECASE,
)


def _strip_junk_and_year(text: str) -> str:
    cert_suffix_match = re.match(
        r"^\s*(\d{7,})\s+CERT(?:IF(?:ICATE)?)?\s*[:.]?\s*[\d][\d\s,./&-]*\s*$",
        text,
        flags=re.IGNORECASE,
    )
    if cert_suffix_match:
        return cert_suffix_match.group(1)

    # BUG FIX #4: "/ SLIP <nominal> BELUM DATANG" (mis. "468,800.86") adalah
    # catatan nominal slip yang belum diterima -- seluruh angka nominalnya
    # (bukan cuma kata SLIP/BELUM DATANG) harus ikut dibuang, supaya
    # pecahan angkanya (468/800/86) tidak disalahartikan sbg suffix polis.
    text = _SLIP_AMOUNT_NOTE_RE.sub("", text)
    text = _CERT_SD_RANGE_RE.sub("", text)
    text = _JUNK_PHRASE_RE.sub("", text)
    text = _YEAR_RE.sub("", text)
    return text


def _strip_attached_noise(text: str) -> str:
    text = re.sub(r"\bP[12345]\b\s*/?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\bTBA\b\s*/?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"(?<=\d)TBA\b", "", text, flags=re.IGNORECASE)
    text = re.sub(r"(?<=\d)VAR\b", "", text, flags=re.IGNORECASE)
    text = re.sub(r"\bVAR\b\s*/?\s*", "", text, flags=re.IGNORECASE)
    text = re.sub(r"^P-(?=\d)", "", text)
    text = text.replace("\x27", "").replace("'", "")
    return text


def _clean_token(tok: str) -> str:
    tok = tok.strip(" :;/")
    if not tok:
        return tok
    if not re.search(r"[A-Za-z]", tok):
        tok = tok.replace(".", "")
    elif re.match(r"^[A-Za-z]+[\d.]+$", tok):
        tok = tok.replace(".", "")
    return tok


def _classify_and_reconstruct(tokens: list) -> list:
    results = []
    current_base = None
    for tok in tokens:
        tok = _clean_token(tok)
        if not tok:
            continue
        digits_only = re.match(r"^\d+$", tok) is not None
        base_has_letter = bool(current_base and re.search(r"[A-Za-z]", current_base))

        if digits_only and current_base is not None and len(tok) < len(current_base):
            if len(current_base) < _MIN_COMPLETE_BASE_LEN and not base_has_letter:
                new_base = current_base + tok
                if results and results[-1] == current_base:
                    results[-1] = new_base
                else:
                    results.append(new_base)
                current_base = new_base
                continue

            is_short_suffix = (
                len(tok) <= _MAX_SUFFIX_LEN
                and (base_has_letter or len(current_base) >= 7)
            )
            if is_short_suffix and len(tok) <= 4:
                results.append(current_base[: -len(tok)] + tok)
                current_base = results[-1]
                continue
            elif is_short_suffix:
                continue

        results.append(tok)
        current_base = tok
    return results


def _dedup_preserve_order(items: list) -> list:
    seen = set()
    deduped = []
    for r in items:
        if r not in seen:
            seen.add(r)
            deduped.append(r)
    return deduped


_POLIS_SD_FULL_RANGE_RE = re.compile(
    r"^\s*(\d{14,18})\s*S\s*[./]?\s*D\s*(\d{14,18})\s*$",
    re.IGNORECASE,
)


def _try_polis_sd_and_plain_segments(val):
    """BUG FIX #10: tangani value yang berisi gabungan segmen dipisah koma,
    di mana tiap segmen adalah salah satu dari (a) pola penuh
    "POLIS1 S/D POLIS2" (dua nomor polis LENGKAP panjang sama), atau
    (b) satu nomor polis polos (14-18 digit). Contoh:
    "012802032200001 S/D 012802032200015,01280203220003" -> polis 1 =
    "012802032200001 SD 012802032200015" (karena rentangnya 15 anggota,
    > MAX_SPLIT_COLS jadi TIDAK di-breakdown penuh, cukup ditulis sbg
    literal "AWAL SD AKHIR"), polis 2 = "01280203220003" apa adanya.
    Kalau rentang S/D-nya kecil (<=MAX_SPLIT_COLS), tetap di-breakdown
    penuh berurutan spt biasa. Return None kalau ada segmen yang tidak
    cocok pola manapun (supaya caller fallback ke proses standar), atau
    kalau tidak ada satupun segmen S/D (bukan kasus khusus ini)."""
    segments = [s.strip() for s in val.split(",")]
    if not segments:
        return None
    results = []
    matched_any_sd = False
    for seg in segments:
        if not seg:
            return None
        m = _POLIS_SD_FULL_RANGE_RE.match(seg)
        if m:
            start_s, end_s = m.group(1), m.group(2)
            if len(start_s) != len(end_s):
                return None
            start_n, end_n = int(start_s), int(end_s)
            count = end_n - start_n + 1
            if count < 1:
                return None
            matched_any_sd = True
            if count == 1:
                results.append(start_s)
            elif count <= MAX_SPLIT_COLS:
                width = len(start_s)
                results.extend(str(n).zfill(width) for n in range(start_n, end_n + 1))
            else:
                results.append(f"{start_s} SD {end_s}")
            continue
        if re.fullmatch(r"\d{14,18}", seg):
            results.append(seg)
            continue
        return None
    if not matched_any_sd:
        return None
    return results


def expand_repetition_chain(val, patterns=None) -> list:
    if pd.isna(val):
        return []
    val = str(val).strip()
    if not val:
        return []

    if _is_standalone_exception(val):
        return [_normalize_spaces(val)]

    # BUG FIX #5 / #10: "POLIS1 S/D POLIS2" (dan gabungannya dgn segmen
    # polis polos lain via koma) -- lihat _try_polis_sd_and_plain_segments().
    _sd_segments_result = _try_polis_sd_and_plain_segments(val)
    if _sd_segments_result is not None:
        return _sd_segments_result

    _revisi_prefix_match = re.search(
        r"/\s*(\d{7,})\s*-\s*(\d+(?:\s*,\s*\d+)*)\s*$",
        val,
        flags=re.IGNORECASE,
    )

    _cert_dash_match = re.search(
        r"(\d{7,})\s*-\s*(\d{6})\s*$",
        val,
        flags=re.IGNORECASE,
    )

    if _cert_dash_match:
        return [_cert_dash_match.group(1)]

    if _revisi_prefix_match:
        base = _revisi_prefix_match.group(1)
        suffix_list = _revisi_prefix_match.group(2)
        suffix_list = re.sub(r"\s*,\s*", ",", suffix_list)
        first_suffix, *remaining_suffixes = suffix_list.split(",")
        first_number = base + first_suffix
        result_parts = [first_number] + remaining_suffixes
        return [",".join(result_parts)]

    _sd_chain_match = re.match(
        r"^\s*(\d{7,})\s*-\s*"
        r"(\d+(?:\s*,\s*\d+)*)"
        r"\s*S\s*/?\s*D\s*"
        r"(\d+)\s*$",
        val,
        flags=re.IGNORECASE,
    )

    if _sd_chain_match:
        base = _sd_chain_match.group(1)
        suffix_list = _sd_chain_match.group(2)
        sd_end = _sd_chain_match.group(3)
        suffix_list = re.sub(r"\s*,\s*", ",", suffix_list)
        first_suffix, *remaining_suffixes = suffix_list.split(",")
        first_number = base + first_suffix
        result_parts = [first_number] + remaining_suffixes + [sd_end]
        return [",".join(result_parts)]

    _dot_suffix_match = re.match(
        r"^\s*(\d{7,})\s*-\s*((?:\.\d+)+)\s*$",
        val,
        flags=re.IGNORECASE,
    )

    if _dot_suffix_match:
        base = _dot_suffix_match.group(1)
        suffix_text = _dot_suffix_match.group(2)
        suffixes = re.findall(r"\.(\d+)", suffix_text)
        results = [base]
        for suffix in suffixes:
            if len(suffix) <= len(base):
                reconstructed = base[:-len(suffix)] + suffix
                results.append(reconstructed)
        return list(dict.fromkeys(results))

    _tba_suffix_m = re.match(r"^(.+?)\s*\+?\s*TBA\s*$", val, re.IGNORECASE)
    if _tba_suffix_m:
        val = _normalize_spaces(_tba_suffix_m.group(1))
        if not val:
            return []

    if re.match(r"^\d{1,4}\.\d{1,4}$", val):
        return [_normalize_spaces(val)]

    _had_narrative = bool(_JUNK_PHRASE_RE.search(val))

    cleaned_text = _strip_junk_and_year(val)
    cleaned_text = _strip_attached_noise(cleaned_text)
    if cleaned_text.count("/") <= 1:
        cleaned_text = _YEAR_2DIGIT_RE.sub("", cleaned_text)
    if cleaned_text.count(".") == 1:
        m = _DOT_YEAR_RE.match(cleaned_text)
        if m:
            cleaned_text = m.group(1)

    if "." in cleaned_text and not re.search(r"[A-Za-z]", cleaned_text):
        _dot_parts = cleaned_text.split(".")
        if len(_dot_parts) > 1 and all(p.isdigit() for p in _dot_parts):
            _rest_lens = {len(p) for p in _dot_parts[1:]}
            if len(_rest_lens) == 1 and next(iter(_rest_lens)) < len(_dot_parts[0]):
                cleaned_text = "+".join(_dot_parts)
            elif len(_rest_lens) == 1 and next(iter(_rest_lens)) == len(_dot_parts[0]) and len(_dot_parts[0]) <= 4:
                return [_normalize_spaces(val)]
    cleaned_text = re.sub(r"[/\-+,]\s*[A-Za-z]\s*$", "", cleaned_text)
    cleaned_text = cleaned_text.replace(":", "").replace("(", "").replace(")", "").replace("¿", " ")
    cleaned_text = cleaned_text.strip()
    cleaned_text = cleaned_text.strip(" /-,&")

    if not cleaned_text:
        return [_normalize_spaces(val)]

    cleaned_text = re.sub(r"\b([A-Za-z]{2,5})\s+(\d{5,})\b", r"\1\2", cleaned_text)
    cleaned_text = cleaned_text.replace("*", "")

    raw_tokens = [t for t in re.split(r"\s+|[+\-,/&;]", cleaned_text) if t.strip()]
    raw_tokens = [_clean_token(t) for t in raw_tokens]
    raw_tokens = [t for t in raw_tokens if t]
    if not raw_tokens:
        return [_normalize_spaces(val)]

    # BUG FIX #3: sisa kata pendek murni huruf (mis. "SA" dari "/ SLIP SA"
    # setelah kata "SLIP" dibuang oleh _JUNK_PHRASE_RE) bukan bagian dari
    # nomor polis -- buang sebagai noise, bukan dianggap token literal.
    raw_tokens = [t for t in raw_tokens if not re.fullmatch(r"[A-Za-z]{1,3}", t)]
    if not raw_tokens:
        return [_normalize_spaces(val)]

    _all_digit_short = (
        len(raw_tokens) > 1
        and all(re.match(r"^\d+$", t) for t in raw_tokens)
        and len(raw_tokens[0]) < 7
    )
    _has_real_list_sep = (
        ("&" in cleaned_text or "+" in cleaned_text)
        and len(raw_tokens[0]) >= 5
    )
    _equal_length = len({len(t) for t in raw_tokens}) == 1
    if _all_digit_short and not _has_real_list_sep:
        if _equal_length or len(raw_tokens) == 2:
            return [_normalize_spaces(val)]
        return ["".join(raw_tokens)]

    reconstructed = _classify_and_reconstruct(raw_tokens)

    if patterns:
        reconstructed = [refine_with_known_pattern(r, patterns) for r in reconstructed]

    reconstructed = _dedup_preserve_order(reconstructed)

    if len(reconstructed) == 1:
        if reconstructed[0].upper() in _CURRENCY_CODES and _had_narrative:
            return [_normalize_spaces(val)]
        return [_normalize_spaces(reconstructed[0])]

    if len(reconstructed) <= MAX_SPLIT_COLS:
        return [_normalize_spaces(r) for r in reconstructed]

    # BUG FIX #12: kalau breakdown-nya terlalu banyak (>5) dan dibiarkan
    # sbg satu literal, DAN bentuknya "base-suffix1,suffix2,..." (dash lalu
    # KOMA -- ciri list suffix, bukan rantai dash-ke-dash), buang dash
    # penghubung base->suffix pertama (mis.
    # "01010311180005-023,024,025,026,028,031,032,033,034" ->
    # "01010311180005023,024,025,026,028,031,032,033,034"). Rantai murni
    # dash-ke-dash (mis. "469-470-467-473-468-014", tanpa koma) TIDAK
    # disentuh -- itu sudah benar dibiarkan literal apa adanya.
    _bail_text = re.sub(r"^(\d{14,18})-(?=\d+,)", r"\1", cleaned_text, count=1)
    return [_normalize_spaces(_bail_text)]


# ─────────────────────────────────────────────────────────────────────────────
# CLEAN POLIS
# ─────────────────────────────────────────────────────────────────────────────


def _remove_known_certificate_parts(val: str) -> str:
    if not val:
        return val

    text = str(val)
    records = _split_policy_records_for_certificate(text)

    if not records:
        return text

    is_multi_record = len(records) >= 2
    record_certs_for_removal = _certs_for_records(records)

    pieces = []
    for record, cert in zip(records, record_certs_for_removal):
        if cert:
            embedded = re.match(
                r"^(\d{15})(0\d{5})(.*)$",
                record,
            )

            if embedded:
                kept = embedded.group(1)
                pieces.append(kept.strip())
            else:
                base_end = re.match(r"\d{14,18}", record)
                if base_end:
                    kept = record[:base_end.end()]
                    rep = re.match(
                        r"\s*/\s*\d{1,2}\b",
                        record[base_end.end():],
                    )
                    if rep:
                        kept += rep.group(0)
                    pieces.append(kept.strip())
                else:
                    pieces.append(record)
        else:
            if re.match(r"^\d{14,18}\s*[-–—]?\s*0{3}\s*$", record):
                base_only = re.match(r"^(\d{14,18})", record)
                pieces.append(base_only.group(1) if base_only else record)
            else:
                pieces.append(record)

    return ",".join(pieces)


def clean_polis(val) -> list:
    if pd.isna(val):
        return []
    val = str(val).strip()
    if not val:
        return []

    # BUG FIX #9: literal whitelist (mis. rantai perulangan >5 anggota yang
    # sengaja dibiarkan apa adanya) juga harus dicek SEBELUM masuk ke
    # _remove_known_certificate_parts(), karena fungsi itu bisa memecah &
    # menyisipkan koma/plus ke dalam value sebelum sempat dicek ulang oleh
    # _is_standalone_exception() di dalam expand_repetition_chain.
    if val.strip() in _LITERAL_WHITELIST:
        return [_normalize_spaces(val)]

    # BUG FIX #5 / #10: "POLIS1 S/D POLIS2" (dan gabungannya dgn segmen
    # polis polos lain via koma) -- cek lebih dulu SEBELUM masuk ke
    # _remove_known_certificate_parts(), karena fungsi itu memecah value
    # berdasarkan kemunculan nomor polis dan bisa menyisipkan koma/plus di
    # antara "S/D" dan polis kedua (merusak pola ini sebelum sempat dicek
    # lagi di expand_repetition_chain). Lihat _try_polis_sd_and_plain_segments().
    _sd_segments_result = _try_polis_sd_and_plain_segments(val)
    if _sd_segments_result is not None:
        return _sd_segments_result

    _trigger = (
        "+" in val
        or "/" in val
        or "-" in val
        or "&" in val
        or ";" in val
        or "," in val
        or "." in val
        or " " in val
        or _YEAR_RE.search(val)
        or re.search(r"\bP[12345]\b", val, re.IGNORECASE)
        or re.search(r"\bTBA\b", val, re.IGNORECASE)
        or re.search(r"\bVAR\b", val, re.IGNORECASE)
        or _JUNK_PHRASE_RE.search(val)
    )
    if _trigger:
        policy_for_clean = _remove_known_certificate_parts(val)

        policy_for_clean = re.sub(
            r"(\d{14,18})\s*/\s*(\d{1,2})\b",
            r"\1+\2",
            policy_for_clean,
        )

        return expand_repetition_chain(policy_for_clean, KNOWN_POLIS_PATTERNS)

    if POLIS_EXCEPTION_RE.search(val):
        return [_normalize_spaces(val)]

    val_upper = val.upper().strip()
    if re.match(r"^VARIOUS\s*$", val_upper):
        return [val.strip()]
    if re.match(r"^VARIOUS\s*-\s*SEE\s+ATTACH", val_upper):
        return [val.strip()]
    if re.match(r"^TBA\s*$", val_upper):
        return [val.strip()]

    return [_normalize_spaces(val)] if val.strip() else []


# ─────────────────────────────────────────────────────────────────────────────
# CLEAN SLIP
# ─────────────────────────────────────────────────────────────────────────────

def clean_slip(val) -> list:
    if pd.isna(val):
        return []
    val = str(val).strip()
    if not val:
        return []

    _trigger = (
        "+" in val
        or "/" in val
        or "-" in val
        or "&" in val
        or ";" in val
        or "," in val
        or "." in val
        or " " in val
        or _YEAR_RE.search(val)
        or re.search(r"\bP[12345]\b", val, re.IGNORECASE)
        or re.search(r"\bTBA\b", val, re.IGNORECASE)
        or re.search(r"\bVAR\b", val, re.IGNORECASE)
        or _JUNK_PHRASE_RE.search(val)
    )
    if _trigger:
        return expand_repetition_chain(val, KNOWN_SLIP_PATTERNS)

    blocks = re.split(r"\s{2,}", val.strip())
    if len(blocks) >= 2:
        left = blocks[0]
        left_check = re.sub(
            r"\b(?:SUMMARY|BORDERO?|BORDERA|BORDRO|SINGGLESHIPMENT|VARIOUS)\b",
            "", left, flags=re.IGNORECASE,
        )
        left_check = _SLIP_NOISE_RE.sub("", left_check)
        left_check = re.sub(r"[\s\+\-/]+", "", left_check).strip()
        if not _is_valid_slip_token(left_check):
            tokens = _extract_slip_tokens("  ".join(blocks[1:]))
            if tokens:
                return _cap_or_join(tokens)

    tokens = _extract_slip_tokens(val)
    if tokens:
        return _cap_or_join(tokens)

    cleaned = re.sub(r"\bVARIOUS\b\s*", "", val, flags=re.IGNORECASE)
    cleaned = re.sub(r"^\s*-\s*", "", cleaned).strip()
    cleaned = _normalize_spaces(cleaned)
    return [cleaned] if cleaned else []


# ─────────────────────────────────────────────────────────────────────────────
# CLEAN INSURED
# ─────────────────────────────────────────────────────────────────────────────

_INSURED_PREFIX_RE = re.compile(r"^\s*(?:PT|CV)\.?\s*", re.IGNORECASE)
_INSURED_GELAR_RE = re.compile(
    r"""
    \b(?:BAPAK|IBU|SDRI?|NYONYA|NY|TUAN|MR|MRS|MS|MBA|
        DR|IR|PROF|HJ|H)\b\.?\s*
    """,
    re.IGNORECASE | re.VERBOSE,
)
_INSURED_BOILERPLATE_RE = re.compile(
    r"""
    \bANY\s+SUBSIDIARY\s+COMPANY\b
  | \bAND\s+ALL\s+SUBSIDIAR(?:Y|IES)\b
  | INCLUDING\s+ANY\s+SUBSIDIARIES.*$
  | \bAS\s+OPERATOR\b
    """,
    re.IGNORECASE | re.VERBOSE,
)
_INSURED_SPLIT_RE = r"\bQQ\b|\bAND\s*/\s*OR\b|/|,|–|¿"


def _remove_polis_slip_from_insured(text: str, polis_ori, slip_ori) -> str:
    for token in (polis_ori, slip_ori):
        if pd.notna(token):
            t = str(token).strip()
            if t and t != "-":
                text = text.replace(t, "")
    return text


def clean_insured(val, polis_ori=None, slip_ori=None) -> list:
    if pd.isna(val):
        return []
    val = str(val).strip()
    if not val:
        return []

    val = _remove_polis_slip_from_insured(val, polis_ori, slip_ori)

    val = re.sub(r",\s*(?:PT|CV)\.?(?=[\s/]|$)", " ", val, flags=re.IGNORECASE)
    val = re.sub(r",\s*KSO\.?\s*$", " KSO", val, flags=re.IGNORECASE)

    raw_parts = re.split(_INSURED_SPLIT_RE, val, flags=re.IGNORECASE)
    raw_parts = _merge_digit_start_parts(raw_parts)
    raw_parts = _merge_branch_keyword_parts(raw_parts)

    cleaned = []
    for p in raw_parts:
        p = _normalize_spaces(p.strip())
        if len(p) <= 2:
            continue
        if p.upper().strip() in INSURED_JUNK_WORDS:
            continue
        if re.match(r"^[^a-zA-Z0-9]+$", p):
            continue
        p_clean = _clean_insured_name(p)
        if not p_clean or len(p_clean) <= 2:
            continue
        if p_clean.upper().strip() in INSURED_JUNK_WORDS:
            continue
        cleaned.append(p_clean)

    if not cleaned:
        fallback = _clean_insured_name(_normalize_spaces(val))
        return [fallback] if fallback else []

    final = []
    for name in cleaned:
        name = _strip_trailing_abbrev_word(name, final)
        if any(_is_abbreviation_of(name, prior) for prior in final):
            continue
        final.append(name)

    return _cap_or_join(final)


# ─────────────────────────────────────────────────────────────────────────────
# KOLOM CERTIFICATE (BARU) -- rule sama persis dengan Data 2/Osbal
# ─────────────────────────────────────────────────────────────────────────────

_CFMT_SD_PHRASE_RE = re.compile(r"(\d{1,7})\s*(?:S\s*[./]?\s*D|¿)\s*(\d{1,7})", re.IGNORECASE)
_CFMT_DASH_PHRASE_RE = re.compile(r"(\d{1,7})\s*-\s*(\d{1,7})")
_CFMT_BARE_NUM_RE = re.compile(r"^\d{1,7}$")


def _cert_pad6(s: str) -> str:
    return s.strip().zfill(6)


def _cert_reconstruct(base: str, suf: str) -> str:
    suf = suf.strip()
    if len(suf) >= len(base):
        return _cert_pad6(suf)
    return base[:-len(suf)] + suf


def format_certificate_segment(raw_text, base_ctx=None, keyword_context=False, force_comma_pairs=False):
    if not raw_text:
        return "", base_ctx
    text = raw_text.strip()

    placeholders = []

    def _protect_sd(m):
        placeholders.append(("sd", m.group(1), m.group(2)))
        return f"\x00{len(placeholders) - 1}\x00"
    text = _CFMT_SD_PHRASE_RE.sub(_protect_sd, text)

    def _protect_dash(m):
        placeholders.append(("dash", m.group(1), m.group(2)))
        return f"\x00{len(placeholders) - 1}\x00"
    text = _CFMT_DASH_PHRASE_RE.sub(_protect_dash, text)

    raw_segments = [s.strip() for s in re.split(r"[,/&\s]+", text) if s.strip()]
    if not raw_segments:
        return "", base_ctx

    sole_segment = len(raw_segments) == 1
    out_parts = []
    current_base = base_ctx

    for seg in raw_segments:
        ph_m = re.fullmatch(r"\x00(\d+)\x00", seg)
        if ph_m:
            kind, left_raw, right_raw = placeholders[int(ph_m.group(1))]
            if current_base and len(left_raw) < 6:
                left_full = _cert_reconstruct(current_base, left_raw)
            else:
                left_full = _cert_pad6(left_raw)
            right_full = _cert_reconstruct(left_full, right_raw) if len(right_raw) < 6 else _cert_pad6(right_raw)

            if force_comma_pairs:
                out_parts.append(left_full)
                out_parts.append(right_full)
            elif keyword_context:
                span = abs(int(right_full) - int(left_full)) + 1
                if span <= 3:
                    lo, hi = sorted([int(left_full), int(right_full)])
                    width = len(left_full)
                    for n in range(lo, hi + 1):
                        out_parts.append(str(n).zfill(width))
                else:
                    out_parts.append(f"{left_full} SD {right_full}")
            elif kind == "sd":
                out_parts.append(f"{left_full} SD {right_full}")
            else:
                if sole_segment:
                    out_parts.append(f"{left_full} SD {right_full}")
                else:
                    out_parts.append(left_full)
                    out_parts.append(right_full)
            current_base = right_full
            continue

        if _CFMT_BARE_NUM_RE.match(seg):
            if current_base and len(seg) < 6:
                full = _cert_reconstruct(current_base, seg)
            else:
                full = _cert_pad6(seg)
            out_parts.append(full)
            current_base = full
            continue

        continue

    if not out_parts:
        return "", current_base
    return ", ".join(out_parts), current_base


_CERT_KEYWORD_RE = re.compile(
    r"CERT(?:IF(?:ICATE)?)?\s*(?:NO\.?)?\s*[:.]?\s*(?=[\d(])",
    re.IGNORECASE,
)
_CERT_TAIL_JUNK_RE = re.compile(r"\s*/\s*VAR(?:IOUS)?\s*$", re.IGNORECASE)

_CERT_BASE_EXACT_LEN_RE = re.compile(r"^\s*(\d{15}|\d{14}|\d{18}|\d{17})(?:\s*/\s*\d{1,2}(?!\d))?")
_CERT_BASE_GREEDY_RE = re.compile(r"^\s*(\d{10,})(?:\s*/\s*\d{1,2}(?!\d))?")

_CERT_HAS_DASH_PAIR_RE = re.compile(r"^(\d{1,7})\s*-\s*(\d{1,7})$")
_CERT_EMBEDDED_POLICY_RE = re.compile(r"\d{10,}")
_CERT_SD_INTERRUPTED_RE = re.compile(
    r"\d{1,7}\s*(?:S\s*[./]?\s*D|¿)\s*\d{10,}\s+\d{1,7}", re.IGNORECASE
)

_CERT_MULTI_RECORD_RE = re.compile(r"(\d{14,18})\s*(?:¿|-)\s*(\d{5,7})(?!\d)")


def _cert_strip_paren(text: str) -> str:
    return text.strip().strip("()").strip()


def _cert_paren_is_narrative(inner: str) -> bool:
    """True kalau isi kurung adalah catatan narasi (ada kata asli), bukan
    list/range angka murni. Angka + pemisah (,&/-. spasi) + huruf S/D
    (notasi S/D) TIDAK dihitung sbg narasi."""
    letters_only = re.sub(r"[\d,&/\-.\s]", "", inner)
    letters_only = re.sub(r"(?i)S\s*\.?\s*D", "", inner)
    letters_only = re.sub(r"[\d,&/\-.\s]", "", letters_only)
    return len(letters_only) > 0


def _cert_looks_like_cert(tail: str) -> bool:
    t = tail.strip()
    has_sd_marker = bool(re.search(r"\d\s*(?:S\s*[./]?\s*D|¿)\s*\d", t, re.IGNORECASE))
    has_list_marker = bool(re.search(r"[,&]", t)) or t.count("/") >= 1
    has_space_number_list = bool(re.fullmatch(r"(?:\d{1,7}\s+)+\d{1,7}", t))
    is_single_long_num = bool(re.fullmatch(r"\d{5,7}", t))
    dash_m = _CERT_HAS_DASH_PAIR_RE.match(t)
    has_dash_pair = bool(dash_m and (len(dash_m.group(1)) >= 5 or len(dash_m.group(2)) >= 5))
    return (has_sd_marker or has_list_marker or has_space_number_list
            or is_single_long_num or has_dash_pair)


def _cert_extract_tail_for_base(val, base_m):
    remainder = val[base_m.end():]
    if remainder.startswith(" - ") or re.match(r"^\s+-\s+", remainder):
        tail = re.sub(r"^\s*-\s*", "", remainder).strip()
    elif remainder[:1] in ("-", "/"):
        tail = remainder[1:].strip()
    elif remainder.strip():
        tail = remainder.strip()
    else:
        return None
    tail = _CERT_TAIL_JUNK_RE.sub("", tail)
    return tail


def _cert_clean_chunk(chunk: str) -> str:
    if _CERT_SD_INTERRUPTED_RE.search(chunk):
        return chunk
    if not _CERT_EMBEDDED_POLICY_RE.search(chunk):
        return chunk if _cert_looks_like_cert(chunk) or re.fullmatch(r"\d{1,7}", chunk.strip()) else ""

    cleaned = _CERT_EMBEDDED_POLICY_RE.sub(" ", chunk)
    cleaned = re.sub(r"^[\s/\-]+|[\s/\-]+$", "", cleaned).strip()
    if cleaned and (_cert_looks_like_cert(cleaned) or re.fullmatch(r"\d{1,7}", cleaned)):
        return cleaned
    return ""


def _cert_clean_tail(tail: str) -> str:
    chunks = [c.strip() for c in tail.split(",") if c.strip()]
    kept = [c for c in (_cert_clean_chunk(c) for c in chunks) if c]
    return ",".join(kept)


def extract_certificate_csi(raw_value) -> str:
    """Ekstrak & format kandidat CERTIFICATE standalone (mis. "CERTIFICATE
    NO 000023 S/D 000026/31" tanpa nomor polis di depannya) -- dipakai
    sbg fallback oleh extract_certificate_slots_csi() saat tidak ada
    nomor polis 14-18 digit sama sekali dalam value ORI. Rule sama
    persis dgn Data 2/Osbal."""
    if raw_value is None:
        return ""
    val = str(raw_value).strip()
    if not val:
        return ""

    multi_matches = _CERT_MULTI_RECORD_RE.findall(val)
    if len(multi_matches) >= 2:
        seen = []
        for _base, cert in multi_matches:
            padded = cert.strip().zfill(6)
            if padded not in seen:
                seen.append(padded)
        if seen:
            return ", ".join(seen)

    kw_m = _CERT_KEYWORD_RE.search(val)
    if kw_m:
        cert_text = val[kw_m.end():]
        cert_text = _cert_strip_paren(cert_text)
        cert_text = _CERT_TAIL_JUNK_RE.sub("", cert_text)
        cert_text = _cert_clean_tail(cert_text)
        had_interrupt = bool(_CERT_SD_INTERRUPTED_RE.search(cert_text))
        cert_text = _CERT_EMBEDDED_POLICY_RE.sub(" ", cert_text)
        formatted, _ = format_certificate_segment(
            cert_text, keyword_context=True, force_comma_pairs=had_interrupt
        )
        return formatted

    paren_m = re.search(r"\(([^()]*)\)", val)
    if paren_m and re.search(r"\d", paren_m.group(1)) and not _cert_paren_is_narrative(paren_m.group(1)):
        inner = paren_m.group(1)
        kw_inner = _CERT_KEYWORD_RE.search(inner)
        is_kw = bool(kw_inner)
        if kw_inner:
            inner = inner[kw_inner.end():]
        if re.search(r"\d", inner):
            formatted, _ = format_certificate_segment(inner, keyword_context=is_kw)
            if formatted:
                return formatted

    for base_re in (_CERT_BASE_EXACT_LEN_RE, _CERT_BASE_GREEDY_RE):
        base_m = base_re.match(val)
        if not base_m:
            continue
        tail = _cert_extract_tail_for_base(val, base_m)
        if tail is None:
            continue

        tail = _cert_clean_tail(tail)
        had_interrupt = bool(_CERT_SD_INTERRUPTED_RE.search(tail))
        tail_clean = _CERT_EMBEDDED_POLICY_RE.sub(" ", tail).strip()
        if tail_clean and _cert_looks_like_cert(tail_clean):
            formatted, _ = format_certificate_segment(
                tail_clean, keyword_context=False, force_comma_pairs=had_interrupt
            )
            if formatted:
                return formatted

    return ""


def _format_cert_tail_precise(tail: str) -> str:
    tail = tail.strip(" ,-/")
    if not tail:
        return ""

    placeholders = []

    def protect_sd(m):
        placeholders.append(("sd", m.group(1), m.group(2)))
        return f"__CERT_PART_{len(placeholders)-1}__"

    work = re.sub(
        r"(\d{1,7})\s*S\s*[./]?\s*D\s*(\d{1,7})",
        protect_sd,
        tail,
        flags=re.IGNORECASE,
    )

    def protect_range(m):
        placeholders.append(("range", m.group(1), m.group(2)))
        return f"__CERT_PART_{len(placeholders)-1}__"

    work = re.sub(
        r"(\d{1,7})\s*-\s*(\d{1,7})",
        protect_range,
        work,
    )

    parts = [x.strip() for x in re.split(r"[,/&\s]+", work) if x.strip()]
    result = []
    previous = None

    for part in parts:
        pm = re.fullmatch(r"__CERT_PART_(\d+)__", part)
        if pm:
            kind, left_raw, right_raw = placeholders[int(pm.group(1))]

            left = _cert_pad6(left_raw)
            right = (
                _cert_reconstruct(left, right_raw)
                if len(right_raw) < 6
                else _cert_pad6(right_raw)
            )

            if kind == "sd":
                result.append(f"{left} SD {right}")
            else:
                span = abs(int(right) - int(left)) + 1

                if span <= 3 and (len(right_raw) == 6 or len(left_raw) == 6):
                    lo, hi = sorted((int(left), int(right)))
                    result.extend(
                        str(n).zfill(6)
                        for n in range(lo, hi + 1)
                    )
                else:
                    result.append(f"{left} SD {right}")

            previous = right
            continue

        if re.fullmatch(r"\d{1,7}", part):
            if previous and len(part) < 6:
                full = _cert_reconstruct(previous, part)
            else:
                full = _cert_pad6(part)

            result.append(full)
            previous = full

    if len(result) >= 4 and all(re.fullmatch(r"\d{6}", x) for x in result):
        nums = [int(x) for x in result]
        if nums == list(range(nums[0], nums[0] + len(nums))):
            return f"{result[0]} SD {result[-1]}"

    return ", ".join(result)


_POLICY_REPETITION_CERT_OVERRIDE_RE = re.compile(
    r"""
    ^\s*
    (?:
        010103112600028\s*-\s*000003\s*,\s*4\s*,\s*6
      | 012408052300003\s*,\s*4
      | P[1-5]\s*/\s*010103012600053\s*,\s*54\s*,\s*55\s*,\s*56
    )
    \s*$
    """,
    re.IGNORECASE | re.VERBOSE,
)


def _extract_certificate_from_policy_record(record: str, multi_record: bool = False, group_has_strong_cert: bool = False) -> str:
    if not record:
        return ""

    text = str(record).strip()

    if _POLICY_REPETITION_CERT_OVERRIDE_RE.match(text):
        return ""

    embedded_m = re.match(
        r"^\s*(\d{15})(0\d{5})(?P<rest>.*)$",
        text,
    )

    if embedded_m:
        base = embedded_m.group(1)
        tail = embedded_m.group(2) + embedded_m.group("rest")
    else:
        base_m = re.match(r"^\s*(\d{14,18})(?!\d)", text)
        if not base_m:
            return ""

        base = base_m.group(1)
        tail = text[base_m.end():].strip()

    tail = re.sub(r"^\s*/\s*\d{1,2}\b", "", tail, count=1).strip()

    # BUG FIX #2: "POLIS1-CERT1 S/D POLIS2-CERT2" (atau varian spasi
    # "POLIS1 CERT1 S/D POLIS2 CERT2") -- karena _split_policy_records_for_certificate
    # memecah per kemunculan nomor polis, S/D di tengah bisa nyangkut jadi
    # EKOR record pertama (mis. "-000001 S/D") tanpa pasangan angka di
    # kanan-nya dalam record itu. S/D yang menggantung begini bukan range
    # sertifikat -- itu cuma penghubung "sampai dengan" antar 2 polis yang
    # berbeda, jadi masing-masing cert (CERT1 & CERT2) tetap berdiri
    # sendiri-sendiri. Buang S/D yang menggantung di ekor SEBELUM parsing
    # cert supaya CERT1 tidak ikut hilang.
    tail = re.sub(r"\s*S\s*[./]?\s*D\s*$", "", tail, flags=re.IGNORECASE).strip()

    if not tail:
        return ""

    kw = re.search(
        r"\bCERT(?:IF(?:ICATE)?)?\s*\.?\s*(?:NO\.?)?\s*[:.]?\s*",
        tail,
        flags=re.IGNORECASE,
    )
    if kw:
        cert_tail = tail[kw.end():].strip(" ,-/")
        return _format_cert_tail_precise(cert_tail)

    special = re.match(
        r"^\s*-\s*0{2,5}\s*-\s*(\d{1,7})"
        r"(?P<rest>.*)$",
        tail,
        flags=re.IGNORECASE,
    )
    if special:
        special_tail = _cert_pad6(special.group(1)) + special.group("rest")
        formatted_special = _format_cert_tail_precise(special_tail)
        if formatted_special:
            return formatted_special

        return _cert_pad6(special.group(1))

    # BUG FIX #3: kurung yang isinya catatan narasi dgn 1 angka pendek
    # nyempil (mis. "(SLIP 1 LAGI BLM ADA)") BUKAN sertifikat -- beda dgn
    # kurung berisi list/range angka murni yang MEMANG valid sbg cert
    # (mis. "(16-24)", "(01&07)"). Bedanya: narasi punya kata2 asli,
    # sedangkan list/range cert cuma angka + pemisah (,&/-. spasi) +
    # paling banyak huruf "S"/"D" dari notasi S/D. Tanpa cek ini, angka
    # apa pun di dalam kurung (termasuk angka nyempil di tengah narasi)
    # salah dianggap nomor sertifikat.
    paren_m = re.search(r"\(([^()]*)\)", tail)
    if paren_m and re.search(r"\d", paren_m.group(1)) and not _cert_paren_is_narrative(paren_m.group(1)):
        inner = paren_m.group(1)
        formatted = _format_cert_tail_precise(inner)
        if formatted:
            return formatted

    candidate = re.sub(r"^\s*[-–—]\s*", "", tail, count=1).strip()
    candidate = candidate.strip(" ,/&")

    if not candidate:
        return ""

    sd_matches = re.findall(
        r"(\d{1,7})\s*S\s*[./]?\s*D\s*(\d{1,7})",
        candidate,
        flags=re.IGNORECASE,
    )
    if sd_matches:
        if any(
            len(a) >= 5 or len(b) >= 5
            for a, b in sd_matches
        ):
            return _format_cert_tail_precise(candidate)

    dash_matches = re.findall(
        r"(\d{1,7})\s*-\s*(\d{1,7})",
        candidate,
    )
    if dash_matches:
        if any(
            len(a) >= 5 or len(b) >= 5
            for a, b in dash_matches
        ):
            return _format_cert_tail_precise(candidate)

    list_numbers = re.findall(r"\d{1,7}", candidate)
    has_list_separator = bool(re.search(r"[,/&]", candidate))
    # BUG FIX #7: list angka pendek dgn anggota LEBIH DARI MAX_SPLIT_COLS
    # DAN tidak satupun anggotanya "lengkap" (>=5 digit, mis. "000001")
    # (spt "023,024,025,026,028,031,032,033,034" -- 9 anggota, semuanya
    # cuma 3 digit berurutan) adalah list suffix PERULANGAN POLIS, bukan
    # sertifikat. Kalau ADA anggota >=5 digit (mis. "000001,000002,...,
    # 8,9,11" atau "00003,00004,...,00011"), itu tetap sertifikat asli
    # walau anggotanya banyak -- heuristik zero-padded/>=5-digit lama
    # tetap berlaku untuk kasus itu.
    if has_list_separator and len(list_numbers) >= 2:
        if len(list_numbers) > MAX_SPLIT_COLS and not any(len(n) >= 5 for n in list_numbers):
            pass
        elif any(len(n) >= 5 or n.startswith("0") for n in list_numbers):
            return _format_cert_tail_precise(candidate)

    single = re.fullmatch(r"0\d{4,6}", candidate)
    if single:
        return _format_cert_tail_precise(candidate)

    if multi_record:
        short_sd = re.search(
            r"\b(\d{1,4})\s*S\s*[./]?\s*D\s*(\d{1,4})\b",
            candidate,
            flags=re.IGNORECASE,
        )
        if short_sd:
            return _format_cert_tail_precise(candidate)

        # BUG FIX #11: "short_range" (dash pendek tanpa syarat panjang, mis.
        # "53-56") dan angka pendek TANPA leading zero (mis. "36") cuma
        # dianggap sertifikat kalau ADA record LAIN dalam value yang sama
        # yang sudah punya bukti sertifikat KUAT (keyword CERT, pola
        # khusus, kurung, S/D eksplisit, dash/list dgn sisi >=5 digit, atau
        # angka tunggal ber-leading-zero). Tanpa bukti kuat dari sibling,
        # angka pendek begini adalah suffix PERULANGAN POLIS biasa (mis.
        # "601" pada "010201211900977-601" yang berdiri sendiri tanpa
        # konteks sertifikat apa pun) -- BUKAN sertifikat.
        if group_has_strong_cert:
            short_range = re.search(r"\b(\d{1,4})\s*-\s*(\d{1,4})\b", candidate)
            if short_range:
                return _format_cert_tail_precise(candidate)

        # BUG FIX #8: single 1-4 digit SELALU dianggap sertifikat kalau
        # DIAWALI ANGKA 0 (mis. "0001") -- itu ciri khas penomoran cert
        # yang di-pad, jadi berdiri sendiri (tidak perlu bukti sibling).
        if re.fullmatch(r"0\d{0,3}", candidate) and not re.fullmatch(r"0+", candidate):
            return _format_cert_tail_precise(candidate)

        # BUG FIX #11 (lanjutan): angka pendek TANPA leading zero (mis.
        # "36") hanya dianggap sertifikat kalau sibling record sudah
        # membuktikan konteks sertifikat KUAT.
        if group_has_strong_cert and re.fullmatch(r"\d{1,4}", candidate) and not re.fullmatch(r"0+", candidate):
            return _format_cert_tail_precise(candidate)

    return ""


def _certs_for_records(records: list) -> list:
    """Hitung CERTIFICATE utk tiap record dgn 2 tahap (lihat BUG FIX #11 di
    _extract_certificate_from_policy_record): tahap 1 cuma bukti kuat,
    tahap 2 (kalau ada sibling yg lolos tahap 1) mengizinkan sinyal lemah
    (short_range / angka pendek tanpa leading zero) utk record yg masih
    kosong. Dipakai bersama oleh extract_certificate_slots_csi() dan
    _remove_known_certificate_parts() supaya konsisten."""
    is_multi_record = len(records) >= 2
    if not is_multi_record:
        return [
            _extract_certificate_from_policy_record(record, multi_record=False)
            for record in records
        ]

    strong_certs = [
        _extract_certificate_from_policy_record(
            record, multi_record=True, group_has_strong_cert=False,
        )
        for record in records
    ]
    if not any(strong_certs):
        return strong_certs
    return [
        c if c else _extract_certificate_from_policy_record(
            record, multi_record=True, group_has_strong_cert=True,
        )
        for c, record in zip(strong_certs, records)
    ]


def _split_policy_records_for_certificate(val: str) -> list:
    if not val:
        return []

    text = str(val).strip()
    starts = []

    embedded_re = re.compile(
        r"(?<!\d)(\d{15})(0\d{5})(?=\s*(?:[-,/&]|$))"
    )
    for m in embedded_re.finditer(text):
        starts.append(m.start())

    normal_re = re.compile(r"(?<!\d)\d{14,18}(?!\d)")
    for m in normal_re.finditer(text):
        if not any(s == m.start() for s in starts):
            starts.append(m.start())

    if not starts:
        return []

    starts = sorted(set(starts))

    records = []
    for i, start_pos in enumerate(starts):
        end_pos = starts[i + 1] if i + 1 < len(starts) else len(text)
        record = text[start_pos:end_pos].strip(" ,")
        if record:
            records.append(record)

    def _record_policy_base(record: str) -> str:
        m = re.search(
            r"(?<!\d)(\d{15})(?=0\d{5}(?=\s*(?:[-,/&]|$)))",
            record,
        )
        if m:
            return m.group(1)

        m = re.search(r"(?<!\d)\d{14,18}(?!\d)", record)
        return m.group(0) if m else ""

    merged = []
    base_to_index = {}

    for record in records:
        base = _record_policy_base(record)

        if base and base in base_to_index:
            idx = base_to_index[base]

            tail = record
            base_pos = tail.find(base)
            if base_pos >= 0:
                tail = tail[base_pos + len(base):]
            tail = tail.strip(" ,/-")

            if tail:
                if re.search(r"S\s*[./]?\s*D\s*$", merged[idx], flags=re.IGNORECASE):
                    merged[idx] = f"{merged[idx]} {tail}"
                else:
                    merged[idx] = f"{merged[idx].rstrip(' ,/')} , {tail}"
        else:
            if base:
                base_to_index[base] = len(merged)
            merged.append(record)

    return merged


def extract_certificate_slots_csi(raw_value, clean_polis_values=None) -> list:
    if raw_value is None or pd.isna(raw_value):
        return []

    val = str(raw_value).strip()
    if not val:
        return []

    if _is_non_cert_standalone(val):
        return []

    records = _split_policy_records_for_certificate(val)
    if not records:
        # BUG FIX #6: kalau sama sekali tidak ada nomor polis 14-18 digit
        # (mis. "CERTIFICATE NO 000023 S/D 000026/31" berdiri sendiri tanpa
        # polis di depannya), tetap coba ekstrak sertifikatnya lewat
        # extract_certificate_csi() -- TAPI hanya kalau ada kata kunci
        # CERT/CERTIF/CERTIFICATE eksplisit. Tanpa syarat ini, value yang
        # sebenarnya adalah POLIS belum lengkap (base <14 digit, mis.
        # "0104011018004-013 / 005 - 013" yang seharusnya di-glue jadi
        # polis oleh mesin reconstruct di clean_polis) akan salah
        # dianggap sertifikat.
        if _CERT_KEYWORD_RE.search(val):
            fallback_cert = extract_certificate_csi(val)
            return [fallback_cert] if fallback_cert else []
        return []

    record_certs = _certs_for_records(records)

    if len(records) == 1:
        cert = record_certs[0] if record_certs else ""
        return [cert] if cert else []

    return record_certs


# ─────────────────────────────────────────────────────────────────────────────
# MITRA BISNIS
# ─────────────────────────────────────────────────────────────────────────────

def get_business_partner(comp_name2, comp_name) -> str:
    broker = "" if pd.isna(comp_name2) else str(comp_name2).strip()
    if not broker or broker.upper() == "DIRECT":
        result = "" if pd.isna(comp_name) else str(comp_name).strip()
    else:
        result = broker
    result = re.sub(r"\.", " ", result)
    return re.sub(r"\s+", " ", result).upper().strip()


# ─────────────────────────────────────────────────────────────────────────────
# PROSES UTAMA
# ─────────────────────────────────────────────────────────────────────────────

def _insert_clean_columns(df, all_lists, prefix, max_cols):
    added = []
    for i in range(1, max_cols + 1):
        col_name = f"{prefix}_{i}"
        df[col_name] = [lst[i - 1] if i - 1 < len(lst) else None for lst in all_lists]
        added.append(col_name)
    return added


def _force_text_format(output_file, prefixes):
    import openpyxl
    wb = openpyxl.load_workbook(output_file)
    ws = wb.active
    header = [c.value for c in ws[1]]
    target_cols = [i for i, h in enumerate(header, start=1) if h and str(h).startswith(prefixes)]
    for col_idx in target_cols:
        for row in ws.iter_rows(min_row=2, min_col=col_idx, max_col=col_idx):
            for cell in row:
                cell.number_format = "@"
    wb.save(output_file)


def process_data(input_file: str, output_file: str) -> None:
    print(f"[1/5] Membaca data dari: {input_file} ...")
    df = pd.read_excel(input_file, header=0)
    print(f"      Total baris keseluruhan: {len(df):,}")

    if CEDANT_COL not in df.columns:
        print(f"\n[ERROR] Kolom '{CEDANT_COL}' tidak ditemukan!")
        print(f"        Kolom tersedia: {list(df.columns)}")
        return

    df[CEDANT_COL] = df[CEDANT_COL].astype(str).str.strip()
    df = df[df[CEDANT_COL] == CEDANT_VALUE].copy()
    print(f"[2/5] Filter cedant '{CEDANT_VALUE}': {len(df):,} baris ditemukan.")

    if df.empty:
        print("\n[WARN] Tidak ada data setelah filter. Proses dihentikan.")
        return

    for col in [POLIS_COL, SLIP_COL, INSURED_COL, BROKER_COL]:
        if col not in df.columns:
            print(f"\n[ERROR] Kolom '{col}' tidak ditemukan di file!")
            print(f"        Kolom tersedia: {list(df.columns)}")
            return

    business_partner_values = [get_business_partner(row[BROKER_COL], row[CEDANT_COL]) for _, row in df.iterrows()]

    df.rename(columns={INSURED_COL: "FAC_INSURED_ORI"}, inplace=True)

    insert_pos = list(df.columns).index(BROKER_COL) + 1 if BROKER_COL in df.columns else len(df.columns)
    df.insert(insert_pos, "BUSINESS_PARTNER", business_partner_values)

    print("[3/5] Menjalankan proses cleaning ...")

    all_clean_polis, all_clean_slip, all_clean_ins, all_certificate_slots = [], [], [], []
    max_polis = max_slip = max_ins = 1

    for _, row in df.iterrows():
        c_polis = _refine_list(clean_polis(row.get(POLIS_COL, "")), KNOWN_POLIS_PATTERNS)
        c_slip  = _refine_list(clean_slip(row.get(SLIP_COL, "")), KNOWN_SLIP_PATTERNS)
        c_ins   = clean_insured(row.get("FAC_INSURED_ORI", ""), row.get(POLIS_COL, ""), row.get(SLIP_COL, ""))
        cert_slots = extract_certificate_slots_csi(row.get(POLIS_COL, ""), c_polis)

        max_polis = max(max_polis, len(c_polis))
        max_slip  = max(max_slip, len(c_slip))
        max_ins   = max(max_ins, len(c_ins))

        all_clean_polis.append(c_polis)
        all_clean_slip.append(c_slip)
        all_clean_ins.append(c_ins)
        all_certificate_slots.append(cert_slots)

    print("[4/5] Menyusun kolom output ...")

    new_columns = []
    for col in df.columns:
        new_columns.append(col)
        if col == POLIS_COL:
            # ========================================================
            # POLICY CLEAN + CERTIFICATE
            # CERTIFICATE_i selalu tepat di kanan POLICY_CLEAN_i.
            #
            # Certificate yang FULL KOSONG tidak dibuat.
            # ========================================================
            max_cert_slots = max(
                (len(x) for x in all_certificate_slots),
                default=0
            )
            max_slots = max(max_polis, max_cert_slots)

            for i in range(1, max_slots + 1):
                policy_col = f"{POLIS_COL}_{i}"
                df[policy_col] = [
                    x[i - 1]
                    if isinstance(x, list)
                    and i - 1 < len(x)
                    and x[i - 1]
                    else None
                    for x in all_clean_polis
                ]
                new_columns.append(policy_col)

                cert_values = [
                    x[i - 1]
                    if isinstance(x, list)
                    and i - 1 < len(x)
                    and x[i - 1]
                    else None
                    for x in all_certificate_slots
                ]

                if any(
                    v is not None and str(v).strip() != ""
                    for v in cert_values
                ):
                    cert_col = f"CERTIFICATE_{i}"
                    df[cert_col] = cert_values
                    new_columns.append(cert_col)

        elif col == SLIP_COL:
            new_columns += _insert_clean_columns(df, all_clean_slip, SLIP_COL, max_slip)
        elif col == "FAC_INSURED_ORI":
            new_columns += _insert_clean_columns(df, all_clean_ins, "FAC_INSURED", max_ins)

    df = df[new_columns]

    os.makedirs(os.path.dirname(output_file), exist_ok=True)
    print(f"[5/5] Menyimpan hasil ke: {output_file} ...")
    df.to_excel(output_file, index=False)
    _force_text_format(output_file, prefixes=(f"{POLIS_COL}_", f"{SLIP_COL}_", "FAC_INSURED_", "CERTIFICATE"))

    print(f"\n{'=' * 55}")
    print("  [OK] Selesai!")
    print(f"  Total baris output  : {len(df):,}")
    print(f"  Total kolom output  : {len(df.columns)}")
    print(f"  File disimpan di    : {output_file}")
    print(f"{'=' * 55}")


if __name__ == "__main__":
    process_data(INPUT_FILE, OUTPUT_FILE)