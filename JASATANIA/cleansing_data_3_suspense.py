"""
================================================================================
 SCRIPT CLEANSING DATA 3 - SUSPENSE (KHUSUS CEDANT: JASA TANIA)
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
INPUT_FILE = "3b. Database Suspense.xlsx"
SHEET_NAME = "Sheet1"
HEADER_ROW = 2
OUTPUT_FILE = "jasa_tania_output_suspense_V3.xlsx"
CEDANT_FILTER = "JASA TANIA"
STATUS_FILTER = "SUSPENSE"

MAX_BREAKDOWN_CODES = 5
MAX_BREAKDOWN_INSURED = 5
TEXT_FORMAT_COLUMN_KEYWORDS = ("POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO", "CERTIFICATE")
RE_I = re.IGNORECASE

CERTIFICATE_PREFIX_PATTERNS = []

# ==============================================================================
# 1. CLEANSING INSURED
# ==============================================================================
_ENTITY_TITLE_RE = re.compile(
    r"\b(?:P\.?T\.?|C\.?V\.?|\(PERSERO\)|PERSERO|TBK\.?|LTD\.?|PTE\.?|N\.?V\.?|"
    r"BAPAK|IBU|NY\.?|NYONYA|MR\.?|MRS\.?|MS\.?)\b", RE_I)

_KNOWN_BANK_KEYWORDS = ["BANK"]
_ENTITY_SPLIT_QQ_RE = re.compile(r"\bQQ\.?\b", RE_I)

def _is_bank_entity(text: str) -> bool:
    return any(kw in text.upper() for kw in _KNOWN_BANK_KEYWORDS)

def _split_insured_entities(raw: str) -> list:
    parts = [p.strip() for p in _ENTITY_SPLIT_QQ_RE.split(raw.strip()) if p.strip()]
    return parts[1:] if len(parts) > 1 and _is_bank_entity(parts[0]) else (parts or [raw.strip()])

def _clean_single_insured(name: str) -> str:
    text = _ENTITY_TITLE_RE.sub("", name.strip())
    text = re.sub(r"[,;:.()]+", " ", text)
    text = re.sub(r"^[-/]+|[-/]+$", "", text)
    return re.sub(r"\s+", " ", text).strip().upper()

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not (original := text.strip()) or original.lower() in ("nan", "none", "-"):
        return []
    cleaned = [c for c in (_clean_single_insured(e) for e in _split_insured_entities(original)) if c]
    return [original] if len(cleaned) > MAX_BREAKDOWN_INSURED or not cleaned else cleaned

# ==============================================================================
# 2. CLEANSING POLIS & SLIP NO
# ==============================================================================
_POLIS_INSTALLMENT_SUFFIX_RE = re.compile(r"-\d{1,2}/\d{1,2}$")
_POLIS_DOUBLE_SLASH_SPLIT_RE = re.compile(r"\s*//\s*")

def clean_polis_code(text: str) -> list:
    if not isinstance(text, str) or not (t := text.strip()) or t.lower() in ("nan", "none"):
        return []

    t = re.sub(r"\s+", " ", t)
    if t == "-": return ["-"]

    if _POLIS_DOUBLE_SLASH_SPLIT_RE.search(t):
        codes = [_POLIS_INSTALLMENT_SUFFIX_RE.sub("", p).upper() for p in _POLIS_DOUBLE_SLASH_SPLIT_RE.split(t) if p.strip()]
        return codes[:1] if len(codes) > MAX_BREAKDOWN_CODES else codes

    return [_POLIS_INSTALLMENT_SUFFIX_RE.sub("", t).upper()]

def clean_slip_code(text: str) -> list:
    if not isinstance(text, str) or not (t := text.strip()) or t.lower() in ("nan", "none"):
        return []
    return [re.sub(r"\s+", " ", t).upper()]

# ==============================================================================
# 3. EKSTRAKSI CERTIFICATE
# ==============================================================================
_SD_NORMALIZE_RE = re.compile(r"\bs\s*/\s*d\b", RE_I)
_CERT_DIGIT_TAIL_RE = re.compile(r"(\d{5,7})$")
_CERT_PREFIX_RES = [re.compile(p, RE_I) for p in CERTIFICATE_PREFIX_PATTERNS]

def _normalize_sd(text: str) -> str:
    return _SD_NORMALIZE_RE.sub("SD", text.strip())

def _has_valid_certificate_prefix(prefix: str) -> bool:
    if not _CERT_PREFIX_RES:
        return False
    return any(p.search(prefix) for p in _CERT_PREFIX_RES)

def _to_six_digit(num_str: str) -> str | None:
    if len(num_str) == 5:
        return "0" + num_str
    if len(num_str) == 6:
        return num_str
    return None

def _extract_single_certificate(raw: str) -> str | None:
    text = raw.strip()
    if not text:
        return None

    match = _CERT_DIGIT_TAIL_RE.search(text)
    if not match:
        return None

    digits = match.group(1)
    prefix = text[: match.start()]

    six_digit = _to_six_digit(digits)
    if six_digit is None:
        return None

    if not prefix or not _has_valid_certificate_prefix(prefix):
        return None

    return six_digit

def extract_certificate(text: str) -> str:
    if not isinstance(text, str) or not (t := text.strip()) or t.lower() in ("nan", "none", "-"):
        return ""

    t = _normalize_sd(t)

    range_match = re.match(r"^(.*?)(\d{5,7})\s*(?:SD|-)\s*(?:\1)?(\d{5,7})$", t, RE_I)
    if range_match:
        prefix, start_digits, end_digits = range_match.groups()
        start_six = _to_six_digit(start_digits)
        end_six = _to_six_digit(end_digits)
        if start_six and end_six and _has_valid_certificate_prefix(prefix):
            start_num, end_num = int(start_six), int(end_six)
            if end_num < start_num:
                return ""
            count = end_num - start_num + 1
            if count > 3:
                return f"{start_six} SD {end_six}"
            return ", ".join(f"{n:06d}" for n in range(start_num, end_num + 1))
        return ""

    single = _extract_single_certificate(t)
    return single if single else ""

# ==============================================================================
# 4. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_status_suspense(df_raw: pd.DataFrame, keyword: str) -> pd.DataFrame:
    status_col = next((c for c in df_raw.columns if str(c).upper().strip() == "STATUS"), None)
    if not status_col:
        raise ValueError("Kolom STATUS tidak ditemukan di source file.")

    mask = df_raw[status_col].astype(str).str.strip().str.upper() == keyword.upper()
    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter STATUS = '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered

def filter_cedant(df_raw: pd.DataFrame, keyword: str) -> pd.DataFrame:
    cedant_cols = [c for c in df_raw.columns if "CEDANT" in str(c).upper()]
    if not cedant_cols:
        raise ValueError("Kolom CEDANT tidak ditemukan di source file.")

    mask = pd.Series(False, index=df_raw.index)
    for col in cedant_cols:
        mask |= df_raw[col].astype(str).str.contains(keyword, case=False, na=False)

    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant mengandung '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered

# ==============================================================================
# 5. BUILD OUTPUT
# ==============================================================================
def _detect_key_columns(df: pd.DataFrame) -> tuple:
    col_insured = next((c for c in df.columns if "INSURED" in str(c).upper()), None)
    col_polis = next((c for c in df.columns if str(c).upper().strip() == "POLIS"), None)
    col_slip = next((c for c in df.columns if "SLIP" in str(c).upper()), None)

    if None in (col_insured, col_polis, col_slip):
        raise ValueError(f"Kolom kunci tidak lengkap. INSURED={col_insured}, POLIS={col_polis}, SLIP={col_slip}")
    print(f"[INFO] Kolom terdeteksi -> INSURED='{col_insured}', POLIS='{col_polis}', SLIP='{col_slip}'")
    return col_insured, col_polis, col_slip

def _breakdown_all_rows(df: pd.DataFrame, col_insured: str, col_polis: str, col_slip: str) -> tuple:
    insured_cln_all, polis_cln_all, slip_cln_all, certificate_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for ins_val, pol_val, slp_val in zip(df[col_insured].to_numpy(), df[col_polis].to_numpy(), df[col_slip].to_numpy()):
        ins_cln = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []
        pol_cln = clean_polis_code(str(pol_val)) if pd.notna(pol_val) else []
        slp_cln = clean_slip_code(str(slp_val)) if pd.notna(slp_val) else []
        cert = extract_certificate(str(pol_val)) if pd.notna(pol_val) else ""

        insured_cln_all.append(ins_cln)
        polis_cln_all.append(pol_cln)
        slip_cln_all.append(slp_cln)
        certificate_all.append(cert)

        if len(ins_cln) > max_ins: max_ins = len(ins_cln)
        if len(pol_cln) > max_pol: max_pol = len(pol_cln)
        if len(slp_cln) > max_slp: max_slp = len(slp_cln)

    return insured_cln_all, polis_cln_all, slip_cln_all, certificate_all, max_ins, max_pol, max_slp

def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_insured, col_polis, col_slip = _detect_key_columns(df_filtered)
    insured_cln_all, polis_cln_all, slip_cln_all, certificate_all, max_ins, max_pol, max_slp = _breakdown_all_rows(
        df_filtered, col_insured, col_polis, col_slip
    )

    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)

    output_cols_data = {}
    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        if col == col_insured:
            for i in range(1, max_ins + 1):
                output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in insured_cln_all]
        if col == col_polis:
            if max_pol >= 1:
                output_cols_data["POLIS_CLN_1"] = [lst[0] if len(lst) >= 1 else "" for lst in polis_cln_all]
            output_cols_data["CERTIFICATE"] = certificate_all
            for i in range(2, max_pol + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in polis_cln_all]
        if col == col_slip:
            for i in range(1, max_slp + 1):
                output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slip_cln_all]

    df_master = pd.DataFrame(output_cols_data)
    drop_cols = [c for c in df_master.columns if "_CLN_" in c and df_master[c].astype(str).str.strip().eq("").all()]
    return df_master.drop(columns=drop_cols)

# ==============================================================================
# 6. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    workbook, text_col_indices = load_workbook(output_path), [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]
    worksheet = workbook.active

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None:
                cell.value = str(cell.value)
    workbook.save(output_path)

# ==============================================================================
# 7. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print("=" * 70 + "\n CLEANSING DATA 3 - SUSPENSE | CEDANT: JASA TANIA\n" + "=" * 70)
    input_path = INPUT_FILE

    if not os.path.exists(input_path):
        candidates = [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]
        if not candidates:
            raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")
        input_path = candidates[0]
        print(f"[AUTO-DETECT] File input tidak ditemukan di path default, memakai: '{input_path}'")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    print(f"[INFO] Total baris source: {len(df_raw)}")

    df_suspense = filter_status_suspense(df_raw, STATUS_FILTER)
    if df_suspense.empty:
        print(f"[WARNING] Tidak ada baris dengan STATUS = '{STATUS_FILTER}'. Proses dihentikan.")
        return

    df_jasa = filter_cedant(df_suspense, CEDANT_FILTER)
    if df_jasa.empty:
        print("[WARNING] Tidak ada baris yang cocok dengan filter cedant. Proses dihentikan.")
        return

    df_hasil = build_output(df_jasa)
    save_with_text_format(df_hasil, OUTPUT_FILE)
    print("=" * 70 + f"\n[SUCCESS] Selesai. Total baris output: {len(df_hasil)}\n[SUCCESS] File hasil: '{OUTPUT_FILE}'\n" + "=" * 70)

if __name__ == "__main__":
    main()