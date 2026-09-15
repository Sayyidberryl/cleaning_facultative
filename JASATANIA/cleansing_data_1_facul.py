"""
================================================================================
 SCRIPT CLEANSING DATA 1 - FACULTATIVE (KHUSUS CEDANT: JASA TANIA)
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
INPUT_FILE = "1b. Transaksi Facul.xlsx"
SHEET_NAME = "Query result"
HEADER_ROW = 0
OUTPUT_FILE = "jasatania_output_facul_V2.xlsx"
CEDANT_FILTER = "JASA TANIA"

CEDANT_COLUMN = "COMP_NAME"
BROKER_NAME_COLUMN = "COMP_NAME.1"
DIRECT_MARKER = "DIRECT"
INSURED_COLUMN = "FAC_INSURED"
POLIS_COLUMN = "FAC_POLICY_NO"
SLIP_COLUMN = "FAC_SLIP"

MAX_BREAKDOWN_CODES = 5
TEXT_FORMAT_COLUMN_KEYWORDS = (
    "POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
    "FAC_POLICY_NO", "FAC_SLIP", "CERTIFICATE"
)
RE_I = re.IGNORECASE

# ==============================================================================
# 1. PRE-COMPILED REGEXES (INSURED & NOISES)
# ==============================================================================
_TITLE_RE = re.compile(
    r"\b(?:S\.?H\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|"
    r"Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu(?!\s+KOTA)|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b", RE_I
)
_LEGAL_ENTITY_RE = re.compile(
    r"\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO|Perseroan\s+Terbatas)\b", RE_I
)
_BANK_INSTITUTION_RE = re.compile(
    r"\bBANK\s+(?:MANDIRI|BRI|BNI|BCA|BUKOPIN|BTN|BTPN|PERMATA|DANAMON|CIMB|MAYBANK|"
    r"PANIN|OCBC|UOB|MEGA|SYARIAH)\b|\bBRI\b|\bBNI\b|\bBCA\b", RE_I
)
_QQ_RE = re.compile(r"\bQQ\b", RE_I)
_LEGAL_ENTITY_WITH_LEADING_PUNCT_RE = re.compile(
    r"[,.]?\s*\b(?:PT\.?(?!\w)|P\.T\.?|P\s+T(?!\w)|CV\.?(?!\w)|C\.V\.?|TBK\.?(?!\w)|Tbk\.?(?!\w)|LTD\.?(?!\w)|LIMITED|PTE\.?(?!\w)|INC\.?(?!\w)|PERSERO(?!\w)|Perseroan\s+Terbatas)", RE_I
)
_ENTITY_SPLIT_RE = re.compile(r"\bQQ+\.?\b|\band\s*/\s*or\b|\bdan\s*/\s*atau\b|;|/|,", RE_I)
_BARE_SUFFIX_FRAGMENT_RE = re.compile(
    r"^\s*(?:PT\.?|P\.T\.?|P\s+T|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO|KSO)\s*$", RE_I
)
_BARE_SUFFIX_WITH_PAREN_RE = re.compile(
    r"^\s*(?:PT\.?|P\.T\.?|P\s+T|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO)?\s*"
    r"(\uE100\d+\uE101)\s*$", RE_I
)
_SUFFIX_THEN_DASH_NEW_ENTITY_RE = re.compile(
    r"^\s*(PT\.?|P\.T\.?|P\s+T|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO)"
    r"\s+-\s+(.+)$", RE_I
)
_PAREN_GROUP_RE = re.compile(r"\([^()]*\)")
_PAREN_TOKEN_RE = re.compile(r"\uE100(\d+)\uE101")
_FINAL_PUNCT_RE = re.compile(r"[,;:\"']")
_FINAL_DASH_RE = re.compile(r"(?<!\s)-(?!\s)")
_FINAL_SLASH_RE = re.compile(r"/")
_FINAL_DOT_RE = re.compile(r"\.")
_FINAL_LPAREN_RE = re.compile(r"\(\s+")
_FINAL_RPAREN_RE = re.compile(r"\s+\)")
_MULTISPACE_RE = re.compile(r"\s+")
_FINAL_EDGE_RE = re.compile(r"^[.\s]+|[.\s]+$")
_FINAL_EMPTY_PAREN_RE = re.compile(r"\(\s*\)")
_DANGLING_PATTERN_RE = re.compile(
    r"^\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b\s*"
    r"|\s*\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b$", RE_I
)
_BARE_CLIENT_FRAGMENT_RE = re.compile(r"^\s*CLIENTS?\s*$", RE_I)
_INSURED_DESCRIPTIVE_RE = re.compile(r"\bVARIOUS\s+INSUREDS?\b", RE_I)
_REPEATED_PG_SPLIT_RE = re.compile(r"(?=\bPG\s+)", RE_I)
_AS_OWNER_NOISE_RE = re.compile(r"\bAS\s+OWNER\s+AND\s*/\s*OR\.?\s*", RE_I)
_BARE_ROMAN_NUMERAL_RE = re.compile(r"^[IVXLCDM]{1,6}$", RE_I)
_CLEAN_COMMA_RE = re.compile(r",\s*(?=[(\s]|$)", RE_I)
_CLEAN_WHOLE_DOT_RE = re.compile(r"\.")
_CLEAN_WHOLE_SLASH_RE = re.compile(r"\s*/\s*")
_CLEAN_WHOLE_EDGE_RE = re.compile(r"^[/\s]+|[/\s]+$")
_HAS_UPPER_RE = re.compile(r"[A-Z]")

_STANDALONE_KEEP_RE = re.compile(r"^\s*(VAR|VARIOUS|TBA|P\d{0,2})\s*$", RE_I)
_PENYELESAIAN_RE = re.compile(
    r"PENYELESAIAN\s+(SUSPEND(?:SE|E)?|SUSPENSE|\(?H\)?UTANG\s+PIUTANG)", RE_I
)
_SEE_ATTACHMENT_RE = re.compile(r"^\s*SEE\s+ATTACHMENT\s*$", RE_I)
_ADMIN_STATUS_NOISE_RE = re.compile(r"/?\s*SLIP\s+BELUM\s+LENGKAP\s*$", RE_I)
_YEAR_CHANGE_NOTE_RE = re.compile(r"TAHUN\s+(?:KE\s+\d+\s+)?\d{4}\s+KE\s+\d{4}", RE_I)
_MONTH_WORDS_ID = ("JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI",
                   "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER")
_MONTH_ALONE_RE = re.compile(r"\b(?:" + "|".join(_MONTH_WORDS_ID) + r")\b", RE_I)
_PLUS_NOISE_TOKEN_RE = re.compile(r"\+\s*(?:P\d{0,2}|VAR(?:IOUS)?|TBA)\b", RE_I)
_SLASH_NOISE_TOKEN_RE = re.compile(r"/\s*(?:VAR(?:IOUS)?)\b", RE_I)
_DASH_NOISE_TOKEN_RE = re.compile(r"\s*-\s*(?:P\d{0,2}|VAR(?:IOUS)?|TBA)\s*$", RE_I)
_LEADING_P_LABEL_RE = re.compile(r"^\s*P\d{1,2}\s*/\s*", RE_I)
_TBA_WITH_SEPARATOR_RE = re.compile(r"(?:^|(?<=[/.\-\s]))\s*TBA\s*(?=[/.\-]|\s|$)", RE_I)
_DASH_SHORT_SUFFIX_DISCARD_RE = re.compile(r"^([A-Za-z]+\d{6,})-\d{1,3}/\d{1,3}\s*$")
_BASE_CODE_RE = re.compile(r"^([A-Za-z]*)(\d{4,})$")
_FULL_CODE_TOKEN_RE = re.compile(r"^[A-Za-z]*\d{6,}$")
_SD_RANGE_RE = re.compile(r"^(.*?)(\d+)\s+S\s*/?\s*D\s+(\d+)\s*$", RE_I)
_P_REVS_RE = re.compile(r"\s*P\d{0,2}\s*/\s*P\d{0,2}\s+[A-Za-z].*", RE_I)
_P_PLUS_SEQ_RE = re.compile(r"\s*P\d{0,2}\s*(?:\+\s*P\d{0,2}\s*)+$", RE_I)
_DIGIT_DOT_MATCH_RE = re.compile(r"^[\d.]+$")
_DOT_COMMA_TAIL_RE = re.compile(r"[.,]\s*$")
_NON_DIGIT_RE = re.compile(r"[^0-9]")
_NON_ALNUM_RE = re.compile(r"[^A-Za-z0-9]")
_DIGIT_SEARCH_RE = re.compile(r"\d")
_POLIS_CERTIFICATE_SUFFIX_RE = re.compile(r"-\s*(?P<cert>\d+)\s*$", RE_I)

_PAREN_STASH: dict[str, str] = {}


# ==============================================================================
# 2. FUNGSI EKSTRAKSI DAN FORMAT CERTIFICATE
# ==============================================================================
def _format_certificate_string(cert_raw: str) -> str:
    if not cert_raw or pd.isna(cert_raw): return ""
    c = str(cert_raw).upper().strip()
    c = re.sub(r"\s+", "", c)
    if not c.isdigit(): return ""

    if len(c) == 6:
        return c
    elif len(c) == 5:
        return "0" + c
    else:
        return ""

def _extract_and_format_certificate(text: str) -> tuple:
    cert_match = _POLIS_CERTIFICATE_SUFFIX_RE.search(text)
    if not cert_match:
        return text, ""

    cert_raw = cert_match.group("cert")
    c = re.sub(r"\s+", "", str(cert_raw))

    if c.isdigit() and len(c) in (5, 6):
        polis_base = text[:cert_match.start()].strip()
        cert_formatted = _format_certificate_string(c)
        return polis_base, cert_formatted

    return text, ""


# ==============================================================================
# 3. CLEANSING INSURED
# ==============================================================================
def _strip_bank_qq_segment(text: str) -> str:
    qq_match = _QQ_RE.search(text)
    if not qq_match: return text
    before_qq = text[: qq_match.start()].strip()
    if _BANK_INSTITUTION_RE.search(before_qq): return text[qq_match.end():].strip()
    return text

def _strip_legal_entity_with_punct(text: str) -> str:
    t = _LEGAL_ENTITY_WITH_LEADING_PUNCT_RE.sub("", text)
    t = _CLEAN_COMMA_RE.sub(" ", t)
    return t

def _strip_titles(text: str) -> str:
    return _TITLE_RE.sub(" ", text)

def _merge_bare_suffix_fragments(parts: list) -> list:
    merged = []
    for p in parts:
        p_stripped = p.strip()
        if merged and (_BARE_SUFFIX_FRAGMENT_RE.match(p_stripped) or _BARE_SUFFIX_WITH_PAREN_RE.match(p_stripped)):
            merged[-1] = merged[-1] + " " + p_stripped
            continue
        suffix_dash_match = _SUFFIX_THEN_DASH_NEW_ENTITY_RE.match(p_stripped) if merged else None
        if suffix_dash_match:
            merged[-1] = merged[-1] + " " + suffix_dash_match.group(1).strip()
            merged.append(suffix_dash_match.group(2).strip())
            continue
        merged.append(p)
    return merged

def _protect_parens_content(text: str) -> str:
    def _mask(match: "re.Match") -> str:
        token = f"\uE100{len(_PAREN_STASH)}\uE101"
        _PAREN_STASH[token] = match.group(0)
        return token
    return _PAREN_GROUP_RE.sub(_mask, text)

def _restore_protected_parens(text: str) -> str:
    def _unmask(match: "re.Match") -> str:
        return _PAREN_STASH.get(f"\uE100{match.group(1)}\uE101", match.group(0))
    return _PAREN_TOKEN_RE.sub(_unmask, text)

def _final_polish(text: str) -> str:
    t = _FINAL_PUNCT_RE.sub(" ", text)
    t = _FINAL_DASH_RE.sub(" ", t)
    t = _FINAL_SLASH_RE.sub(" ", t)
    t = _FINAL_DOT_RE.sub(" ", t)
    t = _FINAL_LPAREN_RE.sub("(", t)
    t = _FINAL_RPAREN_RE.sub(")", t)
    t = _MULTISPACE_RE.sub(" ", t).strip()
    t = _FINAL_EDGE_RE.sub("", t)
    t = _FINAL_EMPTY_PAREN_RE.sub("", t).strip()
    previous = None
    while previous != t:
        previous, t = t, _DANGLING_PATTERN_RE.sub("", t).strip()
        t = _FINAL_EDGE_RE.sub("", t).strip()
    return t.upper()

def _split_repeated_pg_prefix(text: str):
    if len(re.findall(r"\bPG\s+", text, RE_I)) < 2: return None
    parts = [p.strip() for p in _REPEATED_PG_SPLIT_RE.split(text) if p.strip()]
    return parts if len(parts) >= 2 else None

def _has_bare_roman_numeral_slash_fragment(text: str) -> bool:
    if "/" not in text: return False
    parts = [p.strip() for p in text.split("/") if p.strip()]
    return any(_BARE_ROMAN_NUMERAL_RE.match(p) for p in parts)

def _is_bio_farma_cluster_pattern(text: str) -> bool:
    return bool(re.search(r"\bBIO\s+FARMA\b", text, RE_I)) and bool(
        re.search(r"\bCLUSTER\b", text, RE_I)
    )

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return []

    original = text.strip()
    if _INSURED_DESCRIPTIVE_RE.search(original):
        return [original]

    if re.sub(r"\s+", " ", original.upper()).strip() == (
        "CHEVRON PACIFIC INDONESIA : MAKASSAR , RAPAK , GANAL , APLIKANUSA LINTASARTA"
    ):
        return ["CHEVRON PACIFIC INDONESIA MAKASSAR RAPAK GANAL", "APLIKANUSA LINTASARTA"]

    if _is_bio_farma_cluster_pattern(original):
        if re.sub(r"\s+", " ", original.upper()).strip() == "BIO FARMA CLUSTER E&I/ BIO FARMA":
            return ["BIO FARMA CLUSTER E&I", "BIO FARMA"]

        cleaned_whole = _strip_legal_entity_with_punct(original)
        cleaned_whole = _FINAL_PUNCT_RE.sub(" ", cleaned_whole)
        cleaned_whole = re.sub(r"[./]", " ", cleaned_whole)
        cleaned_whole = _MULTISPACE_RE.sub(" ", cleaned_whole).strip()
        return [cleaned_whole.upper()] if cleaned_whole else [original.upper()]

    _PAREN_STASH.clear()
    t = _strip_bank_qq_segment(original)
    t = _AS_OWNER_NOISE_RE.sub(" / ", t)
    t = _strip_titles(t)

    if _has_bare_roman_numeral_slash_fragment(t):
        cleaned_whole = _strip_legal_entity_with_punct(t)
        cleaned_whole = _FINAL_PUNCT_RE.sub(" ", cleaned_whole)
        cleaned_whole = _CLEAN_WHOLE_DOT_RE.sub(" ", cleaned_whole)
        cleaned_whole = _CLEAN_WHOLE_SLASH_RE.sub("/", cleaned_whole)
        cleaned_whole = _MULTISPACE_RE.sub(" ", cleaned_whole).strip()
        cleaned_whole = _CLEAN_WHOLE_EDGE_RE.sub("", cleaned_whole)
        return [cleaned_whole.upper()] if cleaned_whole else [original.upper()]

    t = _protect_parens_content(t)
    raw_parts = _merge_bare_suffix_fragments(_ENTITY_SPLIT_RE.split(t))
    results = []

    for part in raw_parts:
        cleaned = _restore_protected_parens(part)
        cleaned = _strip_legal_entity_with_punct(cleaned)
        cleaned = _final_polish(cleaned)
        cleaned = _MULTISPACE_RE.sub(" ", cleaned).strip()

        if not cleaned or len(cleaned) < 2 or not _HAS_UPPER_RE.search(cleaned):
            continue
        if _BARE_CLIENT_FRAGMENT_RE.match(cleaned):
            continue

        pg_split = _split_repeated_pg_prefix(cleaned)
        if pg_split is not None:
            for sub in pg_split:
                if sub and sub not in results:
                    results.append(sub)
            continue

        if cleaned not in results:
            results.append(cleaned)

    if len(results) > MAX_BREAKDOWN_CODES: return [original]
    if not results: return [original.upper()]
    return results


# ==============================================================================
# 4. CLEANSING POLIS & SLIP
# ==============================================================================
def _strip_tba_noise(text: str) -> str:
    if _STANDALONE_KEEP_RE.match(text): return text
    if not re.search(r"\bTBA\b", text, RE_I): return text
    t = _TBA_WITH_SEPARATOR_RE.sub(" ", text)
    t = re.sub(r"^[/.\-\s]+|[/.\-\s]+$", "", t)
    t = re.sub(r"\s*/\s*/\s*", "/", t)
    t = _MULTISPACE_RE.sub(" ", t).strip()
    return t if t else text

def _strip_leading_p_label(text: str) -> str:
    return _LEADING_P_LABEL_RE.sub("", text)

def _strip_noise(text: str) -> str:
    t = text
    t = _ADMIN_STATUS_NOISE_RE.sub("", t)
    t = _PLUS_NOISE_TOKEN_RE.sub("", t)
    t = _SLASH_NOISE_TOKEN_RE.sub("", t)
    t = _DASH_NOISE_TOKEN_RE.sub("", t)
    t = re.sub(r"[/,]\s*(?=[/,]|$)", " ", t)
    t = re.sub(r"^[/,\s]+|[/,\s]+$", "", t)
    t = re.sub(r"\s*\+\s*$", "", t).strip()
    t = _MULTISPACE_RE.sub(" ", t).strip()
    return t

def _looks_irregular(text: str) -> bool:
    return not _DIGIT_SEARCH_RE.search(text)

def _has_real_code(text: str) -> bool:
    return bool(re.search(r"[A-Za-z]{1,6}\d{5,}", text))

def _strip_dash_short_suffix(text: str):
    match = _DASH_SHORT_SUFFIX_DISCARD_RE.match(text)
    if match: return match.group(1)
    return None

def _split_full_codes_by_slash_amp(text: str):
    if "/" not in text and "&" not in text: return None
    parts = [p.strip() for p in re.split(r"[/&]", text) if p.strip()]
    if len(parts) < 2: return None
    if not all(_FULL_CODE_TOKEN_RE.match(p) for p in parts): return None
    seen, result = set(), []
    for p in parts:
        if p not in seen:
            seen.add(p)
            result.append(p)
    return result if len(result) >= 2 else None

def _reuse_suffix_with_shared_trailing_path(text: str):
    if "+" not in text or "/" not in text: return None
    slash_idx = text.find("/")
    if slash_idx == -1: return None
    before_slash = text[:slash_idx]
    trailing_path = text[slash_idx:]

    if "+" not in before_slash: return None
    plus_parts = [p.strip() for p in before_slash.split("+") if p.strip()]
    if len(plus_parts) < 2: return None
    if not all(re.fullmatch(r"\d+", p) for p in plus_parts): return None

    anchor = plus_parts[0]
    codes = [anchor]
    for p in plus_parts[1:]:
        if len(p) >= len(anchor):
            reconstructed = p
        else:
            reconstructed = anchor[: -len(p)] + p
        codes.append(reconstructed)

    seen, result = set(), []
    for c in codes:
        full = f"{c}{trailing_path}"
        if full not in seen:
            seen.add(full)
            result.append(full)
    return result if len(result) >= 2 else None

def _split_full_codes_by_plus(text: str):
    if "+" not in text: return None
    parts = [p.strip() for p in text.split("+") if p.strip()]
    if len(parts) < 2: return None
    if not any("/" in p for p in parts): return None
    if not all(len(_NON_DIGIT_RE.sub("", p)) >= 4 for p in parts): return None
    seen, result = set(), []
    for p in parts:
        if p not in seen:
            seen.add(p)
            result.append(p)
    return result if len(result) >= 2 else None

def _reuse_suffix_breakdown(text: str):
    dash_prefix_match = re.match(r"^([A-Za-z]*\d{4,})\s*-\s*(\d{1,6})(.*)$", text)
    if dash_prefix_match:
        anchor_prefix_m = re.match(r"^([A-Za-z]*)(\d+)$", dash_prefix_match.group(1))
        anchor_prefix = anchor_prefix_m.group(1) if anchor_prefix_m else ""
        anchor_digits = anchor_prefix_m.group(2) if anchor_prefix_m else dash_prefix_match.group(1)
        first_suffix = dash_prefix_match.group(2)
        remainder = dash_prefix_match.group(3)
        rest_parts = [p.strip() for p in re.split(r"[+,_&\-]", remainder) if p.strip()]
        segments = [first_suffix] + rest_parts
        has_dash_base = True
    else:
        segments = [p.strip() for p in re.split(r"[+,_&]", text) if p.strip()]
        has_dash_base = False
        anchor_prefix, anchor_digits = "", ""

    if not segments: return None
    if not has_dash_base and len(segments) < 2: return None

    codes = []
    short_base_concat_mode = False
    if has_dash_base:
        if len(anchor_digits) < 10: short_base_concat_mode = True
        else: codes.append(f"{anchor_prefix}{anchor_digits}")

    for idx, p in enumerate(segments):
        m_full = _BASE_CODE_RE.match(p)
        if m_full and len(m_full.group(2)) >= 4 and not (anchor_digits and re.fullmatch(r"\d{1,6}", p) and len(p) < len(anchor_digits)):
            anchor_prefix, anchor_digits = m_full.group(1), m_full.group(2)
            codes.append(f"{anchor_prefix}{anchor_digits}")
            short_base_concat_mode = False
            continue
        if re.fullmatch(r"\d{1,6}", p) and anchor_digits:
            if short_base_concat_mode:
                reconstructed = anchor_digits + p
                anchor_digits = reconstructed
                short_base_concat_mode = False
                codes.append(f"{anchor_prefix}{reconstructed}")
                continue
            width = len(p)
            if width >= len(anchor_digits): reconstructed = p
            else: reconstructed = anchor_digits[: -width] + p
            codes.append(f"{anchor_prefix}{reconstructed}")
            continue
        return None

    seen, result = set(), []
    for c in codes:
        if c not in seen:
            seen.add(c)
            result.append(c)
    if not codes: return None
    return result

def _expand_sd_range(text: str):
    match = _SD_RANGE_RE.match(text)
    if not match: return None
    prefix, start_s, end_s = match.groups()
    if not (start_s.isdigit() and end_s.isdigit()): return None

    width = len(end_s)
    if width < len(start_s):
        base_prefix_digits = start_s[:-width]
        start_tail = int(start_s[-width:])
        end_n = int(end_s)
        if start_tail > end_n or (end_n - start_tail + 1) > MAX_BREAKDOWN_CODES:
            return "__KEEP_AS_IS__"
        prefix_clean = prefix.strip()
        joiner = "" if (prefix_clean and not prefix.endswith(" ")) else (" " if prefix_clean else "")
        codes = [
            f"{prefix_clean}{joiner}{base_prefix_digits}{n:0{width}d}".strip()
            for n in range(start_tail, end_n + 1)
        ]
        return codes

    start_n, end_n = int(start_s), int(end_s)
    if start_n > end_n or (end_n - start_n + 1) > MAX_BREAKDOWN_CODES:
        return "__KEEP_AS_IS__"
    prefix_clean = prefix.strip()
    codes = [f"{prefix_clean} {i:0{width}d}".strip() for i in range(start_n, end_n + 1)]
    return codes


def clean_split_code_impl(text: str):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return [], False
    original = text.strip()

    if _STANDALONE_KEEP_RE.match(original): return [original.upper()], True
    if _PENYELESAIAN_RE.search(original): return [original], True
    if _SEE_ATTACHMENT_RE.match(original): return [original], True
    if _YEAR_CHANGE_NOTE_RE.search(original): return [original], True
    if _P_REVS_RE.fullmatch(original): return [original], True

    if _MONTH_ALONE_RE.search(original) and not _has_real_code(original):
        return [original], True

    t = _strip_leading_p_label(original)
    t = _strip_tba_noise(t)
    if not t: return [original], True

    if _P_PLUS_SEQ_RE.fullmatch(t): return [t.strip()], True

    t = _strip_noise(t)
    if not t: return [original], True
    if _STANDALONE_KEEP_RE.match(t): return [t.upper()], True
    if _looks_irregular(t): return [t], True
    if _DIGIT_DOT_MATCH_RE.fullmatch(t) and "." in t: return [t], True

    sd_result = _expand_sd_range(t)
    if sd_result == "__KEEP_AS_IS__": return [t], True
    if sd_result is not None: return sd_result, False

    dash_strip_result = _strip_dash_short_suffix(t)
    if dash_strip_result is not None: return [dash_strip_result], False

    full_code_split_result = _split_full_codes_by_slash_amp(t)
    if full_code_split_result is not None:
        if len(full_code_split_result) > MAX_BREAKDOWN_CODES: return [t], True
        return full_code_split_result, False

    shared_trailing_result = _reuse_suffix_with_shared_trailing_path(t)
    if shared_trailing_result is not None:
        if len(shared_trailing_result) > MAX_BREAKDOWN_CODES: return [t], True
        return shared_trailing_result, True

    full_code_plus_result = _split_full_codes_by_plus(t)
    if full_code_plus_result is not None:
        if len(full_code_plus_result) > MAX_BREAKDOWN_CODES: return [t], True
        return full_code_plus_result, True

    if re.search(r"[+,_&]", t) or re.search(r"^[A-Za-z]*\d{4,}\s*-\s*\d{1,6}", t):
        reuse_result = _reuse_suffix_breakdown(t)
        if reuse_result is not None:
            if len(reuse_result) > MAX_BREAKDOWN_CODES: return [t], True
            return reuse_result, False

    t_final = _DOT_COMMA_TAIL_RE.sub("", t).strip()
    if not t_final: return [original], True
    if not _has_real_code(t_final): return [t_final], True
    return [t_final], False

def clean_split_code(text: str) -> list:
    results, _ = clean_split_code_impl(text)
    return [r.upper() if isinstance(r, str) else r for r in results]


# ==============================================================================
# 5. FUNGSI COMPACT
# ==============================================================================
def _compact_value(value: str) -> str:
    if not isinstance(value, str): return value
    return _NON_ALNUM_RE.sub("", value)

def _compact_split_code_results(results: list, is_kept_as_is: bool) -> list:
    if is_kept_as_is:
        return results
    return [_compact_value(v) for v in results]


# ==============================================================================
# 6. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_cedant(df_raw: pd.DataFrame, keyword: str, cedant_column: str) -> pd.DataFrame:
    if cedant_column not in df_raw.columns:
        raise ValueError(f"Kolom '{cedant_column}' tidak ditemukan. Tersedia: {list(df_raw.columns)}")
    mask = df_raw[cedant_column].astype(str).str.strip().str.upper() == keyword.upper()
    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant '{keyword}': {len(df_filtered)} baris ditemukan.")
    return df_filtered


# ==============================================================================
# 7. BUILD OUTPUT
# ==============================================================================
def _breakdown_all_rows(df: pd.DataFrame, col_insured: str, col_polis: str, col_slip: str):
    ins_arr = df[col_insured].to_numpy()
    pol_arr = df[col_polis].to_numpy()
    slp_arr = df[col_slip].to_numpy()

    insured_cln_all, polis_cln_all, slip_cln_all, cert_cln_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for ins_val, pol_val, slp_val in zip(ins_arr, pol_arr, slp_arr):
        ins_cln = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []

        if pd.notna(pol_val) and str(pol_val).strip():
            polis_base, cert_formatted = _extract_and_format_certificate(str(pol_val))
            pol_results, pol_kept = clean_split_code_impl(polis_base)

            pol_cln = _compact_split_code_results(
                [r.upper() if isinstance(r, str) else r for r in pol_results], pol_kept
            )
            cert_cln = [""] * len(pol_cln)
            if cert_formatted and pol_cln:
                cert_cln[-1] = cert_formatted
        else:
            pol_cln = []
            cert_cln = []

        if pd.notna(slp_val) and str(slp_val).strip():
            slp_results, slp_kept = clean_split_code_impl(str(slp_val))
            slp_cln = _compact_split_code_results(
                [r.upper() if isinstance(r, str) else r for r in slp_results], slp_kept
            )
        else:
            slp_cln = []

        insured_cln_all.append(ins_cln)
        polis_cln_all.append(pol_cln)
        slip_cln_all.append(slp_cln)
        cert_cln_all.append(cert_cln)

        max_ins, max_pol, max_slp = max(max_ins, len(ins_cln)), max(max_pol, len(pol_cln)), max(max_slp, len(slp_cln))

    return insured_cln_all, polis_cln_all, slip_cln_all, cert_cln_all, max_ins, max_pol, max_slp

def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_insured, col_polis, col_slip = INSURED_COLUMN, POLIS_COLUMN, SLIP_COLUMN
    for col, label in ((col_insured, "INSURED"), (col_polis, "POLIS"), (col_slip, "SLIP")):
        if col not in df_filtered.columns:
            raise ValueError(f"Kolom {label}='{col}' tidak ditemukan.")

    ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp = _breakdown_all_rows(
        df_filtered, col_insured, col_polis, col_slip
    )

    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)

    if BROKER_NAME_COLUMN in df_filtered.columns and CEDANT_COLUMN in df_filtered.columns:
        is_direct = df_filtered[BROKER_NAME_COLUMN].astype(str).str.strip().str.upper() == DIRECT_MARKER
        bp_series = df_filtered[CEDANT_COLUMN].where(is_direct, df_filtered[BROKER_NAME_COLUMN])
        business_partner_values = bp_series.astype(str).str.replace(r"^(\s*PT)\.\s*", r"\1 ", regex=True, flags=RE_I).values
    else:
        business_partner_values = None

    output_cols_data = {}
    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        if col == BROKER_NAME_COLUMN and business_partner_values is not None:
            output_cols_data["BUSINESS_PARTNER"] = business_partner_values

        if col == col_insured:
            for i in range(1, max_ins + 1):
                output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in ins_all]

        if col == col_polis:
            for i in range(1, max_pol + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in pol_all]
                output_cols_data[f"CERTIFICATE_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in cert_all]

        if col == col_slip:
            for i in range(1, max_slp + 1):
                output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slp_all]

    df_master = pd.DataFrame(output_cols_data)
    drop_cols = [c for c in df_master.columns if ("_CLN_" in c or c.startswith("CERTIFICATE_")) and df_master[c].astype(str).str.strip().eq("").all()]
    return df_master.drop(columns=drop_cols)


# ==============================================================================
# 8. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    workbook = load_workbook(output_path)
    worksheet = workbook.active
    text_col_indices = [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None: cell.value = str(cell.value)
    workbook.save(output_path)


# ==============================================================================
# 9. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print("=" * 70 + "\n CLEANSING DATA 1 - FACULTATIVE | CEDANT: JASA TANIA\n" + "=" * 70)
    input_path = INPUT_FILE

    if not os.path.exists(input_path):
        candidates = [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]
        if candidates:
            input_path = candidates[0]
            print(f"[AUTO-DETECT] Memakai file terdeteksi: '{input_path}'")
        else:
            raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    print(f"[INFO] Total baris source: {len(df_raw)}")

    df_jasa = filter_cedant(df_raw, CEDANT_FILTER, CEDANT_COLUMN)
    if df_jasa.empty:
        print("[WARNING] Tidak ada baris yang cocok dengan filter cedant. Proses dihentikan.")
        return

    df_hasil = build_output(df_jasa)
    save_with_text_format(df_hasil, OUTPUT_FILE)

    print("=" * 70 + f"\n[SUCCESS] Selesai. Total baris output: {len(df_hasil)}\n[SUCCESS] File hasil: '{OUTPUT_FILE}'\n" + "=" * 70)

if __name__ == "__main__":
    main()