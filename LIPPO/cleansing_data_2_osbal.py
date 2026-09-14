"""
================================================================================
 SCRIPT CLEANSING DATA 2 - OSBAL (CEDANT: PT LIPPO GENERAL INSURANCE)
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
INPUT_FILE = "2b. transaksi Osbal.xlsx"
SHEET_NAME = "Query result"
HEADER_ROW = 0
OUTPUT_FILE = "lippo_output_osbal_V2.xlsx"
CEDANT_FILTER = "LIPPO"

CEDANT_COLUMN = "CCOS_COMP_NAME"
INSURED_COLUMN = "FAC_INSURED"
POLIS_COLUMN = "FAC_POLICY_NO"
SLIP_COLUMN = "FAC_SLIP"

CLSDT_POLIS_COLUMN = "CLSDT_POLICY_NO"
CLSDT_SLIP_COLUMN = "CLSDT_SLIP_NO"
CLSDT_SERTF_COLUMN = "CLSDT_SERTF_NO"
CERTIFICATE_OUTPUT_COL = "CERTIFICATE"

MAX_BREAKDOWN_CODES = 5
CERT_DIGIT_LEN = 6
CERT_RANGE_JOIN_THRESHOLD = 3

TEXT_FORMAT_COLUMN_KEYWORDS = (
    "POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
    "FAC_POLICY_NO", "FAC_SLIP", "CLSDT", "CERTIFICATE",
)

# ==============================================================================
# 1. CLEANSING INSURED
# ==============================================================================
_BYPASS_ALL_RE = re.compile(r"(?i)\bPENYELESAIAN\s+HUTANG\s+PIUTANG\b")
_TRANSACTION_HEADER_RE = re.compile(r"(?i)^\s*(?:Pembayaran|Penagihan)?\s*(?:dan\s+Penagihan)?\s*Premi\s*(?:Fakultatif|SOA|Pensesian)?\s*(?:PAR\s*&?\s*EQ|BIT|GIT|CECR|QS)?\s*")
_TRUNCATE_TRIGGER_RE = re.compile(r"(?i)\bsubsidiar|\bassociat|\baffiliat|\bfiliated\b|related\s+compan|respective\s+right|\bincluding\s+an[yd]|\bindu?ding\s+any|as\s+the\s+owner|as\s+owners?\b|as\s+main\s+contractor|\(as\s+mortgagee\s+and\s+loss\s+payee\)")
_REMOVE_ONLY_RE = re.compile(r"(?i)\bas\s+principal\b|\bsebagai\s+principal\b|\bsebagai\s+kontraktor\b")
_ONLY_WORD_RE = re.compile(r"(?i)\bONLY\b")
_GENERIC_BOILERPLATE_TAIL_RE = re.compile(r"(?i)\bsub\s*-?\s*kontraktor\b.*$|\banak\s+perusahaan\b.*$|\bafiliasi\s+perusahaan\b.*$|\bdan\s+semua\b\s*[–-]?\s*$")
_TITLE_RE = re.compile(r"(?i)\b(?:S\.?H\.?|S\.?I\.?Kom\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b")
_LEGAL_ENTITY_RE = re.compile(r"(?i)\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO|Perseroan\s+Terbatas)\b")
_ENTITY_SPLIT_RE = re.compile(r"(?i)\b(?:QQ|Q\.Q\.)\b|\band\s*/\s*or\b|\bdan\s*/\s*atau\b|\bin\s+this\s+case\b|\bdalam\s+hal\s+ini\b|,|/")
_LEGAL_ENTITY_DASH_SIGNAL_RE = re.compile(r"(?i),\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO)\s*[\r\n]*\s*-\s*")
_GENERIC_DASH_SEPARATOR_RE = re.compile(r"\s*[\r\n]+\s*-\s*|\s+-\s+")
_TRAILING_PAREN_MERGE_RE = re.compile(r"(?i),?\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?)?\s*\(([^)]+)\)\s*$")
_INLINE_PAREN_MERGE_RE = re.compile(r"(?i),?\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?)?\s*\(([^)]+)\)")
_QQ_PAREN_UNWRAP_RE = re.compile(r"(?i)\bQQ\s*\(\s*([^)]*)\)")
_INSURED_DESCRIPTIVE_RE = re.compile(r"(?i)\bVARIOUS\s+INSUREDS?\b|\bACCEPTED\s+BY\b")
_SUSPENSE_WORD_RE = re.compile(r"(?i)\bSUSPENSE\b")

_MONTH_WORDS = ("JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI", "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER", "DES", "JANUARY", "FEBRUARY", "MARCH", "MAY", "JUNE", "JULY", "AUGUST", "OCTOBER", "DECEMBER", "JAN", "FEB", "MAR", "APR", "JUN", "JUL", "AUG", "SEP", "SEPT", "OCT", "NOV", "DEC")
_DATE_BATCH_TOKEN_RE = re.compile(r"(?i)\b(?:" + "|".join(_MONTH_WORDS) + r")\b|\b(?:19|20)\d{2}\b|\bBATCH\s*\d*\b")

_LIPPO_KNOWN_ENTITIES = ["MATAHARI PUTRA PRIMA", "MATAHARI BOSTON DRIGSTORE", "MATAHARI PUSAKA TAMA", "LIPPO GROUP"]
_LIPPO_KNOWN_ENTITIES_SORTED = sorted(_LIPPO_KNOWN_ENTITIES, key=len, reverse=True)
_KNOWN_ENTITY_RE = re.compile("|".join(re.escape(name) for name in _LIPPO_KNOWN_ENTITIES_SORTED))
_ENTITY_SLASH_NORMALIZE_PATTERNS = [(name, re.compile(r"(?i)\b" + r"(?:\s*/\s*|\s+)".join(re.escape(w) for w in name.split()) + r"\b")) for name in _LIPPO_KNOWN_ENTITIES_SORTED if len(name.split()) > 1]

_INSURED_DASH_SPLIT_EXCEPTIONS = ("EBDI VAUNDRI - OSCAR OMEGA", "ADHI - HUTAMA - NINDYA - ABIPRAYA KSO", "PEMERINTAH PROVINSI SULAWESI SELATAN - DINAS BINA MARGA DAN BINA KONSTRUKSI")
_INSURED_DASH_MERGE_SUBSTRINGS = ("PP - WASKITA - WIJAYA KARYA",)

def _unwrap_qq_paren(text: str) -> str:
    def _replace(match):
        inner = re.sub(r"\s+", " ", re.sub(r"[:;]", " ", match.group(1).strip())).strip()
        return f"QQ , {inner}" if inner else "QQ"
    return _QQ_PAREN_UNWRAP_RE.sub(_replace, text)

def _merge_trailing_paren(text: str) -> str:
    match = _TRAILING_PAREN_MERGE_RE.search(text)
    if not match: return text
    inner = match.group(1).strip()
    return text[: match.start()] + text[match.end():] if not inner or len(inner.split()) > 4 or _TRUNCATE_TRIGGER_RE.search(inner) else text[: match.start()] + " " + inner

def _merge_inline_paren(text: str) -> str:
    def _replace(match):
        inner = match.group(1).strip()
        return "" if not inner or len(inner.split()) > 4 or _TRUNCATE_TRIGGER_RE.search(inner) else f" {inner}"
    return _INLINE_PAREN_MERGE_RE.sub(_replace, text)

def _strip_date_batch_info(text: str) -> str: return _DATE_BATCH_TOKEN_RE.sub(" ", text)

def _truncate_boilerplate_tail(text: str) -> str:
    if match := _TRUNCATE_TRIGGER_RE.search(text): text = text[: match.start()]
    if match := _GENERIC_BOILERPLATE_TAIL_RE.search(text): text = text[: match.start()]
    return _REMOVE_ONLY_RE.sub(" ", text)

def _final_polish(text: str) -> str:
    t = re.sub(r"\s+", " ", re.sub(r"[-/().;:\"']", " ", text)).strip()
    prev = None
    while prev != t:
        prev, t = t, re.sub(r"(?i)^\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b\s*|\s*\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b$", "", t).strip()
    return t.upper()

def _apply_dash_merge_substrings(text: str) -> str:
    for sub in _INSURED_DASH_MERGE_SUBSTRINGS:
        if sub.upper() in text.upper():
            merged = re.sub(r"\s*-\s*", " ", sub)
            idx = text.upper().find(sub.upper())
            if idx != -1: text = text[:idx] + merged + text[idx + len(sub):]
    return text

def _try_split_insured_by_dash(text: str):
    stripped = text.strip()
    if _SUSPENSE_WORD_RE.search(stripped): return None
    for exception in _INSURED_DASH_SPLIT_EXCEPTIONS:
        if stripped.upper() == exception.upper():
            return [p.strip() for p in re.split(r"\s*-\s*", stripped) if p.strip()]
    return None

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"): return []
    original = text.strip()

    if _BYPASS_ALL_RE.search(original) or _INSURED_DESCRIPTIVE_RE.search(original) or _SUSPENSE_WORD_RE.search(original):
        return [original]

    if dash_split := _try_split_insured_by_dash(original):
        results = []
        for part in dash_split:
            cleaned = _final_polish(_strip_date_batch_info(_LEGAL_ENTITY_RE.sub(" ", _truncate_boilerplate_tail(part))))
            if cleaned and cleaned not in results: results.append(cleaned)
        return [original] if len(results) > MAX_BREAKDOWN_CODES or not results else results

    t = _apply_dash_merge_substrings(original)
    t = _unwrap_qq_paren(_ONLY_WORD_RE.sub(" ", t))
    t = _merge_trailing_paren(t)
    t = _merge_inline_paren(t)
    t = _TRANSACTION_HEADER_RE.sub("", t)
    t = _TITLE_RE.sub(" ", t)
    if _LEGAL_ENTITY_DASH_SIGNAL_RE.search(t): t = _GENERIC_DASH_SEPARATOR_RE.sub(", ", t)
    for canonical, pat in _ENTITY_SLASH_NORMALIZE_PATTERNS: t = pat.sub(canonical, t)
    t = _truncate_boilerplate_tail(t)

    results = []
    for part in _ENTITY_SPLIT_RE.split(t):
        cleaned = _final_polish(_strip_date_batch_info(_LEGAL_ENTITY_RE.sub(" ", _truncate_boilerplate_tail(part))))
        if not cleaned or len(cleaned) < 2 or not re.search(r"[A-Z]", cleaned): continue

        pos, n, known_found = 0, len(cleaned), []
        while pos < n:
            while pos < n and cleaned[pos] == " ": pos += 1
            if pos >= n or not (match := _KNOWN_ENTITY_RE.match(cleaned, pos)): break
            known_found.append(match.group(0))
            pos = match.end()

        if len(known_found) > 1:
            for name in known_found:
                if name not in results: results.append(name)
            continue
            
        if cleaned not in results: results.append(cleaned)

    return [original] if len(results) > MAX_BREAKDOWN_CODES else results

# ==============================================================================
# 2. CLEANSING POLIS & SLIP NO
# ==============================================================================
_MAIN_DIGITS_PATTERN = r"\d{12,13}"
_TOK_MAIN_RE = re.compile(rf"^{_MAIN_DIGITS_PATTERN}$")
_TOK_CERT_RE = re.compile(r"^\d{6}$")
_TOK_SUFFIX_REPLACE_RE = re.compile(r"^\d{1,11}$")
_TOK_ENDORSE_RE = re.compile(r"^\d{1,2}[/\x00]\d{1,2}$")
_TOK_PLACEHOLDER_RE = re.compile(r"(?i)^P\d{1,2}$")
_TOK_STATUS_RE = re.compile(r"(?i)^(?:New|End(?:orsement)?|Cancel\s*All|Adjustment|Cancellation|Reinstatement|TBA\.?|Attachment|DN)$")
_VARIOUS_WORD_RE = re.compile(r"(?i)\bVAR(?:IOUS)?\b")
_SD_WORD_RE = re.compile(r"(?i)\bSD\b")
_AMPERSAND_SUFFIX_AMBIGUOUS_RE = re.compile(r"\d+\s*&\s*\d+")
_CURRENCY_RE = re.compile(r"(?i)\b(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)\b")
_SEE_ATTACHMENT_RE = re.compile(r"(?i)\bSEE\s+ATTACHMENT\b")
_PENDING_SLIP_RE = re.compile(r"(?i)\bPENDING\s+SLIP\b")
_NO_LABEL_RE = re.compile(r"(?i)\bNO\s*[:.]")
_LETTER_REF_CODE_RE = re.compile(r"(?i)\b(?:\d{1,6}\s*/\s*)?[A-Z]{1,6}-[A-Z]{1,6}\s*/\s*[IVXLCM]{1,6}\s*/\s*\d{2,4}\b")
_BATCH_TOKEN_RE = re.compile(r"(?i)\bBATCH\s*\d*\b")
_BORDERO_WORD_RE = re.compile(r"(?i)\bBORDERO\b|\bBORD\b")
_POLIS_MONTH_RE = re.compile(r"(?i)\b(?:" + "|".join(_MONTH_WORDS) + r")\.?\s*(?:\d{2,4})?\b")
_POLIS_DATE_SLASH_RE = re.compile(r"\b\d{1,2}/\d{1,2}/\d{2,4}\b")
_POLIS_YEAR_RE = re.compile(r"\b(?:19|20)\d{2}\b")
_POLIS_CN_RE = re.compile(r"(?i)\bCN\b")
_POLIS_NARRATIVE_WORDS_RE = re.compile(r"(?i)\bSANY\b|\bNO\s+SLIP\b|\bNO\s+POLIS\b|\bPOLIS\s+VARIOUS\b")
_POLIS_STATUS_WORDS_RE = re.compile(r"(?i)\bCancel\s*All\b|\bNew\b|\bEnd(?:orsement)?\b|\bAdjustment\b|\bCancellation\b|\bReinstatement\b")
_EDGE_SEPARATOR_RE = re.compile(r"^[\s\-/,]+|[\s\-/,]+$")
_NARRATIVE_PREFIX_RE = re.compile(r"(?i)^\s*(?:JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\s+\d{4}\s*-\s*(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)?\s*/?\s*")
_MAIN_CODE_FIRST_RE = re.compile(_MAIN_DIGITS_PATTERN)
_STRONG_SPLIT_CAPTURE_RE = re.compile(r"(\+|-|[,/]|\s+)")
_ALPHA_WORD_RE = re.compile(r"[A-Za-z][A-Za-z0-9]*")
_ALLOWED_ALPHA_TOKENS = {"NEW", "END", "ENDORSEMENT", "CANCEL", "ALL", "ADJUSTMENT", "CANCELLATION", "REINSTATEMENT", "VARIOUS", "VAR", "SD", "ATTACHMENT", "TBA", "DN"}
_FALLBACK_NOISE_WORDS_RE = re.compile(r"(?i)\bTBA\.?\b|\bATTACHMENT\b|\bDN\b")
_BASE_SUFFIX_BASE_RE = re.compile(rf"^({_MAIN_DIGITS_PATTERN})-(\d{{1,11}})\s*-\s*({_MAIN_DIGITS_PATTERN})$")
_E_HASH_SUFFIX_RE = re.compile(rf"(?i)^({_MAIN_DIGITS_PATTERN})\s*-\s*E#\s*\d+\s*$")
_SHORT_BASE_MULTI_SUFFIX_RE = re.compile(r"^(\d{9,11})-(\d{1,6}(?:\s*[,+]\s*\d{1,6})+)\s*$")
_HAS_VALID_MAIN_DIGITS_RE = re.compile(_MAIN_DIGITS_PATTERN)

def _strip_polis_slip_noise(text: str) -> str:
    for pat in (_NARRATIVE_PREFIX_RE, _LETTER_REF_CODE_RE, _NO_LABEL_RE, _PENDING_SLIP_RE, _SEE_ATTACHMENT_RE, _POLIS_NARRATIVE_WORDS_RE, _POLIS_STATUS_WORDS_RE, _BATCH_TOKEN_RE, _BORDERO_WORD_RE, _POLIS_DATE_SLASH_RE, _POLIS_MONTH_RE, _POLIS_YEAR_RE, _POLIS_CN_RE, _CURRENCY_RE):
        text = pat.sub(" ", text)
    match = _MAIN_CODE_FIRST_RE.search(text)
    t = text[match.start():] if match and re.search(r"[A-Za-z]", text[:match.start()]) else text
    return _EDGE_SEPARATOR_RE.sub("", re.sub(r"\s+", " ", t).strip()).strip()

def _split_preserving_continuity(t: str):
    segments, is_soft_list, pending_soft = [], [], False
    for part in _STRONG_SPLIT_CAPTURE_RE.split(t):
        if part in ("+", "-"): pending_soft = True
        elif part is not None and part.strip() == "" and _STRONG_SPLIT_CAPTURE_RE.fullmatch(part or ""): continue
        elif part is not None and _STRONG_SPLIT_CAPTURE_RE.fullmatch(part): pending_soft = False
        elif (seg := (part or "").strip()):
            segments.append(seg)
            is_soft_list.append(pending_soft)
            pending_soft = False
    return segments, is_soft_list

def _clean_fallback_original(original: str) -> str:
    t = _FALLBACK_NOISE_WORDS_RE.sub("", original)
    return re.sub(r"\s{2,}", " ", re.sub(r"[+\-/,]\s*$", "", t.strip())).strip()

def _has_descriptive_narrative(text: str) -> bool:
    return any(not (_TOK_PLACEHOLDER_RE.match(w) or w.upper() in _ALLOWED_ALPHA_TOKENS) for w in _ALPHA_WORD_RE.findall(text))

def _expand_numeric_range(main_a: str, main_b: str):
    if len(main_a) != len(main_b) or not (main_a.isdigit() and main_b.isdigit()): return None
    diff_len = next((i for i in range(1, len(main_a) + 1) if main_a[-i] != main_b[-i]), 0)
    if diff_len == 0 or main_a[:-diff_len] != main_b[:-diff_len]: return None
    try: start, end = int(main_a[-diff_len:]), int(main_b[-diff_len:])
    except ValueError: return None
    return [f"{main_a[:-diff_len]}{str(n).zfill(diff_len)}" for n in range(start, end + 1)] if start <= end and (end - start + 1) <= MAX_BREAKDOWN_CODES else None

def _split_sd_range(text: str):
    if not _SD_WORD_RE.search(text): return None
    parts = _SD_WORD_RE.split(text, maxsplit=1)
    if len(parts) != 2: return "__FALLBACK__"

    left_all, right_all = re.findall(r"\d+", parts[0]), re.findall(r"\d+", parts[1])
    if not left_all or not right_all or not _TOK_MAIN_RE.match(left_all[-1]) or not _TOK_MAIN_RE.match(right_all[0]): return "__FALLBACK__"

    expanded = _expand_numeric_range(left_all[-1], right_all[0])
    return expanded + [t for t in right_all[1:] if _TOK_MAIN_RE.match(t) and t not in expanded] if expanded else "__FALLBACK__"

def _try_reconstruct_suffix23(tokens: list, tokens_is_soft: list):
    if not tokens or not _TOK_MAIN_RE.match(tokens[0]) or not all(_TOK_MAIN_RE.match(t) or _TOK_SUFFIX_REPLACE_RE.match(t) for t in tokens): return None
    if any(_TOK_SUFFIX_REPLACE_RE.match(tok) and not _TOK_MAIN_RE.match(tok) and not soft for tok, soft in zip(tokens, tokens_is_soft)): return None
    if not any(_TOK_SUFFIX_REPLACE_RE.match(tok) and not _TOK_MAIN_RE.match(tok) for tok in tokens): return None

    codes, current_base = [], None
    for tok in tokens:
        if _TOK_MAIN_RE.match(tok):
            current_base = tok
            if current_base not in codes: codes.append(current_base)
        elif current_base and (rec := current_base[:-len(tok)] + tok) not in codes:
            codes.append(rec)
    return codes

def clean_split_code(text: str) -> list:
    codes, _ = _clean_split_code_impl(text)
    return codes

def _clean_split_code_impl(text: str):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"): return [], False
    original = text.strip()

    if _BYPASS_ALL_RE.search(original): return [original], False

    if e_hash_match := _E_HASH_SUFFIX_RE.match(original): return [e_hash_match.group(1)], True

    if short_base_match := _SHORT_BASE_MULTI_SUFFIX_RE.match(original):
        base, suffix_blob = short_base_match.groups()
        suffixes = [s.strip() for s in re.split(r"[,+]", suffix_blob) if s.strip()]
        if suffixes:
            codes = [base]
            for suf in suffixes:
                if len(suf) < len(base): codes.append(base[: -len(suf)] + suf)
            deduped = []
            for c in codes:
                if c not in deduped: deduped.append(c)
            if len(deduped) <= MAX_BREAKDOWN_CODES: return deduped, True

    if not _HAS_VALID_MAIN_DIGITS_RE.search(original): return [original], False

    t = _strip_polis_slip_noise(original)
    if not t or _has_descriptive_narrative(t) or _AMPERSAND_SUFFIX_AMBIGUOUS_RE.search(t): return [original], False

    if (sd_res := _split_sd_range(t)) is not None: return ([original], False) if sd_res == "__FALLBACK__" else (sd_res, True)

    t = re.sub(r"(-\s*\d{1,2})\s*/\s*(\d{1,2})\b", lambda m: f"{m.group(1)}\x00{m.group(2)}", re.sub(r"[\u2012\u2013\u2014\u2015]", "-", _VARIOUS_WORD_RE.sub(" ", t)))
    strong_segments, segment_is_soft = _split_preserving_continuity(t)

    filtered_pairs = [(re.sub(r"[.]", "", seg).strip(), soft) for seg, soft in zip(strong_segments, segment_is_soft) if seg and not bool(_TOK_PLACEHOLDER_RE.match(re.sub(r"[.]", "", seg).strip()) or _TOK_STATUS_RE.match(re.sub(r"[.]", "", seg).strip()))]
    if not filtered_pairs: return [original], False

    tokens, tokens_is_soft = [p[0] for p in filtered_pairs], [p[1] for p in filtered_pairs]
    if (recon := _try_reconstruct_suffix23(tokens, tokens_is_soft)) is not None:
        return ([_clean_fallback_original(original)], False) if len(recon) > MAX_BREAKDOWN_CODES else (recon, True)

    codes, current_base = [], None
    for tok, soft in zip(tokens, tokens_is_soft):
        if current_base and soft and re.match(r"^\d{1,2}\x00\d{1,2}$", tok): continue
        if current_base and soft and _TOK_SUFFIX_REPLACE_RE.match(tok) and not _TOK_MAIN_RE.match(tok) and len(tok) < len(current_base):
            if (rec := current_base[: -len(tok)] + tok) not in codes: codes.append(rec)
            continue

        if "\x00" in tok and not _TOK_MAIN_RE.match(tok):
            if base_match := re.match(rf"^({_MAIN_DIGITS_PATTERN})-?(\d{{1,2}})\x00(\d{{1,2}})$", tok):
                current_base = base_match.group(1)
                if current_base not in codes: codes.append(current_base)
                continue

        anchor, *suffix_parts = [p.replace("\x00", "/") for p in tok.split("-")]
        if not _TOK_MAIN_RE.match(anchor): continue
        current_base = anchor

        while suffix_parts and _TOK_STATUS_RE.match(suffix_parts[-1]): suffix_parts.pop()
        for p in suffix_parts:
            if _TOK_ENDORSE_RE.match(p.replace("/", "\x00")) or re.match(r"^\d{1,2}/\d{1,2}$", p): continue
            if _TOK_SUFFIX_REPLACE_RE.match(p) and len(p) < len(anchor) and (rec := anchor[: -len(p)] + p) not in codes:
                codes.append(rec)
        if anchor not in codes: codes.append(anchor)

    return ([_clean_fallback_original(original)], False) if not codes or len(codes) > MAX_BREAKDOWN_CODES else (codes, True)

# ==============================================================================
# 2B. CLEANSING CERTIFICATE
# ==============================================================================
_CERT_SD_NORMALIZE_RE = re.compile(r"(?i)\bS\s*/\s*D\b")
_CERT_TOKEN_RE = re.compile(r"(?<![A-Za-z0-9])(\d{5,6})(?![A-Za-z0-9])")
_CERT_SD_PAIR_RE = re.compile(r"(?i)(\d{5,6})\s+SD\s+(\d{5,6})")

def _pad_cert_digits(digits: str) -> str: return digits.zfill(CERT_DIGIT_LEN)

def detect_polis_cert_groups(text: str):
    if not text: return None

    tokens = re.findall(r"(?i)\d+|SD", text)
    groups, current_polis, current_certs, current_base_cert, pending_sd = [], None, [], None, False

    for tok in tokens:
        if tok.upper() == "SD":
            pending_sd = True; continue

        n = len(tok)
        if 10 <= n <= 13:
            if current_polis is not None: groups.append((current_polis, current_certs))
            current_polis, current_certs, current_base_cert, pending_sd = tok, [], None, False
            continue

        if current_polis is None: continue

        if pending_sd:
            if current_base_cert is None: pending_sd = False; continue
            endpoint = _pad_cert_digits(tok) if n in (5, 6) else (current_base_cert[:-n] + tok if 1 <= n <= 4 else None)
            
            if endpoint is not None:
                try: start_n, end_n = int(current_base_cert), int(endpoint)
                except ValueError: start_n = end_n = None
                
                if start_n is not None and start_n <= end_n:
                    for x in range(start_n, end_n + 1):
                        c = _pad_cert_digits(str(x))
                        if c not in current_certs: current_certs.append(c)
                current_base_cert = endpoint
            pending_sd = False
        elif n in (5, 6):
            padded = _pad_cert_digits(tok)
            current_certs.append(padded)
            current_base_cert = padded
        elif 1 <= n <= 4 and current_base_cert is not None:
            rebuilt = current_base_cert[:-n] + tok
            current_certs.append(rebuilt)
            current_base_cert = rebuilt

    if current_polis is not None: groups.append((current_polis, current_certs))
    return groups if groups and any(certs for _, certs in groups) else None

def _extract_cert_numbers_from_raw(text: str):
    if not text: return []
    found, consumed_spans = [], []

    for m in _CERT_SD_PAIR_RE.finditer(text):
        if not re.search(r"[A-Za-z0-9]", text[:m.start(1)]): continue
        try: start_n, end_n = int(m.group(1)), int(m.group(2))
        except ValueError: continue
        
        if start_n <= end_n and (end_n - start_n + 1) <= MAX_BREAKDOWN_CODES:
            for n in range(start_n, end_n + 1): found.append(_pad_cert_digits(str(n)))
            consumed_spans.append((m.start(), m.end()))

    scan_text = text
    if consumed_spans:
        chars = list(text)
        for s, e in consumed_spans:
            for i in range(s, e): chars[i] = "#"
        scan_text = "".join(chars)

    for m in _CERT_TOKEN_RE.finditer(scan_text):
        if re.search(r"[A-Za-z0-9]", scan_text[:m.start()]):
            found.append(_pad_cert_digits(m.group(1)))

    return found

def _single_certificate_from_clsdt(clsdt_sertf_value) -> str:
    if clsdt_sertf_value is None or (isinstance(clsdt_sertf_value, float) and pd.isna(clsdt_sertf_value)): return ""
    text = str(clsdt_sertf_value).strip()
    if not text or text.lower() in ("nan", "none", "-"): return ""

    text = _CERT_SD_NORMALIZE_RE.sub("SD", text)
    _, is_polis_pattern = _clean_split_code_impl(text)
    if is_polis_pattern: return ""
    if re.fullmatch(r"\d{5,6}", text): return _pad_cert_digits(text)

    candidates = _extract_cert_numbers_from_raw(text)
    return candidates[0] if candidates else ""

def _join_certificates(cert_list: list) -> str:
    if not cert_list: return ""
    
    seen = []
    for c in cert_list:
        if c not in seen: seen.append(c)
    if len(seen) == 1: return seen[0]

    try: nums = [int(c) for c in seen]
    except ValueError: return ", ".join(seen)

    segments, current_seg = [], [seen[0]]
    for i in range(1, len(seen)):
        if nums[i] - nums[i - 1] == 1: current_seg.append(seen[i])
        else:
            segments.append(current_seg)
            current_seg = [seen[i]]
    segments.append(current_seg)

    parts = []
    for seg in segments:
        if len(seg) > CERT_RANGE_JOIN_THRESHOLD: parts.append(f"{seg[0]} SD {seg[-1]}")
        else: parts.extend(seg)
    return ", ".join(parts)

def build_certificates_for_row(clsdt_sertf_value, fac_policy_value, num_polis: int):
    slots = [""] * max(num_polis, 1)
    if clsdt_cert := _single_certificate_from_clsdt(clsdt_sertf_value):
        slots[0] = clsdt_cert
        return slots

    if _is_blank_value(fac_policy_value): return slots
    fac_text = _CERT_SD_NORMALIZE_RE.sub("SD", str(fac_policy_value).strip())

    if groups := detect_polis_cert_groups(fac_text):
        for i, (_, certs) in enumerate(groups):
            if i < len(slots): slots[i] = _join_certificates(certs)
        return slots

    if single_candidates := _extract_cert_numbers_from_raw(fac_text):
        slots[0] = _join_certificates(single_candidates)

    return slots

# ==============================================================================
# 3. PEMILIHAN SUMBER TERBAIK: CLSDT vs FAC (POLIS & SLIP)
# ==============================================================================
_PENYELESAIAN_RE = re.compile(r"(?i)\bPENYELESAIAN\s+(?:HUTANG|UTANG)\s+PIUTANG\b")
_TRAILING_CURRENCY_WORD_RE = re.compile(r"(?i)\s+(?:IDR|USD|SGD|EUR|JPY|GBP|AUD|MYR)\s*$")

def _is_blank_value(value) -> bool:
    if value is None or (isinstance(value, float) and pd.isna(value)): return True
    return str(value).strip().lower() in ("", "nan", "none")

def _is_valid_clean_result(codes: list) -> bool:
    return all(re.fullmatch(r"\d{9,13}", c) for c in codes) if codes else False

def _pick_best_source_codes(clsdt_value, fac_value):
    clsdt_codes = clean_split_code(str(clsdt_value)) if not _is_blank_value(clsdt_value) else []
    if _is_valid_clean_result(clsdt_codes): return clsdt_codes, "CLSDT"
    fac_codes = clean_split_code(str(fac_value)) if not _is_blank_value(fac_value) else []
    return (fac_codes, "FAC") if fac_codes else (clsdt_codes, "CLSDT")

def _resolve_polis_or_slip(clsdt_value, fac_value):
    fac_str = "" if _is_blank_value(fac_value) else str(fac_value).strip()
    if _PENYELESAIAN_RE.search(fac_str):
        clsdt_codes = clean_split_code(str(clsdt_value)) if not _is_blank_value(clsdt_value) else []
        if _is_valid_clean_result(clsdt_codes):
            keterangan = re.sub(r"\s{2,}", " ", _TRAILING_CURRENCY_WORD_RE.sub("", fac_str).strip())
            return [f"{code} {keterangan}" for code in clsdt_codes], False
        fac_codes = clean_split_code(fac_str)
        return fac_codes, _is_valid_clean_result(fac_codes)

    codes, _ = _pick_best_source_codes(clsdt_value, fac_value)
    return codes, _is_valid_clean_result(codes)

# ==============================================================================
# 4. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_cedant(df_raw: pd.DataFrame, keyword: str, cedant_column: str = "") -> pd.DataFrame:
    if cedant_column:
        if cedant_column not in df_raw.columns: raise ValueError(f"Kolom {cedant_column} tidak ada.")
        cedant_cols = [cedant_column]
    else:
        if not (cedant_cols := [c for c in df_raw.columns if "CEDANT" in str(c).upper()]): raise ValueError("Kolom CEDANT tidak ditemukan lewat auto-detect.")

    mask = pd.Series(False, index=df_raw.index)
    for col in cedant_cols: mask |= df_raw[col].astype(str).str.contains(keyword, case=False, na=False)
    
    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant '{keyword}': {len(df_filtered)} dari total {len(df_raw)} baris.")
    return df_filtered

# ==============================================================================
# 5. BUILD OUTPUT
# ==============================================================================
def _detect_key_columns(df: pd.DataFrame, insured_col: str, polis_col: str, slip_col: str):
    def _res(col: str, kw_or_fn, lbl: str):
        if col:
            if col not in df.columns: raise ValueError(f"{lbl}='{col}' tidak ada.")
            return col
        return next((c for c in df.columns if (kw_or_fn(str(c).upper()) if callable(kw_or_fn) else kw_or_fn in str(c).upper())), None)

    col_ins = _res(insured_col, "INSURED", "INSURED_COLUMN")
    col_pol = _res(polis_col, lambda c: c.strip() == "POLIS", "POLIS_COLUMN")
    col_slp = _res(slip_col, "SLIP", "SLIP_COLUMN")
    
    if None in (col_ins, col_pol, col_slp): raise ValueError("Kolom kunci tidak lengkap.")
    print(f"[INFO] Kolom terdeteksi -> INSURED='{col_ins}', POLIS='{col_pol}', SLIP='{col_slp}'")
    return col_ins, col_pol, col_slp

_NON_ALNUM_RE = re.compile(r"[^A-Za-z0-9]")
def _compact_list(values: list) -> list: return [_NON_ALNUM_RE.sub("", v) if isinstance(v, str) else v for v in values]

def _breakdown_all_rows(df: pd.DataFrame, col_ins: str, col_pol: str, col_slp: str, col_clsdt_pol: str = None, col_clsdt_slp: str = None, col_clsdt_sertf: str = None):
    ins_arr, pol_arr, slp_arr = df[col_ins].to_numpy(), df[col_pol].to_numpy(), df[col_slp].to_numpy()
    clsdt_pol_arr = df[col_clsdt_pol].to_numpy() if col_clsdt_pol and col_clsdt_pol in df.columns else [None] * len(df)
    clsdt_slp_arr = df[col_clsdt_slp].to_numpy() if col_clsdt_slp and col_clsdt_slp in df.columns else [None] * len(df)
    clsdt_sertf_arr = df[col_clsdt_sertf].to_numpy() if col_clsdt_sertf and col_clsdt_sertf in df.columns else [None] * len(df)

    ins_all, pol_all, slp_all, cert_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for ins_val, pol_val, slp_val, clsdt_pol_val, clsdt_slp_val, clsdt_sertf_val in zip(ins_arr, pol_arr, slp_arr, clsdt_pol_arr, clsdt_slp_arr, clsdt_sertf_arr):
        ins = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []
        raw_pol, raw_slp = pol_val if pd.notna(pol_val) else "", slp_val if pd.notna(slp_val) else ""

        clsdt_pol_is_usable = col_clsdt_pol and pd.notna(clsdt_pol_val) and _is_valid_clean_result(clean_split_code(str(clsdt_pol_val)) if not _is_blank_value(clsdt_pol_val) else [])
        polis_cert_groups = detect_polis_cert_groups(_CERT_SD_NORMALIZE_RE.sub("SD", str(raw_pol).strip())) if not clsdt_pol_is_usable and not _is_blank_value(raw_pol) else None

        if polis_cert_groups:
            pol = [p for p, _ in polis_cert_groups]
            pol_should_compact = True
            cert = [_join_certificates(certs) for _, certs in polis_cert_groups]
        else:
            if col_clsdt_pol and pd.notna(clsdt_pol_val):
                pol, pol_should_compact = _resolve_polis_or_slip(clsdt_pol_val, raw_pol)
            else:
                pol = clean_split_code(str(raw_pol).strip()) if not _is_blank_value(raw_pol) else []
                pol_should_compact = _is_valid_clean_result(pol)
            cert = None

        if col_clsdt_slp and pd.notna(clsdt_slp_val): slp, slp_should_compact = _resolve_polis_or_slip(clsdt_slp_val, raw_slp)
        else:
            slp = clean_split_code(str(raw_slp).strip()) if not _is_blank_value(raw_slp) else []
            slp_should_compact = _is_valid_clean_result(slp)

        if pol_should_compact: pol = _compact_list(pol)
        if slp_should_compact: slp = _compact_list(slp)

        if cert is None: cert = build_certificates_for_row(clsdt_sertf_val, raw_pol, num_polis=len(pol) if pol else 1)
        if len(cert) < len(pol): cert = cert + [""] * (len(pol) - len(cert))
        elif len(cert) > len(pol): cert = cert[: len(pol)] if pol else cert

        ins_all.append(ins)
        pol_all.append(pol)
        slp_all.append(slp)
        cert_all.append(cert)
        
        max_ins, max_pol, max_slp = max(max_ins, len(ins)), max(max_pol, len(pol)), max(max_slp, len(slp))
        
    return ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp

def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_ins, col_pol, col_slp = _detect_key_columns(df_filtered, INSURED_COLUMN, POLIS_COLUMN, SLIP_COLUMN)
    col_clsdt_pol = CLSDT_POLIS_COLUMN if CLSDT_POLIS_COLUMN in df_filtered.columns else None
    col_clsdt_slp = CLSDT_SLIP_COLUMN if CLSDT_SLIP_COLUMN in df_filtered.columns else None
    col_clsdt_sertf = CLSDT_SERTF_COLUMN if CLSDT_SERTF_COLUMN in df_filtered.columns else None

    ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp = _breakdown_all_rows(df_filtered, col_ins, col_pol, col_slp, col_clsdt_pol, col_clsdt_slp, col_clsdt_sertf)

    df_filtered = df_filtered.reset_index(drop=True)
    output_cols_data = {col: df_filtered[col].values for col in df_filtered.columns}

    for col in df_filtered.columns:
        if col == col_ins:
            for i in range(1, max_ins + 1): output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in ins_all]
        if col == col_pol:
            for i in range(1, max_pol + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in pol_all]
                output_cols_data[f"{CERTIFICATE_OUTPUT_COL}_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in cert_all]
        if col == col_slp:
            for i in range(1, max_slp + 1): output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slp_all]

    df_master = pd.DataFrame(output_cols_data)
    drop_cols = [c for c in df_master.columns if ("_CLN_" in c) and df_master[c].astype(str).str.strip().eq("").all()]
    return df_master.drop(columns=drop_cols)

# ==============================================================================
# 6. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    workbook = load_workbook(output_path)
    worksheet = workbook.active
    text_col_indices = [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None: cell.value = str(cell.value)
    workbook.save(output_path)

# ==============================================================================
# 7. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print(f"{'='*70}\n CLEANSING DATA 2 - OSBAL | CEDANT: PT LIPPO GENERAL INSURANCE\n{'='*70}")
    
    input_path = INPUT_FILE
    if not os.path.exists(input_path):
        candidates = [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]
        if not candidates: raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")
        input_path = candidates[0]
        print(f"[AUTO-DETECT] Memakai file terdeteksi: '{input_path}'")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    df_lippo = filter_cedant(df_raw, CEDANT_FILTER, CEDANT_COLUMN)

    if not df_lippo.empty:
        df_hasil = build_output(df_lippo)
        save_with_text_format(df_hasil, OUTPUT_FILE)
        print(f"{'='*70}\n[SUCCESS] Selesai. Output: '{OUTPUT_FILE}' ({len(df_hasil)} baris)\n{'='*70}")
    else:
        print("[WARNING] Tidak ada baris yang cocok dengan filter.")

if __name__ == "__main__":
    main()