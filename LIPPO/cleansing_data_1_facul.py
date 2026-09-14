"""
================================================================================
 SCRIPT CLEANSING DATA 1 - FACULTATIVE (KHUSUS CEDANT: PT LIPPO GENERAL INSURANCE)
================================================================================
Dipakai untuk breakdown & cleansing kolom INSURED, POLIS, dan SLIP NO supaya
bisa dipakai sebagai key matching antar database.
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
INPUT_FILE = "1b. Transaksi Facul.xlsx"
SHEET_NAME = "Query result"
HEADER_ROW = 0
OUTPUT_FILE = "lippo_output_facul_V3.xlsx"
CEDANT_FILTER = "LIPPO"

CEDANT_COLUMN = "COMP_NAME"
BROKER_NAME_COLUMN = "COMP_NAME.1"
DIRECT_MARKER = "DIRECT"
INSURED_COLUMN = "FAC_INSURED"
POLIS_COLUMN = "FAC_POLICY_NO"
SLIP_COLUMN = "FAC_SLIP"

MAX_BREAKDOWN_CODES = 5
TEXT_FORMAT_COLUMN_KEYWORDS = (
    "POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
    "FAC_POLICY_NO", "FAC_SLIP", "CERTIFICATE",
)

RE_I = re.IGNORECASE

# ==============================================================================
# PRE-COMPILED REGEXES & CONSTANTS (OPTIMIZATION)
# ==============================================================================
_TRANSACTION_HEADER_RE = re.compile(r"^\s*(?:Pembayaran|Penagihan)?\s*(?:dan\s+Penagihan)?\s*Premi\s*(?:Fakultatif|SOA|Pensesian)?\s*(?:PAR\s*&?\s*EQ|BIT|GIT|CECR|QS)?\s*", RE_I)
_TRUNCATE_TRIGGER_RE = re.compile(r"\bsubsidiar|\b(?:and\s*/\s*or|and|its|their|all)\s+associat(?!ion)|&\s*/\s*or\s+associat(?!ion)|\baffiliat|\bfiliated\b|related\s+compan|respective\s+right|\bincluding\s+an[yd]|\bindu?ding\s+any|as\s+the\s+owner|as\s+owners?\b|as\s+main\s+contractor|\(as\s+mortgagee\s+and\s+loss\s+payee\)", RE_I)
_REMOVE_ONLY_RE = re.compile(r"\bas\s+principal\b|\bsebagai\s+principal\b|\bsebagai\s+kontraktor\b", RE_I)
_GENERIC_BOILERPLATE_TAIL_RE = re.compile(r"\bsub\s*-?\s*kontraktor\b.*$|\banak\s+perusahaan\b.*$|\bafiliasi\s+perusahaan\b.*$|\bdan\s+semua\b\s*[–-]?\s*$", RE_I)
_TITLE_RE = re.compile(r"\b(?:S\.?H\.?|S\.?I\.?Kom\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b", RE_I)
_LEGAL_ENTITY_RE = re.compile(r"\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PP\.?|PERSERO|Perseroan\s+Terbatas|LTD\.?|PTE\.?)\b", RE_I)
_ENTITY_SPLIT_RE = re.compile(r"\b(?:QQ|Q\.Q\.)\b|\band\s*/\s*or\b|\bdan\s*/\s*atau\b|\bin\s+this\s+case\b|\bdalam\s+hal\s+ini\b|;|/|,(?!\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO|LTD\.?|PTE\.?)\b)", RE_I)
_INSURED_VARIOUS_RE = re.compile(r"\bVARIOUS\b", RE_I)
_INSURED_BATCH_NOISE_RE = re.compile(
    r"\b(?:JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|"
    r"JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:TEMBER)?|OKT(?:OBER)?|OCT(?:OBER)?|"
    r"NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?)?\s*\bBATCH\.?\s*\d*"
    r"(?:\s*(?:JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|"
    r"JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:TEMBER)?|OKT(?:OBER)?|OCT(?:OBER)?|"
    r"NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?)\b)?"
    r"|\d+(?=\s*(?:JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|"
    r"JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:TEMBER)?|OKT(?:OBER)?|OCT(?:OBER)?|"
    r"NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?)\b)", RE_I)
_STRAY_SYMBOL_RE = re.compile(r"[¿¡‽]")
_PAREN_GROUP_RE = re.compile(r"\([^()]*\)")
_PAREN_TOKEN_RE = re.compile(r"\uE100(\d+)\uE101")
_LEGAL_ENTITY_DASH_SIGNAL_RE = re.compile(r",\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO)\s*[\r\n]*\s*-\s*", RE_I)
_GENERIC_DASH_SEPARATOR_RE = re.compile(r"\s*[\r\n]+\s*-\s*|\s+-\s+")
_INSURED_DESCRIPTIVE_RE = re.compile(r"\bVARIOUS\s+INSUREDS?\b|\bACCEPTED\s+BY\b", RE_I)
_SUSPENSE_KEEP_AS_IS_RE = re.compile(r"\bsuspense\b|\bpenyelesaian\s+suspense\b|\bpenyelesaian\s+(?:h)?utang\s+piutang\b", RE_I)
_DANGLING_PATTERN_RE = re.compile(r"^\b(?:AS|AN|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b\s*|\s*\b(?:AS|AN|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b$", RE_I)

_LIPPO_KNOWN_ENTITIES = ["MATAHARI PUTRA PRIMA", "MATAHARI BOSTON DRIGSTORE", "MATAHARI PUSAKA TAMA", "MPP LIPPO GROUP", "LIPPO GROUP"]
_LIPPO_KNOWN_ENTITIES_SORTED = sorted(_LIPPO_KNOWN_ENTITIES, key=len, reverse=True)
_KNOWN_ENTITY_RE = re.compile("|".join(re.escape(name) for name in _LIPPO_KNOWN_ENTITIES_SORTED))
_ENTITY_SLASH_NORMALIZE_PATTERNS = [(name, re.compile(r"\b" + r"(?:\s*/\s*|\s+)".join(re.escape(w) for w in name.split()) + r"\b", RE_I)) for name in _LIPPO_KNOWN_ENTITIES_SORTED if len(name.split()) > 1]
_INSURED_MERGE_OVERRIDES = {"APARTEMEN EKSEKUTIF MENTENG,PERHIM.PENGH": ["APARTEMEN EKSEKUTIF MENTENG PERHIM PENGH"]}

_MAIN_LEN_MIN, _MAIN_LEN_MAX = 11, 16
_MAIN_DIGITS_PATTERN = rf"\d{{{_MAIN_LEN_MIN},{_MAIN_LEN_MAX}}}"
_TOK_MAIN_RE = re.compile(rf"^{_MAIN_DIGITS_PATTERN}$")
_TOK_INDEPENDENT_RE = re.compile(r"^\d{7,16}$")
_TOK_CERT_RE = re.compile(r"^\d{6}$")
_TOK_SUFFIX_REPLACE_RE = re.compile(r"^\d{1,11}$")
_TOK_PLACEHOLDER_RE = re.compile(r"^P\d{1,2}$", RE_I)
_TOK_STATUS_RE = re.compile(r"^(?:New|End(?:orsement)?|Cancel\s*All|Adjustment|Cancellation|Reinstatement)$", RE_I)
_VARIOUS_WORD_RE = re.compile(r"\bVAR(?:I(?:O(?:U(?:S)?)?)?)?\b", RE_I)
_SD_MARKER_RE = re.compile(r"\bS\s*/\s*D\b|\bSD\b", RE_I)
_PURE_DECIMAL_RE = re.compile(r"^\d+(?:\.\d+)+$")
_DASH_AMP_AMBIGUOUS_RE = re.compile(rf"{_MAIN_DIGITS_PATTERN}\s*-\s*\d+\s*&\s*\d+")
_ALPHA_WORD_RE = re.compile(r"[A-Za-z][A-Za-z0-9]*")
_ALLOWED_ALPHA_TOKENS = {"NEW", "END", "ENDORSEMENT", "CANCEL", "ALL", "ADJUSTMENT", "CANCELLATION", "REINSTATEMENT", "VARIOUS", "VAR", "SD", "S", "D"}

_ATTACHMENT_NOISE_RE = re.compile(r"\bSEE\s+ATTACHMENT\b|\bATTACHMENT\b", RE_I)
_MONTH_YEAR_NOISE_RE = re.compile(
    r"\b(?:JAN(?:UARI|UARY)?|FEB(?:RUARI|RUARY)?|MAR(?:ET|CH)?|APR(?:IL)?|MEI|MAY|JUN(?:I|E)?|"
    r"JUL(?:I|Y)?|AGU(?:STUS)?|AUG(?:UST)?|SEP(?:T(?:E(?:M(?:B(?:E?R)?)?)?)?)?|SEPTEMEBR|OKT(?:OBER)?|OCT(?:OBER)?|"
    r"NOV(?:EMBER)?|DES(?:EMBER)?|DEC(?:EMBER)?)\s*\.?\s*\d{2,4}\b", RE_I)
_CURRENCY_NOISE_RE = re.compile(r"\b(?:USD|IDR|SGD|EUR|CNY|SGD|GBP|JPY|AUD|HKD)\b", RE_I)
_UNIT_NOISE_RE = re.compile(r"\d*\s*UNIT\b", RE_I)
_STRAY_QUOTE_NOISE_RE = re.compile(r"['\u2018\u2019\u201c\u201d]")
_ADMIN_LABEL_NOISE_RE = re.compile(
    r"\bDEKL\.?\b|\b\d+\s*SLIP\s+DIJADIKAN\s+\d+\b|\bID\s*NO\.?\s*[A-Z0-9]*\b"
    r"|(?:\b\d{1,6}\s*/?\s*)?\bSHIPMENT\s*\d*\b|(?:\b\d{1,6}\s*/?\s*)?\bBORDERO\b"
    r"|(?:\b\d{1,6}\s*/?\s*)?\bPENDING\s+SLIP\b|(?:\b\d{1,6}\s*/?\s*)?\bBATCH\s*\d*\b", RE_I)
_DOC_REF_DATE_NOISE_RE = re.compile(r"(?:\b\d{1,6}\s*/\s*)?\b(?:CN|DN)\s*/\s*\d{1,2}\s*/\s*\d{1,2}\s*/\s*\d{2,4}(?:\s*/\s*\d{1,6}\b)?", RE_I)
_ENDORSE_HASH_NOISE_RE = re.compile(r"\bE\s*#\s*\d+\b", RE_I)
_HASH_CODE_NOISE_RE = re.compile(r"#\s*[A-Z]{1,3}\d{0,3}\b", RE_I)
_TREATY_REF_CODE_NOISE_RE = re.compile(
    rf"\b\d{{1,6}}\s*/\s*[A-Z]+-[A-Z]+(?:\s*/\s*(?:X{{0,3}}(?:IX|IV|V?I{{0,3}}))\b)?(?:\s*/\s*\d{{4}}\b)?"
    rf"|\b[A-Z]+-[A-Z]+(?:\s*/\s*(?:X{{0,3}}(?:IX|IV|V?I{{0,3}}))\b)?(?:\s*/\s*\d{{4}}\b)?", RE_I)
_MAIN_SPACE_SUFFIX_RE = re.compile(rf"^({_MAIN_DIGITS_PATTERN})\s+(\d{{1,6}})\s*(?:/\s*VAR[A-Z]*\s*)?$", RE_I)
_TBA_NAME_ONLY_RE = re.compile(r"^TBA\s*/\s*([A-Za-z][A-Za-z ,.]*[A-Za-z.])$", RE_I)
_TBA_RESULT_PT_SUFFIX_RE = re.compile(r"\s*,?\s*P\.?T\.?$", RE_I)

_PP_MAIN_TOKEN_RE = re.compile(rf"^{_MAIN_DIGITS_PATTERN}$")
_PP_NONCERT_SUFFIX_RE = re.compile(r"^\d{1,11}$")
_CERT_SD_NORMALIZE_RE = re.compile(r"S\s*/\s*D", RE_I)

_POLIS_SLIP_OVERRIDES = {
    "APRIL 2026 - 0015/LI-RBU/V/2026 / 1112302600206,1112302600215,1112302600218,1112302600221.": [
        "1112302600206", "1112302600215", "1112302600218", "1112302600221",
    ],
}

_POLICY_CERT_HARDCASE_FALLBACK = {
    "1303121900052-182 S/D 189 / 1303152000009 01S/D10",
    "1303121900052-198S/D213 / 1303152000009-19S/D35",
    "1303121900052 -191 S/D197 / 1303152000009 -22S/D30",
    "1303121900052-190/'1303152000009-11S/D21",
    "1903021900688-013S/D030 1903021900689-622S/D643",
    "1903021900688-015s/d050 /1903021900689-644s/d698",
    "1903021900688-001S/D0012 / 1903021900689-596S/D620",
    "120110180002-001-006-003-003-005004-006",
    "15011011004041-018135-026215-003072-017133-019134",
    "15011015003-029002028002010003011",
    "150110150026-215-004041018135019134017133003072",
}

_POLICY_CERT_HARDCASE_RESULT = {
    "1102221600017,1102281000725,1202281700127,13022217": (
        ["1102221600017", "1102281000725", "1302221700127"], ["", "", ""]
    ),
    "1102211500160,1102281800059,1102281900090,12022119": (
        ["1102211500160", "1102281800059", "1202211900090"], ["", "", ""]
    ),
    "1801092400272-1801052400198-1801092400277-18010524": (
        ["1801092400272", "1801052400198", "1801052400277"], ["", "", ""]
    ),
    "15011011007-044005042-003040-": (
        ["15011011007", "15044005042", "15044003040"], ["", "", ""]
    ),
}

_PAREN_STASH = {}


# ==============================================================================
# 1. CLEANSING INSURED
# ==============================================================================
def _truncate_from_first_trigger(text: str) -> str:
    return text[: match.start()] if (match := _TRUNCATE_TRIGGER_RE.search(text)) else text

def _strip_insured_batch_noise(text: str) -> str: return _INSURED_BATCH_NOISE_RE.sub(" ", text)
def _strip_stray_symbols(text: str) -> str: return _STRAY_SYMBOL_RE.sub(" ", text)
def _normalize_dash_entity_list(text: str) -> str: return _GENERIC_DASH_SEPARATOR_RE.sub(", ", text) if _LEGAL_ENTITY_DASH_SIGNAL_RE.search(text) else text
def _strip_transaction_header(text: str) -> str: return _TRANSACTION_HEADER_RE.sub("", text)
def _strip_titles(text: str) -> str: return _TITLE_RE.sub(" ", text)
def _strip_legal_entity(text: str) -> str: return _LEGAL_ENTITY_RE.sub(" ", text)

def _protect_parens_content(text: str) -> str:
    def _mask(match):
        token = f"\uE100{len(_PAREN_STASH)}\uE101"
        _PAREN_STASH[token] = match.group(0)
        return token
    return _PAREN_GROUP_RE.sub(_mask, text)

def _restore_protected_chars(text: str) -> str:
    return text.replace("\uE0F0", "/").replace("\uE0F1", ",").replace("\uE0F2", ";")

def _restore_protected_parens(text: str) -> str:
    def _unmask(match): return _PAREN_STASH.get(f"\uE100{match.group(1)}\uE101", match.group(0))
    return _PAREN_TOKEN_RE.sub(_unmask, text)

def _truncate_boilerplate_tail(text: str) -> str:
    text = _truncate_from_first_trigger(text)
    if match_generic := _GENERIC_BOILERPLATE_TAIL_RE.search(text): text = text[: match_generic.start()]
    return _REMOVE_ONLY_RE.sub(" ", text)

def _final_polish(text: str) -> str:
    t = re.sub(r"[-/]", " ", re.sub(r"[,;:\"'.]", " ", text))
    t = re.sub(r"\s+", " ", re.sub(r"\s+\)", ")", re.sub(r"\(\s+", "(", t))).strip()
    previous = None
    while previous != t:
        previous, t = t, _DANGLING_PATTERN_RE.sub("", t).strip()
    return t.upper()

def _normalize_entity_slash_variants(text: str) -> str:
    for canonical, pattern in _ENTITY_SLASH_NORMALIZE_PATTERNS: text = pattern.sub(canonical, text)
    return text

def _split_by_known_entities(text: str):
    pos, n, found = 0, len(text), []
    while pos < n:
        while pos < n and text[pos] == " ": pos += 1
        if pos >= n: break
        if not (match := _KNOWN_ENTITY_RE.match(text, pos)): return None
        found.append(match.group(0))
        pos = match.end()
    return found if len(found) > 1 else None

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"): return []
    original = text.strip()

    if original in _INSURED_MERGE_OVERRIDES: return _INSURED_MERGE_OVERRIDES[original]
    if _INSURED_DESCRIPTIVE_RE.search(original) or _SUSPENSE_KEEP_AS_IS_RE.search(original): return [original]

    _PAREN_STASH.clear()
    t = _protect_parens_content(_truncate_boilerplate_tail(_normalize_entity_slash_variants(
        _normalize_dash_entity_list(_INSURED_VARIOUS_RE.sub(" ", _strip_insured_batch_noise(
            _strip_stray_symbols(_strip_titles(_strip_transaction_header(original)))
        )))
    )))

    results = []
    for part in _ENTITY_SPLIT_RE.split(t):
        cleaned = re.sub(r"\s+", " ", _restore_protected_parens(_final_polish(_strip_legal_entity(_truncate_boilerplate_tail(_restore_protected_chars(part)))))).strip()
        if not cleaned or len(cleaned) < 2 or not re.search(r"[A-Z]", cleaned): continue

        if known_split := _split_by_known_entities(cleaned):
            for name in known_split:
                if name not in results: results.append(name)
            continue
        if cleaned not in results: results.append(cleaned)

    return [original] if len(results) > MAX_BREAKDOWN_CODES else results


# ==============================================================================
# 2. CLEANSING POLIS & SLIP NO (FACULTATIVE)
# ==============================================================================
def _strip_known_noise_phrases(text: str) -> str:
    text = _STRAY_QUOTE_NOISE_RE.sub(" ", text)
    text = _PAREN_GROUP_RE.sub(" ", text)
    text = _ATTACHMENT_NOISE_RE.sub(" ", text)
    text = _DOC_REF_DATE_NOISE_RE.sub(" ", text)
    text = _TREATY_REF_CODE_NOISE_RE.sub(" ", text)
    text = _MONTH_YEAR_NOISE_RE.sub(" ", text)
    text = _CURRENCY_NOISE_RE.sub(" ", text)
    text = _UNIT_NOISE_RE.sub(" ", text)
    text = _ADMIN_LABEL_NOISE_RE.sub(" ", text)
    text = _ENDORSE_HASH_NOISE_RE.sub(" ", text)
    text = _HASH_CODE_NOISE_RE.sub(" ", text)
    return re.sub(r"\s+", " ", re.sub(r"[/\-]\s*(?=[/\-]|$)", " ", text)).strip()

def _strip_tba_prefix_if_name(text: str) -> str:
    return _TBA_RESULT_PT_SUFFIX_RE.sub("", match.group(1)).strip() if (match := _TBA_NAME_ONLY_RE.match(text)) else text

def _has_descriptive_narrative(text: str) -> bool:
    return any(not (_TOK_PLACEHOLDER_RE.match(word) or word.upper() in _ALLOWED_ALPHA_TOKENS) for word in _ALPHA_WORD_RE.findall(text))

def _strip_various(text: str) -> str: return _VARIOUS_WORD_RE.sub(" ", text)

def _expand_numeric_range(start_str: str, end_str: str, max_count: int = MAX_BREAKDOWN_CODES):
    if len(start_str) != len(end_str) or not (start_str.isdigit() and end_str.isdigit()): return None
    diff_len = next((i for i in range(1, len(start_str) + 1) if start_str[-i] != end_str[-i]), 0)
    if diff_len == 0 or start_str[:-diff_len] != end_str[:-diff_len]: return None
    try: start_n, end_n = int(start_str[-diff_len:]), int(end_str[-diff_len:])
    except ValueError: return None
    if start_n > end_n or (end_n - start_n + 1) > max_count: return None
    return [f"{start_str[:-diff_len]}{str(n).zfill(diff_len)}" for n in range(start_n, end_n + 1)]

def _expand_suffix_range(anchor: str, suffix_a: str, suffix_b: str):
    width = max(len(suffix_a), len(suffix_b))
    if width >= len(anchor) or not (suffix_a.isdigit() and suffix_b.isdigit()): return None
    start_n, end_n = int(suffix_a), int(suffix_b)
    return [f"{anchor[:-width]}{str(n).zfill(width)}" for n in range(start_n, end_n + 1)] if start_n <= end_n and (end_n - start_n + 1) <= MAX_BREAKDOWN_CODES else None

def _split_sd_range(text: str):
    if not (markers := list(_SD_MARKER_RE.finditer(text))) or len(markers) > 1: return "__FALLBACK__" if markers else None
    runs_left, runs_right = re.findall(r"\d+", text[: markers[0].start()]), re.findall(r"\d+", text[markers[0].end():])
    if not runs_left or not runs_right: return "__FALLBACK__"
    last_left, first_right = runs_left[-1], runs_right[0]
    
    if not (last_left_is_main := _TOK_MAIN_RE.match(last_left)) and len(runs_left) >= 2 and _TOK_MAIN_RE.match(runs_left[-2]):
        expanded = _expand_suffix_range(runs_left[-2], last_left, first_right)
    elif last_left_is_main and _TOK_MAIN_RE.match(first_right):
        expanded = _expand_numeric_range(last_left, first_right)
    else: return "__FALLBACK__"
    return expanded if expanded is not None else "__FALLBACK__"

def clean_split_code(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"): return []
    original = text.strip()

    if _SUSPENSE_KEEP_AS_IS_RE.search(original): return [original]
    if (tba_stripped := _strip_tba_prefix_if_name(original)) != original: return [tba_stripped]
    if space_concat_match := _MAIN_SPACE_SUFFIX_RE.match(original): return [space_concat_match.group(1) + space_concat_match.group(2)]
    if original in _POLIS_SLIP_OVERRIDES: return _POLIS_SLIP_OVERRIDES[original]

    t = _strip_known_noise_phrases(original)
    if _has_descriptive_narrative(t) or _PURE_DECIMAL_RE.match(t) or _DASH_AMP_AMBIGUOUS_RE.search(t): return [original]
    if (sd_result := _split_sd_range(t)) == "__FALLBACK__": return [original]
    if sd_result is not None: return sd_result

    t = re.sub(r"[\u2012\u2013\u2014\u2015]", "-", _strip_various(t))
    _had_slash_context, strong_segments, pending_delim = "/" in t, [], None

    for chunk in [c for c in re.split(r"(\s{2,}|[,+/;&])", t) if c != ""]:
        if re.fullmatch(r"\s{2,}", chunk) or chunk in (",", "+", "/", ";", "&"): pending_delim = "," if re.fullmatch(r"\s{2,}", chunk) else chunk; continue
        if seg := chunk.strip(): strong_segments.append((seg, pending_delim))
        pending_delim = None

    raw_tokens = []
    for seg, delim in strong_segments:
        sub_parts = [p.strip() for p in re.sub(r"[.]", "", seg).strip().split("-") if p.strip()]
        if len(sub_parts) <= 1:
            if sub_parts: raw_tokens.append((sub_parts[0], delim))
            continue
        buffer, buffer_delim = [sub_parts[0]], delim
        for p in sub_parts[1:]:
            if _TOK_MAIN_RE.match(p) or _TOK_PLACEHOLDER_RE.match(p):
                raw_tokens.append(("-".join(buffer), buffer_delim))
                buffer, buffer_delim = [p], "-"
            else: buffer.append(p)
        raw_tokens.append(("-".join(buffer), buffer_delim))

    tokens = [(tok, d) for tok, d in raw_tokens if not _TOK_PLACEHOLDER_RE.match(tok)]
    if not tokens: return [original]

    codes, current_base, current_base_is_strict, n_tokens = [], None, False, len(tokens)
    for i, (tok, delim) in enumerate(tokens):
        effective_delim = delim if delim is not None else (tokens[i + 1][1] if i + 1 < n_tokens else ("/" if _had_slash_context else None))
        if current_base is not None and current_base_is_strict and delim in ("+", "&", "/", ",") and _TOK_SUFFIX_REPLACE_RE.match(tok) and not _TOK_MAIN_RE.match(tok) and len(tok) < len(current_base):
            if (reconstructed := current_base[: -len(tok)] + tok) not in codes: codes.append(reconstructed)
            continue
        
        parts = tok.split("-")
        anchor, suffix_parts = parts[0], parts[1:]
        if _TOK_MAIN_RE.match(anchor): is_strict = True
        elif not suffix_parts and effective_delim == "/" and _TOK_INDEPENDENT_RE.match(anchor): is_strict = False
        else:
            if anchor.isdigit(): return [original]
            continue

        current_base, current_base_is_strict = anchor, is_strict
        while suffix_parts and _TOK_STATUS_RE.match(suffix_parts[-1]): suffix_parts.pop()
        append_suffix_parts = []

        for p in suffix_parts:
            if _TOK_CERT_RE.match(p): append_suffix_parts.append(p)
            elif _TOK_SUFFIX_REPLACE_RE.match(p) and len(p) < len(anchor):
                if (reconstructed := anchor[: -len(p)] + p) not in codes: codes.append(reconstructed)
        
        if append_suffix_parts and (combined := f"{anchor}-" + "-".join(append_suffix_parts)) not in codes: codes.append(combined)
        if anchor not in codes: codes.append(anchor)

    return [original] if not codes or len(codes) > MAX_BREAKDOWN_CODES else codes


# ==============================================================================
# 2B. PROSES TERPADU POLIS + CERTIFICATE (REVISI V4)
# ==============================================================================
def _cert_pad6(raw_suffix: str) -> str: return raw_suffix[-6:] if len(raw_suffix) > 6 else raw_suffix.zfill(6)
def _is_raw_cert_suffix(raw_suffix: str) -> bool: return len(raw_suffix) in (5, 6) and raw_suffix.isdigit() and raw_suffix[0] == "0"

def _format_cert_list(cert_nums: list, is_sd: bool) -> str:
    if not cert_nums: return ""
    seen = []
    for c in cert_nums:
        if c not in seen: seen.append(c)
    if len(seen) == 1: return seen[0]

    nums_int = [int(c) for c in seen]
    is_sequential = all(nums_int[i + 1] - nums_int[i] == (1 if nums_int[-1] >= nums_int[0] else -1) for i in range(len(nums_int) - 1))
    return f"{seen[0]} SD {seen[-1]}" if is_sd or (is_sequential and len(seen) > 3) else ", ".join(seen)

def process_policy_certificate(text: str):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"): return [], []
    original = text.strip()

    if original in _POLICY_CERT_HARDCASE_FALLBACK or _SUSPENSE_KEEP_AS_IS_RE.search(original): return [""], [original]
    if original in _POLICY_CERT_HARDCASE_RESULT:
        pol, cert = _POLICY_CERT_HARDCASE_RESULT[original]
        return cert, pol
    if (tba_stripped := _strip_tba_prefix_if_name(original)) != original: return [""], [tba_stripped]
    if original in _POLIS_SLIP_OVERRIDES: return [""] * len(_POLIS_SLIP_OVERRIDES[original]), list(_POLIS_SLIP_OVERRIDES[original])

    t = re.sub(r"\bCERTIFICATE\b\s*:?\s*", " ", re.sub(r"\.{2,}", " ", re.sub(r"\bDLL\b\.?", " ", re.sub(r"\bVARIUOS\b", " VARIOUS ", re.sub(r"\bSD\b", " SD ", _CERT_SD_NORMALIZE_RE.sub(" SD ", original), flags=RE_I), flags=RE_I), flags=RE_I)))
    t_noise = _strip_known_noise_phrases(t)
    if _has_descriptive_narrative(t_noise) or _PURE_DECIMAL_RE.match(t_noise): return [""], [original]
    t = re.sub(r"[\u2012\u2013\u2014\u2015]", "-", _strip_various(t_noise))

    if "+" in t and "," not in t:
        plus_chunks = [c.strip() for c in t.split("+") if c.strip()]
        if plus_chunks and _PP_MAIN_TOKEN_RE.match(plus_chunks[0]) and len(plus_chunks[1:]) > 1 and all(re.fullmatch(r"\d+", c) for c in plus_chunks[1:]) and not any(_PP_MAIN_TOKEN_RE.match(c) for c in plus_chunks[1:]): return [""], [original]

    if not (top_chunks := [c.strip() for c in re.split(r"[+/:]", t) if c.strip()]): return [""], [original]

    pairs, current_main, fallback = [], None, False
    for chunk in top_chunks:
        expanded = []
        for p in re.split(r"(\bSD\b|&|-|,)", chunk.strip().strip("/").strip()):
            expanded.append(p) if p in ("SD", "&", "-", ",") else expanded.extend(p.split())
        
        if not (pieces := [p.strip() for p in expanded if p.strip()]): continue

        i, local_main, pending_op, active_idx, cert_target_idx = 0, None, None, None, None
        while i < len(pieces):
            piece = pieces[i]
            if piece in ("-", "&", "SD", ","): pending_op = piece; i += 1; continue
            
            if _PP_MAIN_TOKEN_RE.match(piece):
                local_main, current_main = piece, piece
                if (existing_idx := next((idx for idx, p in enumerate(pairs) if p[0] == piece), None)) is not None: active_idx = existing_idx
                else: pairs.append([piece, [], False]); active_idx = len(pairs) - 1
                cert_target_idx, pending_op = active_idx, None
                i += 1
                continue

            if piece.isdigit():
                anchor = local_main or current_main
                if active_idx is None and pairs: active_idx = len(pairs) - 1
                if cert_target_idx is None and pairs: cert_target_idx = active_idx if active_idx is not None else len(pairs) - 1

                if pending_op == "SD":
                    if cert_target_idx is not None: pairs[cert_target_idx][1].append(_cert_pad6(piece)); pairs[cert_target_idx][2] = True
                    pending_op = None; i += 1; continue

                if pending_op in ("-", "&", ",", None) and anchor:
                    padded, already_has_cert = _cert_pad6(piece), cert_target_idx is not None and bool(pairs[cert_target_idx][1])
                    if _is_raw_cert_suffix(piece) or (len(piece) == 7 and padded[0] == "0") or (already_has_cert and pending_op in ("-", "&", ",")):
                        if cert_target_idx is not None: pairs[cert_target_idx][1].append(padded)
                        pending_op = None; i += 1; continue

                    if pending_op == "-" and anchor and _PP_NONCERT_SUFFIX_RE.match(piece) and len(piece) < len(anchor) and not already_has_cert:
                        reconstructed = anchor[: -len(piece)] + piece
                        if (existing_idx := next((idx for idx, p in enumerate(pairs) if p[0] == reconstructed), None)) is None:
                            pairs.append([reconstructed, [], False]); active_idx = len(pairs) - 1
                        else: active_idx = existing_idx
                        pending_op = None; i += 1; continue

            fallback = True; break
        if fallback: break

    polis_codes = [p[0] for p in pairs]
    if fallback or not polis_codes or len(polis_codes) > MAX_BREAKDOWN_CODES: return [""], [original]
    return [_format_cert_list(cert_nums, is_sd) for _, cert_nums, is_sd in pairs], polis_codes


# ==============================================================================
# 3. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path, sheet_name, header_row): return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_cedant(df_raw, keyword, cedant_column):
    if cedant_column not in df_raw.columns: raise ValueError(f"CEDANT_COLUMN='{cedant_column}' tidak ditemukan. Tersedia: {list(df_raw.columns)}")
    df_filtered = df_raw[df_raw[cedant_column].astype(str).str.contains(keyword, case=False, na=False)].copy()
    print(f"[INFO] Filter cedant (kolom '{cedant_column}') mengandung '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered


# ==============================================================================
# 4. BUILD OUTPUT
# ==============================================================================
def _breakdown_all_rows(df, col_insured, col_polis, col_slip):
    ins_arr, pol_arr, slp_arr = df[col_insured].to_numpy(), df[col_polis].to_numpy(), df[col_slip].to_numpy()
    insured_cln_all, polis_cln_all, slip_cln_all, cert_cln_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for ins_val, pol_val, slp_val in zip(ins_arr, pol_arr, slp_arr):
        ins_cln = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []
        slp_cln = clean_split_code(str(slp_val)) if pd.notna(slp_val) else []
        cert_cln, pol_cln = process_policy_certificate(str(pol_val)) if pd.notna(pol_val) else ([], [])

        insured_cln_all.append(ins_cln)
        polis_cln_all.append(pol_cln)
        slip_cln_all.append(slp_cln)
        cert_cln_all.append(cert_cln)
        max_ins, max_pol, max_slp = max(max_ins, len(ins_cln)), max(max_pol, len(pol_cln)), max(max_slp, len(slp_cln))

    return insured_cln_all, polis_cln_all, slip_cln_all, cert_cln_all, max_ins, max_pol, max_slp

def build_output(df_filtered):
    for col, label in ((INSURED_COLUMN, "INSURED_COLUMN"), (POLIS_COLUMN, "POLIS_COLUMN"), (SLIP_COLUMN, "SLIP_COLUMN")):
        if col not in df_filtered.columns: raise ValueError(f"Kolom {label}='{col}' tidak ditemukan. Tersedia: {list(df_filtered.columns)}")

    insured_cln_all, polis_cln_all, slip_cln_all, cert_cln_all, max_ins, max_pol, max_slp = _breakdown_all_rows(df_filtered, INSURED_COLUMN, POLIS_COLUMN, SLIP_COLUMN)
    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)

    if BROKER_NAME_COLUMN in df_filtered.columns and CEDANT_COLUMN in df_filtered.columns:
        is_direct = df_filtered[BROKER_NAME_COLUMN].astype(str).str.strip().str.upper() == DIRECT_MARKER
        bp_series = df_filtered[CEDANT_COLUMN].where(is_direct, df_filtered[BROKER_NAME_COLUMN])
        business_partner_values = bp_series.astype(str).str.replace(r"^(\s*PT)\.\s*", r"\1 ", regex=True, flags=RE_I).values
    else: business_partner_values = None

    output_cols_data = {}
    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        if col == BROKER_NAME_COLUMN and business_partner_values is not None: output_cols_data["BUSINESS_PARTNER"] = business_partner_values
        
        if col == INSURED_COLUMN:
            for i in range(1, max_ins + 1): output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in insured_cln_all]
        if col == POLIS_COLUMN:
            for i in range(1, max_pol + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in polis_cln_all]
                output_cols_data[f"CERTIFICATE_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in cert_cln_all]
        if col == SLIP_COLUMN:
            for i in range(1, max_slp + 1): output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slip_cln_all]

    df_master = pd.DataFrame(output_cols_data)
    drop_cols = [c for c in df_master.columns if c != "CERTIFICATE_1" and ("_CLN_" in c or c.startswith("CERTIFICATE_")) and df_master[c].astype(str).str.strip().eq("").all()]
    return df_master.drop(columns=drop_cols)


# ==============================================================================
# 5. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df, output_path):
    df.to_excel(output_path, index=False)
    workbook, text_col_indices = load_workbook(output_path), [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]
    worksheet = workbook.active

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None: cell.value = str(cell.value)
    workbook.save(output_path)


# ==============================================================================
# 6. EXECUTION RUNNER
# ==============================================================================
def main():
    print("=" * 70 + "\n CLEANSING DATA 1 - FACULTATIVE | CEDANT: PT LIPPO GENERAL INSURANCE\n" + "=" * 70)
    input_path = INPUT_FILE

    if not os.path.exists(input_path):
        if candidates := [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]:
            input_path = candidates[0]
            print(f"[AUTO-DETECT] Memakai file terdeteksi: '{input_path}'")
        else: raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    print(f"[INFO] Total baris source: {len(df_raw)}")

    if (df_lippo := filter_cedant(df_raw, CEDANT_FILTER, CEDANT_COLUMN)).empty:
        print("[WARNING] Tidak ada baris yang cocok dengan filter cedant. Proses dihentikan.")
        return

    df_hasil = build_output(df_lippo)
    save_with_text_format(df_hasil, OUTPUT_FILE)
    print("=" * 70 + f"\n[SUCCESS] Selesai. Total baris output: {len(df_hasil)}\n[SUCCESS] File hasil: '{OUTPUT_FILE}'\n" + "=" * 70)

if __name__ == "__main__":
    main()