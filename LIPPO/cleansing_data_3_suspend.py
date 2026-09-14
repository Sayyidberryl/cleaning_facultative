"""
================================================================================
 SCRIPT CLEANSING DATA 3 - SUSPENSE (KHUSUS CEDANT: PT LIPPO GENERAL INSURANCE)
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
INPUT_FILE = "3b. Database Suspense.xlsx"
SHEET_NAME = "Sheet1"
HEADER_ROW = 2
OUTPUT_FILE = "lippo_output_suspense_V3.xlsx"
CEDANT_FILTER = "LIPPO"
STATUS_FILTER = "SUSPENSE"

MAX_BREAKDOWN_CODES = 5
MAX_BREAKDOWN_INSURED = 5
TEXT_FORMAT_COLUMN_KEYWORDS = ("POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO", "CERTIFICATE")

# ==============================================================================
# 1. CLEANSING INSURED
# ==============================================================================
_ASTERISK_RE = re.compile(r"\*")
_TRANSACTION_HEADER_RE = re.compile(r"(?i)^\s*(?:Pembayaran|Penagihan)?\s*(?:dan\s+Penagihan)?\s*Premi\s*(?:Fakultatif|SOA|Pensesian)?\s*(?:PAR\s*&?\s*EQ|BIT|GIT|CECR|QS)?\s*")
_TRUNCATE_TRIGGER_RE = re.compile(r"(?i)\bsubsidiar|\bassociat|\baffiliat|\bfiliated\b|related\s+compan|respective\s+right|\bincluding\s+an[yd]|\bindu?ding\s+any|as\s+the\s+owner|as\s+owners?\b|as\s+main\s+contractor|\(as\s+mortgagee\s+and\s+loss\s+payee\)")
_REMOVE_ONLY_RE = re.compile(r"(?i)\bas\s+principal\b|\bsebagai\s+principal\b|\bsebagai\s+kontraktor\b|\bas\s+(?:co\s+)?propert(?:y)?\s+owner\b|\bas\s+operating\s+company\b|\bas\s+property\s+manager\b|\bas\s+contractor\b|\bas\s+co\s+ben[ie]ficiary\b|\bas\s+event\s+project\s+owner\b")
_GENERIC_BOILERPLATE_TAIL_RE = re.compile(r"(?i)\bsub\s*-?\s*kontraktor\b.*$|\banak\s+perusahaan\b.*$|\bafiliasi\s+perusahaan\b.*$|\bsemua\b\s*[–-]?\s*$")
_TITLE_RE = re.compile(r"(?i)\b(?:S\.?H\.?|S\.?I\.?Kom\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b")
_LEGAL_ENTITY_RE = re.compile(r"(?i)\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO|Perseroan\s+Terbatas)\b")
_ENTITY_SPLIT_RE = re.compile(r"(?i)\b(?:QQ|Q\.Q\.)\b|\band\s*/\s*or\b|\bdan\s*/\s*atau\b|\bin\s+this\s+case\b|\bdalam\s+hal\s+ini\b|\band\b|\bor\b|,")
_LEGAL_ENTITY_DASH_SIGNAL_RE = re.compile(r"(?i),\s*(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|PERSERO)\s*[\r\n]*\s*-\s*")
_GENERIC_DASH_SEPARATOR_RE = re.compile(r"\s*[\r\n]+\s*-\s*|\s+-\s+")
_MIDDLE_BOILERPLATE_CLAUSE_RE = re.compile(r"(?i)(?:\band\s*/\s*or\s+)?(?:associat\w*|subsidiar\w*|affiliat\w*|related\s+compan\w*)(?:\s+and\s*/\s*or\s+(?:associat\w*|subsidiar\w*|affiliat\w*|related\s+compan\w*))*\s+for\s+(?:their|its)\s+respective\s+rights?\s+and\s+interests?\b")

_COMPRISING_OF_RE = re.compile(r"(?is)^(?P<head>.*?\bgroup)\s+of\s+companies\s+comprising\s+of\s*:\s*(?P<tail>.+)$")
_CONSISTS_OF_RE = re.compile(r"(?i)consists\s+of\s*:")
_NUMBERED_ITEM_RE = re.compile(r"\d+\s*\.\s*")
_AS_PER_LIST_RE = re.compile(r"(?i)as\s+per\s+list\s+attached\s*\((?P<inner>[^)]+)\)")
_GROUP_OF_COMPANIES_SUFFIX_RE = re.compile(r"(?i)\s+group\s+of\s+companies\b")
_BARE_AND_OR_RE = re.compile(r"(?i)\band\b|\bor\b")
_PAREN_PT_RE = re.compile(r"(?i)\(\s*(?P<inner>(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?)\s+[^)]+)\)")
_PAREN_STRIP_PREFIX_RE = re.compile(r"(?i)^(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?)\s*")
_DANGLING_WORD_RE = re.compile(r"(?i)^\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b\s*|\s*\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b$")

_SPECIFIC_SENTENCE_WHITELIST = {
    "PT Palladium Megah Lestari and/ or PERPETUAL (ASIA) LIMITED (in its capacity as Trustee of Lippo Malls Indonesia Retail Trust) Lippo Malls Indonesia and all subsidiary or controlled companies, for their respective rights and": [
        "PALLADIUM MEGAH LESTARI",
        "PERPETUAL ASIA LIMITED",
        "LIPPO MALLS INDONESIA",
    ],
}
_SPECIFIC_SENTENCE_WHITELIST_NORMALIZED = {re.sub(r"\s+", " ", k.strip()).upper(): v for k, v in _SPECIFIC_SENTENCE_WHITELIST.items()}

_AND_NAME_WHITELIST = [
    "MITSUBISHI HC CAPITAL AND FINANCE INDONESIA",
    "MITSUBISHI UFJ LEASE AND FINANCE INDONESIA",
]

def _strip_asterisk(text: str) -> str:
    return _ASTERISK_RE.sub("", text)

def _truncate_from_first_trigger(text: str) -> str:
    match = _TRUNCATE_TRIGGER_RE.search(text)
    return text[: match.start()] if match else text

def _strip_middle_boilerplate_clause(text: str) -> str:
    match = _MIDDLE_BOILERPLATE_CLAUSE_RE.search(text)
    if not match:
        return text
        
    trailing = text[match.end():].strip(" .,;:-")
    if len(trailing) >= 3 and re.search(r"[A-Za-z]{3,}", trailing):
        return f"{text[: match.start()].rstrip()}, {trailing}"
    return text

def _normalize_dash_entity_list(text: str) -> str:
    if _LEGAL_ENTITY_DASH_SIGNAL_RE.search(text):
        return _GENERIC_DASH_SEPARATOR_RE.sub(", ", text)
    return text

def _strip_transaction_header(text: str) -> str:
    return _TRANSACTION_HEADER_RE.sub("", text)

def _try_pattern_h_specific_sentence(text: str):
    key = re.sub(r"\s+", " ", text.strip()).upper()
    return _SPECIFIC_SENTENCE_WHITELIST_NORMALIZED.get(key)

def _is_real_company_name(inner: str) -> bool:
    name = _PAREN_STRIP_PREFIX_RE.sub("", inner).strip()
    words = [w for w in re.split(r"[\s\-]+", name) if w]
    
    if len(words) < 2:
        return False
        
    has_lowercase = any(c.islower() for c in name)
    all_short_upper = all(len(w) <= 3 and w.isupper() for w in words)
    return not (all_short_upper and not has_lowercase)

def _try_pattern_f_paren_pt(text: str):
    m = _PAREN_PT_RE.search(text)
    if not m:
        return None
        
    inner = m.group("inner").strip()
    if not _is_real_company_name(inner):
        return None
        
    head = text[: m.start()].strip()
    tail_after = text[m.end():].strip()
    
    if tail_after:
        head = f"{head} {tail_after}".strip()
    return [head, inner]

def _try_pattern_a_comprising_of(text: str):
    m = _COMPRISING_OF_RE.match(text.strip())
    if m:
        return [m.group("head").strip(), m.group("tail").strip()]
    return None

def _try_pattern_b_consists_of(text: str):
    m = _CONSISTS_OF_RE.search(text)
    if not m:
        return None
        
    tail = _truncate_from_first_trigger(text[m.end():])
    items = [seg.strip() for seg in _NUMBERED_ITEM_RE.split(tail) if seg.strip()]
    return items or None

def _try_pattern_c_as_per_list(text: str):
    m = _AS_PER_LIST_RE.search(text)
    if not m:
        return None
        
    parts = [p.strip() for p in re.split(r"\s*[\u2013\u2014-]\s*", m.group("inner")) if p.strip()]
    return parts or None

def _strip_group_of_companies_suffix(text: str) -> str:
    return _GROUP_OF_COMPANIES_SUFFIX_RE.sub("", text)

def _truncate_boilerplate_tail(text: str) -> str:
    text = _truncate_from_first_trigger(text)
    match_generic = _GENERIC_BOILERPLATE_TAIL_RE.search(text)
    if match_generic:
        text = text[: match_generic.start()]
    return _REMOVE_ONLY_RE.sub(" ", text)

def _strip_titles(text: str) -> str:
    return _TITLE_RE.sub(" ", text)

def _strip_legal_entity(text: str) -> str:
    return _LEGAL_ENTITY_RE.sub(" ", text)

def _protect_and_whitelist(text: str) -> tuple[str, dict]:
    placeholders = {}
    for i, phrase in enumerate(_AND_NAME_WHITELIST):
        amp_variant = re.sub(r"(?i)\bAND\b", "&", phrase)
        for j, candidate in enumerate([phrase, amp_variant]):
            pattern = re.compile(re.escape(candidate).replace(r"\ ", r"\s+"), re.IGNORECASE)
            match = pattern.search(text)
            if match:
                key = f"ANDWL{i}{j}PLACEHOLDER"
                placeholders[key] = phrase
                text = pattern.sub(key, text)
    return text, placeholders

def _restore_and_whitelist(text: str, placeholders: dict[str, str]) -> str:
    for key, phrase in placeholders.items():
        text = re.sub(key, phrase, text, flags=re.IGNORECASE)
    return text

def _final_polish(text: str) -> str:
    t = re.sub(r"[().;:\"']", " ", text)
    t = re.sub(r"[-/]", " ", t)
    t = re.sub(r"\s+", " ", t).strip()
    
    previous = None
    while previous != t:
        previous = t
        t = _DANGLING_WORD_RE.sub("", t).strip()
    return t.upper()

def _finalize_segments(raw_parts: list) -> list:
    results = []
    for part in raw_parts:
        part, placeholders = _protect_and_whitelist(part)
        for sub in _ENTITY_SPLIT_RE.split(part):
            cleaned = _truncate_boilerplate_tail(sub)
            cleaned = _strip_group_of_companies_suffix(cleaned)
            cleaned = _strip_legal_entity(cleaned)
            cleaned = _final_polish(cleaned)
            cleaned = _restore_and_whitelist(cleaned, placeholders)
            
            if cleaned and len(cleaned) >= 2 and re.search(r"[A-Z]", cleaned) and cleaned not in results:
                results.append(cleaned)
    return results

def _clean_insured_name_breakdown(text: str) -> list:
    specific_result = _try_pattern_h_specific_sentence(text)
    if specific_result is not None:
        return specific_result

    t = text.strip()
    t = _strip_transaction_header(t)
    t = _strip_titles(t)
    t = _normalize_dash_entity_list(t)
    t = _strip_middle_boilerplate_clause(t)

    special_segments = _try_pattern_c_as_per_list(t) or _try_pattern_b_consists_of(t)
    if special_segments is not None:
        return _finalize_segments(special_segments)

    special_segments = _try_pattern_a_comprising_of(t)
    if special_segments is not None:
        return _finalize_segments(special_segments)

    special_segments = _try_pattern_f_paren_pt(t)
    if special_segments is not None:
        return _finalize_segments(special_segments)

    return _finalize_segments([_truncate_boilerplate_tail(_strip_group_of_companies_suffix(t))])

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return []
        
    original = text.strip()
    t = _strip_asterisk(original)
    results = _clean_insured_name_breakdown(t)
    
    if len(results) > MAX_BREAKDOWN_INSURED:
        return [original]
    return results


# ==============================================================================
# 2. CLEANSING POLIS & SLIP NO
# ==============================================================================
_TOK_MAIN_RE = re.compile(r"^\d{13}$")
_TOK_CERT_RE = re.compile(r"^\d{6}$")
_TOK_ENDORSE_RE = re.compile(r"^\d{1,2}/\d{1,2}$")
_TOK_STATUS_RE = re.compile(r"(?i)^(?:New|Endorsement|Cancel\s*All|Adjustment|Cancellation|Reinstatement)$")

def _classify_token(tok: str) -> str:
    if _TOK_MAIN_RE.match(tok): return "MAIN"
    if _TOK_CERT_RE.match(tok): return "CERT"
    if _TOK_ENDORSE_RE.match(tok): return "ENDORSE"
    if _TOK_STATUS_RE.match(tok): return "STATUS"
    return "OTHER"

def clean_split_code(text: str, skip_combined_if_cert: bool = False) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return []

    t = text.strip()
    t = re.sub(r"[.,]", "", t)
    t = re.sub(r"\s+", " ", t).strip()
    
    classified = []
    for tok in t.split("-"):
        tok = tok.strip()
        if tok:
            classified.append((_classify_token(tok), tok))

    codes = []
    i = 0
    n = len(classified)
    
    while i < n:
        kind, tok = classified[i]
        
        if kind == "MAIN":
            suffix_parts = []
            j = i + 1
            
            while j < n and classified[j][0] in ("CERT", "ENDORSE"):
                suffix_parts.append(classified[j][1])
                j += 1

            followed_by_status = j < n and classified[j][0] == "STATUS"
            has_cert_suffix = False
            for k in range(i + 1, j):
                if classified[k][0] == "CERT":
                    has_cert_suffix = True
                    break

            if suffix_parts and not followed_by_status:
                if not (skip_combined_if_cert and has_cert_suffix):
                    combined = f"{tok}-" + "-".join(suffix_parts)
                    if combined not in codes:
                        codes.append(combined)
                
                if tok not in codes:
                    codes.append(tok)
            elif tok not in codes:
                codes.append(tok)

            if followed_by_status:
                j += 1
                
            i = j
            continue
            
        i += 1

    if not codes:
        return [t] if t else []
        
    if len(codes) > MAX_BREAKDOWN_CODES:
        return codes[:1]
    return codes


# ==============================================================================
# 3. EKSTRAKSI CERTIFICATE
# ==============================================================================
_SD_NORMALIZE_RE = re.compile(r"(?i)\bs\s*/\s*d\b")
_CERT_SEGMENT_RE = re.compile(r"(?i)(?P<base>\d{13})-(?P<start>\d{5,6})(?!\d)(?:\s*(?:SD|-)\s*(?:(?P=base)-)?(?P<end>\d{5,6})(?!\d))?")

def _normalize_sd(text: str) -> str:
    return _SD_NORMALIZE_RE.sub("SD", text)

def _to_six_digit(num_str: str) -> str | None:
    if len(num_str) == 5:
        return "0" + num_str
    if len(num_str) == 6:
        return num_str
    return None

def _format_certificate_range(start_six: str, end_six: str) -> str:
    start_num = int(start_six)
    end_num = int(end_six)
    
    if end_num < start_num:
        return ""
        
    count = end_num - start_num + 1
    if count > 3:
        return f"{start_six} SD {end_six}"
        
    return ", ".join(f"{n:06d}" for n in range(start_num, end_num + 1))

def extract_certificate(text: str) -> str:
    if not isinstance(text, str):
        return ""
        
    t = text.strip()
    if not t or t.lower() in ("nan", "none", "-"):
        return ""

    t = _normalize_sd(t)
    results = []
    
    for match in _CERT_SEGMENT_RE.finditer(t):
        start_six = _to_six_digit(match.group("start"))
        if start_six is None:
            continue
            
        end_digits = match.group("end")
        if end_digits:
            end_six = _to_six_digit(end_digits)
            if end_six is None:
                continue
            cert = _format_certificate_range(start_six, end_six)
        else:
            cert = start_six
            
        if cert and cert not in results:
            results.append(cert)

    return ", ".join(results)


# ==============================================================================
# 4. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_status_suspense(df_raw: pd.DataFrame, keyword: str) -> pd.DataFrame:
    status_col = None
    for c in df_raw.columns:
        if str(c).upper().strip() == "STATUS":
            status_col = c
            break
            
    if not status_col:
        raise ValueError("Kolom STATUS tidak ditemukan di source file.")

    mask = df_raw[status_col].astype(str).str.strip().str.upper() == keyword.upper()
    df_filtered = df_raw[mask].copy()
    
    print(f"[INFO] Filter STATUS = '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered

def filter_cedant(df_raw: pd.DataFrame, keyword: str) -> pd.DataFrame:
    cedant_cols = [c for c in df_raw.columns if "CEDANT" in str(c).upper()]
    if not cedant_cols:
        raise ValueError("Kolom CEDANT tidak ditemukan di source file.")

    mask = pd.Series(False, index=df_raw.index)
    for col in cedant_cols:
        mask |= df_raw[col].astype(str).str.contains(keyword, case=False, na=False)

    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant mengandung '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered


# ==============================================================================
# 5. BUILD OUTPUT
# ==============================================================================
def _detect_key_columns(df: pd.DataFrame) -> tuple[str, str, str]:
    col_insured = None
    col_polis = None
    col_slip = None
    
    for c in df.columns:
        c_upper = str(c).upper().strip()
        if "INSURED" in c_upper and col_insured is None:
            col_insured = c
        elif c_upper == "POLIS" and col_polis is None:
            col_polis = c
        elif "SLIP" in c_upper and col_slip is None:
            col_slip = c

    if None in (col_insured, col_polis, col_slip):
        raise ValueError(f"Kolom kunci tidak lengkap. INSURED={col_insured}, POLIS={col_polis}, SLIP={col_slip}")
        
    print(f"[INFO] Kolom terdeteksi -> INSURED='{col_insured}', POLIS='{col_polis}', SLIP='{col_slip}'")
    return col_insured, col_polis, col_slip

def _breakdown_all_rows(df: pd.DataFrame, col_insured: str, col_polis: str, col_slip: str) -> tuple[list, list, list, list, int, int, int]:
    insured_cln_all = []
    polis_cln_all = []
    slip_cln_all = []
    certificate_all = []
    
    max_ins = 0
    max_pol = 0
    max_slp = 0

    for _, row in df.iterrows():
        ins_val = row[col_insured]
        pol_val = row[col_polis]
        slp_val = row[col_slip]
        
        ins_cln = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []
        cert = extract_certificate(str(pol_val)) if pd.notna(pol_val) else ""
        pol_cln = clean_split_code(str(pol_val), skip_combined_if_cert=bool(cert)) if pd.notna(pol_val) else []
        slp_cln = clean_split_code(str(slp_val)) if pd.notna(slp_val) else []

        insured_cln_all.append(ins_cln)
        polis_cln_all.append(pol_cln)
        slip_cln_all.append(slp_cln)
        certificate_all.append(cert)
        
        max_ins = max(max_ins, len(ins_cln))
        max_pol = max(max_pol, len(pol_cln))
        max_slp = max(max_slp, len(slp_cln))

    return insured_cln_all, polis_cln_all, slip_cln_all, certificate_all, max_ins, max_pol, max_slp

def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_insured, col_polis, col_slip = _detect_key_columns(df_filtered)
    breakdown_results = _breakdown_all_rows(df_filtered, col_insured, col_polis, col_slip)
    insured_cln_all, polis_cln_all, slip_cln_all, certificate_all, max_ins, max_pol, max_slp = breakdown_results

    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)
    output_cols_data = {}
    
    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        
        if col == col_insured:
            for i in range(1, max_ins + 1):
                output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in insured_cln_all]
                
        if col == col_polis:
            if max_pol >= 1:
                output_cols_data["POLIS_CLN_1"] = [lst[0] if len(lst) >= 1 else "" for lst in polis_cln_all]
                
            output_cols_data["CERTIFICATE"] = certificate_all
            
            for i in range(2, max_pol + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in polis_cln_all]
                
        if col == col_slip:
            for i in range(1, max_slp + 1):
                output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slip_cln_all]

    df_master = pd.DataFrame(output_cols_data)
    
    drop_cols = []
    for c in df_master.columns:
        if "_CLN_" in c and df_master[c].astype(str).str.strip().eq("").all():
            drop_cols.append(c)
            
    return df_master.drop(columns=drop_cols)


# ==============================================================================
# 6. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    
    workbook = load_workbook(output_path)
    worksheet = workbook.active
    text_col_indices = []
    
    for idx, col in enumerate(df.columns, start=1):
        for kw in TEXT_FORMAT_COLUMN_KEYWORDS:
            if kw in str(col).upper():
                text_col_indices.append(idx)
                break

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None:
                cell.value = str(cell.value)
                
    workbook.save(output_path)


# ==============================================================================
# 7. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print("=" * 70)
    print(" CLEANSING DATA 3 - SUSPENSE | CEDANT: PT LIPPO GENERAL INSURANCE")
    print("=" * 70)
    
    input_path = INPUT_FILE
    if not os.path.exists(input_path):
        candidates = []
        for f in os.listdir("."):
            if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_"):
                candidates.append(f)
                
        if not candidates:
            raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")
            
        input_path = candidates[0]
        print(f"[AUTO-DETECT] File input tidak ditemukan di path default, memakai: '{input_path}'")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    print(f"[INFO] Total baris source: {len(df_raw)}")

    df_suspense = filter_status_suspense(df_raw, STATUS_FILTER)
    if df_suspense.empty:
        print(f"[WARNING] Tidak ada baris dengan STATUS = '{STATUS_FILTER}'. Proses dihentikan.")
        return

    df_lippo = filter_cedant(df_suspense, CEDANT_FILTER)
    if df_lippo.empty:
        print("[WARNING] Tidak ada baris yang cocok dengan filter cedant. Proses dihentikan.")
        return

    df_hasil = build_output(df_lippo)
    save_with_text_format(df_hasil, OUTPUT_FILE)
    
    print("=" * 70)
    print(f"[SUCCESS] Selesai. Total baris output: {len(df_hasil)}")
    print(f"[SUCCESS] File hasil: '{OUTPUT_FILE}'")
    print("=" * 70)

if __name__ == "__main__":
    main()