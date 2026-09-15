import os
import re

import numpy as np
import pandas as pd


INPUT_FILE = os.path.join("input", "2b. transaksi Osbal 01.01.23 - 17.08.26.xlsx")
OUTPUT_FILE = os.path.join("output", "tokio_output_osbal.xlsx")

CEDANT_COL   = "CCOS_COMP_NAME"
CEDANT_VALUE = "PT.ASURANSI TOKIO MARINE INDONESIA"

POLIS_COL   = "FAC_POLICY_NO"
SLIP_COL    = "FAC_SLIP"
INSURED_COL = "FAC_INSURED"

MAX_SPLIT_COLS = 5

# Polis: kata/prefix yang menyebabkan nilai dibiarkan apa adanya
POLIS_EXCEPTION_RE = re.compile(
    r"""
    MOP\s*MARINE
  | (?:LINE\s*SLIP|LINESLIP)
  | \b(?:P1|P2|P3|P73)\s*CANCEL
  | \b(?:P1|P2|P3|P73)\b
  | \bCANCEL\b
  | PENYELESAIAN(?:\s+SUSPENSE)?
  | HUTANG
  | UTANG
    """,
    re.IGNORECASE | re.VERBOSE,
)

# Slip: kata yang menyebabkan nilai dibiarkan apa adanya
SLIP_EXCEPTION_RE = re.compile(
    r"""
    \bSUMMARY\b
  | \bBORDER[OA]\b
  | \bBORDRO\b
  | \bSINGGLESHIPMENT\b
  | \bP1\b
  | \bP2\b
  | \bP3\b
  | PENYELESAIAN(?:\s+SUSPENSE)?
  | HUTANG
  | UTANG
    """,
    re.IGNORECASE | re.VERBOSE,
)

INSURED_JUNK_WORDS = frozenset({
    "PT", "CV", "TBK", "PERSERO", "LTD", "INC", "LLC",
    "AND", "OR", "THE", "OF", "AS",
    "NON FOOD", "DIV",
    "MR", "MR.", "MRS", "MRS.", "MS", "MS.",
})

INSURED_SUFFIX_RE = re.compile(
    r"""
    ,?\s*\bTBK\b\s*(?:,?\s*PT\.?)?
  | ,?\s*\bPT\.?\s*$
  | ,?\s*\bCV\.?\s*$
    """,
    re.IGNORECASE | re.VERBOSE,
)

INSURED_PT_QQ_RE = re.compile(
    r"""
    \bPT\.?\b
  | \bQQ\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

_SLIP_NOISE_RE = re.compile(
    r"""
    \b(?:JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS
        |SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\b
  | \b(?:JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST
        |SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)\b
  | \b(?:IDR|USD|ENG)\b
  | \b(?:P1|P2|P3|P73)\b
  | \bNEW\b
  | \bVARIOUS\b
  | \b20[0-9]{2}\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

_SLIP_TOKEN_RE = re.compile(r"[A-Z0-9][A-Z0-9\-]{6,}", re.IGNORECASE)


# ─────────────────────────────────────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def _normalize_spaces(text: str) -> str:
    return re.sub(r"\s{2,}", " ", text).strip()


def _clean_insured_name(name: str) -> str:
    name = _normalize_spaces(name)

    if re.fullmatch(r"AEON\s+MALL\s*\(TENANTS\)", name, re.IGNORECASE):
        return name.upper()

    for _ in range(3):
        cleaned = _normalize_spaces(INSURED_SUFFIX_RE.sub("", name).strip().strip(","))
        if cleaned == name:
            break
        name = cleaned

    name = INSURED_PT_QQ_RE.sub(" ", name)

    if re.fullmatch(r"\s*AEON\s+MALL\s+\(TENANTS\)\s*", name, flags=re.IGNORECASE):
        return "AEON MALL (TENANTS)"

    name = re.sub(r"\(\s*PERSERO\s*\)", "", name, flags=re.IGNORECASE)
    name = re.sub(r"[()]", " ", name)
    name = name.replace("(", " ").replace(")", " ")
    name = _normalize_spaces(name)
    name = re.sub(r"^[\s.,\-]+|[\s.,\-]+$", "", name)

    return name


def _cap_or_join(items: list) -> list:
    if len(items) > MAX_SPLIT_COLS:
        return [",".join(items)]
    return items


def _unique(items):
    out = []
    for x in items:
        if x and x not in out:
            out.append(x)
    return out


# ─────────────────────────────────────────────────────────────────────────────
# CLEAN POLIS TOKIO
# ─────────────────────────────────────────────────────────────────────────────

def clean_polis_raw(val) -> list:
    if pd.isna(val):
        return []

    val = str(val).strip().upper()

    if not val:
        return []

    val = re.sub(r"\bVARIOUS\b", "", val)
    val = re.sub(r"\s*/\s*VAR$", "", val)

    if val != "TBA":
        val = re.sub(r"\bTBA\b", "", val)

    if val == "TBA":
        return ["TBA"]

    val = val.replace("&", "+")
    val = _normalize_spaces(val)
    val = val.replace(".", "")
    val = re.sub(r"-\d{3}-\d{2}$", "", val)

    if POLIS_EXCEPTION_RE.search(val):
        return [_normalize_spaces(val)]

    if re.search(
        r"(JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\s+\d{4}",
        val,
        flags=re.IGNORECASE,
    ):
        return [_normalize_spaces(val)]

    m = re.fullmatch(
        r"([A-Z]{3}/[A-Z]{4}/\d{2}-[A-Z]\d{6,8})(?:-000-\d+)?",
        val,
        flags=re.IGNORECASE,
    )
    if m:
        return [m.group(1).upper()]

    m = re.fullmatch(
        r"([A-Z]\d{7,8})(?:-000-\d+)?",
        val,
        flags=re.IGNORECASE,
    )
    if m:
        return [m.group(1).upper()]

    if "S/D" in val:
        hasil = re.findall(r"[A-Z]\d{7,8}", val)
        if hasil:
            return _cap_or_join(_unique(hasil))

    val = re.sub(
        r"([A-Z]\d{7,8})-\d{3}(?=\s|\+|$)",
        r"\1",
        val,
        flags=re.IGNORECASE,
    )

    val = re.sub(
        r"([A-Z]\d{7,8})\s+(TMD/|TMI/)",
        r"\1+\2",
        val,
        flags=re.IGNORECASE,
    )

    if "+" in val:
        parts = [p.strip() for p in val.split("+") if p.strip()]
        results = []
        tokio_prefix = None
        tokio_company = None

        for p in parts:
            m_company = re.match(r"^(TMD|TMI)/", p, flags=re.IGNORECASE)
            if m_company:
                tokio_company = m_company.group(1).upper() + "/"

            m = re.fullmatch(r"([A-Z]{4}/\d{2}[A-Z]\d{7,8})", p, flags=re.IGNORECASE)
            if m and tokio_company:
                results.append(tokio_company + m.group(1).upper())
                continue

            m = re.fullmatch(
                r"^([A-Z]{3}/[A-Z]{4}/\d{2}-)([A-Z]\d{7,8})((?:-[A-Z]\d{7,8})+)$",
                p,
                flags=re.IGNORECASE,
            )
            if m:
                tokio_prefix = m.group(1).upper()
                results.append(tokio_prefix + m.group(2).upper())
                suffixes = re.findall(r"[A-Z]\d{7,8}", m.group(3), flags=re.IGNORECASE)
                for suffix in suffixes:
                    results.append(tokio_prefix + suffix.upper())
                continue

            m = re.fullmatch(r"([A-Z]{4}/\d{2}[A-Z]\d{7,8})", p, flags=re.IGNORECASE)
            if m and tokio_prefix:
                results.append("TMD/" + m.group(1).upper())
                continue
            
            m = re.match(
                r"^([A-Z]{3}/[A-Z]{4}/\d{2}-)([A-Z]\d{7,8})$",
                p,
                flags=re.IGNORECASE,
            )
            if m:
                tokio_prefix = m.group(1).upper()
                results.append(tokio_prefix + m.group(2).upper())
                continue

            m = re.match(
                r"^([A-Z]{7,8}\d{2})([A-Z]\d{7,8})$",
                p,
                flags=re.IGNORECASE,
            )
            if m:
                tokio_prefix = m.group(1).upper()
                results.append(tokio_prefix + m.group(2).upper())
                continue
                
            m = re.fullmatch(r"[A-Z]\d{7,8}", p, flags=re.IGNORECASE)
            if m:
                code = m.group(0).upper()
                results.append(tokio_prefix + code if tokio_prefix else code)
                continue

            if p.upper() not in ["VAR", "VARIOUS", "TBA"]:
                results.append(p.upper())

        return _cap_or_join(_unique(results))

    m = re.fullmatch(
        r"([A-Z]{3}/[A-Z]{4}/\d{2}-)([A-Z]\d{7,8})((?:-[A-Z]\d{7,8})+)",
        val,
        flags=re.IGNORECASE,
    )
    if m:
        tokio_prefix = m.group(1).upper()
        first_code = m.group(2).upper()
        results = [tokio_prefix + first_code]
        suffixes = re.findall(r"[A-Z]\d{7,8}", m.group(3), flags=re.IGNORECASE)
        for suffix in suffixes:
            results.append(tokio_prefix + suffix.upper())
        return _cap_or_join(_unique(results))

    m = re.fullmatch(
        r"((?:TMD|TMI)[A-Z]{4}\d{2})([A-Z]\d{7,8})-(\d{4})-((?:TMD|TMI)[A-Z]{4}\d{2}[A-Z]\d{7,8})",
        val,
        flags=re.IGNORECASE,
    )
    if m:
        tokio_prefix = m.group(1).upper()
        first_code = m.group(2).upper()
        second_suffix = m.group(3)
        third_policy = m.group(4).upper()
        code_prefix = first_code[:-4]

        return _cap_or_join(_unique([
            tokio_prefix + first_code,
            tokio_prefix + code_prefix + second_suffix,
            third_policy,
        ]))

    m = re.fullmatch(
        r"((?:TMD|TMI)[A-Z]{4}\d{2})([A-Z]\d{7,8})-([A-Z]\d{7,8})",
        val,
        flags=re.IGNORECASE
    )
    if m:
        tokio_prefix = m.group(1).upper()
        return [
            tokio_prefix + m.group(2).upper(),
            tokio_prefix + m.group(3).upper()
        ]

    m = re.fullmatch(
        r"((?:TMD|TMI)[A-Z]{4}\d{2})([A-Z]\d{7,8})((?:-\d{3,4})+)",
        val,
        flags=re.IGNORECASE
    )
    if m:
        tokio_prefix = m.group(1).upper()
        first_code = m.group(2).upper()
        results = [tokio_prefix + first_code]
        suffixes = re.findall(r"\d{3,4}", m.group(3))

        for suffix in suffixes:
            code_prefix = first_code[:-3] if len(suffix) == 3 else first_code[:-4]
            results.append(tokio_prefix + code_prefix + suffix)

        return _cap_or_join(_unique(results))

    if re.fullmatch(
        r"(?:TMD|TMI)[A-Z]{4}\d{2}[A-Z]\d{7,8}(?:-\d+)+",
        val,
        flags=re.IGNORECASE
    ):
        return [val]

    hasil = re.findall(
        r"(?:TMD|TMI)?[A-Z]{3,5}\d{2}[A-Z]\d{7,8}",
        val,
        flags=re.IGNORECASE
    )
    if len(hasil) > 1:
        return _cap_or_join(_unique([x.upper() for x in hasil]))

    hasil = re.findall(
        r"[A-Z]\d{7,8}",
        val,
        flags=re.IGNORECASE
    )
    if len(hasil) > 1:
        return _cap_or_join(_unique([x.upper() for x in hasil]))

    if re.match(r"^(TMD|TMI)[A-Z0-9]+$", val):
        return [val.upper()]

    m = re.search(r"[A-Z]\d{7,8}", val)
    if m:
        return [m.group(0).upper()]

    return []


def clean_polis(val) -> list:
    if pd.notna(val):
        val_str = str(val).strip().upper()
        if re.fullmatch(r"TMD\.FIAR\.22\.F\d{7,8}(?:\+\d{3})+", val_str):
            return [val_str.replace(".", "")]

    hasil = clean_polis_raw(val)
    hasil_clean = []

    for x in hasil:
        x = re.sub(r"[^A-Z0-9]", "", str(x).upper())
        if x:
            hasil_clean.append(x)

    return hasil_clean

# ─────────────────────────────────────────────────────────────────────────────
# CLEAN SLIP
# ─────────────────────────────────────────────────────────────────────────────

def _clean_slip_core(val) -> list:
    if pd.isna(val):
        return []

    val = str(val).strip()
    if not val:
        return []

    val = _normalize_spaces(val)

    val = re.sub(
        r"([A-Z]{3}/[A-Z]{4}/\d{2}-[A-Z]\d{7,8})-\d{3}-\d{2}",
        r"\1",
        val,
        flags=re.IGNORECASE,
    )

    val = re.sub(
        r"([A-Z]\d{7,8})-\d{3}-\d{2}",
        r"\1",
        val,
        flags=re.IGNORECASE,
    )

    m = re.match(r"^(\d+)\s+S/D\s+\d+", val, flags=re.IGNORECASE)
    if m:
        return [m.group(1)]

    m = re.search(r"(TM[DI]/[A-Z]{4}/\d{2}-[A-Z]\d{7,8})", val, flags=re.IGNORECASE)
    if m:
        return [m.group(1)]

    if "FACILITY" in val:
        return [val]

    if POLIS_EXCEPTION_RE.search(val):
        return [_normalize_spaces(val)]

    if re.fullmatch(r"P\d+(?:\s*\+\s*P\d+)*(\s+CANCEL)?", val, flags=re.IGNORECASE):
        return [_normalize_spaces(val)]

    val = re.sub(r"\s*/\s*END\b", "", val, flags=re.IGNORECASE).strip()

    val = re.sub(
        r"\s*/\s*(VAR\s+ACCOUNT|FACILITY\s+AUTOMOBILE|FACILITY)\b",
        "",
        val,
        flags=re.IGNORECASE,
    )

    val = re.sub(r"\bVARIOUS\b", "", val, flags=re.IGNORECASE)
    val = re.sub(r"\bVAR\b", "", val, flags=re.IGNORECASE)
    val = _normalize_spaces(val)

    val = re.sub(r"/\s*VAR\s+ACCOUNT\b", "", val, flags=re.IGNORECASE)
    val = re.sub(r"/\s*FACILITY\s+AUTOMOBILE\b", "", val, flags=re.IGNORECASE)
    val = re.sub(r"/\s*FACILITY\b", "", val, flags=re.IGNORECASE)
    val = _normalize_spaces(val)

    parts = re.split(r"\s{2,}", val)
    parts = [p.strip() for p in parts if p.strip()]

    if len(parts) >= 2:
        results = [_normalize_spaces(p) for p in parts]
        return _cap_or_join(results)

    if SLIP_EXCEPTION_RE.search(val):
        return [_normalize_spaces(val)]

    if re.fullmatch(r"\d+(?:\.\d+)+(?:-\d+){3,}", val):
        return [_normalize_spaces(val)]

    if "+" in val:
        parts = [p.strip() for p in re.split(r"\s*\+\s*", val) if p.strip()]
        parts = [p for p in parts if not re.fullmatch(r"P\d+", p, flags=re.IGNORECASE)]

        if len(parts) > 1:
            first = parts[0]
            rest = parts[1:]

            if all(re.fullmatch(r"\d+", p) for p in parts):
                return _cap_or_join(parts)

            rest_is_digit = all(re.fullmatch(r"\d+", r) for r in rest)
            suffix_lengths = {len(r) for r in rest} if rest_is_digit else set()

            if rest_is_digit and len(suffix_lengths) == 1:
                suffix_len = suffix_lengths.pop()
                m = re.match(rf"^(.*?)(\d{{{suffix_len}}})$", first)

                if m:
                    prefix = m.group(1)
                    results = [first]
                    for s in rest:
                        s = s.zfill(suffix_len)
                        results.append(prefix + s)
                    return _cap_or_join(results)

            if all(re.match(r"^\d+(?:\.\d+)+$", p) for p in parts):
                return _cap_or_join(parts)

            return [_normalize_spaces(val)]

    m = re.match(r"^(.*?)(\d+)-(\d+)$", val)
    if m:
        prefix, first, second = m.group(1), m.group(2), m.group(3)
        results = [prefix + first]
        if len(second) < len(first):
            second = second.zfill(len(first))
            results.append(prefix + second)
        else:
            results.append(second)
        return _cap_or_join(results)

    slips = re.findall(r"\d+(?:\.\d+)+", val)
    if len(slips) >= 2:
        return _cap_or_join(slips)

    if re.search(
        r"(JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\s+\d{4}",
        val,
        flags=re.IGNORECASE,
    ):
        return [_normalize_spaces(val)]

    m = re.search(r"\d+(?:\.\d+)+", val)
    if m:
        return [m.group(0)]

    return [_normalize_spaces(val)]     


def clean_slip_raw(val) -> list:
    results = _clean_slip_core(val)
    cleaned = []

    for r in results:
        if pd.isna(r):
            continue
        r = str(r).upper()
        r = _normalize_spaces(r)
        r = r.replace(".", "")
        r = re.sub(r'(?<=\d)\s*TBA\b', '', r).strip()

        if r and r not in cleaned:
            cleaned.append(r)

    val_str = "" if pd.isna(val) else str(val).strip()
    if re.fullmatch(r"\d{4}", val_str):
        return ["0000" + val_str]

    return cleaned


def clean_slip(val) -> list:
    hasil = clean_slip_raw(val)
    hasil_clean = []

    for x in hasil:
        x = re.sub(r"[^A-Z0-9]", "", str(x).upper())
        if x:
            hasil_clean.append(x)

    return hasil_clean

# ─────────────────────────────────────────────────────────────────────────────
# CLEAN INSURED
# ─────────────────────────────────────────────────────────────────────────────

def clean_insured(val) -> list:
    if pd.isna(val):
        return []

    val = str(val).strip()
    if not val:
        return []

    bracket_parts = re.findall(r"\(([^)]*)\)", val)
    val = re.sub(r"\([^)]*\)", "", val)

    parts = re.split(r"/|,|QQ", val, flags=re.IGNORECASE)
    parts.extend(bracket_parts)

    cleaned = []

    for p in parts:
        p = _normalize_spaces(p.strip())
        p = re.sub(r"^(MR|MRS|MS|DR|IR)\.?\s+", "", p, flags=re.IGNORECASE)

        if len(p) <= 2 or p.upper() in INSURED_JUNK_WORDS:
            continue
        if re.match(r"^[^A-Za-z0-9]+$", p):
            continue

        p_clean = _clean_insured_name(p)
        if not p_clean or p_clean.upper() in INSURED_JUNK_WORDS:
            continue

        if p_clean not in cleaned:
            cleaned.append(p_clean.upper())

    if not cleaned:
        fallback = _clean_insured_name(_normalize_spaces(val))
        return [fallback.upper()] if fallback else []

    return _cap_or_join(cleaned)


# ─────────────────────────────────────────────────────────────────────────────
# CLEAN CERTIFICATE
# ─────────────────────────────────────────────────────────────────────────────

def clean_certificate(polis_ori) -> str:
    """
    Pada Data 2 Osbal Tokio, tidak ada nomor sertifikat (Certificate di-BLANK-kan).
    """
    return ""


# ─────────────────────────────────────────────────────────────────────────────
# PROCESS DATA (OPTIMIZED EXCEL READ/WRITE)
# ─────────────────────────────────────────────────────────────────────────────

def _insert_clean_columns(
    df: pd.DataFrame,
    all_lists: list,
    prefix: str,
    max_cols: int,
) -> list:
    added = []
    for i in range(1, max_cols + 1):
        col_name = f"clean {prefix} {i}"
        df[col_name] = [lst[i - 1] if i - 1 < len(lst) else None for lst in all_lists]
        added.append(col_name)
    return added


def _fast_read_excel(path: str, sheet_name=None, header: int = 0) -> pd.DataFrame:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    for _ in range(header):
        next(rows_iter)
    cols = next(rows_iter)
    data = list(rows_iter)
    wb.close()
    return pd.DataFrame(data, columns=cols)


def _fast_write_excel(df: pd.DataFrame, path: str) -> None:
    import openpyxl
    os.makedirs(os.path.dirname(path), exist_ok=True)

    df2 = df.astype(object).where(pd.notnull(df), None)

    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Sheet1")
    ws.append(list(df2.columns))
    for row in df2.itertuples(index=False, name=None):
        ws.append(row)
    wb.save(path)


def process_data(input_file: str, output_file: str) -> None:
    print(f"[1/5] Membaca data cepat dari: {input_file} ...")
    df = _fast_read_excel(input_file, header=0)
    print(f"      Total baris keseluruhan: {len(df):,}")

    if CEDANT_COL not in df.columns:
        print(f"\n[ERROR] Kolom '{CEDANT_COL}' tidak ditemukan!")
        print(f"        Kolom tersedia: {list(df.columns)}")
        return

    for col in [POLIS_COL, SLIP_COL, INSURED_COL]:
        if col not in df.columns:
            print(f"\n[ERROR] Kolom '{col}' tidak ditemukan di file!")
            print(f"        Kolom tersedia: {list(df.columns)}")
            return

    df[CEDANT_COL] = df[CEDANT_COL].fillna("").astype(str).str.strip()
    is_tokio = df[CEDANT_COL] == CEDANT_VALUE
    print(f"[2/5] Filter cedant '{CEDANT_VALUE}': {is_tokio.sum():,} baris TOKIO dari total {len(df):,} baris.")

    df.rename(columns={POLIS_COL: "polis_ori", SLIP_COL: "slip_ori", INSURED_COL: "insured_ori"},
              inplace=True)

    print("[3/5] Menjalankan proses cleaning cepat hanya untuk baris TOKIO ...")

    n = len(df)
    all_clean_polis = [[] for _ in range(n)]
    all_clean_slip  = [[] for _ in range(n)]
    all_clean_ins   = [[] for _ in range(n)]
    all_certificates = ["" for _ in range(n)]

    max_polis = max_slip = max_ins = 1

    tokio_idx = np.flatnonzero(is_tokio.to_numpy())
    polis_vals   = df["polis_ori"].to_numpy()
    slip_vals    = df["slip_ori"].to_numpy()
    insured_vals = df["insured_ori"].to_numpy()

    total_t = len(tokio_idx)
    for n_done, pos in enumerate(tokio_idx, 1):
        if n_done % 20_000 == 0:
            print(f"      Progress: {n_done:,} / {total_t:,} baris TOKIO diproses...")

        p_ori = polis_vals[pos]
        c_polis = clean_polis(p_ori)
        c_slip  = clean_slip(slip_vals[pos])
        c_ins   = clean_insured(insured_vals[pos])
        c_cert  = clean_certificate(p_ori)

        max_polis = max(max_polis, len(c_polis))
        max_slip  = max(max_slip,  len(c_slip))
        max_ins   = max(max_ins,   len(c_ins))

        all_clean_polis[pos] = c_polis
        all_clean_slip[pos]  = c_slip
        all_clean_ins[pos]   = c_ins
        all_certificates[pos] = c_cert

    print(f"      Selesai diproses!")
    print(f"      -> Jumlah kolom clean polis  : {max_polis}")
    print(f"      -> Jumlah kolom clean slip   : {max_slip}")
    print(f"      -> Jumlah kolom clean insured: {max_ins}")

    print("[4/5] Menyusun kolom output ...")

    df["CERTIFICATE"] = all_certificates

    new_columns = []
    for col in df.columns:
        if col == "CERTIFICATE":
            continue

        new_columns.append(col)
        if col == "polis_ori":
            new_columns.append("CERTIFICATE")
            new_columns += _insert_clean_columns(df, all_clean_polis, "polis",   max_polis)
        elif col == "slip_ori":
            new_columns += _insert_clean_columns(df, all_clean_slip,  "slip",    max_slip)
        elif col == "insured_ori":
            new_columns += _insert_clean_columns(df, all_clean_ins,   "insured", max_ins)

    df = df[new_columns]

    print(f"[5/5] Menyimpan hasil cepat ke: {output_file} ...")
    _fast_write_excel(df, output_file)

    print(f"\n{'=' * 55}")
    print(f"  [OK] Selesai!")
    print(f"  Total baris output  : {len(df):,}")
    print(f"  Total kolom output  : {len(df.columns)}")
    print(f"  File disimpan di    : {output_file}")
    print(f"{'=' * 55}")


# ─────────────────────────────────────────────────────────────────────────────
# ENTRY POINT
# ─────────────────────────────────────────────────────────────────────────────
if __name__ == "__main__":
    process_data(INPUT_FILE, OUTPUT_FILE)