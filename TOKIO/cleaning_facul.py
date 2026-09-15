import os
import re

import numpy as np
import pandas as pd


INPUT_FILE  = os.path.join("input", "1b. Transaksi Facul 01.01.23 - 17.08.2026.xlsx")
OUTPUT_FILE = os.path.join("output", "tokio_output_facul.xlsx")

CEDANT_COL   = "COMP_NAME"
CEDANT_VALUE = "PT.ASURANSI TOKIO MARINE INDONESIA"

POLIS_COL   = "FAC_POLICY_NO"
SLIP_COL    = "FAC_SLIP"
INSURED_COL = "FAC_INSURED"
BROKER_COL  = "COMP_NAME_1"

MAX_SPLIT_COLS = 5

# Polis: kata/prefix yang menyebabkan nilai dibiarkan apa adanya
POLIS_EXCEPTION_RE = re.compile(
    r"""
    MOP\s*MARINE
  | (?:LINE\s*SLIP|LINESLIP)
  | \b(?:P1|P2|P3|P73)\s*CANCEL
  | \b(?:P1|P2|P3|P73)\b
  | \bCANCEL\b
  | PENYELESAIAN(?:\s+SUSPENSE)?
  | HUTANG
  | UTANG
    """,
    re.IGNORECASE | re.VERBOSE,
)

SLIP_EXCEPTION_RE = re.compile(
    r"""
    \bSUMMARY\b
  | \bBORDER[OA]\b
  | \bBORDRO\b
  | \bSINGGLESHIPMENT\b
  | PENYELESAIAN(?:\s+SUSPENSE)?
  | HUTANG
  | UTANG
    """,
    re.IGNORECASE | re.VERBOSE,
)

MONTH_NAMES = frozenset({
    "JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI",
    "JULI", "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER",
    "JANUARY", "FEBRUARY", "MARCH", "MAY", "JUNE", "JULY", "AUGUST",
    "OCTOBER",
})

# =========================
# INSURED CONFIG
# =========================

INSURED_REMOVE_RE = re.compile(
    r"""
    \b(
        PT|CV|TBK|
        PERSERO|
        LTD|PTE|INC|LLC|
        MR|MRS|MS|
        BAPAK|BPK|IBU|NY|
        DR|DRS|DRA|IR|H|HJ
    )\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

INSURED_POLIS_RE = re.compile(
    r"""
    (POLIS\s*NO\.?.*)
    |
    (POLICY\s*NO\.?.*)
    |
    (SLIP\s*NO\.?.*)
    """,
    re.IGNORECASE | re.VERBOSE,
)

INSURED_SPLIT_RE = re.compile(
    r"""
    \s*,\s*
    |
    \s*/\s*
    |
    \s+QQ\s+
    |
    \s+AND/OR\s+
    |
    \s*&\s*
    |
    \s*\+\s*
    """,
    re.IGNORECASE | re.VERBOSE,
)

INSURED_JUNK_WORDS = {
    "",
    "AND",
    "OR",
    "THE",
    "OF",
    "AS"
}

# Noise token dalam slip
_SLIP_NOISE_RE = re.compile(
    r"""
    \b(?:JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS
        |SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\b
  | \b(?:JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST
        |SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)\b
  | \b(?:IDR|USD|ENG)\b
  | \b(?:P1|P2|P3|P73)\b
  | \bNEW\b
  | \bVARIOUS\b
  | \b20[0-9]{2}\b
    """,
    re.IGNORECASE | re.VERBOSE,
)

_SLIP_TOKEN_RE = re.compile(r"[A-Z0-9][A-Z0-9\-]{6,}", re.IGNORECASE)


# ─────────────────────────────────────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────────────────────────────────────

def _normalize_spaces(text: str) -> str:
    return re.sub(r"\s{2,}", " ", text).strip()


def _is_valid_polis_token(tok: str) -> bool:
    tok = tok.strip()
    return (
        len(tok) >= 5
        and not re.search(r"TBA$", tok, re.IGNORECASE)
        and bool(re.search(r"\d", tok))
    )


def _extract_polis_tokens(text: str) -> list:
    tokens = []
    for block in re.split(r"\s{2,}", text.strip()):
        for tok in re.split(r"\+|\s+", block.strip()):
            tok = tok.strip().strip("-/")
            if _is_valid_polis_token(tok):
                tokens.append(tok)
    return tokens


def _expand_dash_chain(segment: str) -> list:
    results = []
    current_base = None

    for part in segment.split("-"):
        part = part.strip()
        if not part:
            continue

        if len(part) >= 8 and re.search(r"\d", part):
            current_base = part
            results.append(part)
        elif (
            current_base
            and len(part) >= 2
            and re.match(r"^\d+$", part)
            and len(part) < len(current_base)
        ):
            n = len(part)
            results.append(current_base[:-n] + part)
        else:
            if _is_valid_polis_token(part):
                results.append(part)

    return results or ([segment] if _is_valid_polis_token(segment) else [])


def _is_valid_slip_token(tok: str) -> bool:
    tok = tok.strip()
    return (
        len(tok) >= 7
        and bool(re.search(r"\d", tok))
        and not re.match(r"^\d{4}$", tok)
    )


def _strip_slip_suffix(tok: str) -> str:
    m = re.match(r"^(.+?)-(\d{1,6})$", tok)
    if m and len(m.group(1)) > len(m.group(2)):
        return m.group(1)
    return tok


def _extract_slip_tokens(text: str) -> list:
    results = []
    for block in re.split(r"\s{2,}", text.strip()):
        clean = _SLIP_NOISE_RE.sub(" ", block)
        clean = re.sub(r"^[\s\-/+,]+|[\s\-/+,]+$", "", clean).strip()

        for cand in _SLIP_TOKEN_RE.findall(clean):
            cand = cand.strip("-")
            parts = cand.split("-")
            if len(parts) == 2:
                a, b = parts
                if _is_valid_slip_token(a) and _is_valid_slip_token(b) and abs(len(a) - len(b)) <= 2:
                    results.extend([a, b])
                    continue
            if _is_valid_slip_token(cand):
                results.append(_strip_slip_suffix(cand))

    return results


def _clean_insured_name(name: str) -> str:
    if pd.isna(name):
        return ""
    name = str(name).upper()
    name = INSURED_POLIS_RE.sub("", name)
    name = INSURED_REMOVE_RE.sub(" ", name)
    name = re.sub(r"[()]", " ", name)
    name = name.replace("/", " ")
    name = name.replace("-", " ")
    name = _normalize_spaces(name)
    return name

def split_insured(name):
    if pd.isna(name):
        return []
    name = str(name)
    parts = INSURED_SPLIT_RE.split(name)
    hasil = []
    for p in parts:
        p = _clean_insured_name(p)
        if not p or p in INSURED_JUNK_WORDS:
            continue
        hasil.append(p)
    return hasil


def _cap_or_join(items: list) -> list:
    if len(items) > MAX_SPLIT_COLS:
        return [",".join(items)]
    return items


# ============================================================
# NEW: CERTIFICATE EXTRACTION & FORMATTING HELPERS
# ============================================================

def format_cert_range(start_str, end_str):
    """Format 6 digit padding & tentukan pakai koma atau SD"""
    try:
        start_int = int(start_str)
        end_int = int(end_str)
        
        if end_int < start_int:
            return f"{str(start_int).zfill(6)} SD {str(end_int).zfill(6)}"
            
        count = end_int - start_int + 1
        
        if count > 3:
            return f"{str(start_int).zfill(6)} SD {str(end_int).zfill(6)}"
        else:
            return ", ".join(str(i).zfill(6) for i in range(start_int, end_int + 1))
    except ValueError:
        return ""

def extract_cert(p):
    """Mendeteksi dan menarik nilai sertifikat dari teks, mengembalikan (sisa_polis, cert)"""
    cert = ""
    
    m_sd = re.search(r"(?:-|,)?(\d{1,6})_SD_(\d{1,6})\b", p, re.IGNORECASE)
    if m_sd:
        # Jadikan cert HANYA JIKA digit aslinya >= 3
        if len(m_sd.group(1)) >= 3 or len(m_sd.group(2)) >= 3:
            cert = format_cert_range(m_sd.group(1), m_sd.group(2))
        p = p[:m_sd.start()] + p[m_sd.end():]
        return p.strip(" -/,"), cert
        
    m_comma = re.search(r"(?:-|,)((?:\d{1,6},)+\d{1,6})\b", p)
    if m_comma:
        nums = re.findall(r"\d{1,6}", m_comma.group(1))
        if any(len(n) >= 3 for n in nums):
            cert = ", ".join(n.zfill(6) for n in nums)
        p = p[:m_comma.start()] + p[m_comma.end():]
        return p.strip(" -/,"), cert
        
    m_range = re.search(r"-(\d{1,6})-(\d{1,6})\b", p)
    if m_range:
        if len(m_range.group(1)) >= 3 or len(m_range.group(2)) >= 3:
            cert = format_cert_range(m_range.group(1), m_range.group(2))
        p = p[:m_range.start()] + p[m_range.end():]
        return p.strip(" -/,"), cert
        
    m_single = re.search(r"([A-Z]\d{5,8}|\d{8,})(?:-|,)(\d{1,6})\b", p)
    if m_single:
        if len(m_single.group(2)) >= 3:
            cert = m_single.group(2).zfill(6)
        p = p[:m_single.start()] + m_single.group(1) + p[m_single.end():]
        return p.strip(" -/,"), cert
        
    m_end = re.search(r"(?:-|,)(\d{1,6})$", p)
    if m_end:
        if len(m_end.group(1)) >= 3:
            cert = m_end.group(1).zfill(6)
        p = p[:m_end.start()]
        return p.strip(" -/,"), cert

    return p.strip(" -/,"), cert

def clean_polis(val) -> list:
    if pd.isna(val):
        return []
    

    

    val = str(val).strip()
    if not val:
        return []

    # ============================================================
    # TOKIO - COMMA REPETITION
    # ============================================================

    m_comma_rep = re.fullmatch(
        r"^(.*?)-(\d{2,7}(?:,\d{2,7})+)$",
        val
    )

    if m_comma_rep:

        base = m_comma_rep.group(1).strip()

        suffixes = [
            x.strip()
            for x in m_comma_rep.group(2).split(",")
        ]

        hasil_repetition = [
            {
                "polis": f"{base}-{suffix}",
                "cert": ""
            }
            for suffix in suffixes
            if re.fullmatch(r"\d{2,7}", suffix)
        ]

        # > 5 → biarkan POLIS ORI apa adanya
        if len(hasil_repetition) > MAX_SPLIT_COLS:
            return [{
                "polis": val,
                "cert": ""
            }]

        return hasil_repetition

    # ============================================================
    # FIX SUPER CLEANER: HAPUS JUNK DAN LEM SPASI/TITIK/DASH
    # ============================================================
    val = re.sub(r"\bVARIOUS\b", "", val, flags=re.IGNORECASE)
    val = re.sub(r"\bVARIO\b", "", val, flags=re.IGNORECASE)
    val = re.sub(r"\bVARI\b", "", val, flags=re.IGNORECASE)
    val = re.sub(r"\bVAR\b", "", val, flags=re.IGNORECASE)
    
    val = re.sub(r"\(\s*END\s*\)", "", val, flags=re.IGNORECASE)
    val = re.sub(r"\bEND\b", "", val, flags=re.IGNORECASE)
    
    # Hapus P1, P2, P3, dll di AWAL kata sebelum memproses regex yang lain
    val = re.sub(r"\b(?:P1|P2|P3|P73)\b", "", val, flags=re.IGNORECASE)
    
    # 1. FIX: Lem Spasi antara Prefix dan Base
    val = re.sub(r"([A-Z]{3}/[A-Z]{3,4}/\d{2})\s+([A-Z]\d{4,8})", r"\1-\2", val)
    
    # 2. Lem S/D dan Koma Certificate
    val = re.sub(r"(\d{1,6})\s+S/?D\s+(\d{1,6})", r"\1_SD_\2", val, flags=re.IGNORECASE)
    val = re.sub(r"\s+-\s+(\d)", r"-\1", val)
    val = re.sub(r",\s+(\d)", r",\1", val)
    
    # 3. FIX: Ubah pola titik repitisi menjadi '+' 
    val = re.sub(r"(?<=\d{3})\.(?=\d{3,5}\b)", "+", val)
    
    # 4. Ubah spasi slash menjadi plus agar bisa di-split dengan aman
    val = re.sub(r"\s+/\s*|\s*/\s+", " + ", val)
    
    # 5. FIX: Ubah pola dash/koma repitisi menjadi '+' (MENCEGAH REPETISI JADI CERTIFICATE)
    while True:
        new_val = re.sub(
            r"([A-Z0-9\./-]{8,}(?:\+\d{2,7})*)\s*(?:-|,)\s*(\d{2,7})\b",
            # Aturan: Kalau angkanya 6 digit (123456) atau diawali 00 (000001), itu Certificate, jadi jangan diubah jadi '+'
            lambda m: m.group(0) if (m.group(2).startswith("00") or len(m.group(2))==6) else f"{m.group(1)}+{m.group(2)}",
            val
        )
        if new_val == val:
            break
        val = new_val
    
    val = re.sub(r"[\s/]+$", "", val)
    val = _normalize_spaces(val)

    def finalize(chunks):
        res = []
        for chunk in chunks:
            if not chunk.strip():
                continue
                
            p_clean, c_clean = extract_cert(chunk)
            
            # Simulated Certificate extraction for -001, -002 patterns 
            # if extract_cert failed to catch it
            m_end_cert = re.search(r"-(\d{1,3})$", p_clean)
            if m_end_cert and not c_clean:
                # Cek jika angka di ujung itu seperti -001
                if m_end_cert.group(1).startswith("00"):
                     c_clean = m_end_cert.group(1).zfill(6)
                     p_clean = p_clean[:m_end_cert.start()]

            # Sapu bersih junk
            p_clean = re.sub(r"/\d{1,3}$", "", p_clean)
            p_clean = re.sub(r"[,-]\d{1,2}$", "", p_clean) 
            
            if re.fullmatch(r"[0-9.]+", p_clean):
                p_clean = p_clean.replace(".", "")
            elif re.fullmatch(r"[0-9-]+", p_clean):
                p_clean = p_clean.replace("-", "")
                
            if p_clean:
                # Jika certificate mengandung koma, pecah jadi beberapa kolom
                if c_clean and "," in c_clean:
                    for cert_item in c_clean.split(","):
                        res.append({"polis": p_clean, "cert": cert_item.strip()})
                else:
                    res.append({"polis": p_clean, "cert": c_clean})
        
        if len(res) > MAX_SPLIT_COLS:
            comb_p = ",".join(x["polis"] for x in res)
            comb_c = ",".join(x["cert"] for x in res if x["cert"])
            return [{"polis": comb_p, "cert": comb_c}]
        return res

    if POLIS_EXCEPTION_RE.search(val) or re.search(r"\d+TBA\d+", val, re.IGNORECASE):
        p, c = extract_cert(val)
        return [{"polis": _normalize_spaces(p), "cert": c}]

    val = val.replace("&", "+")
    val = re.sub(r"\+\s*\+", "+", val)
    val = re.sub(r"^\s*\+\s*|\s*\+\s*$", "", val)
    val = _normalize_spaces(val)

    if re.fullmatch(r"\d+\+\d+", val):
        return finalize([val])

    if "+" in val:
        hasil = []
        base = None
        prefix = "" 
        for p in val.split("+"):
            p = p.strip()
            
            m_full = re.search(r"^(.+?)([A-Z]\d{4,8})(.*)$", p)
            m_base_only = re.search(r"^([A-Z]\d{4,8})(.*)$", p)
            
            if m_full and not m_base_only:
                prefix = m_full.group(1)
                base = m_full.group(2)
                suffix = m_full.group(3)
                if p not in hasil:
                    hasil.append(p)
                continue
                
            elif m_base_only:
                new_base = m_base_only.group(1)
                suffix = m_base_only.group(2)
                base = new_base
                if prefix:
                    full = prefix + new_base + suffix
                    if full not in hasil:
                        hasil.append(full)
                else:
                    if p not in hasil:
                        hasil.append(p)
                continue

            if base and prefix and re.match(r"^\d+", p):
                m_num = re.match(r"^(\d+)(.*)$", p)
                num = m_num.group(1)
                rest = m_num.group(2)
                if len(num) >= len(base) - 1:
                    code = base[0] + num[-(len(base)-1):]
                else:
                    code = base[:-len(num)] + num
                polis = prefix + code + rest
                if polis not in hasil:
                    hasil.append(polis)
                continue

            if p:
                if p not in hasil:
                    hasil.append(p)
        return finalize(hasil)

    if re.search(r"\s+/\s+", val):
        hasil = [p.strip() for p in re.split(r"\s+/\s+", val) if p.strip()]
        return finalize(hasil)

    if " " in val:
        hasil = [p.strip() for p in val.split() if p.strip()]
        return finalize(hasil)

    return finalize([val])

# ─────────────────────────────────────────────────────────────────────────────
# CLEAN SLIP
# ─────────────────────────────────────────────────────────────────────────────

def clean_slip(val) -> list:
    if pd.isna(val):
        return []
    val = str(val).upper().strip()
    val = val.replace(".", "")
    if not val:
        return []
    val = _normalize_spaces(val)

    if SLIP_EXCEPTION_RE.search(val):
        return [_normalize_spaces(val)]

    if re.search(r"\b(EQ|PAR)/\d+", val, re.IGNORECASE):
        return [_normalize_spaces(val)]

    if re.fullmatch(r"MDD/FCMI/\d{2}-F\d{7}(?:&\d{4,5})+", val, re.IGNORECASE):
        return [_normalize_spaces(val)]

    val = re.sub(r"\bEND(?:\.?1)?\b", "", val, flags=re.IGNORECASE)
    val=re.sub(r"END\.1\.?","",val)
    if ";" in val:
        val=val.replace(";","+")
        
    if re.search(r"[+&]", val):
        parts = [x.strip() for x in re.split(r"[+&]", val) if x.strip()]
        hasil = []
        base = None
        prefix = ""
        for p in parts:
            if re.fullmatch(r"\d{1,3}", p):
                continue
            p = re.sub(r"-\d{2}$", "", p)
            if base and prefix and re.fullmatch(r"\d{4,5}", p):
                code = base[:-len(p)] + p
                hasil.append(prefix + code)
                base = code
                continue
            hasil.append(p)
            m = re.match(r"^(.*?)([A-Z]\d{7}|\d{5,8})$", p)
            if m:
                prefix = m.group(1)
                code = m.group(2)
                base = code
                hasil.append(prefix + code)
                continue
        return _cap_or_join(list(dict.fromkeys(hasil)))

    slips = []
    for s in re.split(r"\s*[;+&]\s*", val):
        s = s.strip()
        if not s:
            continue
        s = re.sub(
            r"\b(USD|IDR|SGD|EUR|JPY|AUD|GBP|ORI|ORI\.|ORIGINAL|COPY|REALISASI|REALIZATION|CANCEL|CANCELLED|ENDT?|ENDORSEMENT|SA|P1|P2|P3|VAR|REVISI|REV)\b.*$",
            "", s, flags=re.IGNORECASE
        )
        s = _normalize_spaces(s)
        s = re.sub(r"-(\d{2})$", "", s)
        s = s.strip("-_,.; ")
        if not s:
            continue
        if s not in slips:
            slips.append(s)

    return _cap_or_join(slips)

# ─────────────────────────────────────────────────────────────────────────────
# CLEAN INSURED & MITRA BISNIS
# ─────────────────────────────────────────────────────────────────────────────

def clean_insured(val) -> list:
    if pd.isna(val):
        return []
    val = str(val).strip()
    if not val:
        return []

    insureds = split_insured(val)
    cleaned = []
    for ins in insureds:
        ins = _normalize_spaces(ins)
        if not ins or len(ins) <= 2 or ins.upper() in INSURED_JUNK_WORDS:
            continue
        ins = ins.upper()
        if ins not in cleaned:
            cleaned.append(ins)

    if not cleaned:
        fallback = _clean_insured_name(val)
        if fallback:
            cleaned.append(fallback.upper())

    return _cap_or_join(cleaned)


def get_mitra_bisnis(comp_name2, comp_name) -> str:
    broker = "" if pd.isna(comp_name2) else str(comp_name2).strip()
    if not broker or broker.upper() == "DIRECT":
        return "" if pd.isna(comp_name) else str(comp_name).strip()
    return broker


def clean_business_partners(val):
    if pd.isna(val):
        return ""
    val = str(val).strip().upper()
    val = re.sub(r"\bPT\.\s*", "PT ", val, flags=re.IGNORECASE)
    return _normalize_spaces(val)


# ─────────────────────────────────────────────────────────────────────────────
# PROCESS DATA
# ─────────────────────────────────────────────────────────────────────────────

def _insert_paired_columns(df: pd.DataFrame, all_clean_data: list, max_cols: int) -> list:
    """Sisipkan kolom 'clean polis' dan maksimal 3 'Certificate' secara berdampingan."""
    added = []
    for i in range(1, max_cols + 1):
        polis_col = f"clean polis {i}"
        cert_col = f"Certificate {i}"
        
        df[polis_col] = [lst[i - 1]["polis"] if i - 1 < len(lst) else None for lst in all_clean_data]
        added.append(polis_col)
        
        # RULE: HANYA BUAT KOLOM CERTIFICATE SAMPAI MAKSIMAL KE-3
        if i <= 3:
            df[cert_col] = [lst[i - 1]["cert"] if i - 1 < len(lst) else None for lst in all_clean_data]
            added.append(cert_col)
            
    return added


def _insert_clean_columns(df: pd.DataFrame, all_lists: list, prefix: str, max_cols: int) -> list:
    added = []
    for i in range(1, max_cols + 1):
        col_name = f"clean {prefix} {i}"
        df[col_name] = [lst[i - 1] if i - 1 < len(lst) else None for lst in all_lists]
        added.append(col_name)
    return added


def _fast_read_excel(path: str, sheet_name=None, header: int = 0) -> pd.DataFrame:
    import openpyxl
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)
    for _ in range(header):
        next(rows_iter)
        
    raw_cols = list(next(rows_iter))
    
    cols = []
    seen = {}
    for col in raw_cols:
        col_str = str(col).strip() if col is not None else ""
        if col_str in seen:
            seen[col_str] += 1
            cols.append(f"{col_str}_{seen[col_str]}")
        else:
            seen[col_str] = 0
            cols.append(col_str)
            
    data = list(rows_iter)
    wb.close()
    return pd.DataFrame(data, columns=cols)


def _fast_write_excel(df: pd.DataFrame, path: str) -> None:
    import openpyxl
    os.makedirs(os.path.dirname(path), exist_ok=True)
    df2 = df.astype(object).where(pd.notnull(df), None)
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Sheet1")
    ws.append(list(df2.columns))
    for row in df2.itertuples(index=False, name=None):
        ws.append(row)
    wb.save(path)


def process_data(input_file: str, output_file: str) -> None:
    print(f"[1/5] Membaca data dari: {input_file} ...")
    df = _fast_read_excel(input_file, header=0)
    print(f"      Total baris keseluruhan: {len(df):,}")

    if CEDANT_COL not in df.columns:
        print(f"\n[ERROR] Kolom '{CEDANT_COL}' tidak ditemukan!")
        return

    required_cols = [POLIS_COL, SLIP_COL, INSURED_COL, BROKER_COL]
    for col in required_cols:
        if col not in df.columns:
            print(f"\n[ERROR] Kolom '{col}' tidak ditemukan di file!")
            return

    df[CEDANT_COL] = df[CEDANT_COL].astype(str).str.strip()
    is_tokio = df[CEDANT_COL] == CEDANT_VALUE
    print(f"[2/5] Filter cedant '{CEDANT_VALUE}': {is_tokio.sum():,} baris TOKIO dari total {len(df):,} baris.")

    broker_s = df[BROKER_COL].fillna("").astype(str).str.strip()
    cedant_s = df[CEDANT_COL].fillna("").astype(str).str.strip()
    use_cedant = (broker_s == "") | (broker_s.str.upper() == "DIRECT")
    mitra_values = [
        clean_business_partners(x)
        for x in np.where(use_cedant, cedant_s, broker_s)
    ]

    df.rename(columns={POLIS_COL: "polis_ori", SLIP_COL: "slip_ori", INSURED_COL: "insured_ori"},
              inplace=True)

    insert_pos = list(df.columns).index(BROKER_COL) + 1 if BROKER_COL in df.columns else len(df.columns)
    df.insert(insert_pos, "BUSINESS PARTNERS", mitra_values)

    print("[3/5] Menjalankan proses cleaning hanya untuk baris TOKIO ...")

    n = len(df)
    all_clean_polis = [[] for _ in range(n)]
    all_clean_slip  = [[] for _ in range(n)]
    all_clean_ins   = [[] for _ in range(n)]
    max_polis = max_slip = max_ins = 1

    tokio_idx = np.flatnonzero(is_tokio.to_numpy())
    polis_vals   = df["polis_ori"].to_numpy()
    slip_vals    = df["slip_ori"].to_numpy()
    insured_vals = df["insured_ori"].to_numpy()

    total_w = len(tokio_idx)
    for n_done, pos in enumerate(tokio_idx, 1):
        if n_done % 5_000 == 0:
            print(f"      Progress: {n_done:,} / {total_w:,} baris TOKIO diproses...")

        c_polis_paired = clean_polis(polis_vals[pos])
        c_slip  = clean_slip(slip_vals[pos])
        c_ins   = clean_insured(insured_vals[pos])

        max_polis = max(max_polis, len(c_polis_paired))
        max_slip  = max(max_slip,  len(c_slip))
        max_ins   = max(max_ins,   len(c_ins))

        all_clean_polis[pos] = c_polis_paired
        all_clean_slip[pos]  = c_slip
        all_clean_ins[pos]   = c_ins

    print(f"      Selesai diproses!")

    print("[4/5] Menyusun kolom output ...")

    new_columns = []
    for col in df.columns:
        new_columns.append(col)
        if col == "polis_ori":
            new_columns += _insert_paired_columns(df, all_clean_polis, max_polis)
        elif col == "slip_ori":
            new_columns += _insert_clean_columns(df, all_clean_slip,  "slip",    max_slip)
        elif col == "insured_ori":
            new_columns += _insert_clean_columns(df, all_clean_ins,   "insured", max_ins)

    df = df[new_columns]

    print(f"[5/5] Menyimpan hasil ke: {output_file} ...")
    _fast_write_excel(df, output_file)

    print(f"\n{'=' * 55}")
    print(f"  [OK] Selesai!")
    print(f"{'=' * 55}")

if __name__ == "__main__":
    process_data(INPUT_FILE, OUTPUT_FILE)