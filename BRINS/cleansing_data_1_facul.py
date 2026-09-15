"""
================================================================================
 SCRIPT CLEANSING DATA 1 - FACULTATIVE (KHUSUS CEDANT: BRINS GENERAL INSURANCE)
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
INPUT_FILE = "1b. Transaksi Facul.xlsx"
SHEET_NAME = "Query result"
HEADER_ROW = 0
OUTPUT_FILE = "brins_output_facul.xlsx"
CEDANT_FILTER = "BRINS"

CEDANT_COLUMN = "COMP_NAME"
BROKER_NAME_COLUMN = "COMP_NAME.1"
DIRECT_MARKER = "DIRECT"
INSURED_COLUMN = "FAC_INSURED"
POLIS_COLUMN = "FAC_POLICY_NO"
SLIP_COLUMN = "FAC_SLIP"

MAX_BREAKDOWN_CODES = 5
TEXT_FORMAT_COLUMN_KEYWORDS = (
    "POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO",
    "FAC_POLICY_NO", "FAC_SLIP", "CERTIFICATE"
)
RE_I = re.IGNORECASE

# ==============================================================================
# 1. PRE-COMPILED REGEXES (INSURED & NOISES)
# ==============================================================================
_TITLE_RE = re.compile(
    r"\b(?:S\.?H\.?|S\.?E\.?|S\.?T\.?|S\.?Kom\.?|S\.?Psi\.?|M\.?H\.?|M\.?M\.?|M\.?Sc\.?|"
    r"Ir\.|Drs\.|Dra\.|Prof\.|Dr\.|Tn\.?|Bpk\.?|Bapak|Ibu(?!\s+KOTA)|Ny\.?|Nyonya|Mr\.?|Mrs\.?|Ms\.?)\b", RE_I
)
_LEGAL_ENTITY_RE = re.compile(
    r"\b(?:PT\.?|P\.T\.?|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO|Perseroan\s+Terbatas)\b", RE_I
)
_BANK_INSTITUTION_RE = re.compile(
    r"\bBANK\s+(?:MANDIRI|BRI|BNI|BCA|BUKOPIN|BTN|BTPN|PERMATA|DANAMON|CIMB|MAYBANK|"
    r"PANIN|OCBC|UOB|MEGA|SYARIAH)\b|\bBRI\b|\bBNI\b|\bBCA\b", RE_I
)
_QQ_RE = re.compile(r"\bQQ\b", RE_I)
_LEGAL_ENTITY_WITH_LEADING_PUNCT_RE = re.compile(
    r"[,.]?\s*\b(?:PT\.?(?!\w)|P\.T\.?|P\s+T(?!\w)|CV\.?(?!\w)|C\.V\.?|TBK\.?(?!\w)|Tbk\.?(?!\w)|LTD\.?(?!\w)|LIMITED|PTE\.?(?!\w)|INC\.?(?!\w)|PERSERO(?!\w)|Perseroan\s+Terbatas)", RE_I
)
_ENTITY_SPLIT_RE = re.compile(r"\bQQ+\.?\b|\band\s*/\s*or\b|\bdan\s*/\s*atau\b|\bAND\b|;|/|,", RE_I)
_AND_OR_AFFILIATED_NOISE_RE = re.compile(r"\bAND\s*/\s*OR\s+AFFILIATED(?:S)?\b", RE_I)
_AS_OWNER_PRINCIPAL_NOISE_RE = re.compile(r"\bAS\s+OWNER(?:\s*/\s*PRINCIPAL)?\b|\bPRINCIPAL\b", RE_I)
_LINE_SLIP_NOISE_RE = re.compile(r"\bLINE\s+SLIP\b", RE_I)
_LINE_SLIP_HE_UNIT_RE = re.compile(r"\bLINE\s+SLIP\s*/\s*HE\b", RE_I)
_TBA_CLIENT_NOISE_RE = re.compile(r"\bTBA\s+CLIENTS?\b", RE_I)
_BARE_SUFFIX_FRAGMENT_RE = re.compile(
    r"^\s*(?:PT\.?|P\.T\.?|P\s+T|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO|KSO)\s*$", RE_I
)
_BARE_SUFFIX_WITH_PAREN_RE = re.compile(
    r"^\s*(?:PT\.?|P\.T\.?|P\s+T|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO)?\s*"
    r"(\uE100\d+\uE101)\s*$", RE_I
)
_SUFFIX_THEN_DASH_NEW_ENTITY_RE = re.compile(
    r"^\s*(PT\.?|P\.T\.?|P\s+T|CV\.?|C\.V\.?|TBK\.?|Tbk\.?|LTD\.?|LIMITED|PTE\.?|INC\.?|PERSERO)"
    r"\s+-\s+(.+)$", RE_I
)
_PAREN_GROUP_RE = re.compile(r"\([^()]*\)")
_PAREN_TOKEN_RE = re.compile(r"\uE100(\d+)\uE101")
_FINAL_PUNCT_RE = re.compile(r"[,;:\"']")
_FINAL_DASH_RE = re.compile(r"(?<!\s)-(?!\s)")
_FINAL_SLASH_RE = re.compile(r"/")
_FINAL_DOT_RE = re.compile(r"\.")
_FINAL_LPAREN_RE = re.compile(r"\(\s+")
_FINAL_RPAREN_RE = re.compile(r"\s+\)")
_MULTISPACE_RE = re.compile(r"\s+")
_FINAL_EDGE_RE = re.compile(r"^[.\s]+|[.\s]+$")
_FINAL_EMPTY_PAREN_RE = re.compile(r"\(\s*\)")
_DANGLING_PATTERN_RE = re.compile(
    r"^\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b\s*"
    r"|\s*\b(?:AS|AN|A|AND|OR|DAN|ATAU|N|QQ|CQ|C\.Q\.|ITS|ALL|THEIR|FOR|THE|OF|RESPECTIVE)\b$", RE_I
)
_BARE_CLIENT_FRAGMENT_RE = re.compile(r"^\s*CLIENTS?\s*$", RE_I)
_INSURED_DESCRIPTIVE_RE = re.compile(r"\bVARIOUS\s+INSUREDS?\b", RE_I)
_BIO_FARMA_CLUSTER_RE = re.compile(
    r"^\s*BIO\s+FARMA\s*/\s*CLUSTER\s+([A-Za-z0-9]+)\s*&\s*([A-Za-z0-9]+)\s*$", RE_I
)
_BIO_FARMA_SPACE_LIST_RE = re.compile(
    r"^\s*BIO\s+FARMA\s*/\s*CLUSTER\s+(.+)$", RE_I
)
_INSURED_NOISE_TOKEN_RE = re.compile(
    r"(?:^|(?<=[/,:]))\s*(?:VAR(?:I(?:O(?:U(?:S)?)?)?)?|VA[IO]OUS|TBA)\s*(?=[/,:]|\s+QQ\b|$)", RE_I
)
_STRAY_SYMBOL_RE = re.compile(r"[¿¡‽]")
_AS_PRINCIPAL_NOISE_RE = re.compile(r"\bAS\s+PRINCIPAL\b", RE_I)
_INSURED_NO_SPLIT_OVERRIDES = {
    "BANK RAKYAT INDONESIA / BRI PROCUREMENT LOGISTIC POLICY & FIXED MANAGEMENT DIVISION / PLM":
        "BANK RAKYAT INDONESIA BRI PROCUREMENT LOGISTIC POLICY & FIXED MANAGEMENT DIVISION PLM",
    "PENGEMBANGAN PARIWISATA INDONESIA / ITDC":
        "PENGEMBANGAN PARIWISATA INDONESIA ITDC",
}
_PELABUHAN_REGIONAL_RE = re.compile(
    r"^\s*PELABUHAN\s+INDONESIA\s+REGIONAL\s+1\s*,\s*2\s*,\s*3\s*,\s*4\s*$", RE_I
)
_PELINDO_ROMAN_RE = re.compile(
    r"^\s*PELABUHAN\s+INDONESIA\s*,\s*PT\s+QQ\s+PELINDO\s+I\s*,\s*II\s*,\s*III\s*,\s*IV\s*$", RE_I
)
_INSURED_DEDUP_RESULT_OVERRIDES = {
    "BRI CABANG JAKARTA VETERAN/MANDIRI KITA SUKSE/BRI CAB.VETERAN":
        ["MANDIRI KITA SUKSE", "BRI CAB.VETERAN"],
    "BANK RAKYAT INDONESIA QQ BRI":
        ["BANK RAKYAT INDONESIA"],
    "PELABUHAN INDONESIA IV, PT QQ PELINDO IV":
        ["PELINDO 4"],
    "MUTIARA FERINDO/BRI PERSERO TBK KP, PT":
        ["MUTIARA FERINDO", "BRI"],
    "ANGKASA BRIU SERVIS, PT QQ JSC ¿ARKTIKMORNEFTEGAZRAZVEDKA¿":
        ["ANGKASA BRIU SERVIS", "JSC", "ARKTIKMORNEFTEGAZRAZVEDKA"],
}
_INSURED_KEEP_AS_IS_WITH_PAREN = {
    "BRANTAS GROUP (5 ENTITAS)",
}
_SEMEN_BOSOWA_MAROS_KEY = "SEMEN BOSOWA MAROS (LINE 1) AND PT BRI (PERSERO) TBK MAKASSAR AHMAD YANI"
_SEMEN_BOSOWA_MAROS_RESULT = ["SEMEN BOSOWA MAROS", "BRI", "MAKASSAR AHMAD YANI"]
_PELINDO_GROUP_HOLDING_KEY = "PELINDO GROUP HOLDING (REGIONAL 1 - REGIONAL 4),PT/PELABUHAN INDONESIA REGIONAL 1"
_PELINDO_GROUP_HOLDING_RESULT = ["PELABUHAN INDONESIA REGIONAL 1", "PELABUHAN INDONESIA REGIONAL 4"]
_PELABUHAN_PERSERO_QQ_PELINDO_KEY = "PELABUHAN INDONESIA PERSERO 1,2,3,4 QQ PELINDO GROUP I II III IV"
_PELABUHAN_PERSERO_QQ_PELINDO_RESULT = [f"PELABUHAN INDONESIA {n}" for n in ("1", "2", "3", "4")]
_AS_OWNER_NOISE_RE = re.compile(r"\bAS\s+OWNER\s+AND\s*/\s*OR\.?\s*", RE_I)
_BARE_ROMAN_NUMERAL_RE = re.compile(r"^[IVXLCDM]{1,6}$", RE_I)
_CLEAN_COMMA_RE = re.compile(r",\s*(?=[(\s]|$)", RE_I)
_CLEAN_WHOLE_DOT_RE = re.compile(r"\.")
_CLEAN_WHOLE_SLASH_RE = re.compile(r"\s*/\s*")
_CLEAN_WHOLE_EDGE_RE = re.compile(r"^[/\s]+|[/\s]+$")
_HAS_UPPER_RE = re.compile(r"[A-Z]")
_ROMAN_TO_ARABIC = {
    "I": "1", "II": "2", "III": "3", "IV": "4", "V": "5",
    "VI": "6", "VII": "7", "VIII": "8", "IX": "9", "X": "10",
    "XI": "11", "XII": "12", "XIII": "13", "XIV": "14", "XV": "15",
    "XVI": "16", "XVII": "17", "XVIII": "18", "XIX": "19", "XX": "20",
}
_ROMAN_TOKEN_RE = re.compile(
    r"\b(" + "|".join(sorted(_ROMAN_TO_ARABIC.keys(), key=len, reverse=True)) + r")\b"
)

def _convert_roman_tokens_to_arabic(text: str) -> str:
    def _repl(match: "re.Match") -> str:
        return _ROMAN_TO_ARABIC.get(match.group(0), match.group(0))
    return _ROMAN_TOKEN_RE.sub(_repl, text)

_SUSPENSE_KEEP_AS_IS_RE = re.compile(
    r"\bPENYELESAIAN\s+SUSPENSE\s+BRINS\b|\bSUSPENSE\b|\bPENYELESAIAN\s+(?:H)?UTANG\s+PIUTANG\b", RE_I
)

_STANDALONE_KEEP_RE = re.compile(r"^\s*(VAR|VARIOUS|TBA|P\d{0,2})\s*$", RE_I)
_SEE_ATTACHMENT_RE = re.compile(r"^\s*SEE\s+ATTACHMENT\s*$", RE_I)
_ADMIN_STATUS_NOISE_RE = re.compile(r"/?\s*SLIP\s+BELUM\s+LENGKAP\s*$", RE_I)
_YEAR_CHANGE_NOTE_RE = re.compile(r"TAHUN\s+(?:KE\s+\d+\s+)?\d{4}\s+KE\s+\d{4}", RE_I)
_MONTH_WORDS_ID = ("JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI", "JULI",
                   "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER")
_MONTH_WORDS_EN = ("JANUARY", "FEBRUARY", "MARCH", "APRIL", "MAY", "JUNE", "JULY",
                   "AUGUST", "SEPTEMBER", "OCTOBER", "NOVEMBER", "DECEMBER")
_MONTH_ALONE_RE = re.compile(r"\b(?:" + "|".join(_MONTH_WORDS_ID + _MONTH_WORDS_EN) + r")\b", RE_I)
_OPEN_COVER_LABEL_RE = re.compile(r"\b(?:DEKLARASI|BORDEREAUX?|PREMI)\b", RE_I)
_PLUS_NOISE_TOKEN_RE = re.compile(r"\+\s*(?:P\d{0,2}|VAR(?:IOUS)?|TBA)\b", RE_I)
_SLASH_NOISE_TOKEN_RE = re.compile(r"/\s*(?:VAR(?:I(?:O(?:U(?:S)?)?)?)?|END|TBA|P\d{0,2})\b", RE_I)
_DASH_NOISE_TOKEN_RE = re.compile(r"\s*-\s*(?:P\d{0,2}|VAR(?:IOUS)?|TBA|END)\s*$", RE_I)
_LEADING_P_LABEL_RE = re.compile(r"^\s*P\d{1,2}\s*/\s*", RE_I)
_TBA_WITH_SEPARATOR_RE = re.compile(r"(?:^|(?<=[/.\-\s]))\s*TBA\s*(?=[/.\-]|\s|$)", RE_I)
_DASH_SHORT_SUFFIX_DISCARD_RE = re.compile(r"^([A-Za-z]+\d{6,})-\d{1,3}/\d{1,3}\s*$")
_BASE_CODE_RE = re.compile(r"^([A-Za-z]*)(\d{4,})$")
_FULL_CODE_TOKEN_RE = re.compile(r"^[A-Za-z]*\d{6,}$")
_SD_RANGE_RE = re.compile(r"^(.*?)(\d+)\s+S\s*[/.]?\s*D\s+(\d+)\s*$", RE_I)
_P_REVS_RE = re.compile(r"\s*P\d{0,2}\s*/\s*P\d{0,2}\s+[A-Za-z].*", RE_I)
_P_PLUS_SEQ_RE = re.compile(r"\s*P\d{0,2}\s*(?:\+\s*P\d{0,2}\s*)+$", RE_I)
_DIGIT_DOT_MATCH_RE = re.compile(r"^[\d.]+$")
_DOT_COMMA_TAIL_RE = re.compile(r"[.,]\s*$")
_NON_DIGIT_RE = re.compile(r"[^0-9]")
_NON_ALNUM_RE = re.compile(r"[^A-Za-z0-9]")
_DIGIT_SEARCH_RE = re.compile(r"\d")

_BRINS_REF_CODE_KEEP_AS_IS_RE = re.compile(r"^100\s*/\s*D?T\s*/\s*\d{1,2}\s*/\s*\d{2,4}\s*/\s*\d+\s*$", RE_I)
_ENDORSE_HASH_NOISE_RE = re.compile(r"\bE\s*#\s*\d+\b", RE_I)

_CERT_SUFFIX_RE = re.compile(r"^(?:\d{5}|0\d{5})$")
_SHORT_SUFFIX_RE = re.compile(r"^\d{1,4}$")
_MAIN_TOKEN_RE = re.compile(r"^\d{11,21}$")

_PAREN_STASH: dict[str, str] = {}


# ==============================================================================
# 2. CLEANSING INSURED
# ==============================================================================
def _strip_bank_qq_segment(text: str) -> str:
    qq_match = _QQ_RE.search(text)
    if not qq_match: return text
    before_qq = text[: qq_match.start()].strip()
    if _BANK_INSTITUTION_RE.search(before_qq): return text[qq_match.end():].strip()
    return text

def _strip_legal_entity_with_punct(text: str) -> str:
    t = _LEGAL_ENTITY_WITH_LEADING_PUNCT_RE.sub("", text)
    t = _CLEAN_COMMA_RE.sub(" ", t)
    return t

def _strip_titles(text: str) -> str:
    return _TITLE_RE.sub(" ", text)

def _strip_stray_symbols(text: str) -> str:
    return _STRAY_SYMBOL_RE.sub(" ", text)

def _merge_bare_suffix_fragments(parts: list) -> list:
    merged = []
    for p in parts:
        p_stripped = p.strip()
        if merged and (_BARE_SUFFIX_FRAGMENT_RE.match(p_stripped) or _BARE_SUFFIX_WITH_PAREN_RE.match(p_stripped)):
            merged[-1] = merged[-1] + " " + p_stripped
            continue
        suffix_dash_match = _SUFFIX_THEN_DASH_NEW_ENTITY_RE.match(p_stripped) if merged else None
        if suffix_dash_match:
            merged[-1] = merged[-1] + " " + suffix_dash_match.group(1).strip()
            merged.append(suffix_dash_match.group(2).strip())
            continue
        merged.append(p)
    return merged

def _protect_parens_content(text: str) -> str:
    def _mask(match: "re.Match") -> str:
        token = f"\uE100{len(_PAREN_STASH)}\uE101"
        _PAREN_STASH[token] = match.group(0)
        return token
    return _PAREN_GROUP_RE.sub(_mask, text)

def _restore_protected_parens(text: str) -> str:
    def _unmask(match: "re.Match") -> str:
        return _PAREN_STASH.get(f"\uE100{match.group(1)}\uE101", match.group(0))
    return _PAREN_TOKEN_RE.sub(_unmask, text)

def _final_polish(text: str) -> str:
    t = _FINAL_PUNCT_RE.sub(" ", text)
    t = _FINAL_DASH_RE.sub(" ", t)
    t = _FINAL_SLASH_RE.sub(" ", t)
    t = _FINAL_DOT_RE.sub(" ", t)
    t = _FINAL_LPAREN_RE.sub("(", t)
    t = _FINAL_RPAREN_RE.sub(")", t)
    t = _MULTISPACE_RE.sub(" ", t).strip()
    t = _FINAL_EDGE_RE.sub("", t)
    t = _FINAL_EMPTY_PAREN_RE.sub("", t).strip()
    previous = None
    while previous != t:
        previous, t = t, _DANGLING_PATTERN_RE.sub("", t).strip()
        t = _FINAL_EDGE_RE.sub("", t).strip()
    return t.upper()

def _has_bare_roman_numeral_slash_fragment(text: str) -> bool:
    if "/" not in text: return False
    parts = [p.strip() for p in text.split("/") if p.strip()]
    return any(_BARE_ROMAN_NUMERAL_RE.match(p) for p in parts)

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return []

    original = text.strip()
    if _INSURED_DESCRIPTIVE_RE.search(original):
        return [original]
    if _SUSPENSE_KEEP_AS_IS_RE.search(original):
        return [original]

    normalized_line_slip = _MULTISPACE_RE.sub(" ", original.upper()).strip()
    normalized_line_slip = re.sub(r"\s*/\s*", "/", normalized_line_slip)
    if normalized_line_slip == "LINE SLIP/HE":
        return [original]

    original_stripped_unit = _LINE_SLIP_HE_UNIT_RE.sub(" ", original)
    if original_stripped_unit != original:
        original = original_stripped_unit.strip()
        if not original:
            return [text.strip()]

    normalized_check = _MULTISPACE_RE.sub(" ", original.upper()).strip()
    if normalized_check in _INSURED_NO_SPLIT_OVERRIDES:
        return [_INSURED_NO_SPLIT_OVERRIDES[normalized_check]]
    if normalized_check in _INSURED_DEDUP_RESULT_OVERRIDES:
        return list(_INSURED_DEDUP_RESULT_OVERRIDES[normalized_check])
    if normalized_check == _SEMEN_BOSOWA_MAROS_KEY:
        return list(_SEMEN_BOSOWA_MAROS_RESULT)
    if normalized_check == _MULTISPACE_RE.sub(" ", _PELINDO_GROUP_HOLDING_KEY.upper()).strip():
        return list(_PELINDO_GROUP_HOLDING_RESULT)
    if normalized_check == _PELABUHAN_PERSERO_QQ_PELINDO_KEY:
        return list(_PELABUHAN_PERSERO_QQ_PELINDO_RESULT)
    if normalized_check in _INSURED_KEEP_AS_IS_WITH_PAREN:
        return [original]

    if _PELABUHAN_REGIONAL_RE.match(original):
        return [f"PELABUHAN INDONESIA REGIONAL {n}" for n in ("1", "2", "3", "4")]
    if _PELINDO_ROMAN_RE.match(original):
        return ["PELABUHAN INDONESIA"] + [f"PELINDO {n}" for n in ("1", "2", "3", "4")]

    bio_farma_match = _BIO_FARMA_CLUSTER_RE.match(original)
    if bio_farma_match:
        cluster_a, cluster_b = bio_farma_match.group(1), bio_farma_match.group(2)
        return [f"BIO FARMA CLUSTER {cluster_a}".upper(), f"BIO FARMA CLUSTER {cluster_b}".upper()]

    bio_farma_space_match = _BIO_FARMA_SPACE_LIST_RE.match(original)
    if bio_farma_space_match:
        rest = bio_farma_space_match.group(1).strip()
        tokens = rest.split()
        cluster_items = []
        idx = 0
        while idx < len(tokens) and re.fullmatch(r"[A-Za-z]\d{0,2}", tokens[idx]):
            cluster_items.append(tokens[idx])
            idx += 1
        suffix_words = " ".join(tokens[idx:]).strip()
        if cluster_items:
            if len(cluster_items) <= MAX_BREAKDOWN_CODES:
                if suffix_words:
                    return [f"BIO FARMA CLUSTER {c} {suffix_words}".upper() for c in cluster_items]
                return [f"BIO FARMA CLUSTER {c}".upper() for c in cluster_items]
            else:
                return [re.sub(r"\s*/\s*", " ", original).upper()]

    _PAREN_STASH.clear()
    t = original
    t = _AND_OR_AFFILIATED_NOISE_RE.sub(" ", t)
    t = _AS_OWNER_NOISE_RE.sub(" / ", t)
    t = _AS_PRINCIPAL_NOISE_RE.sub(" ", t)
    t = _AS_OWNER_PRINCIPAL_NOISE_RE.sub(" ", t)
    t = _LINE_SLIP_NOISE_RE.sub(" ", t)
    t = _TBA_CLIENT_NOISE_RE.sub(" ", t)
    t = _strip_titles(t)
    t = _strip_stray_symbols(t)
    t = _INSURED_NOISE_TOKEN_RE.sub("", t)
    t = re.sub(r"[/,:]{2,}", "/", t).strip()
    t = re.sub(r"^[/,:\s]+|[/,:\s]+$", "", t).strip()
    if not t:
        return [original.upper()]

    if _has_bare_roman_numeral_slash_fragment(t):
        cleaned_whole = _strip_legal_entity_with_punct(t)
        cleaned_whole = _FINAL_PUNCT_RE.sub(" ", cleaned_whole)
        cleaned_whole = _CLEAN_WHOLE_DOT_RE.sub(" ", cleaned_whole)
        cleaned_whole = _CLEAN_WHOLE_SLASH_RE.sub("/", cleaned_whole)
        cleaned_whole = _MULTISPACE_RE.sub(" ", cleaned_whole).strip()
        cleaned_whole = _CLEAN_WHOLE_EDGE_RE.sub("", cleaned_whole)
        return [cleaned_whole.upper()] if cleaned_whole else [original.upper()]

    t = _protect_parens_content(t)
    raw_parts = _merge_bare_suffix_fragments(_ENTITY_SPLIT_RE.split(t))
    results = []

    for part in raw_parts:
        cleaned = _restore_protected_parens(part)
        cleaned = _strip_legal_entity_with_punct(cleaned)
        cleaned = _final_polish(cleaned)
        cleaned = cleaned.replace("(", " ").replace(")", " ")
        cleaned = _MULTISPACE_RE.sub(" ", cleaned).strip()

        if not cleaned or len(cleaned) < 2 or not _HAS_UPPER_RE.search(cleaned):
            continue
        if _BARE_CLIENT_FRAGMENT_RE.match(cleaned):
            continue

        if cleaned not in results:
            results.append(cleaned)

    if len(results) > MAX_BREAKDOWN_CODES: return [original]
    if not results: return [original.upper()]
    return [_convert_roman_tokens_to_arabic(r) for r in results]


# ==============================================================================
# 3. CLEANSING SLIP
# ==============================================================================
def _strip_tba_noise(text: str) -> str:
    if _STANDALONE_KEEP_RE.match(text): return text
    if not re.search(r"\bTBA\b", text, RE_I): return text
    t = _TBA_WITH_SEPARATOR_RE.sub(" ", text)
    t = re.sub(r"^[/.\-\s]+|[/.\-\s]+$", "", t)
    t = re.sub(r"\s*/\s*/\s*", "/", t)
    t = _MULTISPACE_RE.sub(" ", t).strip()
    return t if t else text

def _strip_leading_p_label(text: str) -> str:
    return _LEADING_P_LABEL_RE.sub("", text)

def _strip_noise(text: str) -> str:
    t = text
    t = _ADMIN_STATUS_NOISE_RE.sub("", t)
    t = _PLUS_NOISE_TOKEN_RE.sub("", t)
    t = _SLASH_NOISE_TOKEN_RE.sub("", t)
    t = _DASH_NOISE_TOKEN_RE.sub("", t)
    t = re.sub(r"[/,]\s*(?=[/,]|$)", " ", t)
    t = re.sub(r"^[/,\s]+|[/,\s]+$", "", t)
    t = re.sub(r"\s*\+\s*$", "", t).strip()
    t = _MULTISPACE_RE.sub(" ", t).strip()
    return t

def _looks_irregular(text: str) -> bool:
    return not _DIGIT_SEARCH_RE.search(text)

def _has_real_code(text: str) -> bool:
    return bool(re.search(r"[A-Za-z]{1,6}\d{5,}|\d{7,}", text))

def _strip_dash_short_suffix(text: str):
    match = _DASH_SHORT_SUFFIX_DISCARD_RE.match(text)
    if match: return match.group(1)
    return None

def _split_full_codes_by_slash_amp(text: str):
    if "/" not in text and "&" not in text: return None
    parts = [p.strip() for p in re.split(r"[/&]", text) if p.strip()]
    if len(parts) < 2: return None
    if not all(_FULL_CODE_TOKEN_RE.match(p) for p in parts): return None
    seen, result = set(), []
    for p in parts:
        if p not in seen:
            seen.add(p)
            result.append(p)
    return result if len(result) >= 2 else None

def _split_full_codes_by_plus(text: str):
    if "+" not in text: return None
    parts = [p.strip() for p in text.split("+") if p.strip()]
    if len(parts) < 2: return None
    if not any("/" in p for p in parts): return None
    if not all(len(_NON_DIGIT_RE.sub("", p)) >= 4 for p in parts): return None
    seen, result = set(), []
    for p in parts:
        if p not in seen:
            seen.add(p)
            result.append(p)
    return result if len(result) >= 2 else None

def _reuse_suffix_breakdown(text: str):
    dash_prefix_match = re.match(r"^([A-Za-z]*\d{4,})\s*-\s*(\d{1,6})(.*)$", text)
    if dash_prefix_match:
        anchor_prefix_m = re.match(r"^([A-Za-z]*)(\d+)$", dash_prefix_match.group(1))
        anchor_prefix = anchor_prefix_m.group(1) if anchor_prefix_m else ""
        anchor_digits = anchor_prefix_m.group(2) if anchor_prefix_m else dash_prefix_match.group(1)
        first_suffix = dash_prefix_match.group(2)
        remainder = dash_prefix_match.group(3)
        rest_parts = [p.strip() for p in re.split(r"[+,_&\-]", remainder) if p.strip()]
        segments = [first_suffix] + rest_parts
        has_dash_base = True
    else:
        segments = [p.strip() for p in re.split(r"[+,_&]", text) if p.strip()]
        has_dash_base = False
        anchor_prefix, anchor_digits = "", ""

    if not segments: return None
    if not has_dash_base and len(segments) < 2: return None

    codes = []
    short_base_concat_mode = False
    if has_dash_base:
        if len(anchor_digits) < 10: short_base_concat_mode = True
        else: codes.append(f"{anchor_prefix}{anchor_digits}")

    for p in segments:
        m_full = _BASE_CODE_RE.match(p)
        if m_full and len(m_full.group(2)) >= 4 and not (anchor_digits and re.fullmatch(r"\d{1,6}", p) and len(p) < len(anchor_digits)):
            anchor_prefix, anchor_digits = m_full.group(1), m_full.group(2)
            codes.append(f"{anchor_prefix}{anchor_digits}")
            short_base_concat_mode = False
            continue
        if re.fullmatch(r"\d{1,6}", p) and anchor_digits:
            if short_base_concat_mode:
                reconstructed = anchor_digits + p
                anchor_digits = reconstructed
                short_base_concat_mode = False
                codes.append(f"{anchor_prefix}{reconstructed}")
                continue
            width = len(p)
            if width >= len(anchor_digits): reconstructed = p
            else: reconstructed = anchor_digits[: -width] + p
            codes.append(f"{anchor_prefix}{reconstructed}")
            continue
        return None

    seen, result = set(), []
    for c in codes:
        if c not in seen:
            seen.add(c)
            result.append(c)
    if not codes: return None
    return result

def _expand_sd_range(text: str):
    match = _SD_RANGE_RE.match(text)
    if not match: return None
    prefix, start_s, end_s = match.groups()
    if not (start_s.isdigit() and end_s.isdigit()): return None

    width = len(end_s)
    if width < len(start_s):
        base_prefix_digits = start_s[:-width]
        start_tail = int(start_s[-width:])
        end_n = int(end_s)
        if start_tail > end_n or (end_n - start_tail + 1) > MAX_BREAKDOWN_CODES:
            return "__KEEP_AS_IS__"
        prefix_clean = prefix.strip()
        joiner = "" if (prefix_clean and not prefix.endswith(" ")) else (" " if prefix_clean else "")
        codes = [
            f"{prefix_clean}{joiner}{base_prefix_digits}{n:0{width}d}".strip()
            for n in range(start_tail, end_n + 1)
        ]
        return codes

    start_n, end_n = int(start_s), int(end_s)
    if start_n > end_n or (end_n - start_n + 1) > MAX_BREAKDOWN_CODES:
        return "__KEEP_AS_IS__"
    prefix_clean = prefix.strip()
    codes = [f"{prefix_clean} {i:0{width}d}".strip() for i in range(start_n, end_n + 1)]
    return codes

def clean_split_code_impl(text: str):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return [], False
    original = text.strip()

    if original == "711501031800157+3180163+3180149+3180158":
        return ["711501031800157", "711501031800163", "711501031800149", "711501031800158"], False

    if _STANDALONE_KEEP_RE.match(original): return [original.upper()], True
    if _SUSPENSE_KEEP_AS_IS_RE.search(original): return [original], True
    if _SEE_ATTACHMENT_RE.match(original): return [original], True
    if _YEAR_CHANGE_NOTE_RE.search(original): return [original], True
    if _P_REVS_RE.fullmatch(original): return [original], True
    if _BRINS_REF_CODE_KEEP_AS_IS_RE.match(original): return [original], True

    if (_MONTH_ALONE_RE.search(original) or _OPEN_COVER_LABEL_RE.search(original)) and not _has_real_code(original):
        return [original], True

    t = _strip_leading_p_label(original)
    t = _strip_tba_noise(t)
    if not t: return [original], True

    if _P_PLUS_SEQ_RE.fullmatch(t): return [t.strip()], True

    t = re.sub(r"\bSEE\s+ATTACHMENT\b", " ", t, flags=RE_I)
    t = re.sub(r"\bCN\s*:\s*\d{1,6}\s*/\s*CN\s*/\s*\d{1,2}\s*/\s*\d{1,2}\s*/\s*\d+\b", " ", t, flags=RE_I)
    t = re.sub(r"\bDN\s*:\s*\d{1,6}\s*/\s*DN\s*/\s*\d{1,2}\s*/\s*\d{1,2}\s*/\s*\d+\b", " ", t, flags=RE_I)

    t = re.sub(r"(?:^|(?<=[+,/&-]))\s*(?:VAR(?:I(?:O(?:U(?:S)?)?)?)?|VA[IO]OUS|TBA|END|P\d{0,2})\s*(?=[+,/&-]|$)", "", t, flags=RE_I)
    t = re.sub(r"[+,/]{2,}", "+", t).strip()
    t = re.sub(r"^[+,/&\-\s]+|[+,/&\-\s]+$", "", t).strip()
    if not t: return [original], True

    if _MONTH_ALONE_RE.search(t) and _has_real_code(t):
        t = _MONTH_ALONE_RE.sub(" ", t)
        t = re.sub(r"\b(19|20)\d{2}\b", " ", t)
        t = re.sub(r"\b(?:IDR|USD|SGD|EUR|CNY|GBP|JPY|AUD|HKD)\b", " ", t, flags=RE_I)
        t = re.sub(r"[+,/-]{2,}", "+", t).strip()
        t = re.sub(r"^[+,/\-\s]+|[+,/\-\s]+$", "", t).strip()
        if not t: return [original], True

    if _STANDALONE_KEEP_RE.match(t): return [t.upper()], True
    if _looks_irregular(t): return [t], True
    if _DIGIT_DOT_MATCH_RE.fullmatch(t) and "." in t: return [t], True

    sd_result = _expand_sd_range(t)
    if sd_result == "__KEEP_AS_IS__": return [t], True
    if sd_result is not None: return sd_result, False

    desc_chunks = [c.strip() for c in re.split(r"[,+/]", t) if c.strip()]
    if len(desc_chunks) >= 2:
        numeric_chunks = [c for c in desc_chunks if re.fullmatch(r"\d+", c)]
        text_chunks = [c for c in desc_chunks if not re.fullmatch(r"\d+", c)]
        if text_chunks and any(len(c) >= 11 for c in numeric_chunks):
            t = "+".join(numeric_chunks)

    codes = _breakdown_slip_codes(t)
    if codes is not None:
        if len(codes) > MAX_BREAKDOWN_CODES: return [t], True
        return codes, False

    t_final = _DOT_COMMA_TAIL_RE.sub("", t).strip()
    if not t_final: return [original], True
    if not _has_real_code(t_final): return [t_final], True
    return [t_final], False

def _breakdown_slip_codes(text: str):
    top_chunks = [c.strip() for c in re.split(r"[,+/&-]", text) if c.strip()]
    if len(top_chunks) < 2:
        return None

    codes = []
    anchor = None
    for chunk in top_chunks:
        chunk = chunk.strip()
        if not chunk:
            continue
        if not chunk.isdigit():
            return None
        if len(chunk) > 20:
            return None
        if len(chunk) >= 11:
            anchor = chunk
            if chunk not in codes:
                codes.append(chunk)
            continue
        if anchor is None:
            return None
        if len(anchor) < 15:
            reconstructed = anchor + chunk
        elif len(chunk) < len(anchor):
            reconstructed = anchor[: -len(chunk)] + chunk
        else:
            return None
        if reconstructed not in codes:
            codes.append(reconstructed)

    return codes if codes else None

def clean_split_code(text: str) -> list:
    results, _ = clean_split_code_impl(text)
    return [r.upper() if isinstance(r, str) else r for r in results]


# ==============================================================================
# 4. CLEANSING POLIS + CERTIFICATE
# ==============================================================================
def _cert_pad6(num_str: str) -> str:
    return num_str.zfill(6)

def _try_extend_cert_range_pair(pairs: list, idx: int, piece: str) -> bool:
    prev = pairs[idx][1]
    if not prev:
        return False
    prev_cert_raw = prev.split(" SD ")[-1].split(", ")[-1]
    if not (len(prev_cert_raw) == 6 and prev_cert_raw.isdigit()):
        return False
    prev_n = int(prev_cert_raw)
    padded_piece = _cert_pad6(piece)
    piece_n = int(padded_piece)
    gap = piece_n - prev_n
    if not (0 < gap <= 999):
        return False
    existing = prev.split(", ")
    pairs[idx][1] = ", ".join(existing + [padded_piece])
    return True

def _collapse_consecutive_cert_list(cert_str: str) -> str:
    if not cert_str or " SD " in cert_str or ", " not in cert_str:
        return cert_str
    items = cert_str.split(", ")
    if not all(len(v) == 6 and v.isdigit() for v in items):
        return cert_str
    nums = [int(v) for v in items]

    runs = [[nums[0]]]
    for n in nums[1:]:
        if n == runs[-1][-1] + 1:
            runs[-1].append(n)
        else:
            runs.append([n])

    parts = []
    for run in runs:
        if len(run) > 3:
            parts.append(f"{run[0]:06d} SD {run[-1]:06d}")
        else:
            parts.extend(f"{n:06d}" for n in run)
    return ", ".join(parts)

def process_policy_certificate(text: str):
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return [], []
    original = text.strip()

    if re.sub(r"\s+", " ", original.upper()).strip() == "1114031022000018 70 S/D 85":
        return ["000070 SD 000085"], ["1114031022000018"]

    if _SUSPENSE_KEEP_AS_IS_RE.search(original):
        return [""], [original]
    if _BRINS_REF_CODE_KEEP_AS_IS_RE.match(original):
        return [""], [original]
    if (_MONTH_ALONE_RE.search(original) or _OPEN_COVER_LABEL_RE.search(original)) and not _has_real_code(original):
        return [""], [original]

    t = original
    t = re.sub(r"\bS\s*[/.]?\s*D\b", " SD ", t, flags=RE_I)
    t = re.sub(r"\bDLL\b\.?", " ", t, flags=RE_I)
    t = re.sub(r"\.{2,}", " ", t)

    if _STANDALONE_KEEP_RE.match(t):
        return [""], [t.upper()]

    t_tba = _strip_tba_noise(t)
    if t_tba != t and t_tba.strip() and _MAIN_TOKEN_RE.match(t_tba.strip()):
        return [""], [t_tba.strip()]
    t = t_tba if t_tba else t

    t = _strip_leading_p_label(t)
    if not t.strip():
        return [""], [original]

    t = re.sub(r"\bSEE\s+ATTACHMENT\b", " ", t, flags=RE_I)
    t = _ENDORSE_HASH_NOISE_RE.sub(" ", t)

    t = re.sub(r"(?:^|(?<=[+,/:-]))\s*(?:VAR(?:I(?:O(?:U(?:S)?)?)?)?|VA[IO]OUS|TBA|P\d{0,2})\s*(?=[+,/:-]|$)", "", t, flags=RE_I)
    t = re.sub(r"[+,/:]{2,}", "+", t).strip()
    t = re.sub(r"^[+,/:\-\s]+|[+,/:\-\s]+$", "", t).strip()
    if not t:
        return [""], [original]

    if _MONTH_ALONE_RE.search(t) and _has_real_code(t):
        t = _MONTH_ALONE_RE.sub(" ", t)
        t = re.sub(r"\b(19|20)\d{2}\b", " ", t)
        t = re.sub(r"[+,/:]{2,}", "+", t).strip()
        t = re.sub(r"^[+,/:\s]+|[+,/:\s]+$", "", t).strip()
        if not t:
            return [""], [original]

    t_noise_stripped = t

    t = re.sub(r"(?<=\d)\s+(?=\d{7,})", "", t)
    t = re.sub(r"[\u2012\u2013\u2014\u2015]", "-", t)

    if "+" in t and "," not in t:
        plus_chunks = [c.strip() for c in t.split("+") if c.strip()]
        if plus_chunks and _MAIN_TOKEN_RE.match(plus_chunks[0]):
            trailing = plus_chunks[1:]
            if len(trailing) > MAX_BREAKDOWN_CODES - 1 and all(re.fullmatch(r"\d+", c) for c in trailing) and not any(_MAIN_TOKEN_RE.match(c) for c in trailing):
                return [""], [original]

    top_chunks = [c.strip() for c in re.split(r"[,+/:]", t) if c.strip()]
    if not top_chunks:
        return [""], [original]

    pairs: list[list] = []
    current_main = None
    fallback = False

    for chunk in top_chunks:
        chunk = chunk.strip().strip("/").strip()
        chunk = re.sub(rf"^(\d{{11,21}})\s+(?=\d{{1,6}}\b)", r"\1-", chunk)
        pieces = re.split(r"(\bSD\b|&|-)", chunk)
        pieces = [p.strip() for p in pieces if p.strip()]
        if not pieces:
            continue

        i = 0
        local_main = None
        pending_op = None
        active_idx = None

        while i < len(pieces):
            piece = pieces[i]

            if piece in ("-", "&", "SD"):
                pending_op = piece
                i += 1
                continue

            if _MAIN_TOKEN_RE.match(piece):
                local_main, current_main = piece, piece
                existing_idx = next((idx for idx, p in enumerate(pairs) if p[0] == piece), None)
                if existing_idx is not None:
                    active_idx = existing_idx
                else:
                    pairs.append([piece, ""])
                    active_idx = len(pairs) - 1
                pending_op = None
                i += 1
                continue

            if piece.isdigit():
                anchor = local_main or current_main
                if active_idx is None and pairs:
                    active_idx = len(pairs) - 1

                if pending_op == "SD":
                    if _CERT_SUFFIX_RE.match(piece):
                        end_padded = _cert_pad6(piece)
                        if active_idx is not None:
                            prev = pairs[active_idx][1]
                            prev_start = prev.split(" SD ")[0] if " SD " in prev else prev
                            start_raw = prev_start.split(", ")[-1] if prev_start else None
                            if start_raw and len(start_raw) == 6 and start_raw.isdigit():
                                start_n, end_n = int(start_raw), int(end_padded)
                                if start_n <= end_n:
                                    nums = [str(n).zfill(6) for n in range(start_n, end_n + 1)]
                                    existing = prev_start.split(", ") if ", " in prev_start else [prev_start]
                                    pairs[active_idx][1] = ", ".join(existing[:-1] + nums) if len(existing) > 1 else ", ".join(nums)
                                else:
                                    pairs[active_idx][1] = f"{prev_start}, {end_padded}"
                            else:
                                pairs[active_idx][1] = f"{prev_start}, {end_padded}" if prev_start else end_padded
                        pending_op = None
                        i += 1
                        continue

                    if anchor and active_idx is not None and pairs:
                        width = len(piece)
                        last_polis = pairs[-1][0]
                        start_tail = last_polis[-width:] if len(last_polis) >= width else None
                        if start_tail and start_tail.isdigit():
                            start_n, end_n = int(start_tail), int(piece)
                            if start_n <= end_n and (end_n - start_n + 1) <= MAX_BREAKDOWN_CODES:
                                for n in range(start_n + 1, end_n + 1):
                                    suf = str(n).zfill(width)
                                    recon = anchor[: -width] + suf
                                    if recon not in [p[0] for p in pairs]:
                                        pairs.append([recon, ""])
                            else:
                                fallback = True
                                break
                    pending_op = None
                    i += 1
                    continue

                if anchor and (_CERT_SUFFIX_RE.match(piece) or (pending_op == "&" and _SHORT_SUFFIX_RE.match(piece))):
                    if pending_op == "-" and active_idx is not None and _try_extend_cert_range_pair(pairs, active_idx, piece):
                        pending_op = None
                        i += 1
                        continue
                    padded = _cert_pad6(piece)
                    if active_idx is not None:
                        prev = pairs[active_idx][1]
                        existing = [c for c in prev.split(", ") if c] if prev and " SD " not in prev else ([prev] if prev else [])
                        if padded not in existing:
                            existing.append(padded)
                            pairs[active_idx][1] = ", ".join(existing)
                    pending_op = None
                    i += 1
                    continue

                if anchor and (_SHORT_SUFFIX_RE.match(piece) or (re.fullmatch(r"\d{1,11}", piece) and len(piece) < len(anchor))):
                    if active_idx is not None and (pending_op == "-" or pairs[active_idx][1]) and _try_extend_cert_range_pair(pairs, active_idx, piece):
                        pending_op = None
                        i += 1
                        continue
                    if len(anchor) < 16:
                        reconstructed = anchor + piece
                    else:
                        reconstructed = anchor[: -len(piece)] + piece
                    existing_idx = next((idx for idx, p in enumerate(pairs) if p[0] == reconstructed), None)
                    if existing_idx is None:
                        pairs.append([reconstructed, ""])
                        active_idx = len(pairs) - 1
                    else:
                        active_idx = existing_idx
                    pending_op = None
                    i += 1
                    continue

                fallback = True
                break

            fallback = True
            break

        if fallback:
            break

    polis_codes = [p[0] for p in pairs]
    if fallback or not polis_codes:
        return [""], [t_noise_stripped]
    if len(polis_codes) > MAX_BREAKDOWN_CODES:
        return [""], [t_noise_stripped]

    cert_list = [_collapse_consecutive_cert_list(p[1]) for p in pairs]
    return cert_list, polis_codes


# ==============================================================================
# 5. FUNGSI COMPACT
# ==============================================================================
def _compact_value(value: str) -> str:
    if not isinstance(value, str): return value
    return _NON_ALNUM_RE.sub("", value)

def _compact_split_code_results(results: list, is_kept_as_is: bool) -> list:
    if is_kept_as_is:
        return results
    return [_compact_value(v) for v in results]


# ==============================================================================
# 6. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_cedant(df_raw: pd.DataFrame, keyword: str, cedant_column: str) -> pd.DataFrame:
    if cedant_column not in df_raw.columns:
        raise ValueError(f"Kolom '{cedant_column}' tidak ditemukan. Tersedia: {list(df_raw.columns)}")
    mask = df_raw[cedant_column].astype(str).str.contains(keyword, case=False, na=False)
    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant (kolom '{cedant_column}') mengandung '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered


# ==============================================================================
# 7. BUILD OUTPUT
# ==============================================================================
def _breakdown_all_rows(df: pd.DataFrame, col_insured: str, col_polis: str, col_slip: str):
    ins_arr = df[col_insured].to_numpy()
    pol_arr = df[col_polis].to_numpy()
    slp_arr = df[col_slip].to_numpy()

    insured_cln_all, polis_cln_all, slip_cln_all, cert_cln_all = [], [], [], []
    max_ins = max_pol = max_slp = 0

    for ins_val, pol_val, slp_val in zip(ins_arr, pol_arr, slp_arr):
        ins_cln = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []

        if pd.notna(pol_val) and str(pol_val).strip():
            cert_cln, pol_cln = process_policy_certificate(str(pol_val))
            pol_cln = [v.upper() if isinstance(v, str) else v for v in pol_cln]
        else:
            pol_cln, cert_cln = [], []

        if pd.notna(slp_val) and str(slp_val).strip():
            slp_results, slp_kept = clean_split_code_impl(str(slp_val))
            slp_cln = _compact_split_code_results(
                [r.upper() if isinstance(r, str) else r for r in slp_results], slp_kept
            )
        else:
            slp_cln = []

        insured_cln_all.append(ins_cln)
        polis_cln_all.append(pol_cln)
        slip_cln_all.append(slp_cln)
        cert_cln_all.append(cert_cln)

        max_ins, max_pol, max_slp = max(max_ins, len(ins_cln)), max(max_pol, len(pol_cln)), max(max_slp, len(slp_cln))

    return insured_cln_all, polis_cln_all, slip_cln_all, cert_cln_all, max_ins, max_pol, max_slp

def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_insured, col_polis, col_slip = INSURED_COLUMN, POLIS_COLUMN, SLIP_COLUMN
    for col, label in ((col_insured, "INSURED"), (col_polis, "POLIS"), (col_slip, "SLIP")):
        if col not in df_filtered.columns:
            raise ValueError(f"Kolom {label}='{col}' tidak ditemukan.")

    ins_all, pol_all, slp_all, cert_all, max_ins, max_pol, max_slp = _breakdown_all_rows(
        df_filtered, col_insured, col_polis, col_slip
    )

    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)

    if BROKER_NAME_COLUMN in df_filtered.columns and CEDANT_COLUMN in df_filtered.columns:
        is_direct = df_filtered[BROKER_NAME_COLUMN].astype(str).str.strip().str.upper() == DIRECT_MARKER
        bp_series = df_filtered[CEDANT_COLUMN].where(is_direct, df_filtered[BROKER_NAME_COLUMN])
        business_partner_values = bp_series.astype(str).str.replace(r"^(\s*PT)\.\s*", r"\1 ", regex=True, flags=RE_I).values
    else:
        business_partner_values = None

    output_cols_data = {}
    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        if col == BROKER_NAME_COLUMN and business_partner_values is not None:
            output_cols_data["BUSINESS_PARTNER"] = business_partner_values

        if col == col_insured:
            for i in range(1, max_ins + 1):
                output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in ins_all]

        if col == col_polis:
            for i in range(1, max_pol + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in pol_all]
                output_cols_data[f"CERTIFICATE_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in cert_all]

        if col == col_slip:
            for i in range(1, max_slp + 1):
                output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slp_all]

    df_master = pd.DataFrame(output_cols_data)
    drop_cols = [c for c in df_master.columns if ("_CLN_" in c or c.startswith("CERTIFICATE_")) and df_master[c].astype(str).str.strip().eq("").all()]
    return df_master.drop(columns=drop_cols)


# ==============================================================================
# 8. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    workbook = load_workbook(output_path)
    worksheet = workbook.active
    text_col_indices = [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None: cell.value = str(cell.value)
    workbook.save(output_path)


# ==============================================================================
# 9. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print("=" * 70 + "\n CLEANSING DATA 1 - FACULTATIVE | CEDANT: BRINS GENERAL INSURANCE\n" + "=" * 70)
    input_path = INPUT_FILE

    if not os.path.exists(input_path):
        candidates = [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]
        if candidates:
            input_path = candidates[0]
            print(f"[AUTO-DETECT] Memakai file terdeteksi: '{input_path}'")
        else:
            raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    print(f"[INFO] Total baris source: {len(df_raw)}")

    df_brins = filter_cedant(df_raw, CEDANT_FILTER, CEDANT_COLUMN)
    if df_brins.empty:
        print("[WARNING] Tidak ada baris yang cocok dengan filter cedant. Proses dihentikan.")
        return

    df_hasil = build_output(df_brins)
    save_with_text_format(df_hasil, OUTPUT_FILE)

    print("=" * 70 + f"\n[SUCCESS] Selesai. Total baris output: {len(df_hasil)}\n[SUCCESS] File hasil: '{OUTPUT_FILE}'\n" + "=" * 70)

if __name__ == "__main__":
    main()