"""
================================================================================
 SCRIPT CLEANSING DATA 3 - SUSPENSE (KHUSUS CEDANT: BRINS)
================================================================================
"""

import os
import re
import pandas as pd
from openpyxl import load_workbook

# ==============================================================================
# 0. KONFIGURASI
# ==============================================================================
INPUT_FILE = "3b. Database Suspense.xlsx"
SHEET_NAME = "Sheet1"
HEADER_ROW = 2
OUTPUT_FILE = "brins_output_suspend.xlsx"
CEDANT_FILTER = "BRINS"
STATUS_FILTER = "SUSPENSE"

MAX_BREAKDOWN_CODES = 5
MAX_BREAKDOWN_INSURED = 5
TEXT_FORMAT_COLUMN_KEYWORDS = ("POLIS", "SLIP", "RECEIPT NO", "CREDIT NOTES", "DETAIL RINCIAN NO", "CERTIFICATE")
RE_I = re.IGNORECASE

# ==============================================================================
# 1. CLEANSING INSURED
# ==============================================================================
_ENTITY_TITLE_RE = re.compile(
    r"\b(?:P\.?T\.?|C\.?V\.?|PERSERO|TBK\.?|LTD\.?|PTE\.?|N\.?V\.?|"
    r"BAPAK|IBU|NY\.?|NYONYA|MR\.?|MRS\.?|MS\.?)\b", RE_I
)
_PAREN_CONTENT_RE = re.compile(r"\([^)]*\)")
_OPEN_COVER_PREFIX_RE = re.compile(r"^\s*(?:PREMI\s+)?AF\s+MARINE\s+CARGO\s+", RE_I)
_OPEN_COVER_SUFFIX_RE = re.compile(
    r"\s+GROUP\s+(?:JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER|"
    r"JAN|FEB|MAR|APR|JUN|JUL|AUG|AGT|SEP|OKT|OCT|NOV|DEC|DES)\s+\d{4}\s*$", RE_I
)
_ENTITY_SPLIT_QQ_RE = re.compile(r"\bQQ\.?\b", RE_I)
_KNOWN_BANK_KEYWORDS = ["BANK", "BRI", "BNI", "BCA", "BTN", "MANDIRI", "CIMB", "PERMATA", "DANAMON", "OCBC", "MAYBANK", "PANIN"]

def _is_bank_entity(text: str) -> bool:
    upper = text.upper()
    return any(re.search(rf"\b{kw}\b", upper) for kw in _KNOWN_BANK_KEYWORDS)

def _split_insured_entities(raw: str) -> list:
    text = raw.strip()
    if not text:
        return []

    parts = [p.strip() for p in _ENTITY_SPLIT_QQ_RE.split(text) if p.strip()]
    if len(parts) < 2:
        return [text]

    non_bank_parts = [p for p in parts if not _is_bank_entity(p)]
    return non_bank_parts if non_bank_parts else parts

def _clean_single_insured(name: str) -> str:
    text = name.strip()
    text = _ENTITY_TITLE_RE.sub("", text)
    text = _PAREN_CONTENT_RE.sub("", text)
    text = _OPEN_COVER_PREFIX_RE.sub("", text)
    text = _OPEN_COVER_SUFFIX_RE.sub("", text)
    text = re.sub(r"[,;:.]+", " ", text)
    text = re.sub(r"^[-/]+|[-/]+$", "", text)
    return re.sub(r"\s+", " ", text).strip().upper()

def clean_insured_name(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none", "-"):
        return []

    original = text.strip()
    entities = _split_insured_entities(original)
    if not entities:
        return []

    cleaned = [c for c in (_clean_single_insured(e) for e in entities) if c]
    if not cleaned:
        return []
    return [original] if len(cleaned) > MAX_BREAKDOWN_INSURED else cleaned

# ==============================================================================
# 2. CLEANSING POLIS & SLIP NO
# ==============================================================================
_POLIS_INSTALLMENT_SUFFIX_RE = re.compile(r"-\d{1,2}/\d{1,2}$")
_POLIS_CERTIFICATE_SUFFIX_RE = re.compile(r"-(\d{6})$")
_POLIS_CERTIFICATE_PLUS_INSTALLMENT_RE = re.compile(r"^(.+)-(\d{6})-\d{1,2}/\d{1,2}$")

def clean_polis_code(text: str) -> tuple:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none"):
        return [], []

    t = re.sub(r"\s+", " ", text.strip())
    if t == "-":
        return ["-"], []

    combo_match = _POLIS_CERTIFICATE_PLUS_INSTALLMENT_RE.match(t)
    if combo_match:
        return [combo_match.group(1).upper()], [combo_match.group(2).upper()]

    cert_match = _POLIS_CERTIFICATE_SUFFIX_RE.search(t)
    if cert_match:
        certificate = cert_match.group(1)
        polis_base = _POLIS_CERTIFICATE_SUFFIX_RE.sub("", t)
        return [polis_base.upper()], [certificate.upper()]

    t = _POLIS_INSTALLMENT_SUFFIX_RE.sub("", t)
    return ([t.upper()] if t else []), []

def clean_slip_code(text: str) -> list:
    if not isinstance(text, str) or not text.strip() or text.strip().lower() in ("nan", "none"):
        return []
    t = re.sub(r"\s+", " ", text.strip())
    return [t.upper()] if t else []

# ==============================================================================
# 3. LOAD & FILTER DATA
# ==============================================================================
def load_source(file_path: str, sheet_name: str, header_row: int) -> pd.DataFrame:
    return pd.read_excel(file_path, sheet_name=sheet_name, header=header_row)

def filter_cedant(df_raw: pd.DataFrame, keyword: str) -> pd.DataFrame:
    cedant_cols = [c for c in df_raw.columns if "CEDANT" in str(c).upper()]
    if not cedant_cols:
        raise ValueError("Kolom CEDANT tidak ditemukan di source file.")

    mask = pd.Series(False, index=df_raw.index)
    for col in cedant_cols:
        mask |= df_raw[col].astype(str).str.contains(keyword, case=False, na=False)

    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter cedant mengandung '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered

def filter_status(df_raw: pd.DataFrame, keyword: str) -> pd.DataFrame:
    status_col = next((c for c in df_raw.columns if str(c).upper().strip() == "STATUS"), None)
    if status_col is None:
        raise ValueError("Kolom STATUS tidak ditemukan di source file.")

    mask = df_raw[status_col].astype(str).str.strip().str.upper() == keyword.upper()
    df_filtered = df_raw[mask].copy()
    print(f"[INFO] Filter status = '{keyword}': {len(df_filtered)} baris ditemukan dari total {len(df_raw)} baris.")
    return df_filtered

# ==============================================================================
# 4. BUILD OUTPUT
# ==============================================================================
def _detect_key_columns(df: pd.DataFrame) -> tuple:
    col_insured = next((c for c in df.columns if "INSURED" in str(c).upper()), None)
    col_polis = next((c for c in df.columns if str(c).upper().strip() == "POLIS"), None)
    col_slip = next((c for c in df.columns if "SLIP" in str(c).upper()), None)

    if None in (col_insured, col_polis, col_slip):
        raise ValueError(f"Kolom kunci tidak lengkap. INSURED={col_insured}, POLIS={col_polis}, SLIP={col_slip}")
    print(f"[INFO] Kolom terdeteksi -> INSURED='{col_insured}', POLIS='{col_polis}', SLIP='{col_slip}'")
    return col_insured, col_polis, col_slip

def _breakdown_all_rows(df: pd.DataFrame, col_insured: str, col_polis: str, col_slip: str) -> tuple:
    ins_arr = df[col_insured].to_numpy()
    pol_arr = df[col_polis].to_numpy()
    slp_arr = df[col_slip].to_numpy()
    
    insured_cln_all, polis_cln_all, cert_cln_all, slip_cln_all = [], [], [], []
    max_ins = max_pol = max_cert = max_slp = 0

    for ins_val, pol_val, slp_val in zip(ins_arr, pol_arr, slp_arr):
        ins_cln = clean_insured_name(str(ins_val)) if pd.notna(ins_val) else []
        pol_cln, cert_cln = clean_polis_code(str(pol_val)) if pd.notna(pol_val) else ([], [])
        slp_cln = clean_slip_code(str(slp_val)) if pd.notna(slp_val) else []

        insured_cln_all.append(ins_cln)
        polis_cln_all.append(pol_cln)
        cert_cln_all.append(cert_cln)
        slip_cln_all.append(slp_cln)
        
        max_ins, max_pol = max(max_ins, len(ins_cln)), max(max_pol, len(pol_cln))
        max_cert, max_slp = max(max_cert, len(cert_cln)), max(max_slp, len(slp_cln))

    return insured_cln_all, polis_cln_all, cert_cln_all, slip_cln_all, max_ins, max_pol, max_cert, max_slp

def build_output(df_filtered: pd.DataFrame) -> pd.DataFrame:
    col_insured, col_polis, col_slip = _detect_key_columns(df_filtered)
    insured_cln_all, polis_cln_all, cert_cln_all, slip_cln_all, max_ins, max_pol, max_cert, max_slp = _breakdown_all_rows(df_filtered, col_insured, col_polis, col_slip)

    original_cols = list(df_filtered.columns)
    df_filtered = df_filtered.reset_index(drop=True)
    output_cols_data = {}
    
    for col in original_cols:
        output_cols_data[col] = df_filtered[col].values
        if col == col_insured:
            for i in range(1, max_ins + 1):
                output_cols_data[f"INSURED_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in insured_cln_all]
        if col == col_polis:
            for i in range(1, max_pol + 1):
                output_cols_data[f"POLIS_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in polis_cln_all]
            for i in range(1, max_cert + 1):
                output_cols_data[f"CERTIFICATE_{i}" if max_cert > 1 else "CERTIFICATE"] = [lst[i - 1] if i <= len(lst) else "" for lst in cert_cln_all]
        if col == col_slip:
            for i in range(1, max_slp + 1):
                output_cols_data[f"SLIP_NO_CLN_{i}"] = [lst[i - 1] if i <= len(lst) else "" for lst in slip_cln_all]

    df_master = pd.DataFrame(output_cols_data)
    drop_cols = [c for c in df_master.columns if ("_CLN_" in c or c.startswith("CERTIFICATE")) and df_master[c].astype(str).str.strip().eq("").all()]
    return df_master.drop(columns=drop_cols)

# ==============================================================================
# 5. SIMPAN OUTPUT
# ==============================================================================
def save_with_text_format(df: pd.DataFrame, output_path: str) -> None:
    df.to_excel(output_path, index=False)
    workbook = load_workbook(output_path)
    worksheet = workbook.active
    text_col_indices = [idx for idx, col in enumerate(df.columns, start=1) if any(kw in str(col).upper() for kw in TEXT_FORMAT_COLUMN_KEYWORDS)]

    for col_idx in text_col_indices:
        for row in range(2, worksheet.max_row + 1):
            cell = worksheet.cell(row=row, column=col_idx)
            cell.number_format = "@"
            if cell.value is not None:
                cell.value = str(cell.value)
    workbook.save(output_path)

# ==============================================================================
# 6. EXECUTION RUNNER
# ==============================================================================
def main() -> None:
    print("=" * 70 + "\n CLEANSING DATA 3 - SUSPENSE | CEDANT: BRINS\n" + "=" * 70)
    input_path = INPUT_FILE

    if not os.path.exists(input_path):
        candidates = [f for f in os.listdir(".") if f.endswith(".xlsx") and not f.startswith("Hasil_") and not f.startswith("Clean_")]
        if not candidates:
            raise FileNotFoundError("File Excel sumber tidak ditemukan di folder kerja.")
        input_path = candidates[0]
        print(f"[AUTO-DETECT] File input tidak ditemukan di path default, memakai: '{input_path}'")

    df_raw = load_source(input_path, SHEET_NAME, HEADER_ROW)
    print(f"[INFO] Total baris source: {len(df_raw)}")

    df_brins = filter_cedant(df_raw, CEDANT_FILTER)
    if df_brins.empty:
        print("[WARNING] Tidak ada baris yang cocok dengan filter cedant. Proses dihentikan.")
        return

    df_brins = filter_status(df_brins, STATUS_FILTER)
    if df_brins.empty:
        print("[WARNING] Tidak ada baris yang cocok dengan filter status. Proses dihentikan.")
        return

    df_hasil = build_output(df_brins)
    save_with_text_format(df_hasil, OUTPUT_FILE)
    print("=" * 70 + f"\n[SUCCESS] Selesai. Total baris output: {len(df_hasil)}\n[SUCCESS] File hasil: '{OUTPUT_FILE}'\n" + "=" * 70)

if __name__ == "__main__":
    main()