import os
import re

import numpy as np
import pandas as pd


# ============================================================
# CONFIG
# ============================================================

INPUT_FILE = os.path.join(
    "input",
    "1b. Transaksi Facul 01.01.23 - 17.08.2026.xlsx"
)

OUTPUT_FILE = os.path.join(
    "output",
    "tripa_output_facul.xlsx"
)


# ============================================================
# HEADER FACUL DATA 1
# ============================================================

CEDANT_COL = "COMP_NAME"
CEDANT_VALUE = "PT ASURANSI TRIPAKARTA"

POLIS_COL = "FAC_POLICY_NO"
SLIP_COL = "FAC_SLIP"
INSURED_COL = "FAC_INSURED"

BROKER_NAME_COL = "COMP_NAME_1"
BROKER_CODE_COL = "FAC_BROKER"

MAX_SPLIT_COLS = 5


# ============================================================
# POLIS EXCEPTION
# ============================================================

POLIS_EXCEPTION_RE = re.compile(
    r"""
    MOP\s*MARINE
  | (?:LINE\s*SLIP|LINESLIP)
  | \b(?:P1|P2|P3|P73)\s*CANCEL
  | \b(?:P1|P2|P3|P73)\b
  | \bCANCEL\b
  | PENYELESAIAN(?:\s+SUSPENSE)?
  | HUTANG
  | UTANG
    """,
    re.IGNORECASE | re.VERBOSE,
)


# ============================================================
# SLIP EXCEPTION
# ============================================================

SLIP_EXCEPTION_RE = re.compile(
    r"""
    \bSUMMARY\b
  | \bBORDER[OA]\b
  | \bBORDRO\b
  | \bSINGGLESHIPMENT\b
  | PENYELESAIAN(?:\s+SUSPENSE)?
  | HUTANG
  | UTANG
    """,
    re.IGNORECASE | re.VERBOSE,
)


# ============================================================
# MONTH
# ============================================================

MONTH_NAMES = frozenset({
    "JANUARI", "FEBRUARI", "MARET", "APRIL", "MEI", "JUNI",
    "JULI", "AGUSTUS", "SEPTEMBER", "OKTOBER", "NOVEMBER", "DESEMBER",
    "JANUARY", "FEBRUARY", "MARCH", "MAY", "JUNE", "JULY", "AUGUST",
    "OCTOBER",
})


# ============================================================
# INSURED CONFIG
# ============================================================

INSURED_REMOVE_RE = re.compile(
    r"""
    \b(
        PT|CV|TBK|
        PERSERO|
        LTD|PTE|INC|LLC|
        MR|MRS|MS|
        BAPAK|BPK|IBU|NY|
        DR|DRS|DRA|IR|H|HJ
    )\b
    """,
    re.IGNORECASE | re.VERBOSE,
)


INSURED_POLIS_RE = re.compile(
    r"""
    (POLIS\s*NO\.?.*)
    |
    (POLICY\s*NO\.?.*)
    |
    (SLIP\s*NO\.?.*)
    """,
    re.IGNORECASE | re.VERBOSE,
)


INSURED_SPLIT_RE = re.compile(
    r"""
    \s*,\s*
    |
    \s*/\s*
    |
    \s+QQ\s+
    |
    \s+AND/OR\s+
    |
    \s*&\s*
    |
    \s*\+\s*
    """,
    re.IGNORECASE | re.VERBOSE,
)


INSURED_JUNK_WORDS = {
    "",
    "AND",
    "OR",
    "THE",
    "OF",
    "AS"
}


# ============================================================
# SLIP NOISE
# ============================================================

_SLIP_NOISE_RE = re.compile(
    r"""
    \b(?:JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS
        |SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER)\b
  | \b(?:JANUARY|FEBRUARY|MARCH|APRIL|MAY|JUNE|JULY|AUGUST
        |SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)\b
  | \b(?:IDR|USD|ENG)\b
  | \b(?:P1|P2|P3|P73)\b
  | \bNEW\b
  | \bVARIOUS\b
  | \b20[0-9]{2}\b
    """,
    re.IGNORECASE | re.VERBOSE,
)


_SLIP_TOKEN_RE = re.compile(
    r"[A-Z0-9][A-Z0-9\-]{6,}",
    re.IGNORECASE
)


# ============================================================
# HELPER
# ============================================================

def _normalize_spaces(text: str) -> str:
    return re.sub(r"\s{2,}", " ", text).strip()


def _is_valid_polis_token(tok: str) -> bool:
    tok = tok.strip()
    return (
        len(tok) >= 5
        and not re.search(r"TBA$", tok, re.IGNORECASE)
        and bool(re.search(r"\d", tok))
    )


def _clean_insured_name(name: str) -> str:
    if pd.isna(name):
        return ""

    name = str(name).upper()

    name = re.sub(
        r"\s*/\s*BORD(?:ER|ERO)?\b.*$",
        "",
        name,
        flags=re.IGNORECASE
    )

    name = INSURED_POLIS_RE.sub("", name)
    name = INSURED_REMOVE_RE.sub(" ", name)
    name = re.sub(r"[()]", " ", name)
    name = name.replace("/", " ")
    name = name.replace("-", " ")

    return _normalize_spaces(name)


def split_insured(name):
    if pd.isna(name):
        return []

    name = str(name)
    parts = INSURED_SPLIT_RE.split(name)

    hasil = []
    for p in parts:
        p = _clean_insured_name(p)
        if not p or p in INSURED_JUNK_WORDS:
            continue
        hasil.append(p)

    return hasil


def _cap_or_join(items: list) -> list:
    if len(items) > MAX_SPLIT_COLS:
        return [",".join(items)]
    return items


# ============================================================
# CLEAN CERTIFICATE (TERUPDATE)
# ============================================================
def clean_certificate(polis_ori):
    """
    Mengambil & merapikan CERTIFICATE dari POLIS ORI.
    """
    if pd.isna(polis_ori):
        return ""

    polis = str(polis_ori).strip().upper()

    if not polis:
        return ""

    # Hapus data kotor / suspense
    if re.search(r"\bSUSPENSE\b|\bENDORSEMENT\b", polis):
        return ""

    # 1. Handling Pola S/D Eksplisit (misal: 00001 S/D 00008)
    if re.search(r"\bS\s*/?\s*D\b", polis):
        parts_sd = re.split(r"\s*S\s*/?\s*D\s*", polis)
        if len(parts_sd) >= 2:
            left = parts_sd[0].strip()
            # Potong bagian kanan sampai menemukan koma/spasi pemisah polis berikutnya
            right = re.split(r"\s*,\s*|\s+/\s+", parts_sd[1])[0].strip()

            m_left = re.search(r"(?:^|[\s-])(\d{1,6})$", left)
            cert_left = m_left.group(1).zfill(6) if m_left else ""

            m_right = re.search(r"(?:^|[\s-])(\d{1,6})$", right)
            cert_right = m_right.group(1).zfill(6) if m_right else ""

            if cert_left and cert_right:
                return f"{cert_left} SD {cert_right}"

    # Bersihkan suffix / VAR / VARIOUS
    polis_clean = re.sub(r"\s*/\s*(?:VAR|VARIOUS|TBA).*$", "", polis, flags=re.IGNORECASE)

    # 2. Pola Multi Polis Dipisah Slash (contoh: 10303071700120 - 000001 / 10303071700120 - 000002)
    if "/" in polis_clean and "-" in polis_clean:
        sub_polis = polis_clean.split("/")
        certs = []
        for sp in sub_polis:
            m = re.search(r"-\s*(\d{1,6})\b", sp.strip())
            if m:
                certs.append(m.group(1).zfill(6))
        if certs:
            return ", ".join(dict.fromkeys(certs))

    # 3. Pola Dipisah Koma (contoh: 10412032000260-000033 , 000062 ,10512032000350-001)
    if "," in polis_clean:
        # Pisahkan segmen berdasarkan koma
        segments = [s.strip() for s in polis_clean.split(",")]
        certs = []
        for seg in segments:
            # Jika segmen mengandung polis baru dengan dash pendek (-001 / -02), abaikan suffix tersebut
            if re.search(r"\d{8,}-\d{1,3}$", seg):
                continue
            
            # Cari angka cert (bisa berupa -000033 atau angka berdiri sendiri 000062)
            m = re.search(r"(?:^|-|\s)(\d{4,6})$", seg)
            if m:
                certs.append(m.group(1).zfill(6))
                
        if certs:
            # Jika rentang banyak angka dengan koma (contoh No 5) -> buat A SD B
            if len(certs) > 2:
                return f"{certs[0]} SD {certs[-1]}"
            # Jika hanya 2 angka dipisah koma (contoh No 4) -> buat A, B
            return ", ".join(dict.fromkeys(certs))

    # 4. Pola Dash Bertumpuk Tanpa Koma (contoh: 10209051600001-000048-00059)
    if re.search(r"\d+-\d{4,6}-\d{4,6}", polis_clean):
        first_dash = polis_clean.find("-")
        if first_dash != -1:
            suffix_part = polis_clean[first_dash + 1:]
            tokens = re.findall(r"\b\d{4,6}\b", suffix_part)
            valid_certs = [t.zfill(6) for t in tokens]
            if valid_certs:
                return ", ".join(dict.fromkeys(valid_certs))

    # 5. Handling Single Certificate Standar
    m_single = re.search(r"-\s*(\d{4,6})$", polis_clean)
    if m_single:
        return m_single.group(1).zfill(6)

    return ""
# ============================================================
# CLEAN POLIS
# ============================================================

def clean_polis(val) -> list:
    if pd.isna(val):
        return []

    val = str(val).strip()
    if not val:
        return []

    original = val
    val = val.upper()

    if re.fullmatch(r"\d{1,6}(?:\s*[+&-]\s*\d{1,6})+", val):
        return [val]

    if re.fullmatch(r"\s*VARIOUS\s*", val, re.IGNORECASE):
        return ["VARIOUS"]

    val = re.sub(r"\s*/\s*VARI(?:OUS)?\b.*$", "", val, flags=re.IGNORECASE)
    val = re.sub(r"\s+VARI(?:OUS)?\b.*$", "", val, flags=re.IGNORECASE)
    val = _normalize_spaces(val)

    m = re.fullmatch(r"(\d+)\s*S/D\s*(\d+)", val, re.IGNORECASE)
    if m:
        start = m.group(1)
        end = m.group(2)
        if len(start) == len(end):
            start_num = int(start)
            end_num = int(end)
            jumlah = abs(end_num - start_num) + 1

            if jumlah > 5:
                return [original]

            hasil = []
            step = 1 if start_num <= end_num else -1
            for num in range(start_num, end_num + step, step):
                polis = str(num).zfill(len(start))
                if polis not in hasil:
                    hasil.append(polis)
            return hasil

    if re.search(r"\bVARIOUS\b", original, re.IGNORECASE) and re.search(r"\d{8,}", original):
        candidates = re.findall(r"\b\d{8,}\b", original)
        hasil = []
        for p in candidates:
            if p not in hasil:
                hasil.append(p)
        if hasil:
            return _cap_or_join(hasil)

    if re.fullmatch(r"\d+(?:\s*-\s*\d+)+", val):
        parts = [x.strip() for x in re.split(r"\s*-\s*", val) if x.strip()]
        base = parts[0]
        suffixes = parts[1:]

        if len(suffixes) == 1:
            return [base]

        hasil = [base]
        for suffix in suffixes:
            if len(suffix) < len(base):
                polis = base[:len(base) - len(suffix)] + suffix
            else:
                polis = suffix
            if polis not in hasil:
                hasil.append(polis)
        return _cap_or_join(hasil)

    if re.search(r"-\s*00000\d", val):
        bases = re.findall(r"\b\d{8,}\b", val)
        hasil = []
        for p in bases:
            if p not in hasil:
                hasil.append(p)
        if hasil:
            return _cap_or_join(hasil)

    if "&" in val:
        parts = re.split(r"\s*&\s*", val)
        hasil = []
        for p in parts:
            p = re.sub(r"[^A-Z0-9]", "", p.strip())
            if p:
                hasil.append(p)
        return _cap_or_join(list(dict.fromkeys(hasil)))

    if "/" in val:
        parts = re.split(r"\s*/\s*", val)
        hasil = []
        for p in parts:
            p = p.strip()
            if not p or re.fullmatch(r"\d{1,3}", p) or p.upper() in {"VARIOUS", "VAR"}:
                continue
            p = re.sub(r"[^A-Z0-9]", "", p)
            if len(p) >= 8 and re.search(r"\d", p):
                if p not in hasil:
                    hasil.append(p)
        if hasil:
            return _cap_or_join(hasil)

    m = re.fullmatch(r"(\d{7,})-(\d{1,6})(?:\+\d{1,6})+", val)
    if m:
        parts = val.split("+")
        first = parts[0]
        base, first_suffix = first.split("-", 1)
        suffixes = [first_suffix, *parts[1:]]
        hasil = []
        for suffix in suffixes:
            if suffix.upper() in {"VAR", "VARIOUS", "TBA"} or not re.fullmatch(r"\d{1,6}", suffix):
                continue
            polis = base[:len(base) - len(suffix)] + suffix if len(suffix) < len(base) else suffix
            if polis not in hasil:
                hasil.append(polis)
        return hasil

    if "+" in val:
        parts = [p.strip() for p in val.split("+") if p.strip()]
        hasil = []
        base = None

        for p in parts:
            clean = re.sub(r"[^A-Z0-9]", "", p)
            if len(clean) >= 8:
                hasil.append(clean)
                base = clean
                continue

            if base and re.fullmatch(r"\d{6}", p):
                suffix = p[-6:]
                polis = base[:-6] + suffix
                if polis not in hasil:
                    hasil.append(polis)
                base = polis
                continue

            if base and re.fullmatch(r"\d{3,4}", p):
                suffix = p
                polis = base[:len(base) - len(suffix)] + suffix
                if polis not in hasil:
                    hasil.append(polis)
                base = polis
                continue

            if re.fullmatch(r"\d{5,6}", p):
                hasil.append(p)
                base = p

        if hasil:
            return _cap_or_join(list(dict.fromkeys(hasil)))

    if re.fullmatch(r"[0-9.]+", val):
        return [val.replace(".", "")]

    if re.fullmatch(r"[0-9-]+", val):
        return [val.replace("-", "")]

    val = _normalize_spaces(val)
    return [val] if val else []


# ============================================================
# CLEAN SLIP
# ============================================================

def clean_slip(val) -> list:
    if pd.isna(val):
        return []

    val = str(val).upper().strip().replace(".", "")
    if not val:
        return []

    val = _normalize_spaces(val)

    m = re.match(r"^(\d{7,})\s*/\s*0\b", val)
    if m:
        return [m.group(1)]

    m = re.fullmatch(r"(\d{7,})\s*/\s*(\d{7,})", val)
    if m:
        return [m.group(1), m.group(2)]

    m = re.match(r"^(\d{7,})\s*-\s*\d+/CN/", val, re.IGNORECASE)
    if m:
        return [m.group(1)]

    m = re.fullmatch(r"P[123]\s*/\s*(\d{7,})", val, re.IGNORECASE)
    if m:
        return [m.group(1)]

    if re.search(r"\bSUMMARY\b|\bSINGGLESHIPMENT\b|PENYELESAIAN(?:\s+SUSPENSE)?|\bHUTANG\b|\bUTANG\b", val, re.IGNORECASE):
        return [_normalize_spaces(val)]

    m = re.fullmatch(r"(JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER|JANUARY|FEBRUARY|MARCH|MAY|JUNE|JULY|AUGUST|SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)\s+(20[0-9]{2})\s*/\s*VARIOUS", val, re.IGNORECASE)
    if m:
        return [f"{m.group(1).upper()} {m.group(2)}"]

    m = re.search(r"(\d{7,})\s*S/D\s*(\d{7,})", val, re.IGNORECASE)
    if m:
        start, end = m.group(1), m.group(2)
        if len(start) == len(end):
            start_num, end_num = int(start), int(end)
            jumlah = abs(end_num - start_num) + 1
            if jumlah > 5:
                return [f"{start} S/D {end}"]
            hasil = []
            step = 1 if start_num <= end_num else -1
            for num in range(start_num, end_num + step, step):
                slip = str(num).zfill(len(start))
                if slip not in hasil:
                    hasil.append(slip)
            return hasil

    m = re.fullmatch(r"(\d{7,})\s*/\s*(\d{1,6})", val)
    if m:
        base, suffix = m.group(1), m.group(2)
        if len(suffix) < len(base):
            slip = base[:len(base) - len(suffix)] + suffix
            return [base, slip]

    if "+" in val:
        parts = [x.strip() for x in val.split("+") if x.strip()]
        if len(parts) >= 2 and all(re.match(r"^\d{5,}/", p) for p in parts):
            hasil = []
            for p in parts:
                p = _normalize_spaces(p)
                if p not in hasil:
                    hasil.append(p)
            return _cap_or_join(hasil)

    if re.fullmatch(r"\d{1,7}(?:\s*[+&-]\s*\d{1,7})+", val):
        first_part = re.split(r"\s*[+&-]\s*", val)[0]
        if len(first_part) <= 7:
            return [val]

    if re.fullmatch(r"\d{7,}(?:\s*-\s*\d{1,6})+", val):
        parts = [p.strip() for p in re.split(r"\s*-\s*", val) if p.strip()]
        base = parts[0]
        hasil = [base]
        for suffix in parts[1:]:
            if not re.fullmatch(r"\d{1,6}", suffix):
                continue
            polis = base[:len(base) - len(suffix)] + suffix if len(suffix) < len(base) else suffix
            if polis not in hasil:
                hasil.append(polis)
        return _cap_or_join(hasil)

    if re.fullmatch(r"\d{7,}(?:\s*\+\s*\d{1,6})+", val):
        parts = [p.strip() for p in re.split(r"\s*\+\s*", val) if p.strip()]
        base = parts[0]
        hasil = [base]
        for suffix in parts[1:]:
            if not re.fullmatch(r"\d{1,6}", suffix):
                continue
            slip = base[:len(base) - len(suffix)] + suffix if len(suffix) < len(base) else suffix
            if slip not in hasil:
                hasil.append(slip)
        return _cap_or_join(hasil)

    candidates = re.findall(r"\b\d{7,}\b", val)
    if candidates:
        hasil = list(dict.fromkeys(candidates))
        return hasil if len(hasil) <= MAX_SPLIT_COLS else [val]

    val = re.sub(r"^SEE\s+ATTACHMENT\s*/\s*", "", val, flags=re.IGNORECASE).strip()

    if re.search(r"\b(EQ|PAR)/\d+", val, re.IGNORECASE) or re.fullmatch(r"MDD/FCMI/\d{2}-F\d{7}(?:&\d{4,5})+", val, re.IGNORECASE):
        return [_normalize_spaces(val)]

    val = re.sub(r"\bEND(?:\.?1)?\b", "", val, flags=re.IGNORECASE)
    val = re.sub(r"END\.?1\.?", "", val, flags=re.IGNORECASE)
    val = _normalize_spaces(val)

    if re.search(r"/CN/", val, re.IGNORECASE):
        return [val]

    if "," in val:
        parts = [x.strip() for x in val.split(",") if x.strip()]
        hasil = list(dict.fromkeys(parts))
        if hasil:
            return _cap_or_join(hasil)

    m = re.fullmatch(r"(\d{7,})/\d{1,3}", val)
    if m:
        return [m.group(1)]

    m = re.fullmatch(r"(\d{7,})-(\d{1,6})", val)
    if m:
        base, suffix = m.group(1), m.group(2)
        hasil = [base]
        slip = base[:len(base) - len(suffix)] + suffix if len(suffix) < len(base) else suffix
        if slip not in hasil:
            hasil.append(slip)
        return _cap_or_join(hasil)

    if re.fullmatch(r"\d{1,6}(?:\s*[&+-]\s*\d{1,6})+", val):
        return [val]

    m = re.fullmatch(r"(\d{7,})-(\d{1,6})(?:\+\d{1,6})+", val)
    if m:
        parts = val.split("+")
        first = parts[0]
        base, first_suffix = first.split("-", 1)
        suffixes = [first_suffix, *parts[1:]]
        hasil = []
        for suffix in suffixes:
            if suffix.upper() in {"VAR", "VARIOUS", "TBA"} or not re.fullmatch(r"\d{1,6}", suffix):
                continue
            slip = base[:len(base) - len(suffix)] + suffix if len(suffix) < len(base) else suffix
            if slip not in hasil:
                hasil.append(slip)
        return hasil

    if "+" in val:
        parts = [x.strip() for x in val.split("+") if x.strip()]
        if len(parts) >= 2 and re.fullmatch(r"\d{7,}", parts[0]):
            base = parts[0]
            hasil = [base]
            for suffix in parts[1:]:
                if re.fullmatch(r"\d{3,6}", suffix):
                    slip = base[:len(base) - len(suffix)] + suffix if len(suffix) < len(base) else suffix
                    if slip not in hasil:
                        hasil.append(slip)
                else:
                    if suffix.upper() not in {"VAR", "VARIOUS", "TBA"} and suffix not in hasil:
                        hasil.append(suffix)
            return _cap_or_join(hasil)

    if re.fullmatch(r"\d+(?:-\d+)+", val):
        parts = [x.strip() for x in val.split("-") if x.strip()]
        if len(parts) >= 2 and len(parts[0]) >= 7:
            base = parts[0]
            hasil = [base]
            for suffix in parts[1:]:
                if not re.fullmatch(r"\d{1,6}", suffix):
                    continue
                max_suffix_len = 16 - len(base)
                if max_suffix_len <= 0:
                    continue
                use_len = min(len(suffix), max_suffix_len)
                suffix_use = suffix[-use_len:]
                slip = base[:-use_len] + suffix_use
                if len(slip) <= 16 and slip not in hasil:
                    hasil.append(slip)
            if len(hasil) > 1:
                return _cap_or_join(hasil)

    if ";" in val:
        val = val.replace(";", "+")

    if re.search(r"[+&]", val):
        parts = [x.strip() for x in re.split(r"[+&]", val) if x.strip()]
        return _cap_or_join(list(dict.fromkeys(parts)))

    slips = []
    for s in re.split(r"\s*[;+&]\s*", val):
        s = s.strip()
        if not s:
            continue
        s = re.sub(r"\b(USD|IDR|SGD|EUR|JPY|AUD|GBP|ORI|ORIGINAL|COPY|REALISASI|REALIZATION|CANCEL|CANCELLED|ENDT?|ENDORSEMENT|SA|P1|P2|P3|VAR|REVISI|REV)\b.*$", "", s, flags=re.IGNORECASE)
        s = _normalize_spaces(s)
        s = re.sub(r"-(\d{2})$", "", s)
        s = s.strip("-_,.; ")
        if s and s not in slips:
            slips.append(s)

    return _cap_or_join(slips)


# ============================================================
# CLEAN INSURED
# ============================================================

def clean_insured(val) -> list:
    if pd.isna(val):
        return []

    val = str(val).strip()
    if not val:
        return []

    val_upper = val.upper()

    if re.match(r"^BIO\s+FARMA\b", val_upper, re.IGNORECASE):
        val_upper = re.sub(r"\s*/\s*", " ", val_upper)
        val_upper = re.sub(r"\s*&\s*", " ", val_upper)
        return [_normalize_spaces(val_upper)]

    insureds = split_insured(val)
    cleaned = []

    for ins in insureds:
        ins = _normalize_spaces(ins)
        if not ins:
            continue

        if re.fullmatch(r"BORD(?:ER|ERO)?\s+(?:JANUARI|FEBRUARI|MARET|APRIL|MEI|JUNI|JULI|AGUSTUS|SEPTEMBER|OKTOBER|NOVEMBER|DESEMBER|JANUARY|FEBRUARY|MARCH|MAY|JUNE|JULY|AUGUST|SEPTEMBER|OCTOBER|NOVEMBER|DECEMBER)(?:\s+20[0-9]{2})?", ins, flags=re.IGNORECASE):
            continue

        if len(ins) <= 2 or ins.upper() in INSURED_JUNK_WORDS:
            continue

        ins = ins.upper()
        if ins not in cleaned:
            cleaned.append(ins)

    if not cleaned:
        fallback = _clean_insured_name(val)
        if fallback:
            cleaned.append(fallback.upper())

    return _cap_or_join(cleaned)


# ============================================================
# BUSINESS PARTNERS
# ============================================================

def clean_business_partners(val):
    if pd.isna(val):
        return ""
    val = str(val).strip().upper()
    val = re.sub(r"\bPT\.\s*", "PT ", val, flags=re.IGNORECASE)
    return _normalize_spaces(val)


def get_mitra_bisnis(broker_name, broker_code, cedant) -> str:
    broker_name = "" if pd.isna(broker_name) else str(broker_name).strip()
    broker_code = "" if pd.isna(broker_code) else str(broker_code).strip()
    cedant = "" if pd.isna(cedant) else str(cedant).strip()

    if broker_name and broker_name.upper() != "DIRECT":
        return clean_business_partners(broker_name)

    if broker_code and broker_code.upper() != "DIRECT":
        return clean_business_partners(broker_code)

    return clean_business_partners(cedant)


# ============================================================
# INSERT CLEAN COLUMNS
# ============================================================

def _insert_clean_columns(df: pd.DataFrame, all_lists: list, prefix: str, max_cols: int) -> list:
    added = []
    for i in range(1, max_cols + 1):
        col_name = f"clean {prefix} {i}"
        df[col_name] = [lst[i - 1] if i - 1 < len(lst) else None for lst in all_lists]
        added.append(col_name)
    return added


# ============================================================
# FAST READ & WRITE EXCEL
# ============================================================

def _fast_read_excel(path: str, sheet_name=None, header: int = 0) -> pd.DataFrame:
    import openpyxl

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[sheet_name] if sheet_name else wb.worksheets[0]
    rows_iter = ws.iter_rows(values_only=True)

    for _ in range(header):
        next(rows_iter)

    cols = list(next(rows_iter))
    seen = {}
    new_cols = []

    for col in cols:
        col = "" if col is None else str(col).strip()
        if col not in seen:
            seen[col] = 0
            new_cols.append(col)
        else:
            seen[col] += 1
            new_cols.append(f"{col}_{seen[col]}")

    data = list(rows_iter)
    wb.close()
    return pd.DataFrame(data, columns=new_cols)


def _fast_write_excel(df: pd.DataFrame, path: str) -> None:
    import openpyxl

    output_dir = os.path.dirname(path)
    if output_dir:
        os.makedirs(output_dir, exist_ok=True)

    df2 = df.astype(object).where(pd.notnull(df), None)
    wb = openpyxl.Workbook(write_only=True)
    ws = wb.create_sheet("Sheet1")
    ws.append(list(df2.columns))

    for row in df2.itertuples(index=False, name=None):
        ws.append(row)

    wb.save(path)


# ============================================================
# PROCESS DATA
# ============================================================

def process_data(input_file: str, output_file: str) -> None:
    print(f"[1/5] Membaca data dari: {input_file} ...")
    df = _fast_read_excel(input_file, header=0)
    print(f"      Total baris keseluruhan: {len(df):,}")

    if CEDANT_COL not in df.columns:
        print(f"\n[ERROR] Kolom '{CEDANT_COL}' tidak ditemukan!")
        return

    required_cols = [POLIS_COL, SLIP_COL, INSURED_COL, BROKER_NAME_COL, BROKER_CODE_COL]
    for col in required_cols:
        if col not in df.columns:
            print(f"\n[ERROR] Kolom '{col}' tidak ditemukan di file!")
            return

    df[CEDANT_COL] = (
        df[CEDANT_COL]
        .fillna("")
        .astype(str)
        .str.upper()
        .str.strip()
        .str.replace(r"\.", "", regex=True)
        .str.replace(r"\s+", " ", regex=True)
    )

    cedant_value_normalized = CEDANT_VALUE.upper().replace(".", "")
    is_tripa = df[CEDANT_COL] == cedant_value_normalized
    total_tripa = int(is_tripa.sum())

    print(f"\n[2/5] Filter cedant '{CEDANT_VALUE}': {total_tripa:,} baris TRIPA dari total {len(df):,} baris.")

    broker_name_s = df[BROKER_NAME_COL].fillna("").astype(str).str.strip()
    broker_code_s = df[BROKER_CODE_COL].fillna("").astype(str).str.strip()
    cedant_s = df[CEDANT_COL].fillna("").astype(str).str.strip()

    mitra_values = [
        get_mitra_bisnis(bn, bc, cd)
        for bn, bc, cd in zip(broker_name_s, broker_code_s, cedant_s)
    ]

    df.rename(
        columns={
            POLIS_COL: "polis_ori",
            SLIP_COL: "slip_ori",
            INSURED_COL: "insured_ori"
        },
        inplace=True
    )

    insert_pos = list(df.columns).index(BROKER_NAME_COL) + 1 if BROKER_NAME_COL in df.columns else len(df.columns)
    df.insert(insert_pos, "BUSINESS PARTNERS", mitra_values)

    print("\n[3/5] Menjalankan proses cleaning hanya untuk baris TRIPA ...")

    n = len(df)
    all_clean_polis = [[] for _ in range(n)]
    all_clean_slip = [[] for _ in range(n)]
    all_clean_ins = [[] for _ in range(n)]
    all_certificates = ["" for _ in range(n)]  # PENAMPUNG CERTIFICATE

    max_polis = 1
    max_slip = 1
    max_ins = 1

    tripa_idx = np.flatnonzero(is_tripa.to_numpy())
    polis_vals = df["polis_ori"].to_numpy()
    slip_vals = df["slip_ori"].to_numpy()
    insured_vals = df["insured_ori"].to_numpy()

    for n_done, pos in enumerate(tripa_idx, 1):
        if n_done % 5_000 == 0:
            print(f"      Progress: {n_done:,} / {total_tripa:,} baris TRIPA diproses...")

        p_ori = polis_vals[pos]
        c_polis = clean_polis(p_ori)
        c_slip = clean_slip(slip_vals[pos])
        c_ins = clean_insured(insured_vals[pos])
        c_cert = clean_certificate(p_ori)  # PANGGIL CLEAN CERTIFICATE

        max_polis = max(max_polis, len(c_polis))
        max_slip = max(max_slip, len(c_slip))
        max_ins = max(max_ins, len(c_ins))

        all_clean_polis[pos] = c_polis
        all_clean_slip[pos] = c_slip
        all_clean_ins[pos] = c_ins
        all_certificates[pos] = c_cert  # SIMPAN KE LIST

    print("\n      Selesai diproses!")
    print(f"      -> Jumlah kolom clean polis  : {max_polis}")
    print(f"      -> Jumlah kolom clean slip   : {max_slip}")
    print(f"      -> Jumlah kolom clean insured: {max_ins}")

    print("\n[4/5] Menyusun kolom output ...")

    # MASUKKAN CERTIFICATE KE DATAFRAME
    df["CERTIFICATE"] = all_certificates

    new_columns = []
    for col in df.columns:
        if col == "CERTIFICATE":
            continue

        new_columns.append(col)

        if col == "polis_ori":
            polis_cols = _insert_clean_columns(
                df, all_clean_polis, "polis", max_polis
            )

            new_columns += polis_cols

            # CERTIFICATE di sebelah kanan clean polis 1
            if "clean polis 1" in polis_cols:
                idx = new_columns.index("clean polis 1") + 1
                new_columns.insert(idx, "CERTIFICATE")

        elif col == "slip_ori":
            new_columns += _insert_clean_columns(df, all_clean_slip, "slip", max_slip)

        elif col == "insured_ori":
            new_columns += _insert_clean_columns(df, all_clean_ins, "insured", max_ins)

    df = df[new_columns]

    print(f"\n[5/5] Menyimpan hasil ke: {output_file} ...")
    _fast_write_excel(df, output_file)

    print(f"\n{'=' * 60}")
    print("  [OK] Selesai!")
    print(f"  Total baris output  : {len(df):,}")
    print(f"  Total kolom output  : {len(df.columns)}")
    print(f"  Total baris TRIPA   : {total_tripa:,}")
    print(f"  File disimpan di    : {output_file}")
    print(f"{'=' * 60}")


# ============================================================
# ENTRY POINT
# ============================================================

if __name__ == "__main__":
    print("MASUK MAIN")
    process_data(INPUT_FILE, OUTPUT_FILE)